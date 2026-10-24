"""GB 19153—2019 第三方复核：第三文本引擎(MuPDF)逐源格比对 + 结构/单调/压力列/边界候选复算。

只读脚本；不写入被复核的任何文件。
依赖：pymupdf(PyMuPDF)、openpyxl（非 harness 自带依赖，需另装）。
用法：python REVIEW_20261013_third_engine_check.py
"""
import hashlib
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

import pymupdf

ROOT = Path(__file__).resolve().parents[1]          # docs/standards/gb19153_2019
REPO = ROOT.parents[2]                              # 仓库根
PDF = REPO / "标准原文" / "5. GB 19153-2019 容积式空气压缩机能效限定值及能效等级.pdf"
XLSX = ROOT / "data" / "GB19153_2019_标准表值.xlsx"

EXPECT_PDF_SHA = "ab4ff8e101da57ba0b5a778fa2882cb5f4eb8c2c5d403829cac83d1cef1c09cc"
EXPECT_XLSX_SHA = "7686bbc675e3f3189b4d21898738dd4f4303e57cf66e0df8db52bbdc89dd7dc7"
TABLES = ["表1", "表2", "表3", "表4", "表5", "表6"]
HEADER_PAGE = {"表1": 5, "表2": 8, "表3": 12, "表4": 12, "表5": 14, "表6": 15}

FULL = {chr(0xFF10 + i): chr(0x30 + i) for i in range(10)}
FULL[chr(0xFF0E)] = "."


def norm(s):
    """全角转半角、统一破折号/方括号、去空白。"""
    s = "".join(FULL.get(c, c) for c in str(s if s is not None else ""))
    for d in "—–―−":
        s = s.replace(d, "-")
    return re.sub(r"\s+", "", s.replace("［", "[").replace("］", "]"))


def sha256(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def main():
    print("=== 0. 文件身份 ===")
    pdf_sha, xlsx_sha = sha256(PDF), sha256(XLSX)
    print("  PDF :", pdf_sha, "OK" if pdf_sha == EXPECT_PDF_SHA else "!! 与记录不一致")
    print("  XLSX:", xlsx_sha, "OK" if xlsx_sha == EXPECT_XLSX_SHA else "!! 与记录不一致")

    doc = pymupdf.open(PDF)
    print("  PDF 页数:", doc.page_count)
    words = {i + 1: doc[i].get_text("words") for i in range(doc.page_count)}
    wb = load_workbook(XLSX, data_only=True)

    print("\n=== 1. 结构计数 ===")
    stat = {}
    for name in TABLES:
        ws = wb[name]
        vals = dash = 0
        keys = Counter()
        groups = defaultdict(dict)
        for r in range(6, ws.max_row + 1):
            g = ws.cell(row=r, column=7).value
            if isinstance(g, (int, float)):
                vals += 1
                groups[(ws.cell(row=r, column=1).value, ws.cell(row=r, column=4).value,
                        ws.cell(row=r, column=5).value)][ws.cell(row=r, column=3).value] = g
            else:
                dash += 1
            keys[tuple(ws.cell(row=r, column=c).value for c in (1, 2, 3, 4, 5))] += 1
        full = sum(1 for v in groups.values() if len(v) == 3)
        nonmono = sum(1 for v in groups.values()
                      if len(v) == 3 and not (v[1] <= v[2] <= v[3]))
        dup = sum(1 for v in keys.values() if v > 1)
        stat[name] = (vals, dash)
        print(f"  {name}: 行={ws.max_row - 5} 数值={vals} 破折号={dash} "
              f"完整三档组={full} 非单调={nonmono} 重复键={dup}")
    tv = sum(v[0] for v in stat.values())
    td = sum(v[1] for v in stat.values())
    print(f"  合计：数值={tv}（期望 2502） 破折号={td}（期望 417）")

    print("\n=== 2. 第三引擎(MuPDF)逐源格字面比对 ===")
    ok = bad = 0
    for name in TABLES:
        ws = wb[name]
        for r in range(6, ws.max_row + 1):
            page = ws.cell(row=r, column=9).value
            bbox = ws.cell(row=r, column=13).value
            lit = norm(ws.cell(row=r, column=14).value)
            if page is None or bbox is None:
                continue
            x0, y0, x1, y1 = [float(v) for v in str(bbox).split(",")]
            got = ""
            for w in words[int(page)]:
                if w[0] < x1 + 1.5 and w[2] > x0 - 1.5 and w[1] < y1 + 2.0 and w[3] > y0 - 2.0:
                    got += w[4]
            if lit and lit in norm(got):
                ok += 1
            else:
                bad += 1
                if bad <= 10:
                    print(f"  !!! {name} 行{r} PDF{page} 期望{lit!r} 实得{norm(got)[:40]!r}")
    print(f"  一致={ok} 不符={bad}（期望 2919/0）")

    print("\n=== 3. 压力→列中心对齐（原页表头数字中心 vs 工作簿列中心）===")
    bands = {"表1": (150, 170), "表2": (150, 170), "表3": (150, 205),
             "表4": (390, 410), "表5": (140, 175), "表6": (140, 175)}
    for name in TABLES:
        pno = HEADER_PAGE[name]
        lo, hi = bands[name]
        labels = sorted((round((w[0] + w[2]) / 2, 1), float(norm(w[4])))
                        for w in words[pno]
                        if lo <= w[1] <= hi and re.fullmatch(r"\d+(\.\d+)?", norm(w[4])))
        ws = wb[name]
        # 先找出「同一源格覆盖多个压力」的横跨合并格（如表4 0.55~1.5 kW 的 1.25/1.4 MPa），
        # 这些 bbox 会污染列跨度，必须剔除；单列内的「—」可以保留。
        owners = defaultdict(set)
        for r in range(6, ws.max_row + 1):
            if ws.cell(row=r, column=9).value != pno:
                continue
            owners[ws.cell(row=r, column=13).value].add(ws.cell(row=r, column=4).value)
        span = {}
        for r in range(6, ws.max_row + 1):
            if ws.cell(row=r, column=9).value != pno:
                continue
            p = ws.cell(row=r, column=4).value
            bb = ws.cell(row=r, column=13).value
            if len(owners[bb]) > 1:
                continue
            x0, _, x1, _ = [float(v) for v in str(bb).split(",")]
            s = span.setdefault(p, [x0, x1])
            s[0], s[1] = min(s[0], x0), max(s[1], x1)
        # 表1/表2 每个压力含风冷＋液冷两列，表头数字居中于「列对」
        centers = sorted((round((a + b) / 2, 1), float(p)) for p, (a, b) in span.items())
        same = len(labels) == len(centers) and all(
            abs(lx - cx) <= 0.5 and lv == cv for (lx, lv), (cx, cv) in zip(labels, centers))
        print(f"  {name}: 表头中心={labels}")
        print(f"       列中心  ={centers}  → {'精确一致(±0.5pt)' if same else '!! 不一致'}")

    print("\n=== 4. 边界候选抽检（值取自工作簿，值本身已由第 2 步对原文核验）===")

    def get(table, power, pressure, branch=None, grade=None):
        ws = wb[table]
        out = []
        for r in range(6, ws.max_row + 1):
            if ws.cell(row=r, column=1).value != power or ws.cell(row=r, column=4).value != pressure:
                continue
            if branch is not None and ws.cell(row=r, column=5).value != branch:
                continue
            if grade is not None and ws.cell(row=r, column=3).value != grade:
                continue
            out.append((ws.cell(row=r, column=3).value, ws.cell(row=r, column=7).value))
        return out

    checks = [
        ("C01 表1 1.5kW/0.3/风冷", get("表1", 1.5, 0.3, "风冷"), [5.8, 6.5, 7.4]),
        ("C05 表2 2.2kW/0.3/风冷", get("表2", 2.2, 0.3, "风冷"), [6.1, 6.8, 7.8]),
        ("C05b 表1 2.2kW/0.3/风冷", get("表1", 2.2, 0.3, "风冷"), [5.6, 6.2, 7.0]),
        ("C07 表3 2.2kW/0.25/3级", get("表3", 2.2, 0.25, None, 3), [6.9]),
        ("C07b 表3 3kW/0.25/3级", get("表3", 3, 0.25, None, 3), ["—"]),
        ("C08 表3 18.5kW/0.5", get("表3", 18.5, 0.5), [6.2, 6.6, 7.3]),
        ("C10 表4 18.5kW/0.7 两分支", get("表4", 18.5, 0.7), [7.3, 6.7, 7.8, 7.1, 8.8, 7.4]),
        ("C14 表5 0.25kW/0.2/1级", get("表5", 0.25, 0.2, None, 1), [5.2]),
        ("C14b 表6 0.25kW/0.2/1级", get("表6", 0.25, 0.2, None, 1), [5.5]),
        ("C15 表5 1.5kW/0.4/3级", get("表5", 1.5, 0.4, None, 3), [9.0]),
        ("C15b 表5 1.8kW/0.4/3级", get("表5", 1.8, 0.4, None, 3), [9.1]),
    ]
    def same_value(a, b):
        if isinstance(a, str) or isinstance(b, str):
            return str(a) == str(b)
        return abs(float(a) - float(b)) < 1e-9

    for label, rows, expect in checks:
        got = [v for _, v in rows]
        flag = "OK" if len(got) == len(expect) and all(
            same_value(x, y) for x, y in zip(got, expect)) else "!! 不符"
        print(f"  {label}: 实际={got} 期望={expect} → {flag}")
    print("  表5!G127 =", wb["表5"].cell(row=127, column=7).value,
          "（字面", wb["表5"].cell(row=127, column=14).value, "）")
    print("  表5!G148 =", wb["表5"].cell(row=148, column=7).value,
          "（字面", wb["表5"].cell(row=148, column=14).value, "）")

    print("\n=== 5. 表4 方括号液冷与合并破折号覆盖 ===")
    ws = wb["表4"]
    print("  分支分布:", Counter((ws.cell(row=r, column=5).value, ws.cell(row=r, column=6).value)
                                for r in range(6, ws.max_row + 1)))
    bbox = Counter()
    for r in range(6, ws.max_row + 1):
        if ws.cell(row=r, column=6).value == "源表破折号":
            bbox[ws.cell(row=r, column=13).value] += 1
    print("  破折号按合并源格统计（组合数, bbox）:",
          [(c, b) for b, c in bbox.items()])

    print("\n=== 6. 公式单元格 ===")
    f = sum(1 for name in wb.sheetnames for row in wb[name].iter_rows()
            for c in row if c.data_type == "f")
    print("  公式单元格数:", f, "（期望 0）")


if __name__ == "__main__":
    main()
