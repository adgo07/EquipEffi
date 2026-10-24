"""Read-only workbook/CSV verification; writes only this directory's audit result."""
import csv, hashlib, json
from collections import Counter
from decimal import Decimal
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parents[2]
def rows(name):
    with (root / name).open(encoding='utf-8-sig', newline='') as stream:
        return list(csv.DictReader(stream))
values = rows('data/table_values.csv')
audit = rows('evidence/cell_audit.csv')
powers = rows('evidence/row_audit.csv')
book = openpyxl.load_workbook(root/'data/GB30253_2024_TABLE_VALUES.xlsx', data_only=False)
checks = []
def check(name, condition, detail):
    checks.append(dict(name=name, status='PASS' if condition else 'FAIL', detail=detail))
check('sheet_order', book.sheetnames == ['使用说明与索引']+[f'表{i:02d}' for i in range(1,30)], len(book.sheetnames))
check('coverage', len(values)==10803 and len(audit)==10803 and len(powers)==707, dict(values=len(values),audit=len(audit),power_rows=len(powers)))
expected = {i:567 if i==1 else 315 if i<=7 else 392 if i<=10 else 375 if i<=28 else 420 for i in range(1,30)}
check('table_cell_counts', dict(Counter(int(r['table']) for r in values))==expected, expected)
mismatch=[]
for r in values:
    actual=book[r['sheet']][r['cell']].value
    source=r['efficiency']
    same=actual==source if source=='—' else isinstance(actual,(int,float)) and Decimal(str(actual))==Decimal(source)
    if not same: mismatch.append([r['sheet'],r['cell'],source,actual])
check('xlsx_efficiency_values', not mismatch, dict(compared=len(values),mismatches=mismatch))
power_mismatch=[]
for r in powers:
    actual=book[r['sheet']][r['power_cell']].value
    expected_power=r['power_kW']
    same=str(actual)==expected_power
    if isinstance(actual,(int,float)): same=Decimal(str(actual))==Decimal(expected_power)
    if not same: power_mismatch.append([r['sheet'],r['power_cell'],expected_power,actual])
check('xlsx_power_rows', not power_mismatch, dict(compared=len(powers),mismatches=power_mismatch))
check('two_pdf_coordinate_reads', all(r['comparison']=='MATCH' for r in audit), Counter(r['comparison'] for r in audit))
formulas=[f'{s.title}!{c.coordinate}' for s in book for row in s for c in row if c.data_type=='f']
check('no_formulas', not formulas, formulas)
check('source_anomaly_preserved', book['表01']['N23'].value==0 and book['表01']['G23'].value=='—' and book['表01']['U23'].value=='—', 'N23=0.0; G23/U23=—')
check('dash_count', sum(r['efficiency']=='—' for r in values)==152, sum(r['efficiency']=='—' for r in values))
files=['data/GB30253_2024_TABLE_VALUES.xlsx','data/table_values.csv','evidence/cell_audit.csv','evidence/row_audit.csv']
result=dict(checks=checks, totals=dict(pass_count=sum(c['status']=='PASS' for c in checks),fail_count=sum(c['status']=='FAIL' for c in checks),error_count=0,skip_count=0), sha256={f:hashlib.sha256((root/f).read_bytes()).hexdigest() for f in files})
(root/'evidence/verification_result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
print(json.dumps(result,ensure_ascii=False,indent=2))
assert result['totals']['fail_count']==0
