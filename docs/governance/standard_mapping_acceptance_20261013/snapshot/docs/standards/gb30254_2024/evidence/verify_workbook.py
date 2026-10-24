"""Read-only XLSX/source verification; openpyxl is used only to inspect saved cells."""
from pathlib import Path
import json,csv,re,hashlib
from decimal import Decimal
from collections import Counter
from pypdf import PdfReader
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

BASE=Path(__file__).resolve().parents[1]
ROOT=BASE.parents[2]
SOURCE=ROOT/'标准原文/3. GB 30254-2024 高压三相笼型异步电动机能效限定值及能效等级.pdf'
XLSX=BASE/'data/GB30254_2024_标准表值.xlsx'
rows=json.loads((BASE/'_work/decoded_rows.json').read_text(encoding='utf-8'))
wb=load_workbook(XLSX,data_only=False)
pdf=PdfReader(SOURCE)
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
tokens={p:re.findall(r'/G[0-9A-F]{2}',pdf.pages[p-1].extract_text()) for p in range(7,19)}
source_sha=sha(SOURCE)
offsets=Counter(); rownums=Counter(); checks=[]; stats={}; second_reader=[]
for r in rows:
    t,p=r['table'],r['pdf_page']; rownums[t]+=1; xr=rownums[t]+7
    raw=['/G'+format(0x27 if c=='’' else ord(c),'02X') for word in r['raw'] for c in word]
    # Compare a full row of original pypdf glyph tokens with independently positioned pdfplumber row.
    stream=tokens[p]; hits=[i for i in range(offsets[p],len(stream)-len(raw)+1) if stream[i:i+len(raw)]==raw]
    second_ok=bool(hits)
    if second_ok: offsets[p]=hits[0]+len(raw)
    else: second_reader.append(dict(pdf_page=p,power=r['values'][0],raw=raw))
    sh=wb[f'表{t}']; assert sh.max_row==7+sum(x['table']==t for x in rows)
    for j,(rawtoken,expected) in enumerate(zip(r['raw'],r['values']),1):
        c=sh.cell(xr,j); actual=c.value
        ok=(actual==expected) if expected=='—' else Decimal(str(actual))==Decimal(expected)
        fmt='0.0' if '.' in expected else '0'
        format_ok=expected=='—' or c.number_format==fmt
        checks.append(dict(source_sha256=source_sha,table=t,pdf_page=p,printed_page=r['printed_page'],page_data_row=r['page_row'],power_kw=r['values'][0],pole='' if j==1 else 2*((j-2)//3+1),grade='' if j==1 else (j-2)%3+1,sheet=sh.title,cell=c.coordinate,source_glyph=rawtoken,source_display=expected,xlsx_value=actual,xlsx_number_format=c.number_format,value_match=ok,display_precision_match=format_ok,second_pdf_reader_row_match=second_ok,method='glyph-decoded PDF text to saved XLSX; not independent raster transcription'))
        assert ok and format_ok,(t,c.coordinate,expected,actual,c.number_format)
    assert [sh.cell(xr,len(r['values'])+k).value for k in (1,2,3)]==[p,r['printed_page'],r['page_row']]
assert list(rownums.values())==[43,42,35,35,25,24],rownums
assert not second_reader, second_reader
for t in range(1,7):
    tr=[r for r in rows if r['table']==t]
    assert len({r['values'][0] for r in tr})==len(tr)
    ps=[int(r['values'][0]) for r in tr]; assert ps==sorted(ps)
    for r in tr:
        for g in range(1,len(r['values']),3):
            v=r['values'][g:g+3]
            assert all(x=='—' for x in v) or all(x!='—' for x in v)
            if v[0]!='—': assert Decimal(v[0])>=Decimal(v[1])>=Decimal(v[2])
    stats[t]=dict(rows=len(tr),numeric_efficiency_cells=sum(v!='—' for r in tr for v in r['values'][1:]),dash_cells=sum(v=='—' for r in tr for v in r['values'][1:]),xlsx_range=f'表{t}!A8:{get_column_letter(len(tr[0]["values"]))}{7+len(tr)}')
assert len(wb.sheetnames)==6
assert not any(c.data_type=='f' for s in wb for row in s for c in row)
with (BASE/'evidence/cell_verification.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(checks[0]));w.writeheader();w.writerows(checks)
summary=dict(date='2026-10-13',executor='Codex GB30254 Stage A session',source_sha256=sha(SOURCE),workbook_sha256=sha(XLSX),tables=stats,checked_cells=len(checks),numeric_efficiency_cells=sum(v['numeric_efficiency_cells'] for v in stats.values()),dash_cells=sum(v['dash_cells'] for v in stats.values()),power_cells=len(rows),xlsx_value_failures=sum(not x['value_match'] for x in checks),display_precision_failures=sum(not x['display_precision_match'] for x in checks),second_reader_unmatched_rows=second_reader,formula_count=0,independent_acceptance='NOT_RUN',visual_numeric_coverage='glyph anchors and visual sample only; not full independent cell transcription')
(BASE/'evidence/verification_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in summary.items() if k!='second_reader_unmatched_rows'},ensure_ascii=False,indent=2)); print('unmatched second-reader rows',len(second_reader))
