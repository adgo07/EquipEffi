"""Read-only PDF extraction; writes only this standard directory. No runtime data import."""
from pathlib import Path
import csv, hashlib, json, re
from pypdf import PdfReader

BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[2]
PDF = ROOT / '标准原文/9. GB 28381-2012 离心鼓风机能效限定值及节能评价值.pdf'
SHA = '07410053c5b190193169686f54b5f86fb243a9a069c3f6c3e0f085637c9b1a75'
assert hashlib.sha256(PDF.read_bytes()).hexdigest() == SHA
pages = [p.extract_text() for p in PdfReader(PDF).pages]
(BASE/'tmp').mkdir(exist_ok=True)
(BASE/'data').mkdir(exist_ok=True)
(BASE/'evidence').mkdir(exist_ok=True)
(BASE/'tmp/extracted.txt').write_text('\n\n'.join(pages),encoding='utf-8')
tables={i:[] for i in range(1,9)}
current=0
for physical, txt in enumerate(pages,1):
    # Table title (including continued title), not narrative references.
    for line in txt.splitlines():
        title=re.match(r'^表\s*([1-8])\s',line)
        if title:
            current=int(title.group(1))
        nums=re.findall(r'(?<![\d.])\d{2}\.\d(?!\d)',line)
        if current and len(nums)==5:
            tables[current].append({'values':nums,'pdf_page':physical,'printed_page':physical-2})
ratios=['<0.020','0.021~0.030','0.031~0.040','0.041~0.050','0.051~0.060','0.061~0.080','0.081~0.090','>0.091']
diameters=['<300','301~400','401~600','601~800','>801']
names=['单级双支撑低速','多级低速','单级双支撑高速','多级高速']
cells=[]
for t,rows in tables.items():
    assert len(rows)==(8 if t%2 else 16),(t,len(rows))
    for r,row in enumerate(rows,1):
        row.update(source_row=r,ratio=ratios[(r-1)//(1 if t%2 else 2)], stages=('1' if t%2 else (['2~3','4~6'] if t in [2,6] else ['2','3~4'])[(r-1)%2]),branch=names[(t-1)%4])
        for c,v in enumerate(row['values'],1):
            cells.append({'source_id':f'GB28381-2012-T{t}-R{r:02}-C{c}', 'table':t,'source_row':r,'source_column':c,'branch':row['branch'],'ratio':row['ratio'],'stages':row['stages'],'diameter_mm':diameters[c-1],'value_lexical':v,'unit':'%','pdf_page':row['pdf_page'],'printed_page':row['printed_page'],'xlsx_sheet':f'表{t}','xlsx_cell':f'{chr(67+c)}{r+5}','source_sha256':SHA})
(BASE/'data/tables.json').write_text(json.dumps({'standard':'GB 28381-2012','source_sha256':SHA,'diameters_mm':diameters,'tables':tables},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
with (BASE/'evidence/source_cells.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(cells[0]));w.writeheader();w.writerows(cells)
print(json.dumps({'tables':len(tables),'physical_rows':sum(map(len,tables.values())),'threshold_cells':len(cells),'source_sha256':SHA}))

if (BASE/'data/GB28381_2012_TABLES.xlsx').exists():
    from openpyxl import load_workbook
    wb=load_workbook(BASE/'data/GB28381_2012_TABLES.xlsx',data_only=False)
    results=[]
    for cell in cells:
        actual=wb[cell['xlsx_sheet']][cell['xlsx_cell']].value
        results.append({**cell,'actual':actual,'result':'PASS' if actual==float(cell['value_lexical']) else 'FAIL','method':'PDF text to exported XLSX (same extraction lineage; not independent visual audit)'})
    with (BASE/'evidence/xlsx_cell_check.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(results[0]));w.writeheader();w.writerows(results)
    checks={'table_sheets':len(wb.sheetnames),'expected_table_sheets':8,'threshold_pass':sum(r['result']=='PASS' for r in results),'threshold_fail':sum(r['result']=='FAIL' for r in results),'formula_count':sum(c.data_type=='f' for ws in wb for row in ws for c in row),'blank_threshold_cells':sum(r['actual'] is None for r in results),'source_sha256':SHA,'workbook_sha256':hashlib.sha256((BASE/'data/GB28381_2012_TABLES.xlsx').read_bytes()).hexdigest()}
    (BASE/'evidence/check_results.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(checks))
    assert checks['threshold_fail']==0 and checks['table_sheets']==8 and checks['formula_count']==0
