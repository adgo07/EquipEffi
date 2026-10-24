"""Read-only XLSX checks against original lexical data and source-cell audit.
openpyxl is used only to inspect the saved artifact-tool workbook, never author it.
"""
from pathlib import Path
from decimal import Decimal
import json, hashlib, re
import openpyxl
ROOT=Path(__file__).resolve().parent
p=ROOT/'data/GB19577_2024_TABLE_VALUES.xlsx'
wb=openpyxl.load_workbook(p,data_only=False)
d=json.loads((ROOT/'data/table_values.json').read_text(encoding='utf-8'))
audit=json.loads((ROOT/'evidence/cell_audit.json').read_text(encoding='utf-8'))
checks=[]
def check(kind,name,actual,expected):
    checks.append(dict(kind=kind,name=name,actual=actual,expected=expected,result='pass' if actual==expected else 'fail'))
check('structure','Sheet names',wb.sheetnames,['表'+str(i) for i in range(1,9)]+['表A.1','表A.2'])
for a in audit:
    c=wb[a['sheet']][a['cell']]
    expected=Decimal(a['expected_lexical'])
    if a['table']=='A.1' and a['cell'].startswith('B'): expected/=100
    check('numeric',f"{a['sheet']}!{a['cell']}",str(Decimal(str(c.value))),str(expected))
    # Decimal lexical equality ignores trailing zeros; formats retain display source precision.
    checks[-1]['result']='pass' if Decimal(str(c.value))==expected else 'fail'
    check('numeric_type',f"{a['sheet']}!{a['cell']}",c.data_type,'n')
    lex=a['expected_lexical']
    fmt='0%' if a['table']=='A.1' and a['cell'].startswith('B') else '0.'+'0'*len(lex.split('.')[1]) if '.' in lex else '0'
    check('format',f"{a['sheet']}!{a['cell']}",c.number_format,fmt)
for t,rows in d['tables'].items():
    s=wb['表'+t]
    for r in rows:
        n=r['excel_row']
        for col,field in [('B','standard'),('C','branch'),('D','capacity'),('E','metric')]: check('routing',f'{s.title}!{col}{n}',s[f'{col}{n}'].value,r[field])
    check('structure',s.title+' freeze',s.freeze_panes,'A6')
    check('structure',s.title+' gridlines',s.sheet_view.showGridLines,False)
for n in range(10,16):check('blank_semantics',f'表7!J{n}',wb['表7'][f'J{n}'].value,None)
for n in range(6,10):check('dash_semantics',f'表A.1!F{n}',wb['表A.1'][f'F{n}'].value,'—')
check('hours','sum',sum(wb['表A.2'][f'C{n}'].value for n in range(6,21)),1007)
formulas=[];errors=[]
for s in wb:
    for row in s:
        for c in row:
            if c.data_type=='f':formulas.append(f'{s.title}!{c.coordinate}')
            if c.data_type=='e':errors.append(f'{s.title}!{c.coordinate}')
check('artifact','business formulas',formulas,[])
check('artifact','Excel errors',errors,[])
text=(ROOT/'BOUNDARY_CANDIDATES.md').read_text(encoding='utf-8')
candidates=re.findall(r'^\|(B-[A-Z0-9-]+)\|',text,re.M)
check('structure','unique boundary candidate IDs',len(set(candidates)),len(candidates))
summary={k:sum(x['result']==k for x in checks) for k in ['pass','fail','error','skip','not_run']}
summary.update(numeric_cells=sum(x['kind']=='numeric' for x in checks),energy_source_rows=69,sheets=len(wb.sheetnames),boundary_candidates=len(candidates),workbook_sha256=hashlib.sha256(p.read_bytes()).hexdigest(),independent_acceptance='not_run',production_tests='not_run')
(ROOT/'evidence/workbook_checks.json').write_text(json.dumps(dict(summary=summary,checks=checks),ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False))
assert summary['fail']==0,[x for x in checks if x['result']!='pass'][:10]
