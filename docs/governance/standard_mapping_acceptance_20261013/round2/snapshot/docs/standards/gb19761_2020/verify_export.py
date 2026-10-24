"""Read-only XLSX verification; does not author or resave the workbook."""
from pathlib import Path
import hashlib,json,re
import openpyxl
root=Path(__file__).resolve().parent
rows=json.loads((root/'data/table_values.json').read_text(encoding='utf-8'))
xlsx=root/'data/GB19761_2020_表1-4核对工作簿.xlsx'
wb=openpyxl.load_workbook(xlsx,data_only=False)
assert wb.sheetnames==['表1','表2','表3','表4']
matches=0
for r in rows:
    s=wb[r['sheet']]; n=r['xlsx_row']
    expected=[r['id'],r['condition'],r['ns'],r['size'],*[r[k] if r[k] is not None else '—' for k in ['g3','g2','g1']],r['metric'],r['source'],r['note']]
    for c,value in enumerate(expected,1):
        actual=s.cell(n,c).value
        assert (actual or '')==(value or ''),(s.title,n,c,actual,value)
        matches+=1
    for c in [5,6,7]:
        if r['g3'] is not None: assert s.cell(n,c).data_type=='n'
formulas=[f'{s.title}!{c.coordinate}' for s in wb for row in s for c in row if c.data_type=='f']
assert not formulas
assert sum(isinstance(c.value,(int,float)) for s in wb for row in s.iter_rows(min_row=5,min_col=5,max_col=7) for c in row)==402
assert wb['表2']['E51'].value==wb['表2']['F51'].value==wb['表2']['G51'].value=='—'
assert len({r['id'] for r in rows})==135
result={'xlsx_sha256':hashlib.sha256(xlsx.read_bytes()).hexdigest(),'sheet_count':4,
        'export_cell_matches':matches,'numeric_cells':402,'dash_cells':3,'business_formulas':len(formulas),
        'fail':0,'error':0,'independent_source_review':'NOT_RUN','native_excel_open':'NOT_RUN'}
(root/'evidence/export_check.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
