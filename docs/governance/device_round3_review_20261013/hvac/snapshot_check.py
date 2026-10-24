from pathlib import Path
import openpyxl,json,hashlib,re,collections
R=Path.cwd();O=R/'docs/governance/device_round3_review_20261013/hvac'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
rows=[]
for key in ['gb19577_2024','gb19576_2019','gb21454_2021','gb37479_2019']:
 p=next((R/'docs/standards'/key/'data').glob('*.xlsx'));snap=next((R/'_acceptance_20261013/round3/evidence_snapshots').glob(key+'__*.txt'))
 wb=openpyxl.load_workbook(p,data_only=True);wbf=openpyxl.load_workbook(p,data_only=False)
 lines=['FILE '+str(p.relative_to(R/'docs/standards'/key))];count=0
 for ws in wb:
  lines+=['',f'==== SHEET {ws.title}  max_row={ws.max_row} max_col={ws.max_column} merged={[str(r) for r in ws.merged_cells.ranges]}']
  for row in ws.iter_rows():
   vals=[]
   for c in row:
    v=c.value;fv=wbf[ws.title][c.coordinate].value
    if v is None and fv is None:continue
    st='' if v is None else str(v)
    if isinstance(fv,str) and fv.startswith('='):st+=' {F=%s}'%fv
    vals.append('%s=%s'%(c.coordinate,st.replace('\n','\\n')));count+=1
   if vals:lines.append('R%-4d | %s'%(row[0].row,' | '.join(vals)))
 current='\n'.join(lines);old=snap.read_text('utf-8-sig');target=O/(key+'_current_grid.txt');target.write_text(current,encoding='utf8')
 rows.append(dict(standard=key,workbook_path=str(p),workbook_sha256=sha(p),snapshot_path=str(snap),snapshot_sha256=sha(snap),nonempty_cell_count=count,normalized_text_equal=current.rstrip()==old.rstrip(),limit='文本值/Sheet尺寸/合并区；不是OOXML和样式全比对；快照时序不由此认证'))
(O/'snapshot_comparison.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf8');print(rows)
owner=json.loads((O/'owner_unresolved_inventory.json').read_text('utf8'))
for x in owner:
 x.update(id=x['existing_issue_id'] if x['review_item_id'].startswith('OWNER') else 'LOCAL-HVAC-'+x['review_item_id'],source_location=x['original_excerpt_and_location'],ambiguity=x['impact'],options=x['options_and_recommendation'],recommendation=x['options_and_recommendation'],source_opinion_ids=[x.get('origin',x['existing_issue_id'])])
(O/'owner_items.json').write_text(json.dumps(owner,ensure_ascii=False,indent=2),encoding='utf8')
