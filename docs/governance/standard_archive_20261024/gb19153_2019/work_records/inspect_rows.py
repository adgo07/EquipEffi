import warnings
warnings.filterwarnings("ignore")
from openpyxl import load_workbook

WB = r"G:\Python Project\EquipEffi\docs\standards\gb19153_2019\data\GB19153_2019_标准表值.xlsx"
wb = load_workbook(WB, data_only=True)

ws = wb["表4"]
print("=== 表4 rows 225..308 ===")
for r in range(225, ws.max_row + 1):
    print(r, [ws.cell(row=r, column=c).value for c in range(1, 15)])

print()
print("=== 表4 colF/colE 分布 ===")
import collections
print(collections.Counter((ws.cell(row=r, column=5).value, ws.cell(row=r, column=6).value) for r in range(6, ws.max_row + 1)))

ws2 = wb["表2"]
print()
print("=== 表2 rows 285..296 ===")
for r in range(285, 297):
    print(r, [ws2.cell(row=r, column=c).value for c in range(1, 15)])
