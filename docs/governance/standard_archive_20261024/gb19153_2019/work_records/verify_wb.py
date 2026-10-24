"""独立核验 GB19153 工作簿结构（只读）。"""
import sys, json, collections
from openpyxl import load_workbook

WB = r"G:\Python Project\EquipEffi\docs\standards\gb19153_2019\data\GB19153_2019_标准表值.xlsx"
wb = load_workbook(WB, data_only=False, read_only=False)
print("SHEETS:", wb.sheetnames)

expected_rows = {"表1": 1013, "表2": 761, "表3": 437, "表4": 308, "表5": 215, "表6": 215}
summary = {}
for name in wb.sheetnames:
    ws = wb[name]
    print("---", name, "dims=", ws.dimensions, "max_row=", ws.max_row, "max_col=", ws.max_column)
    # header row 5
    hdr = [ws.cell(row=5, column=c).value for c in range(1, ws.max_column + 1)]
    print("   row5:", hdr)
    r6 = [ws.cell(row=6, column=c).value for c in range(1, ws.max_column + 1)]
    print("   row6:", r6)
    # count formulas
    f = 0
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                f += 1
    print("   formula cells:", f)

for name in ["表1", "表2", "表3", "表4", "表5", "表6"]:
    ws = wb[name]
    vals = 0
    dashes = 0
    blanks = 0
    other = 0
    keys = collections.Counter()
    last_row = 0
    for r in range(6, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if a is None and ws.cell(row=r, column=7).value is None:
            continue
        last_row = r
        g = ws.cell(row=r, column=7).value
        if g is None:
            blanks += 1
        elif isinstance(g, str) and g.strip() in ("—", "-", "–"):
            dashes += 1
        elif isinstance(g, (int, float)):
            vals += 1
        else:
            other += 1
        key = tuple(ws.cell(row=r, column=c).value for c in (1, 2, 3, 4, 5))
        keys[key] += 1
    dup = {k: v for k, v in keys.items() if v > 1}
    print(f"{name}: rows_to={last_row} values={vals} dash={dashes} blank={blanks} other={other} dup_keys={len(dup)}")
    if dup:
        print("   sample dup:", list(dup.items())[:3])
    print("   colH units:", collections.Counter(ws.cell(row=r, column=8).value for r in range(6, ws.max_row + 1)).most_common(6))
    print("   colE branch:", collections.Counter(ws.cell(row=r, column=5).value for r in range(6, ws.max_row + 1)).most_common(8))
    print("   colF sourcekind:", collections.Counter(ws.cell(row=r, column=6).value for r in range(6, ws.max_row + 1)).most_common(8))
    print("   colI/J pages:", collections.Counter((ws.cell(row=r, column=9).value, ws.cell(row=r, column=10).value) for r in range(6, ws.max_row + 1)).most_common(30))
    pw = collections.Counter(ws.cell(row=r, column=1).value for r in range(6, ws.max_row + 1))
    print("   distinct power:", len(pw), sorted([x for x in pw if isinstance(x, (int, float))]))
    pr = collections.Counter(ws.cell(row=r, column=4).value for r in range(6, ws.max_row + 1))
    print("   distinct pressure:", len(pr), sorted([x for x in pr if isinstance(x, (int, float))]))
    gr = collections.Counter(ws.cell(row=r, column=3).value for r in range(6, ws.max_row + 1))
    print("   grades:", gr.most_common())

ws = wb["表7权重与说明"]
print("表7 dims", ws.dimensions, ws.max_row, ws.max_column)
for r in range(1, ws.max_row + 1):
    print("   ", r, [ws.cell(row=r, column=c).value for c in range(1, min(ws.max_column, 8) + 1)])
