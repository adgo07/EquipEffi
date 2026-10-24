"""GB19153 映射/工作簿综合核验：第三引擎逐格 + 单调性 + 边界候选数值。"""
import re
from collections import Counter, defaultdict
from openpyxl import load_workbook
import pymupdf

PDF = r"G:\Python Project\EquipEffi\标准原文\5. GB 19153-2019 容积式空气压缩机能效限定值及能效等级.pdf"
WB = r"G:\Python Project\EquipEffi\docs\standards\gb19153_2019\data\GB19153_2019_标准表值.xlsx"

FULL = {chr(0xFF10 + i): chr(0x30 + i) for i in range(10)}
FULL.update({chr(0xFF0E): ".", chr(0x3000): "", "．": "."})


def norm(s):
    if s is None:
        return ""
    s = "".join(FULL.get(ch, ch) for ch in str(s))
    for d in "—–―−":
        s = s.replace(d, "-")
    s = s.replace("［", "[").replace("］", "]")
    return re.sub(r"\s+", "", s)


doc = pymupdf.open(PDF)
words = {}
for i in range(doc.page_count):
    words[i + 1] = doc[i].get_text("words")

wb = load_workbook(WB, data_only=True)

# ---------- 1. 第三引擎逐格核对 ----------
print("=== 1. MuPDF 逐格核对（bbox 相交 + 子串包含）===")
tot = [0, 0]
for name in ["表1", "表2", "表3", "表4", "表5", "表6"]:
    ws = wb[name]
    ok = bad = 0
    fails = []
    for r in range(6, ws.max_row + 1):
        page = ws.cell(row=r, column=9).value
        bbox = ws.cell(row=r, column=13).value
        lit = norm(ws.cell(row=r, column=14).value)
        if page is None or bbox is None:
            continue
        x0, y0, x1, y1 = [float(v) for v in str(bbox).split(",")]
        got = []
        for w in words[int(page)]:
            wx0, wy0, wx1, wy1, wt = w[0], w[1], w[2], w[3], w[4]
            if wx0 < x1 + 1.5 and wx1 > x0 - 1.5 and wy0 < y1 + 2.0 and wy1 > y0 - 2.0:
                got.append(wt)
        joined = norm("".join(got))
        if lit and lit in joined:
            ok += 1
        else:
            bad += 1
            if len(fails) < 6:
                fails.append((r, int(page), lit, joined[:40]))
    print(f"  {name}: pass={ok} fail={bad}")
    for f in fails:
        print("     FAIL row=%s page=%s expect=%r got=%r" % f)
    tot[0] += ok
    tot[1] += bad
print("  TOTAL pass=%d fail=%d" % tuple(tot))

# ---------- 2. 三档单调性 ----------
print("\n=== 2. 三档单调性 L1<=L2<=L3 ===")
for name in ["表1", "表2", "表3", "表4", "表5", "表6"]:
    ws = wb[name]
    groups = defaultdict(dict)
    for r in range(6, ws.max_row + 1):
        g = ws.cell(row=r, column=7).value
        key = (ws.cell(row=r, column=1).value, ws.cell(row=r, column=4).value,
               ws.cell(row=r, column=5).value)
        grade = ws.cell(row=r, column=3).value
        if isinstance(g, (int, float)):
            groups[key][grade] = g
    bad = 0
    partial = 0
    for k, v in groups.items():
        if len(v) == 3:
            if not (v[1] <= v[2] <= v[3]):
                bad += 1
                if bad <= 5:
                    print("   NON-MONOTONIC", k, v)
        else:
            partial += 1
    print(f"  {name}: groups={len(groups)} complete3={len(groups)-partial} partial={partial} nonmonotonic={bad}")

# ---------- 3. 边界候选数值核对 ----------
print("\n=== 3. BOUNDARY_CANDIDATES 数值核对 ===")


def got(table, power, pressure, branch=None, grade=None, col=7):
    ws = wb[table]
    res = []
    for r in range(6, ws.max_row + 1):
        if ws.cell(row=r, column=1).value != power:
            continue
        if ws.cell(row=r, column=4).value != pressure:
            continue
        if branch is not None and ws.cell(row=r, column=5).value != branch:
            continue
        if grade is not None and ws.cell(row=r, column=3).value != grade:
            continue
        res.append((ws.cell(row=r, column=3).value, ws.cell(row=r, column=5).value,
                    ws.cell(row=r, column=7).value, r))
    return res


checks = [
    ("C01 表1 1.5kW 0.3MPa 风冷", got("表1", 1.5, 0.3, "风冷"), [5.8, 6.5, 7.4]),
    ("C05 表2 2.2kW 0.3MPa 风冷", got("表2", 2.2, 0.3, "风冷"), [6.1, 6.8, 7.8]),
    ("C05b 表1 2.2kW 0.3MPa 风冷", got("表1", 2.2, 0.3, "风冷"), [5.6, 6.2, 7.0]),
    ("C07 表3 2.2kW 0.25MPa 3级", got("表3", 2.2, 0.25, None, 3), [6.9]),
    ("C07b 表3 3kW 0.25MPa 3级", got("表3", 3, 0.25, None, 3), ["—"]),
    ("C08 表3 18.5kW 0.5MPa 全档", got("表3", 18.5, 0.5), [6.2, 6.6, 7.3]),
    ("C08b 表3 22kW 0.5MPa 全档", got("表3", 22, 0.5), ["—", "—", "—"]),
    ("C10 表4 18.5kW 0.7MPa", got("表4", 18.5, 0.7), ["7.3/6.7", "7.8/7.1", "8.8/7.4"]),
    ("C11 表4 0.55kW 1.25/1.4MPa", got("表4", 0.55, 1.25) + got("表4", 0.55, 1.4), ["—"] * 6),
    ("C14 表5 0.25kW 0.2MPa 1级", got("表5", 0.25, 0.2, None, 1), [5.2]),
    ("C14b 表6 0.25kW 0.2MPa 1级", got("表6", 0.25, 0.2, None, 1), [5.5]),
    ("C15 表5 1.5kW 0.4MPa 3级", got("表5", 1.5, 0.4, None, 3), [9.0]),
    ("C15b 表5 1.8kW 0.4MPa 3级", got("表5", 1.8, 0.4, None, 3), [9.1]),
]
for label, rows, expect in checks:
    vals = [v[2] for v in rows]
    print(f"  {label}: rows={[(v[3], v[1], v[2]) for v in rows]} expect={expect}")
print("  表5!G127 =", wb["表5"].cell(row=127, column=7).value,
      "| 表5!G148 =", wb["表5"].cell(row=148, column=7).value)
print("  表5 row127 =", [wb["表5"].cell(row=127, column=c).value for c in range(1, 15)])
print("  表5 row148 =", [wb["表5"].cell(row=148, column=c).value for c in range(1, 15)])

# ---------- 4. 表1/2 液冷破折号覆盖 ----------
print("\n=== 4. 表1/2 低功率液冷破折号覆盖 ===")
for name, powers in (("表1", [1.5, 2.2, 3, 4, 5.5, 7.5, 11, 15]),
                     ("表2", [2.2, 3, 4, 5.5, 7.5, 11, 15])):
    ws = wb[name]
    cnt = 0
    for r in range(6, ws.max_row + 1):
        if ws.cell(row=r, column=5).value == "液冷" and ws.cell(row=r, column=6).value == "源表破折号" \
                and ws.cell(row=r, column=1).value in powers:
            cnt += 1
    print(f"  {name}: 低功率液冷破折号组合数={cnt}（期望 {len(powers)*6*3}）")

# ---------- 5. 页脚页码 ----------
print("\n=== 5. 各 PDF 页可见页脚 ===")
for pno in range(3, 19):
    pg = doc[pno - 1]
    ws_ = [w for w in pg.get_text("words") if w[1] > 750]
    ws_.sort(key=lambda w: w[0])
    print(f"  PDF{pno}: " + " | ".join("%s@%.0f" % (w[4], w[0]) for w in ws_))
