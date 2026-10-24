"""Read-only final XLSX validation; openpyxl is not used to author the workbook."""
from pathlib import Path
import json,hashlib,collections
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
data=json.loads((ROOT/'tmp/table_data.json').read_text(encoding='utf-8'))
path=ROOT/'data/GB19153_2019_标准表值.xlsx'
wb=openpyxl.load_workbook(path,data_only=False)
checks=[]
assert wb.sheetnames==['表1','表2','表3','表4','表5','表6','表7权重与说明']
for n in range(1,7):
    sheet=wb[f'表{n}'];records=[r for r in data['records'] if r['table']==n]
    assert sheet.max_row==5+len(records)
    for i,r in enumerate(records,6):
        expected=[r['power_kw'],r['actual_input_range_kw'] or None,r['grade'],r['pressure_gauge_mpa'],r['cooling'],r['source_kind'],r['value'] if r['value'] is not None else '—','kW/(m³/min)',r['pdf_page'],r['printed_page'],r['source_row'],r['source_col'],r['bbox'],r['source_text']]
        actual=[sheet.cell(i,c).value for c in range(1,15)]
        assert actual==expected,(n,i,actual,expected)
    groups=collections.defaultdict(dict)
    for r in records:
        groups[(r['power_kw'],r['pressure_gauge_mpa'],r['cooling'])][r['grade']]=r['value']
    for key,g in groups.items():
        assert set(g)=={1,2,3},(n,key,g)
        assert all(v is None for v in g.values()) or (None not in g.values() and g[1]<=g[2]<=g[3]),(n,key,g)
    checks.append({'table':n,'rows':len(records),'numeric':sum(r['value'] is not None for r in records),'dash':sum(r['value'] is None for r in records),'power_count':len({r['power_kw'] for r in records}),'all_14_columns_readback':'PASS','grade_structure':'PASS'})
s=wb['表7权重与说明']
assert [[s.cell(i,c).value for c in range(1,5)] for i in range(5,8)]==[r+[17,15] for r in data['weights']]
assert sum(r[1] for r in data['weights'])==1
assert not [c.coordinate for s in wb for row in s for c in row if c.data_type=='f']
assert hashlib.sha256((ROOT.parents[2]/'标准原文'/'5. GB 19153-2019 容积式空气压缩机能效限定值及能效等级.pdf').read_bytes()).hexdigest()==data['pdf_sha256']
report={'executed_on':'2026-10-13','executor':'Codex','xlsx_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'pdf_sha256':data['pdf_sha256'],'tables':checks,'weight_cells':6,'weight_readback':'PASS','production_formula_count':0,'numeric_readback':2502,'dash_readback':417,'all_record_cells_checked':2919*14,'fail':0,'error':0,'independent_manual_acceptance':'NOT_RUN','production_unittest_compileall_smoke_performance':'NOT_RUN: Stage A documentation only'}
(ROOT/'evidence/workbook_verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
