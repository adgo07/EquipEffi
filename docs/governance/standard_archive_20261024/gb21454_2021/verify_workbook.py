"""Read-only XLSX verification. Checks export fidelity, NOT independent source truth."""
from pathlib import Path
import csv, hashlib, json
from decimal import Decimal
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
root=Path(__file__).resolve().parent
data=json.loads((root/'data/table_values.json').read_text(encoding='utf-8'))
p=root/'data/GB21454_2021_标准表值.xlsx'
wb=load_workbook(p,data_only=False)
assert wb.sheetnames==[t['name'] for t in data['tables']]
evidence=root/'evidence';evidence.mkdir(exist_ok=True)
audit=[]; counts={}; fails=[]
for t in data['tables']:
    s=wb[t['name']]; numeric=0; symbol=0
    for ri,row in enumerate(t['rows'],6):
        for ci,expected in enumerate(row,1):
            cell=s.cell(ri,ci); actual=cell.value
            ok=actual==expected
            if isinstance(expected,(int,float)):
                numeric+=1
                ok=cell.data_type=='n' and Decimal(str(actual))==Decimal(str(expected))
            if expected=='—':symbol+=1
            src=f"{t['name']} 数据行{ri-5} {t['headers'][ci-1]}"
            page=str(t['page']);printed=str(t['page']-2)
            if t['name'] in ['表A.3','表A.4']:
                src=row[3]+' / '+t['headers'][ci-1]
                page='9' if 'PDF9 ' in row[3] else '10';printed=str(int(page)-2)
            if t['name'] in ['表A.1','表A.2'] and ci in [5,6]:
                src=f"{t['name']} 使用侧合并源格A～E / {t['headers'][ci-1]}"
            audit.append([data['source_sha256'],t['name'],page,printed,src,cell.coordinate,str(expected),str(actual),'MATCH' if ok else 'MISMATCH','执行者页图转录自核 + XLSX回读；非独立验收'])
            if not ok:fails.append(f'{t["name"]}!{cell.coordinate}')
    for ci,h in enumerate(t['headers'],1):assert s.cell(5,ci).value==h
    if 'sourceTotal' in t:
        assert sum(r[2] for r in t['rows'])==t['sourceTotal']
        c=s.cell(6+len(t['rows']),3)
        assert c.value==t['sourceTotal']
        audit.append([data['source_sha256'],t['name'],'10','8',t['name']+' 原文总计',c.coordinate,str(t['sourceTotal']),str(c.value),'MATCH','原文总计自核 + 行求和 + XLSX回读'])
    counts[t['name']]={'data_rows':len(t['rows']),'data_cells':sum(map(len,t['rows'])),'numeric_export_cells':numeric,'dash_cells':symbol}
with (evidence/'cell_audit.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.writer(f);w.writerow(['source_sha256','sheet','pdf_physical_page','printed_page','source_cell_locator','xlsx_cell','expected_transcribed','actual_xlsx','export_result','method']);w.writerows(audit)
formulas=[f'{s.title}!{c.coordinate}' for s in wb for row in s for c in row if c.data_type=='f']
errors=[f'{s.title}!{c.coordinate}' for s in wb for row in s for c in row if c.data_type=='e']
summary={'date':'2026-10-13','source_sha256':data['source_sha256'],'xlsx_sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'tables':counts,'audit_cells_including_two_totals':len(audit),'export_mismatches':fails,'formula_cells':formulas,'error_cells':errors,'independent_source_acceptance':'NOT_RUN','native_excel_check':'NOT_RUN','production_tests':'NOT_RUN'}
(evidence/'check_results.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
assert not fails and not formulas and not errors
print(json.dumps(summary,ensure_ascii=False,indent=2))
