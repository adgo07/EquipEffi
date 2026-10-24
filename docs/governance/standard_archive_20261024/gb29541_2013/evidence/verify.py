"""Read-only PDF/XLSX verification. Writes evidence only inside own standard folder."""
from pathlib import Path
from decimal import Decimal
import csv, hashlib, json, re
from pypdf import PdfReader
from openpyxl import load_workbook

base=Path(__file__).resolve().parents[1]
root=base.parents[2]
pdf=root/'标准原文/14. GB 29541 2013 热泵热水机（器）能效限定值及能效等级.pdf'
xlsx=base/'data/GB29541_2013_标准表值.xlsx'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(pdf)=='94d5f6cc5502e76da01fc030a00d2acd54f40dd83cb967886149fb6b158e15df'
d=PdfReader(pdf)
assert len(d.pages)==5
table_text=re.split(r'\n4\.2\s',d.pages[3].extract_text())[0]
pdf_values=re.findall(r'\b\d+\.\d{2}\b',table_text)
assert len(pdf_values)==45, pdf_values
# Same-author visual-check list; not an independent transcription.
visual_rows=[
 '4.60 4.40 4.10 3.90 3.70',
 '4.20 4.00 3.80 3.60 3.40',
 '3.80 3.60 3.40 3.20 3.00',
 '4.60 4.40 4.10 3.90 3.70',
 '4.60 4.40 4.10 3.90 3.70',
 '4.50 4.30 4.00 3.80 3.60',
 '3.90 3.70 3.50 3.30 3.10',
 '3.90 3.70 3.50 3.30 3.10',
 '3.80 3.60 3.40 3.20 3.00',
]
visual=[x for row in visual_rows for x in row.split()]
assert visual==pdf_values
wb=load_workbook(xlsx,data_only=False,read_only=False)
assert wb.sheetnames==['表1_能效限值','条款数值与条件']
records=[]
for i,expected in enumerate(visual):
 r=6+i//5; col=6+i%5
 cell=wb.worksheets[0].cell(r,col)
 actual=Decimal(str(cell.value))
 assert actual==Decimal(expected)
 assert cell.data_type=='n' and cell.number_format=='0.00'
 records.append([sha(pdf),'表1',4,'2',f'T1-R{i//5+1:02}',f'等级{i%5+1}',f'表1_能效限值!{cell.coordinate}',expected,format(actual,'.2f'),pdf_values[i],'原页复看+独立文本提取+保存后坐标复读','MATCH','Codex本会话','2026-10-13'])
clauses=[('表1',4,'2','C6','10'),('4.1',3,'1','C7','5'),('4.2',4,'2','C8','10'),('4.3',4,'2','C9','0.95'),('4.3',4,'2','C10','0.95'),('4.3',4,'2','C11','1.10'),('5',4,'2','C12','5'),('6',4,'2','C13','2'),('7',4,'2','C14','2'),('8.2',5,'无印刷页码','C15','1'),('8.2',5,'无印刷页码','C16','2')]
for clause,page,printed,coord,expected in clauses:
 cell=wb.worksheets[1][coord]
 assert Decimal(str(cell.value))==Decimal(expected) and cell.data_type=='n'
 records.append([sha(pdf),clause,page,printed,'条款数值',cell.offset(column=-1).value,f'条款数值与条件!{coord}',expected,str(cell.value),'原页数字/汉字数词','原页复看+保存后坐标复读','MATCH','Codex本会话','2026-10-13'])
with (base/'evidence/cell_verification.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f);w.writerow(['pdf_sha256','source_clause','pdf_page','printed_page','source_row','source_column_or_item','xlsx_coordinate','source_expected','xlsx_actual','text_check','method','result','reviewer','date']);w.writerows(records)
formulas=[f'{ws.title}!{c.coordinate}' for ws in wb for row in ws for c in row if c.data_type=='f']
assert not formulas
report={'date':'2026-10-13','pdf_sha256':sha(pdf),'xlsx_sha256':sha(xlsx),'pdf_pages':5,'tables':1,'physical_data_rows':9,'grade_columns':5,'table_numeric_cells_matched':45,'clause_numeric_cells_matched':11,'formula_count':0,'independent_acceptance':False,'checks':{'pass':7,'fail':0,'error':0,'skip':0,'not_run':['引用GB/T21362-2008和GB/T23137-2008全文核验','独立验收','生产unittest/compileall/smoke/性能基线（本次无生产改动）','Excel桌面原生打开（本次静态值工作簿）']}}
(base/'evidence/verification_results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
