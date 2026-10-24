"""只读XLSX，逐坐标核对原格审计CSV；不通过openpyxl写工作簿。"""
from pathlib import Path
from decimal import Decimal
import csv, hashlib, json
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
book=ROOT/'data/GB18613_2020_表1-6校对数据_A1.xlsx'
sha=hashlib.sha256(book.read_bytes()).hexdigest()
wb=load_workbook(book,data_only=False)
source=list(csv.DictReader((ROOT/'evidence/source_cell_audit.csv').open(encoding='utf-8-sig')))
checks=[]
for r in source:
    c=wb[r['sheet']][r['cell']]
    raw=r['source_lexeme']; kind=r['source_kind']
    if kind=='NUMBER':
        equal=c.data_type=='n' and Decimal(str(c.value))==Decimal(raw)
        fmt=('0.'+'0'*len(raw.split('.')[1])) if '.' in raw else '0'
        if int(r['source_column'])>1: fmt='0.0'
        equal=equal and c.number_format==fmt
    elif kind=='EMPTY_SOURCE': equal=c.value is None
    else: equal=c.value==raw
    checks.append(dict(sheet=r['sheet'],cell=r['cell'],table=r['table'],pdf_page=r['pdf_page'],
                       source_body_row=r['source_body_row'],source_column=r['source_column'],
                       expected=raw,actual='' if c.value is None else str(c.value),
                       expected_kind=kind,actual_type=c.data_type,number_format=c.number_format,
                       result='PASS' if equal else 'FAIL',workbook_sha256=sha))
with (ROOT/'evidence/workbook_cell_check.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(checks[0]));w.writeheader();w.writerows(checks)
errors=[x for x in checks if x['result']!='PASS']
assert not errors, errors[:5]
formula_count=sum(c.data_type=='f' for s in wb for row in s for c in row)
error_cells=[f'{s.title}!{c.coordinate}' for s in wb for row in s for c in row if c.data_type=='e']
assert formula_count==0 and not error_cells
summary=dict(workbook_sha256=sha,source_sha256=source[0]['source_sha256'],
             actual_sheets=wb.sheetnames,checked_cells=len(checks),pass_count=len(checks),
             fail_count=0,error_count=0,skip_count=0,not_run_count=0,
             formula_count=formula_count,error_cells=error_cells,
             frozen_headers={s.title:s.freeze_panes for s in wb},
             verification='self-check of exported workbook; not independent acceptance',
             independent_acceptance='NOT_RUN',native_excel_open='NOT_RUN')
(ROOT/'evidence/workbook_check.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False))
