"""Read-only XLSX verification and reproducible visual spotcheck reconciliation."""
from pathlib import Path
from decimal import Decimal
import csv, json, hashlib
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parent
book=ROOT/'GB28381_2026_TABLE_VALUES.xlsx'
wb=load_workbook(book,read_only=False,data_only=False)
with (ROOT/'CELL_AUDIT.csv').open(encoding='utf-8-sig') as f: audit=list(csv.DictReader(f))
assert wb.sheetnames==[f'表{i}' for i in range(1,7)]
numeric=dashes=0
for r in audit:
    name,coord=r['xlsx_cell'].split('!');cell=wb[name][coord];expected=r['source_token']
    if expected=='—':
        assert cell.value=='—',(name,coord,cell.value);dashes+=1
    else:
        assert isinstance(cell.value,(int,float)),(name,coord,'not numeric')
        assert Decimal(str(cell.value))==Decimal(expected),(name,coord,cell.value,expected)
        precision=len(expected.split('.')[1]) if '.' in expected else 0
        fmt='0.'+'0'*precision if precision else '0'
        assert cell.number_format==fmt,(name,coord,cell.number_format,fmt)
        numeric+=1
    row=cell.row
    assert Decimal(str(wb[name][f'A{row}'].value))==Decimal(r['flow'])
    if int(r['table'])>=4: assert wb[name][f'B{row}'].value==int(r['grade'])
    page_col={1:5,2:5,3:5,4:13,5:12,6:17}[int(r['table'])]
    assert wb[name].cell(row,page_col).value==int(r['pdf_page'])
    assert wb[name].cell(row,page_col+1).value==int(r['printed_page'])
    assert wb[name].cell(row,page_col+2).value==int(r['pdf_grid_row'])
formula_count=sum(c.data_type=='f' for sh in wb for row in sh for c in row)
error_count=sum(c.data_type=='e' for sh in wb for row in sh for c in row)
assert formula_count==error_count==0
# Expectations manually transcribed from displayed original page images, not JSON.
# This is the SAME agent's source-image spotcheck, not independent acceptance.
samples=[
 (1,'500',1,'','82.10'),(1,'2000',1,'','83.58'),(1,'400000',3,'','81.22'),(1,'800000',1,'','89.96'),
 (2,'60000',3,'','81.30'),(2,'457800',1,'','89.15'),(2,'1000000',3,'','83.90'),
 (3,'500',1,'','70.19'),(3,'100000',1,'','77.54'),(3,'400000',3,'','74.34'),
 (4,'1',3,'68.6','256.1'),(4,'1.25',2,'78.4','240.8'),(4,'4',1,'9.8','19.0'),
 (4,'31.5',3,'98.0','179.6'),(4,'125',1,'88.2','115'),(4,'315',3,'98.0','144.9'),
 (4,'800',1,'98.0','—'),(4,'1250',3,'78.4','104.5'),
 (5,'1',2,'14.7','40.95'),(5,'1',2,'19.6','51.05'),(5,'1',2,'29.4','77.85'),
 (5,'1',1,'39.2','—'),(5,'2.5',3,'44.1','130.1'),(5,'100',2,'39.2','56.4'),
 (5,'1000',3,'49.0','68.1'),(5,'1250',3,'49.0','66.8'),
 (6,'2',1,'39.2','—'),(6,'2',1,'166.6','349.0'),(6,'5',1,'127.4','211'),
 (6,'6.3',1,'147.0','262'),(6,'12.5',3,'68.6','107'),(6,'16',2,'68.6','100'),
 (6,'100',3,'205.8','229.5'),(6,'200',3,'205.8','197.0')]
spot=[]
for table,flow,grade,pressure,expected in samples:
    hits=[r for r in audit if (int(r['table']),r['flow'],int(r['grade']),r['pressure_kPa'])==(table,flow,grade,pressure)]
    assert len(hits)==1
    r=hits[0]; assert r['source_token']==expected,(r,expected)
    spot.append({'table':table,'flow':flow,'grade':grade,'pressure_kPa':pressure,'image_expected':expected,'transcribed_token':r['source_token'],'xlsx_cell':r['xlsx_cell'],'pdf_page':r['pdf_page'],'printed_page':r['printed_page'],'result':'MATCH','review_kind':'same_agent_visual_spotcheck','review_date':'2026-10-13'})
with (ROOT/'VISUAL_SPOTCHECK.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(spot[0]));w.writeheader();w.writerows(spot)
report={'xlsx_sha256':hashlib.sha256(book.read_bytes()).hexdigest(),'sheets':len(wb.sheetnames),'numeric_value_and_format_pass':numeric,'dash_pass':dashes,'coordinate_and_branch_pass':len(audit),'formula_cells':formula_count,'error_cells':error_count,'visual_spotcheck_pass':len(spot),'visual_numeric_samples':sum(x[-1]!='—' for x in samples),'visual_dash_samples':sum(x[-1]=='—' for x in samples),'fail':0,'error':0,'independent_cell_acceptance':'NOT_RUN','native_excel_open':'NOT_RUN','production_tests':'NOT_RUN'}
(ROOT/'CHECK_RESULTS.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
print(json.dumps(report,ensure_ascii=False))
