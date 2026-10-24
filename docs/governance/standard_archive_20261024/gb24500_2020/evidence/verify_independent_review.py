"""GB 24500—2020 独立复核脚本（核验会话 2026-10-13）。

只读复核，不修改任何交付文件：
  1. 从文件重新计算原 PDF / 工作簿 SHA-256，与 evidence/verification_summary.json 比较；
  2. 用 pdfplumber 独立提取原 PDF 数字 token 及其坐标；
  3. 用 openpyxl 读取真实 Sheet/单元格，与 evidence/cell_verification.csv 逐格比较；
  4. 检查 CSV 覆盖是否与工作簿效率单元格集合一致（不多不少）；
  5. 用 PDF 表格线独立确认 4 张表的分组/合并结构（容量列、合并燃料标签、表4 冷凝列）。

本脚本与 tmp/verify.py 是两条相互独立的检查路径（本脚本不使用 tmp/expected_cells.json，
也不使用 tmp/verify.py 的硬编码坐标窗口）。
运行：python docs/standards/gb24500_2020/evidence/verify_independent_review.py
依赖：pdfplumber、openpyxl（执行环境须在复核记录中写明版本）。
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import sys
import unicodedata
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]          # docs/standards/gb24500_2020
repo = root.parents[2]                              # gb24500_2020 -> standards -> docs -> 仓库根
pdf = repo / "标准原文/11. GB+24500-2020 工业锅炉能效限定值及能效等级.pdf"
xlsx = root / "data/GB24500_2020_TABLE_VALUES.xlsx"
summary = json.loads((root / "evidence/verification_summary.json").read_text(encoding="utf-8"))

failures: list[str] = []


def check(ok: bool, label: str, detail: str = "") -> None:
    print(("PASS  " if ok else "FAIL  ") + label + (f"  [{detail}]" if detail else ""))
    if not ok:
        failures.append(label)


# 1. 文件哈希 ---------------------------------------------------------------
pdf_sha = hashlib.sha256(pdf.read_bytes()).hexdigest()
xlsx_sha = hashlib.sha256(xlsx.read_bytes()).hexdigest()
check(pdf_sha == summary["pdf_sha256"], "原 PDF SHA-256 与校验摘要一致", pdf_sha)
check(xlsx_sha == summary["xlsx_sha256"], "工作簿 SHA-256 与校验摘要一致", xlsx_sha)

book = load_workbook(xlsx, data_only=False)
check(book.sheetnames == ["表1 层状燃煤", "表2 流化床燃煤", "表3 生物质", "表4 室燃", "限定值与来源"],
      "5 个 Sheet 名称与顺序", str(book.sheetnames))
formula_cells = sum(1 for s in book for row in s for c in row if c.data_type == "f")
check(formula_cells == 0, "工作簿零公式单元格", f"formula_cells={formula_cells}")
check(book["限定值与来源"]["B5"].value == 97, "5.2 电加热最低热效率 = 97", str(book["限定值与来源"]["B5"].value))

# 2/3. PDF 坐标 vs CSV vs 工作簿 -------------------------------------------
csv_rows = list(csv.DictReader((root / "evidence/cell_verification.csv").open(encoding="utf-8-sig")))
check(len(csv_rows) == summary["threshold_cells"] == 78, "逐格复核 CSV 行数 = 78", str(len(csv_rows)))

with pdfplumber.open(pdf) as doc:
    check(len(doc.pages) == 6, "原 PDF 物理页数 = 6", str(len(doc.pages)))
    tokens: dict[tuple[int, float, float], str] = {}
    for pno, page in enumerate(doc.pages, start=1):
        for w in page.extract_words():
            t = unicodedata.normalize("NFKC", w["text"]).strip()
            if re.fullmatch(r"[0-9]+[ab]?(\([0-9]+[ab]\))?%?", t):
                tokens[(pno, round(w["x0"], 2), round(w["top"], 2))] = t

    mismatches, missing_token = [], []
    for r in csv_rows:
        page_no = int(r["pdf_physical_page"])
        x0, top = (float(v) for v in r["pdf_bbox_pt"].split(",")[:2])
        tok = tokens.get((page_no, round(x0, 2), round(top, 2)))
        if tok is None:
            missing_token.append((r["sheet"], r["cell"], r["pdf_bbox_pt"]))
        elif tok != r["source_token"]:
            mismatches.append((r["sheet"], r["cell"], tok, r["source_token"]))
        actual = book[r["sheet"]][r["cell"]].value
        if str(actual) != r["xlsx_value"]:
            mismatches.append((r["sheet"], r["cell"], actual, r["xlsx_value"]))
    check(not missing_token, "CSV 的 PDF 坐标均能在原 PDF 中命中数字 token", str(missing_token[:3]))
    check(not mismatches, "CSV 源 token / 工作簿取值逐格一致", str(mismatches[:5]))

    # 4. 覆盖完整性：工作簿效率单元格必须与 CSV 完全一致
    covered = {((r["sheet"]), r["cell"]) for r in csv_rows}
    found: set[tuple[str, str]] = set()
    for ws in book.worksheets:
        if ws.title == "限定值与来源":
            continue
        header = [ws.cell(row=5, column=c).value for c in range(1, ws.max_column + 1)]
        for c in range(1, ws.max_column + 1):
            if header[c - 1] and ("效率" in str(header[c - 1]) or str(header[c - 1]).startswith(("D≤", "D>"))):
                for row in range(6, ws.max_row + 1):
                    cell = ws.cell(row=row, column=c)
                    if cell.value is not None:
                        found.add((ws.title, cell.coordinate))
    check(found == covered, "CSV 覆盖集合 = 工作簿效率单元格集合", f"workbook={len(found)} csv={len(covered)}")

    # 5. 结构：用表格线确认容量列 / 合并标签 / 表4 冷凝列
    def vlines(pno: int) -> dict[float, list[tuple[float, float]]]:
        out: dict[float, list[tuple[float, float]]] = {}
        for ln in doc.pages[pno - 1].lines:
            if abs(ln["x0"] - ln["x1"]) < 0.6 and abs(ln["top"] - ln["bottom"]) > 3:
                out.setdefault(round(ln["x0"], 1), []).append((round(ln["top"], 1), round(ln["bottom"], 1)))
        return out

    # 表1：F/G 两个容量列在 x=451.1 分栏，且 451.1 贯穿全部数据行（248.3-654.3）
    p4 = vlines(4)
    x451 = p4.get(451.1, [])
    check(bool(x451) and min(a for a, _ in x451) <= 248.3 and max(b for _, b in x451) >= 654.3,
          "表1 有左右两个容量效率列（x=451.1 分隔，贯穿全部数据行）", str(sorted(x451)))
    # 表1：x=99.2 只出现在烟煤(248.3-383.6)与无烟煤(451.3-586.6)块 => 合并燃料标签按Ⅱ/Ⅲ类拆分
    x992 = p4.get(99.2, [])
    check(bool(x992) and all(248.0 <= a and b <= 586.9 for a, b in x992),
          "表1 Ⅱ/Ⅲ类子栏只覆盖烟煤与无烟煤块（合并标签结构）", str(x992))

    # 表2（同一物理页上半部）：效率列为单一列 467.9-531.9，列内无分隔线 => 无容量分档
    p5 = vlines(5)
    interior = [x for x, spans in p5.items()
                if 467.9 < x < 531.9 and any(a >= 185.0 and b <= 564.2 for a, b in spans)]
    check(not interior, "表2 效率列内无容量分栏（单一效率列）", str(sorted(interior)))
    # 表3：F/G 在 x=440.8 分栏，只覆盖表3 数据行（664.0-754.2）
    x440 = p5.get(440.8, [])
    check(bool(x440) and all(660.0 <= a and b <= 754.3 for a, b in x440),
          "表3 有小/大容量两列（x=440.8 分隔）", str(x440))

    # 表4：第二效率列分隔线 x=441.3 只存在于天然气（冷凝）行组 227.9-282.1
    x441 = vlines(6).get(441.3, [])
    check(bool(x441) and all(227.9 <= a and b <= 282.1 for a, b in x441),
          "表4 的第二效率列分隔线只存在于天然气（冷凝）行组", str(x441))

print()
if failures:
    print(f"结论：FAIL（{len(failures)} 项）")
    for f in failures:
        print("  - " + f)
    sys.exit(1)
print("结论：全部通过（独立复核路径）")
