from pathlib import Path
import json,csv,hashlib,datetime,ast
from decimal import Decimal,InvalidOperation
import openpyxl
b=Path('docs/standards/gbt36561_2018');o=Path('docs/governance/device_full_acceptance_20261013/reverification/gbt36561_batch1')
j=json.loads((b/'data/table_values.json').read_text(encoding='utf8'));c=list(csv.DictReader((b/'data/table_values.csv').open(encoding='utf-8-sig')));e=json.loads((b/'evidence/workbook_coordinates.json').read_text(encoding='utf8'))
w=openpyxl.load_workbook(b/'data/GBT36561_2018_TABLE_VALUES.xlsx',read_only=True,data_only=False);s=w['正文表值']
ids=['GBT36561-2018-T6-0089','GBT36561-2018-T6-0090','GBT36561-2018-T11-0293'];checks=[]
for rid in ids:
 a=next(x for x in j if x['record_id']==rid);d=next(x for x in c if x['record_id']==rid);z=next(x for x in e if x['record_id']==rid);row=int(z['value_cell'][1:]);wv=s.cell(row,6).value;wc=s.cell(row,8).value
 result={'record_id':rid,'json':a,'csv':d,'evidence':z,'xlsx_value':wv,'xlsx_condition':wc,'xlsx_cell':'正文表值!'+('H' if rid in ids[:2] else 'F')+str(row)}
 if rid==ids[0]: ok=all('300 mm' in x and '500' not in x for x in [a['condition'],d['condition'],z['condition'],wc])
 elif rid==ids[1]:ok=all('>φ300 mm' in x and '≥' not in x for x in [a['condition'],d['condition'],z['condition'],wc])
 else:ok=all(Decimal(str(x))==1 for x in [a['value'],d['value'],z['value'],wv]) and a['operator']=='≤' and a['unit']=='h'
 result['pass']=ok;checks.append(result)
src=(b/'build_data.py').read_text(encoding='utf8');tree=ast.parse(src);literals=[]
for n in ast.walk(tree):
 if isinstance(n,ast.Tuple):
  try:literals.append(ast.literal_eval(n))
  except (ValueError,TypeError):pass
build={'FS01':'距主轴端面300 mm处测量' in src,'FS02':'直径>φ300 mm圆盘' in src,'FS03':('井式炉',950,['≤0.2','0.2～1.0','1.0～2.5'],[1,1,2]) in literals}
reg=[]
def equal(a,b):
 try:return Decimal(str(a))==Decimal(str(b))
 except InvalidOperation:return str(a)==str(b)
for x in csv.DictReader(Path('docs/governance/device_full_acceptance_20261013/fans_submersible/gbt36561_2018/independent_visual_cells.csv').open(encoding='utf-8-sig')):
 val=w[x['cell'].split('!')[0]][x['cell'].split('!')[1]].value;reg.append({'record_id':x['record_id'],'cell':x['cell'],'expected':x['visual_source_value'],'actual':val,'pass':equal(x['visual_source_value'],val)})
# Compare complete T6/T11 rows across all four representations, including neighbours.
rows=[]
for a in j:
 if a['table'] not in ['6','11']:continue
 d=next(x for x in c if x['record_id']==a['record_id']);z=next(x for x in e if x['record_id']==a['record_id']);r=int(z['value_cell'][1:]);errors=[]
 for k in a:
  if not equal(a[k],d[k]) or a[k]!=z[k]:errors.append(k)
 if not equal(s.cell(r,6).value,a['value']):errors.append('xlsx_value')
 if (s.cell(r,8).value or '')!=a['condition']:errors.append('xlsx_condition')
 rows.append({'record_id':a['record_id'],'pass':not errors,'errors':errors})
sha=json.loads((b/'evidence/SHA256SUMS.json').read_text(encoding='utf8'));sha_checks=[{'path':p,'pass':hashlib.sha256((b/p).read_bytes()).hexdigest()==v} for p,v in sha.items() if (b/p).is_file()]
reviews=list(csv.DictReader((b/'evidence/numeric_review.csv').open(encoding='utf-8-sig')))
res={'targets':checks,'build_source_literal_checks':build,'full_numeric_regression':{'pass':sum(x['pass'] for x in reg),'fail':sum(not x['pass'] for x in reg)},'T6_T11_cross_asset':rows,'published_sha_checks':sha_checks,'numeric_review_target_rows':[x for x in reviews if x['cell'] in ['F90','F91','F294','H90','H91']],'scope':'FS01/02/03 only; FS04 excluded'}
(o/'checks.json').write_text(json.dumps(res,ensure_ascii=False,indent=2),encoding='utf8');(o/'numeric_regression.json').write_text(json.dumps(reg,ensure_ascii=False,indent=2),encoding='utf8')
start=json.loads((o/'start_hashes.json').read_text(encoding='utf8'));end={'time_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'files':[]};changes=[]
for x in start['files']:
 p=Path(x['path']);h=hashlib.sha256(p.read_bytes()).hexdigest() if p.exists() else None;end['files'].append({'path':x['path'],'sha256':h})
 if h!=x['sha256']:changes.append({'path':x['path'],'start_sha256':x['sha256'],'end_sha256':h})
(o/'end_hashes.json').write_text(json.dumps(end,ensure_ascii=False,indent=2),encoding='utf8');(o/'hash_changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({'target_pass':[x['pass'] for x in checks],'build':build,'regression':res['full_numeric_regression'],'cross_rows':len(rows),'cross_fail':[x for x in rows if not x['pass']],'sha_fail':[x for x in sha_checks if not x['pass']],'changed_files':changes,'review_targets':res['numeric_review_target_rows']},ensure_ascii=False,indent=2))

