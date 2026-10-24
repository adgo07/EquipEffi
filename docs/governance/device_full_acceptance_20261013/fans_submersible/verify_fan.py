import openpyxl,csv,json
from pathlib import Path
out=Path('docs/governance/device_full_acceptance_20261013/fans_submersible')
# Independently transcribed from rendered source pages 4-6 (not copied from mapping data).
raw={
'表1':'''43 58 61 46 59 62 49 60 63 52 61 64 56 64 67 59 65 68
45 62 65 48 63 66 51 64 67 54 65 68 59 68 71 63 69 72
49 65 70 52 66 71 55 67 72 58 68 72 62 70 75 65 73 78
52 66 71 55 67 72 58 68 73 61 69 73 63 71 76 66 75 80''',
'表2':'''62 72 75 65 75 78 69 78 81
65 74 77 68 77 80 72 80 83
68 76 79 71 79 82 75 82 85
62 70 75 65 75 78 68 78 81
65 72 78 68 75 81 70 78 84
68 75 80 71 78 83 72 81 85
62 70 77 63 72 79 64 73 83
65 72 82 66 75 83 67 76 84
64 74 81 70 76 85 73 80 86
69 75 82 73 79 86 75 82 87
67 74 79 69 76 81 71 79 85
71 77 82 73 79 84 75 81 86
73 78 83 75 80 85 77 82 87
70 79 84 72 81 86 75 83 88
63 73 78 66 78 83 73 82 87 76 84 89
— — — 70 79 84 72 81 86''',
'表3':'''55 66 69 58 69 72 60 73 77
59 68 71 61 71 74 63 75 79
61 70 73 64 73 76 66 77 81
63 72 75 67 75 78 69 79 83''',
'表4':'''36 43 46 37 50 54 39 50 53 43 55 60 50 60 63
35 42 45 35 49 53 38 49 52 42 54 59 49 59 62
35 43 46 36 49 52 38 49 52 42 55 59 49 59 62
34 42 45 35 48 51 37 48 51 41 54 58 48 58 61
33 43 46 35 49 52 37 48 51 41 55 58 48 58 61
32 42 45 34 48 51 36 47 50 40 54 57 47 57 61
33 42 45 35 48 51 37 47 51 41 54 57 47 57 61
31 41 44 33 47 50 35 46 50 39 53 56 46 56 60
32 41 44 34 47 50 36 47 51 40 53 56 46 56 60
30 40 43 32 46 49 34 46 50 38 52 55 45 55 59'''}
w=openpyxl.load_workbook(next(Path('docs/standards/gb19761_2020').rglob('*.xlsx')),data_only=False);records=[]
for name,txt in raw.items():
 vals=txt.split(); coords=[w[name].cell(r,c) for r in range(5,w[name].max_row+1) for c in [5,6,7]]
 assert len(vals)==len(coords),(name,len(vals),len(coords))
 for expected,cell in zip(vals,coords):
  actual=str(cell.value);eq=(expected==actual or (expected!='—' and float(expected)==float(actual)))
  records.append(dict(table=name,cell=cell.coordinate,visual_source_value=expected,workbook_value=actual,match=eq))
with (out/'gb19761_2020'/'independent_visual_cells.csv').open('w',encoding='utf-8-sig',newline='') as f:
 writer=csv.DictWriter(f,fieldnames=records[0]);writer.writeheader();writer.writerows(records)
print('fan',len(records),'mismatches',[r for r in records if not r['match']])

