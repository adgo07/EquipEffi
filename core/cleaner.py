# -*- coding: utf-8 -*-
"""core/cleaner.py - 清洗引擎（模块A）
从输入Excel读取设备数据 → 清洗/标准化/校验 → 输出标准params + 问题清单
与判定器解耦：清洗只产出干净的参数字典，判定只看字典。
"""
import re
from pathlib import Path

import openpyxl


def to_float(v):
    """宽松数值转换：'45kW'→45、'1,500'→1500、'—'→None"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    s = s.replace(",", "").replace("，", "").replace(" ", "")
    s = s.replace("kW", "").replace("kw", "").replace("ＫＷ", "")
    s = s.replace("kVA", "").replace("kva", "").replace("KVA", "")
    s = s.replace("m³/h", "").replace("m3/h", "").replace("m³/min", "")
    s = s.replace("MPa", "").replace("mpa", "").replace("ＭＰａ", "")
    s = s.replace("kV", "").replace("kv", "")
    s = s.replace("Pa", "").replace("pa", "")
    s = s.replace("r/min", "").replace("rpm", "")
    s = s.replace("W", "").replace("w", "")
    s = s.replace("℃", "").replace("°C", "")
    s = s.replace("K", "").replace("k", "")
    s = s.replace("t/h", "").replace("MW", "")
    s = s.replace("kJ/kg", "").replace("kJ/m³", "")
    s = s.replace("（", "").replace("）", "").replace("(", "").replace(")", "")
    s = s.replace("—", "").replace("-", "").replace("－", "")
    s = s.replace("／", "").replace("/", "").replace("％", "")
    s = s.replace(">", "").replace("<", "").replace("≥", "").replace("≤", "")
    s = s.replace("≥", "")
    if not s or s in ("/", "一", "无"):
        return None
    try:
        return float(s)
    except (TypeError, ValueError):
        return None


def clean_text(v):
    """文本清理：去空格、全角转半角"""
    if v is None:
        return ""
    s = str(v).strip()
    # 全角转半角
    out = []
    for ch in s:
        code = ord(ch)
        if code == 0x3000:
            out.append(" ")
        elif 0xFF01 <= code <= 0xFF5E:
            out.append(chr(code - 0xFEE0))
        else:
            out.append(ch)
    return "".join(out).strip()


def parse_poles_from_model(model):
    """从型号提取极数：取'-'后的数字（如 YE3-225M-4→4、YXKK400-14→14）
    注意：机座号中的数字（225M）不取"""
    m = clean_text(model)
    if not m:
        return None
    # 找最后一个'-'或空格后的数字
    parts = re.split(r"[-－_ ]+", m)
    if len(parts) < 2:
        return None
    tail = parts[-1]
    # 尾部可能是纯数字（极数）或"数字L"等
    mm = re.match(r"^(\d+)$", tail)
    if mm:
        p = int(mm.group(1))
        if p in (2, 4, 6, 8, 10, 12, 14, 16, 18, 20, 24, 28, 32):
            return p
    return None


# 各设备sheet：列名关键词 → 参数键（支持V1.11与2024监察版两种模板）
COLUMN_MAP = {
    "transformer": {
        "变压器类别": "category", "设备型式": "category", "额定容量": "capacity_kva",
        "铁芯材质": "core_material", "绝缘等级": "insulation", "连接组标号": "connection",
        "空载损耗": "no_load_loss_w", "负载损耗": "load_loss_w",
    },
    "motor_lv": {
        "额定功率": "power_kw", "设备效率": "efficiency_pct", "铭牌": "efficiency_pct",
        "极数": "poles", "型号": "model", "电机型号": "model",
    },
    "motor_hv": {
        "电压等级": "voltage_kv", "冷却方式": "cooling", "额定功率": "power_kw",
        "设备效率": "efficiency_pct", "铭牌": "efficiency_pct", "极数": "poles", "型号": "model",
    },
    "motor_pmsm": {
        "额定功率": "power_kw", "转速": "speed_rpm", "设备效率": "efficiency_pct",
        "铭牌": "efficiency_pct", "极数": "poles", "型号": "model",
    },
    "compressor": {
        "设备型式": "type", "容积流量": "flow_m3min", "额定功率": "power_kw",
        "排气压力": "pressure_mpa", "冷却方式": "cooling", "是否变频": "variable",
        "机组比功率": "specific_power", "型号": "model",
    },
    "pump_water": {
        "流量": "flow_m3h", "扬程": "head_m", "转速": "speed_rpm", "吸入形式": "suction",
        "单吸": "suction", "级数": "stages", "泵轴方向": "shaft", "支撑形式": "support",
        "设备效率": "efficiency_pct", "泵效率": "efficiency_pct", "型号": "model",
    },
    "pump_chem": {
        "流量": "flow_m3h", "扬程": "head_m", "转速": "speed_rpm", "吸入形式": "suction",
        "单吸": "suction", "级数": "stages", "设备效率": "efficiency_pct", "泵效率": "efficiency_pct",
        "型号": "model",
    },
    "fan": {
        "风机类型": "type", "压力": "pressure_pa", "出口滞止压力": "outlet_pa",
        "流量": "flow_m3h", "主轴转速": "speed_rpm", "机号": "no",
        "进口滞止密度": "density", "等熵": "isentropic_k", "风机效率": "efficiency_pct",
        "轮毂比": "hub_ratio",
    },
    "blower": {
        "风机类型": "type", "离心级数": "stages", "叶轮出口宽度": "b2_mm",
        "叶轮出口直径": "d2_mm", "进口绝对压力": "p1_kpa", "出口绝对压力": "p2_kpa",
        "进口温度": "t1_k", "出口温度": "t2_k", "多变效率": "efficiency_pct",
        "能效限定值": "manual_limit", "节能评价值": "manual_save",
    },
    "submersible": {
        "类型": "subtype", "额定": "power_kw", "电泵规定效率": "eta_db",
        "容差": "delta_eta", "设备效率": "efficiency_pct", "泵效率": "efficiency_pct",
        "型号": "model",
    },
    "boiler": {
        "设备类型": "boiler_type", "是否冷凝": "condensing", "蒸发量": "capacity",
        "热功率": "capacity", "燃料收到基低位发热量": "heat_value", "锅炉热效率": "efficiency_pct",
        "热效率": "efficiency_pct", "型号": "model",
    },
    "heat_treatment": {
        "设备类型": "furnace", "总折合重量": "weight_t", "电炉总耗能量": "elec_kwh",
        "燃料总耗量": "fuel_m3", "燃料热值": "fuel_hv", "燃料系数": "fuel_coef",
        "型号": "model",
    },
}

# sheet名 → 设备key + 起始数据行（V1.11布局 + 2024监察版）
SHEET_CONFIG = {
    "变压器": ("transformer", 4),
    "低压电动机": ("motor_lv", 3),
    "高压电动机": ("motor_hv", 4),
    "永磁同步电机": ("motor_pmsm", 4),
    "空压机": ("compressor", 3),
    "清水泵": ("pump_water", 4),
    "清水离心泵 ": ("pump_water", 4),
    "清水离心泵": ("pump_water", 4),
    "化工泵": ("pump_chem", 4),
    "化工离心泵": ("pump_chem", 4),
    "通风机": ("fan", 3),
    "鼓风机": ("blower", 3),
    "离心鼓风机": ("blower", 3),
    "潜水电泵": ("submersible", 3),
    "工业锅炉": ("boiler", 3),
    "热处理设备": ("heat_treatment", 3),
}


class Cleaner:
    def __init__(self, path):
        self.path = Path(path)
        self.wb = openpyxl.load_workbook(path, data_only=True)
        self.devices = {}   # key -> [{params, issues, row, name}]
        self.summary = []   # 每sheet统计

    def run(self):
        for sheet_name, (key, data_start) in SHEET_CONFIG.items():
            if sheet_name not in self.wb.sheetnames:
                continue
            ws = self.wb[sheet_name]
            self._clean_sheet(ws, key, data_start, sheet_name)
        return self.devices

    def _find_header_row(self, ws, key):
        """找含'序号'的表头行"""
        for r in range(1, 6):
            for c in range(1, 8):
                v = ws.cell(r, c).value
                if v and "序号" in str(v):
                    return r
        return None

    def _col_index(self, ws, header_row, keyword):
        """按关键词找列号（表头行模糊匹配），返回(列号, 表头原文)"""
        for c in range(1, ws.max_column + 1):
            v = ws.cell(header_row, c).value
            if v and keyword in str(v):
                return c, str(v)
        return None, None

    def _clean_sheet(self, ws, key, data_start, sheet_name):
        colmap = COLUMN_MAP[key]
        hdr = self._find_header_row(ws, key)
        if hdr is None:
            hdr = data_start - 1
        # 建立列映射（参数 → (列号, 表头原文)）
        col_for = {}
        for kw, param in colmap.items():
            if param in col_for:
                continue  # 已有更早匹配
            c, txt = self._col_index(ws, hdr, kw)
            if c:
                col_for[param] = (c, txt)
        out = []
        for r in range(hdr + 1, ws.max_row + 1):
            seq = ws.cell(r, 1).value
            if not isinstance(seq, (int, float)) or seq == 0:
                continue
            # 核心列检查（有参数才算设备行）
            has_data = any(ws.cell(r, c).value not in (None, "") for c, _ in col_for.values())
            if not has_data:
                continue
            params = {}
            issues = []
            for param, (c, htxt) in col_for.items():
                v = ws.cell(r, c).value
                if param == "model":
                    params["model"] = clean_text(v)
                elif param == "poles":
                    # 无效值（公式错误/斜杠）→None，由型号解析
                    if v in ("#VALUE!", "#REF!", "#DIV/0!", "/", "—", "无"):
                        params["poles"] = None
                    else:
                        params["poles"] = v
                elif param in ("suction", "cooling", "variable", "type", "shaft", "support",
                               "category", "core_material", "insulation", "connection",
                               "subtype", "boiler_type", "condensing", "furnace"):
                    params[param] = clean_text(v)
                else:
                    if v in ("/", "—", "无", "#VALUE!", "#REF!", "#DIV/0!"):
                        params[param] = None
                        continue
                    f = to_float(v)
                    if v not in (None, "") and f is None:
                        issues.append(f"R{r}参数[{param}]非数值[{v}]")
                    # 单位换算：鼓风机压力表头MPa→kPa
                    if param in ("p1_kpa", "p2_kpa") and f is not None and "MPa" in htxt:
                        f = f * 1000
                    # 风机效率小数→百分数
                    if param == "efficiency_pct" and f is not None and f < 1 and key in ("fan", "blower"):
                        f = f * 100
                    params[param] = f
            # 变压器2024版：绝缘等级/连接组合并列解析
            if key == "transformer" and params.get("insulation") and not params.get("connection"):
                ins = params["insulation"]
                if "/" in ins and "℃" in ins:
                    parts = ins.split("/", 1)
                    params["insulation"], params["connection"] = parts[0].strip(), parts[1].strip()
            # 专项清洗
            self._special_clean(key, params, issues, r)
            # 判定必需项缺失检查
            self._required_check(key, params, issues)
            out.append({"params": params, "issues": issues, "row": r,
                        "name": str(ws.cell(r, 3).value or ws.cell(r, 2).value or f"R{r}")})
        self.devices[key] = out
        self.summary.append({"sheet": sheet_name, "key": key, "devices": len(out),
                             "issues": sum(len(i) for i in out)})

    def _special_clean(self, key, p, issues, r):
        """专项清洗规则"""
        if key == "motor_lv" or key == "motor_hv" or key == "motor_pmsm":
            if p.get("poles") is None and p.get("model"):
                p["poles"] = parse_poles_from_model(p["model"])
                if p["poles"]:
                    issues.append(f"R{r}极数由型号自动解析={p['poles']}")
                else:
                    issues.append(f"R{r}极数缺失且型号无法解析[{p.get('model')}]")
        if key == "compressor":
            # 冷却标准化
            if p.get("cooling"):
                p["cooling"] = "液冷" if ("液" in p["cooling"] or "水" in p["cooling"]) else "风冷"
            # 变频标准化
            if p.get("variable"):
                p["variable"] = "是" if ("变频" in p["variable"] or "变" in p["variable"] and "不变" not in p["variable"]) else "否"
            # 比功率校验：功率/流量 > 比功率 → 可疑（额定功率÷流量≠实测比功率）
            if p.get("power_kw") and p.get("flow_m3min") and p.get("specific_power"):
                ratio = p["power_kw"] / p["flow_m3min"]
                if ratio > p["specific_power"]:
                    issues.append(f"R{r}比功率{p['specific_power']}小于功率/流量比值{ratio:.2f}，请核实是否为实测机组输入功率")
        if key == "pump_water":
            # 泵类型推断（支撑形式/吸入形式/级数）
            if p.get("stages") and p.get("support"):
                if "管道" in p["support"]:
                    p["pump_type"] = "管道"
                elif p["stages"] > 1:
                    p["pump_type"] = "多级"
                elif "双" in p.get("suction", ""):
                    p["pump_type"] = "单级双吸"
                else:
                    p["pump_type"] = "单级单吸"
        if key == "fan":
            if p.get("efficiency_pct") is None:
                pass  # 效率缺失由判定器提示

    def _required_check(self, key, p, issues):
        """判定必需项检查"""
        required = {
            "transformer": ["category", "capacity_kva", "no_load_loss_w", "load_loss_w"],
            "motor_lv": ["power_kw", "efficiency_pct", "poles"],
            "motor_hv": ["voltage_kv", "cooling", "power_kw", "efficiency_pct", "poles"],
            "motor_pmsm": ["power_kw", "efficiency_pct", "poles"],
            "compressor": ["type", "power_kw", "pressure_mpa", "specific_power"],
            "pump_water": ["flow_m3h", "head_m", "speed_rpm", "efficiency_pct"],
            "pump_chem": ["flow_m3h", "head_m", "speed_rpm", "efficiency_pct"],
            "fan": ["pressure_pa", "flow_m3h", "speed_rpm", "no"],
            "blower": ["type", "b2_mm", "d2_mm"],
            "submersible": ["power_kw", "eta_db", "efficiency_pct"],
            "boiler": ["efficiency_pct"],
            "heat_treatment": ["furnace"],
        }
        for f in required.get(key, []):
            if p.get(f) in (None, ""):
                issues.append(f"缺少判定必需项[{f}]")
