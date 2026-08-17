# -*- coding: utf-8 -*-
"""gui/main_window.py - 主窗口（tkinter，大众主流蓝白配色）
功能：模板下载 / 设备类别选择 / 拖拽上传 / 分析（清洗→判定→汇总→写回）
"""
import os
import queue
import shutil
import sys
import threading
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
    HAS_DND = True
except ImportError:
    HAS_DND = False

from core.cleaner import Cleaner
from core.evaluator import evaluate_all
from core.summary import build_summary

# ===== 大众主流配色（Windows 11 / Office风格）=====
COLOR_BG = "#F3F3F3"        # 窗口浅灰
COLOR_CARD = "#FFFFFF"      # 白色卡片
COLOR_PRIMARY = "#0078D4"   # 微软蓝（主按钮/标题）
COLOR_TEXT = "#333333"      # 正文深灰
COLOR_SUB = "#666666"       # 次级文字
COLOR_BORDER = "#E1E1E1"    # 边框
COLOR_OK = "#107C10"        # 成功绿
FONT = "Microsoft YaHei"

# 设备类别选项（key, 显示名）
DEVICE_OPTIONS = [
    ("transformer", "变压器"), ("motor_lv", "低压电动机"), ("motor_hv", "高压电动机"),
    ("motor_pmsm", "永磁同步电机"), ("compressor", "空压机"),
    ("pump_water", "清水泵"), ("pump_chem", "化工泵"), ("fan", "通风机"),
    ("blower", "鼓风机"), ("submersible", "潜水电泵"), ("boiler", "锅炉"),
    ("heat_treatment", "热处理"),
]
TEMPLATE_NAME = "设备能效分析模板.xlsx"


class MainWindow:
    def __init__(self, root):
        self.root = root
        self.queue = queue.Queue()
        self.out_path = None
        root.title("设备能效分析工具")
        root.geometry("760x680")
        root.minsize(680, 580)
        root.configure(bg=COLOR_BG)
        self._build_ui()
        self._bind_dnd()
        self._check_queue()

    # ---------- 界面 ----------
    def _build_ui(self):
        # 标题栏
        header = tk.Frame(self.root, bg=COLOR_PRIMARY, height=64)
        header.pack(fill="x")
        header.pack_propagate(False)
        tk.Label(header, text="设备能效分析工具", bg=COLOR_PRIMARY, fg="white",
                 font=(FONT, 16, "bold")).pack(side="left", padx=24, pady=12)
        tk.Label(header, text="v0.2.0", bg=COLOR_PRIMARY, fg="#D0E8FF",
                 font=(FONT, 9)).pack(side="right", padx=24)

        body = tk.Frame(self.root, bg=COLOR_BG)
        body.pack(fill="both", expand=True, padx=20, pady=16)

        # ① 输入文件卡片
        card1 = self._make_card(body)
        tk.Label(card1, text="① 选择/拖入设备台账Excel", bg=COLOR_CARD, fg=COLOR_TEXT,
                 font=(FONT, 11, "bold")).grid(row=0, column=0, sticky="w", pady=(2, 6))
        self.input_var = tk.StringVar()
        self.input_entry = ttk.Entry(card1, textvariable=self.input_var, font=(FONT, 10))
        self.input_entry.grid(row=1, column=0, sticky="ew", padx=(0, 8))
        self.input_entry.insert(0, "请选择或拖入企业填写的台账文件…")
        self.input_entry.bind("<Button-1>", lambda e: self._browse_input())
        btn_style = {"font": (FONT, 10), "bg": "#E8F1FA", "fg": COLOR_PRIMARY,
                     "activebackground": "#D6E8F7", "relief": "flat", "cursor": "hand2"}
        tk.Button(card1, text="浏览…", command=self._browse_input, **btn_style).grid(row=1, column=1)
        card1.columnconfigure(0, weight=1)

        # ② 设备类别卡片
        card2 = self._make_card(body)
        top = tk.Frame(card2, bg=COLOR_CARD)
        top.pack(fill="x")
        tk.Label(top, text="② 分析设备类别", bg=COLOR_CARD, fg=COLOR_TEXT,
                 font=(FONT, 11, "bold")).pack(side="left")
        tk.Button(top, text="全选", command=lambda: self._set_all(True),
                  font=(FONT, 9), bg="#F5F5F5", fg=COLOR_TEXT, relief="flat", cursor="hand2").pack(side="right", padx=4)
        tk.Button(top, text="全不选", command=lambda: self._set_all(False),
                  font=(FONT, 9), bg="#F5F5F5", fg=COLOR_TEXT, relief="flat", cursor="hand2").pack(side="right")
        grid = tk.Frame(card2, bg=COLOR_CARD)
        grid.pack(fill="x", pady=(8, 2))
        self.device_vars = {}
        for i, (key, label) in enumerate(DEVICE_OPTIONS):
            var = tk.BooleanVar(value=True)
            self.device_vars[key] = var
            tk.Checkbutton(grid, text=label, variable=var, bg=COLOR_CARD, fg=COLOR_TEXT,
                           font=(FONT, 10), activebackground=COLOR_CARD,
                           selectcolor="white").grid(row=i // 4, column=i % 4, sticky="w", padx=8, pady=2)

        # ③ 模板下载
        card3 = self._make_card(body)
        tk.Label(card3, text="③ 没有台账模板？", bg=COLOR_CARD, fg=COLOR_TEXT,
                 font=(FONT, 11, "bold")).pack(side="left")
        tk.Button(card3, text="下载空白模板", command=self._download_template,
                  font=(FONT, 10), bg="#E8F1FA", fg=COLOR_PRIMARY, relief="flat",
                  cursor="hand2", activebackground="#D6E8F7").pack(side="right")

        # 开始按钮
        self.start_btn = tk.Button(body, text="开始分析", command=self._start,
                                   bg=COLOR_PRIMARY, fg="white", font=(FONT, 12, "bold"),
                                   relief="flat", cursor="hand2", height=2,
                                   activebackground="#005A9E", activeforeground="white")
        self.start_btn.pack(fill="x", pady=10)

        # 进度条
        self.progress = ttk.Progressbar(body, mode="determinate", maximum=100)
        self.progress.pack(fill="x", pady=(4, 2))

        # 日志区
        log_frame = tk.Frame(body, bg=COLOR_BG)
        log_frame.pack(fill="both", expand=True, pady=(8, 0))
        self.log = tk.Text(log_frame, height=10, font=(FONT, 9), fg=COLOR_TEXT,
                           bg="white", relief="flat", wrap="word", state="disabled")
        self.log.pack(side="left", fill="both", expand=True)
        scroll = ttk.Scrollbar(log_frame, command=self.log.yview)
        scroll.pack(side="right", fill="y")
        self.log.configure(yscrollcommand=scroll.set)

        self._log("欢迎使用设备能效分析工具\n1. 可直接把台账Excel拖入窗口\n2. 勾选要分析的设备类别\n3. 点击“开始分析”")

    def _make_card(self, parent):
        card = tk.Frame(parent, bg=COLOR_CARD, highlightbackground=COLOR_BORDER, highlightthickness=1)
        card.pack(fill="x", pady=(0, 8))
        inner = tk.Frame(card, bg=COLOR_CARD)
        inner.pack(fill="x", padx=14, pady=10)
        inner.columnconfigure(0, weight=1)
        return inner

    # ---------- 拖拽 ----------
    def _bind_dnd(self):
        if not HAS_DND:
            self._log("（拖拽功能不可用，请用浏览按钮选择文件）")
            return
        self.root.drop_target_register(DND_FILES)
        self.root.dnd_bind("<<Drop>>", self._on_drop)
        self.input_entry.drop_target_register(DND_FILES)
        self.input_entry.dnd_bind("<<Drop>>", self._on_drop)

    def _on_drop(self, event):
        files = self.root.tk.splitlist(event.data)
        for f in files:
            if f.lower().endswith((".xlsx", ".xlsm")):
                self.input_var.set(f)
                self._log(f"已拖入：{f}")
                return

    # ---------- 交互 ----------
    def _browse_input(self):
        path = filedialog.askopenfilename(
            title="选择设备台账Excel",
            filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")])
        if path:
            self.input_var.set(path)
            self._log(f"已选择：{path}")

    def _set_all(self, val):
        for var in self.device_vars.values():
            var.set(val)

    def _download_template(self):
        src = self._template_path()
        if src is None:
            messagebox.showwarning("提示", "模板文件缺失，请重新安装工具")
            return
        out = filedialog.asksaveasfilename(
            title="保存空白模板", defaultextension=".xlsx",
            initialfile=TEMPLATE_NAME,
            filetypes=[("Excel文件", "*.xlsx")])
        if out:
            try:
                shutil.copy2(src, out)
                self._log(f"模板已保存：{out}")
                messagebox.showinfo("完成",
                                    f"模板已保存到：\n{out}\n\n将模板发给企业填写，收回后再用本工具分析。")
            except Exception as e:
                messagebox.showerror("错误", f"保存失败：{e}")

    def _template_path(self):
        if getattr(sys, "frozen", False):
            ext = Path(sys.executable).parent / "template" / TEMPLATE_NAME
            if ext.exists():
                return ext
            return Path(getattr(sys, "_MEIPASS", ".")) / "template" / TEMPLATE_NAME
        p = Path(__file__).resolve().parent.parent / "template" / TEMPLATE_NAME
        return p if p.exists() else None

    # ---------- 分析流程 ----------
    def _start(self):
        src = self.input_var.get().strip()
        if not src or not Path(src).exists():
            messagebox.showwarning("提示", "请先选择有效的Excel文件")
            return
        selected = [k for k, v in self.device_vars.items() if v.get()]
        if not selected:
            messagebox.showwarning("提示", "请至少勾选一类设备")
            return
        self.start_btn.configure(state="disabled", text="分析中…")
        self.progress["value"] = 0
        threading.Thread(target=self._worker, args=(src, selected), daemon=True).start()

    def _worker(self, src, selected):
        try:
            self._log("\n[1/4] 正在读取与清洗数据…")
            cleaner = Cleaner(src)
            devices = cleaner.run(keys=selected)
            n = sum(len(v) for v in devices.values())
            self._post("progress", 30)
            self._log(f"清洗完成：{len(devices)}类设备，共{n}台")

            self._log("[2/4] 正在判定能效等级…")
            results = evaluate_all(devices)
            self._post("progress", 65)
            judged = sum(1 for v in results.values() for it in v
                         if it["result_dict"]["result"] not in ("无法判定", "不在范围"))
            self._log(f"判定完成：{judged}台已判定")

            self._log("[3/4] 正在汇总统计…")
            summary = build_summary(results)
            self._post("progress", 80)
            for k, v in summary["total"].items():
                self._log(f"  {k}: {v}台")

            self._log("[4/4] 正在写回结果…")
            out = self._write_back(src, results)
            self._post("progress", 100)
            self._log(f"\n✅ 分析完成！\n输出文件：{out}")
            self._post("done", out)
        except Exception as e:
            self._post("error", f"分析失败：{type(e).__name__}: {e}")

    def _write_back(self, src, results):
        """写回判定结果到输出文件（zip补丁，图片保留）
        写回内容：判定列结论 + 1/2/3级标准限值列 + 过程指标（泵ns/Ci、变压器空载负载、风机ψ/ns）"""
        from core.writer import patch_cells
        src_path = Path(src)
        out = src_path.parent / f"{src_path.stem}_能效分析结果.xlsx"
        changes = {}
        for key, items in results.items():
            sheet_name = self._sheet_name_for(key, src_path)
            if sheet_name is None:
                continue
            scheme = self._judge_scheme(sheet_name, src_path, key)
            if not scheme or not scheme.get("judge"):
                self._log(f"  ⚠ {sheet_name}：未找到判定列，跳过")
                continue
            sc = changes.setdefault(sheet_name, {})
            for it in items:
                rd = it["result_dict"]
                row = it["row"]
                sc[f"{scheme['judge']}{row}"] = rd.get("result", "")
                self._apply_levels(sc, scheme, key, rd, row)
                self._apply_extra(sc, scheme, key, rd, row)
        if not changes:
            import shutil
            shutil.copy2(src, out)
        else:
            res = patch_cells(src, out, changes)
            self._log(f"  已写入 {len(res.get('patched_sheets', {}))} 个sheet（含标准限值/过程指标）")
        return out

    def _judge_scheme(self, sheet_name, src_path, key):
        """动态定位写回列方案：
        {"judge": 判定列, "levels": {"1":[列...],"2":[...],"3":[...]},
         "extra": {"ns":列,"ci":列,"limit":列,"save":列,"psi":列}}"""
        import openpyxl
        wb = openpyxl.load_workbook(src_path, read_only=True)
        ws = wb[sheet_name]
        # 表头行
        hdr = None
        for r in range(1, 6):
            for c in range(1, 8):
                v = ws.cell(r, c).value
                if v and "序号" in str(v):
                    hdr = r
                    break
            if hdr:
                break
        if hdr is None:
            hdr = 2
        scheme = {"hdr": hdr}
        # 判定列
        for c in range(1, min(ws.max_column, 40) + 1):
            v = ws.cell(hdr, c).value
            if v and ("能效判定" in str(v) or "能效等级" in str(v) or "判定" in str(v)):
                scheme["judge"] = openpyxl.utils.get_column_letter(c)
                break
        # 1/2/3级限值列（表头行或其下一行的"1级"等；变压器R3次级表头）
        levels = {"1": [], "2": [], "3": []}
        for rr in (hdr, hdr + 1):
            for c in range(1, min(ws.max_column, 40) + 1):
                v = str(ws.cell(rr, c).value or "").strip()
                if v in ("1级", "1 级") or v.startswith("1级"):
                    levels["1"].append(openpyxl.utils.get_column_letter(c))
                elif v in ("2级", "2 级") or v.startswith("2级"):
                    levels["2"].append(openpyxl.utils.get_column_letter(c))
                elif v in ("3级", "3 级") or v.startswith("3级"):
                    levels["3"].append(openpyxl.utils.get_column_letter(c))
                elif v in ("一等", "二等", "三等") and key == "heat_treatment":
                    levels[{"一等": "1", "二等": "2", "三等": "3"}[v]].append(
                        openpyxl.utils.get_column_letter(c))
            if any(levels.values()):
                break
        scheme["levels"] = levels
        # 特殊指标列（表头行）
        extra = {}
        for c in range(1, min(ws.max_column, 40) + 1):
            v = str(ws.cell(hdr, c).value or "")
            if key in ("pump_water", "pump_chem"):
                if "比转速" in v or "/ns" in v or v.strip() == "ns":
                    extra["ns"] = openpyxl.utils.get_column_letter(c)
                if "Ci" in v or "常数" in v:
                    extra["ci"] = openpyxl.utils.get_column_letter(c)
            elif key == "fan":
                if "压力系数" in v or "/Ψ" in v:
                    extra["psi"] = openpyxl.utils.get_column_letter(c)
                if "比转速" in v or "/ns" in v:
                    extra["ns"] = openpyxl.utils.get_column_letter(c)
            elif key == "blower":
                if "能效限定值" in v:
                    extra["limit"] = openpyxl.utils.get_column_letter(c)
                if "节能评价值" in v:
                    extra["save"] = openpyxl.utils.get_column_letter(c)
        scheme["extra"] = extra
        wb.close()
        return scheme

    def _apply_levels(self, sc, scheme, key, rd, row):
        """写1/2/3级限值到标准值列"""
        levels = scheme["levels"]
        if key == "transformer":
            # 每级2列（空载+负载），按出现顺序配对
            for i, lv in enumerate(("1", "2", "3")):
                cols = levels.get(lv, [])
                nl = rd.get("no_load_levels") or [None] * 3
                ld = rd.get("load_levels") or [None] * 3
                if len(cols) >= 2:
                    if nl[i] is not None:
                        sc[f"{cols[0]}{row}"] = nl[i]
                    if ld[i] is not None:
                        sc[f"{cols[1]}{row}"] = ld[i]
            return
        if key == "blower":
            # level1=评价值→Z列、level2=限定值→Y列（由scheme.extra定位）
            return
        for lv in ("1", "2", "3"):
            v = rd.get("level" + lv)
            cols = levels.get(lv, [])
            if v is not None and cols:
                sc[f"{cols[0]}{row}"] = round(float(v), 2)

    def _apply_extra(self, sc, scheme, key, rd, row):
        """写过程指标（泵ns/Ci、变压器空载负载已由_apply_levels处理、风机ψ/ns、鼓风机限值）"""
        extra = scheme.get("extra", {})
        if key in ("pump_water", "pump_chem"):
            if extra.get("ns") and rd.get("ns"):
                sc[f"{extra['ns']}{row}"] = rd["ns"]
            if extra.get("ci") and rd.get("ci"):
                sc[f"{extra['ci']}{row}"] = "/".join(str(x) for x in rd["ci"])
        elif key == "fan":
            if extra.get("psi") and rd.get("psi"):
                sc[f"{extra['psi']}{row}"] = rd["psi"]
            if extra.get("ns") and rd.get("ns"):
                sc[f"{extra['ns']}{row}"] = rd["ns"]
        elif key == "blower":
            if extra.get("limit") and rd.get("level2") is not None:
                sc[f"{extra['limit']}{row}"] = round(float(rd["level2"]), 2)
            if extra.get("save") and rd.get("level1") is not None:
                sc[f"{extra['save']}{row}"] = round(float(rd["level1"]), 2)

    def _sheet_name_for(self, key, src_path):
        """设备key → 输入文件中的sheet名（按SHEET_CONFIG反向）"""
        from core.cleaner import SHEET_CONFIG
        import openpyxl
        try:
            wb = openpyxl.load_workbook(src_path, read_only=True)
            names = wb.sheetnames
            wb.close()
        except Exception:
            return None
        for name, (k, _) in SHEET_CONFIG.items():
            if k == key and name in names:
                return name
        return None

    def _judge_col(self, sheet_name, src_path):
        """找sheet的能效判定列：必须在表头行（含'序号'的行）中找，
        避免R1标题行（如《XX能效限定值及能效等级》）误导"""
        import openpyxl
        wb = openpyxl.load_workbook(src_path, read_only=True)
        ws = wb[sheet_name]
        hdr = None
        for r in range(1, 6):
            for c in range(1, 8):
                v = ws.cell(r, c).value
                if v and "序号" in str(v):
                    hdr = r
                    break
            if hdr:
                break
        if hdr is None:
            hdr = 2
        col = None
        for c in range(1, min(ws.max_column, 40) + 1):
            v = ws.cell(hdr, c).value
            if v and ("能效判定" in str(v) or "能效等级" in str(v) or "判定" in str(v)):
                col = openpyxl.utils.get_column_letter(c)
                break
        wb.close()
        return col

    # ---------- 队列/日志 ----------
    def _post(self, kind, data):
        self.queue.put((kind, data))

    def _check_queue(self):
        try:
            while True:
                kind, data = self.queue.get_nowait()
                if kind == "log":
                    self._log(data)
                elif kind == "progress":
                    self.progress["value"] = data
                elif kind == "done":
                    self.start_btn.configure(state="normal", text="开始分析")
                    if messagebox.askyesno("完成", f"分析完成！\n输出：{data}\n\n是否现在打开输出文件？"):
                        os.startfile(data)  # noqa: S606
                elif kind == "error":
                    self.start_btn.configure(state="normal", text="开始分析")
                    messagebox.showerror("错误", data)
        except queue.Empty:
            pass
        self.root.after(200, self._check_queue)

    def _log(self, msg):
        self.log.configure(state="normal")
        self.log.insert("end", msg + "\n")
        self.log.see("end")
        self.log.configure(state="disabled")


def run():
    root = (TkinterDnD.Tk() if HAS_DND else tk.Tk())
    MainWindow(root)
    root.mainloop()


if __name__ == "__main__":
    run()
