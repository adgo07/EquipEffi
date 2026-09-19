"""Refresh the human-review workbook without discarding reviewer edits.

The review workbook is generated from the machine standard packages.  When a
standard package gains traceable metadata (for example a newly explicit fixed
gate), the flattened row set can change.  Rebuilding the workbook is safer
than appending rows by hand, but the three reviewer columns must be merged by
stable data ID so an existing review is never silently lost.
"""
from __future__ import annotations

import argparse
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any

from openpyxl import load_workbook

try:
    from .build_human_review_book import _write_book
except ImportError:  # direct execution: ``python tools/refresh_review_book_preserving_edits.py``
    from build_human_review_book import _write_book


REVIEW_SHEET = "标准数据平铺"
REVIEW_HEADERS = ("校对状态", "建议修订值", "校对意见")


def _header_map(sheet: Any) -> dict[str, int]:
    return {
        str(cell.value): index
        for index, cell in enumerate(sheet[3], start=1)
        if cell.value not in (None, "")
    }


def _review_values(sheet: Any) -> dict[str, tuple[Any, ...]]:
    headers = _header_map(sheet)
    id_column = headers.get("数据ID")
    columns = tuple(headers.get(name) for name in REVIEW_HEADERS)
    if id_column is None or any(column is None for column in columns):
        raise ValueError("标准数据平铺缺少数据ID或人工校对列")
    values: dict[str, tuple[Any, ...]] = {}
    for row in sheet.iter_rows(min_row=4, max_col=max((id_column, *columns))):
        identifier = row[id_column - 1].value
        if identifier in (None, ""):
            continue
        values[str(identifier)] = tuple(row[column - 1].value for column in columns)
    return values


def refresh(existing: Path, output: Path) -> dict[str, int | str]:
    """Rebuild *output* from current machine data and merge reviewer columns."""

    existing = Path(existing).resolve()
    output = Path(output).resolve()
    if not existing.is_file():
        raise FileNotFoundError(existing)
    with TemporaryDirectory(prefix="equipeffi-review-refresh-") as temporary:
        rebuilt = Path(temporary) / "rebuilt.xlsx"
        _write_book(rebuilt)
        old_book = load_workbook(existing, read_only=True, data_only=False)
        new_book = load_workbook(rebuilt, read_only=False, data_only=False)
        try:
            old_sheet = old_book[REVIEW_SHEET]
            new_sheet = new_book[REVIEW_SHEET]
            old_values = _review_values(old_sheet)
            new_headers = _header_map(new_sheet)
            id_column = new_headers["数据ID"]
            reviewer_columns = tuple(new_headers[name] for name in REVIEW_HEADERS)
            matched = 0
            new_rows = 0
            for row in new_sheet.iter_rows(min_row=4, max_col=max((id_column, *reviewer_columns))):
                identifier = row[id_column - 1].value
                if identifier in (None, ""):
                    continue
                new_rows += 1
                previous = old_values.get(str(identifier))
                if previous is None:
                    continue
                matched += 1
                for column, value in zip(reviewer_columns, previous):
                    new_sheet.cell(row=row[0].row, column=column).value = value
            output.parent.mkdir(parents=True, exist_ok=True)
            new_book.save(output)
        finally:
            old_book.close()
            new_book.close()
    return {
        "output": str(output),
        "previous_review_rows": len(old_values),
        "rebuilt_rows": new_rows,
        "merged_rows": matched,
        "new_rows_without_previous_review": new_rows - matched,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="重建标准校对册并按数据ID保留人工校对列")
    parser.add_argument("existing", type=Path, help="已有人工校对册")
    parser.add_argument("output", type=Path, help="重建后的输出文件")
    args = parser.parse_args(argv)
    print(refresh(args.existing, args.output))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
