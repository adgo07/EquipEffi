"""Audit V4 workbook data-validation coverage and numeric boundary rules.

This is intentionally independent from the evaluation engine.  It checks the
Excel-facing contract so a regenerated template cannot silently regress to a
small fixed validation range or accept percentage efficiencies as fractions.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from equipeffi.application.services.v4_template_contract import V4_DEVICE_SHEETS
from equipeffi.infrastructure.excel.v4_reader import V4WorkbookReaderImpl


MAX_DATA_ROW = 1_048_576


def _header(value: object) -> str:
    return str(value or "").split("\n", 1)[0].strip()


def _column(sheet, display_name: str) -> int | None:
    for cell in sheet[3]:
        if _header(cell.value) == display_name:
            return cell.column
    return None


def audit(path: Path) -> dict[str, object]:
    path = Path(path)
    contract = V4WorkbookReaderImpl().read_contract(path)
    workbook = load_workbook(path, read_only=False, data_only=False)
    errors: list[str] = []
    checks = {"sheets": len(V4_DEVICE_SHEETS), "validation_ranges": 0, "efficiency_fields": 0}
    try:
        for sheet_name in V4_DEVICE_SHEETS:
            sheet = workbook[sheet_name]
            for validation in sheet.data_validations.dataValidation:
                for token in str(validation.sqref).split():
                    try:
                        min_col, min_row, max_col, max_row = range_boundaries(token)
                    except ValueError:
                        continue
                    if min_row == 4 and min_col == max_col:
                        checks["validation_ranges"] += 1
                        if max_row != MAX_DATA_ROW:
                            errors.append(f"{sheet_name}:{token}未覆盖到Excel最大行")

            for field in contract.fields_for_sheet(sheet_name, editable_only=True):
                if "效率" not in field.display_name:
                    continue
                checks["efficiency_fields"] += 1
                col = _column(sheet, field.display_name)
                if col is None:
                    errors.append(f"{sheet_name}:{field.display_name}找不到表头")
                    continue
                col_letter = get_column_letter(col)
                matches = [
                    item for item in sheet.data_validations.dataValidation
                    if any(str(r).split(":", 1)[0].startswith(col_letter) for r in str(item.sqref).split())
                ]
                if not matches:
                    errors.append(f"{sheet_name}:{field.display_name}缺少数据验证")
                    continue
                formulas = " ".join(
                    str(value or "") for item in matches for value in (item.formula1, item.formula2)
                )
                if sheet_name == "工业锅炉" and field.field_id == "design_efficiency":
                    if ">=1" not in formulas or "<=IF" not in formulas or "110" not in formulas or "E4=\"室燃燃烧锅炉（燃气冷凝）\"" not in formulas:
                        errors.append("工业锅炉:设计热效率未保留1～100及冷凝锅炉110例外")
                elif not any(item.type == "decimal" and item.operator == "between" and item.formula1 == "1" and item.formula2 == "100" for item in matches):
                    errors.append(f"{sheet_name}:{field.display_name}未限制为1～100")

            if sheet_name == "多联式空调":
                col = _column(sheet, "机外静压")
                if col is not None:
                    letter = get_column_letter(col)
                    if not any(
                        item.operator == "greaterThanOrEqual" and item.formula1 == "0"
                        for item in sheet.data_validations.dataValidation
                        if any(str(r).split(":", 1)[0].startswith(letter) for r in str(item.sqref).split())
                    ):
                        errors.append("多联式空调:机外静压未允许0且拒绝负值")
    finally:
        workbook.close()
    return {"path": str(path.resolve()), "checks": checks, "errors": errors, "is_valid": not errors}


def main(argv: list[str] | None = None) -> int:
    args = list(argv if argv is not None else sys.argv[1:])
    if len(args) != 1:
        print("用法: python tools/audit_v4_validations.py 模板.xlsx", file=sys.stderr)
        return 2
    result = audit(Path(args[0]))
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["is_valid"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
