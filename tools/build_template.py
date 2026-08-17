# -*- coding: utf-8 -*-
"""tools/build_template.py - 从V1.11生成干净的空模板
用途：GUI"下载模板"功能分发的标准台账模板
- 保留全部设备sheet的表头+格式（黄色填写区）
- 清空所有数据行（含判定公式、示例数据）
- 删除标准sheet（工具判定用内置JSON，不给企业看）
用法: python tools/build_template.py
输出: template/设备能效分析模板.xlsx
"""
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")

V111 = r"G:\标准  规范\02_能耗限额_终端产品\用能设备\用能设备能效分析表V1.11.xlsx"
OUT = Path(__file__).resolve().parent.parent / "template" / "设备能效分析模板.xlsx"

# 设备sheet保留（企业填写用）；标准sheet/版本变更删除
DEVICE_SHEETS = [
    "变压器", "低压电动机", "高压电动机", "永磁同步电机", "空压机",
    "清水泵", "化工泵", "通风机", "鼓风机", "潜水电泵", "工业锅炉", "热处理设备",
    "热泵和冷水机组", "热泵热水机", "风管送风式空调", "单元式空调", "多联式空调",
    "注意事项",
]

HEADER_ROWS = {  # sheet → 表头行数（数据从这行之后开始）
    "变压器": 3, "低压电动机": 2, "高压电动机": 3, "永磁同步电机": 3,
    "空压机": 2, "清水泵": 3, "化工泵": 3, "通风机": 2, "鼓风机": 2,
    "潜水电泵": 2, "工业锅炉": 2, "热处理设备": 2,
    "热泵和冷水机组": 3, "热泵热水机": 3, "风管送风式空调": 3,
    "单元式空调": 2, "多联式空调": 2, "注意事项": 1,
}


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    src = openpyxl.load_workbook(V111, keep_vba=False)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in DEVICE_SHEETS:
        if sheet not in src.sheetnames:
            print(f"跳过（不存在）: {sheet}")
            continue
        s = src[sheet]
        t = wb.create_sheet(sheet)
        keep = HEADER_ROWS.get(sheet, 2)
        # 复制表头行（含样式）
        for r in range(1, min(keep + 1, s.max_row + 1)):
            for c in range(1, min(s.max_column, 60) + 1):
                src_cell = s.cell(r, c)
                dst_cell = t.cell(r, c)
                dst_cell.value = src_cell.value
                if src_cell.has_style:
                    dst_cell.font = src_cell.font.copy()
                    dst_cell.border = src_cell.border.copy()
                    dst_cell.fill = src_cell.fill.copy()
                    dst_cell.alignment = src_cell.alignment.copy()
        # 复制列宽
        for c in range(1, min(s.max_column, 60) + 1):
            letter = openpyxl.utils.get_column_letter(c)
            w = s.column_dimensions[letter].width
            if w:
                t.column_dimensions[letter].width = w
        # 复制合并单元格（表头区）
        for m in s.merged_cells.ranges:
            if m.min_row <= keep:
                t.merge_cells(str(m))
        # 冻结窗格
        t.freeze_panes = f"A{keep + 1}"
    wb.save(OUT)
    print(f"✅ 模板已生成: {OUT}")

if __name__ == "__main__":
    main()
