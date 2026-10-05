"""Phase 8：其他设备 Sheet「语义不变」检查。

Owner 规则要求：除「离心泵」Sheet 的授权修改外，Writer / 模板处理不得删除 Sheet、
重命名其他 Sheet、调整其他设备业务字段或改变其业务规则。验收应检查
cell values / formulas、defined names、validation、protection、tables、sheet order、
关键 style/structure 的**语义/结构**不变。

**不要求**整个 xlsx ZIP 逐字节不变（正常 OOXML 重写会改变压缩字节）。

用法::

    python tools/check_v6_untouched_sheets.py --baseline <V6 基线> --template <资产>
"""
from __future__ import annotations

import argparse
import json
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_BASELINE = (REPO_ROOT / "specs" / "equipment_efficiency" / "templates"
                    / "设备能效分析空白模板_重构版V6_变压器.xlsx")
DEFAULT_TEMPLATE = (REPO_ROOT / "src" / "equipeffi" / "resources" / "templates"
                    / "设备能效分析空白模板_重构版V6_20261005.xlsx")

#: 本阶段唯一授权修改的 Sheet。
AUTHORIZED_SHEET = "离心泵"
#: 授权修改的间接影响面：`EV_PUMP_CATEGORY` 在「配置」Sheet 上，且 FieldDictionary
#: 的「离心泵」行会被同步。配置 Sheet 只允许这些坐标变化。
CONFIG_SHEET = "配置"
ALLOWED_CONFIG_CELLS: frozenset[str] = frozenset(
    [f"AE{row}" for row in range(2, 62)]
    + [f"{col}{row}" for row in range(2, 400) for col in ("H", "I", "N")]
)


@dataclass
class Difference:
    sheet: str
    kind: str
    where: str
    baseline: str
    template: str

    def as_dict(self) -> dict:
        return {"sheet": self.sheet, "kind": self.kind, "where": self.where,
                "baseline": self.baseline, "template": self.template}


@dataclass
class Report:
    baseline: str
    template: str
    differences: list[Difference] = field(default_factory=list)
    checked_sheets: list[str] = field(default_factory=list)

    @property
    def passed(self) -> bool:
        return not self.differences

    def as_dict(self) -> dict:
        return {"baseline": self.baseline, "template": self.template,
                "passed": self.passed,
                "checked_sheets": self.checked_sheets,
                "difference_count": len(self.differences),
                "differences": [d.as_dict() for d in self.differences]}


def _cell_text(value) -> str:
    return "" if value is None else str(value)


def _compare_sheet(baseline_sheet, template_sheet, name: str,
                   allowed_cells: frozenset[str]) -> list[Difference]:
    out: list[Difference] = []
    rows = max(baseline_sheet.max_row, template_sheet.max_row)
    columns = max(baseline_sheet.max_column, template_sheet.max_column)
    for row in range(1, rows + 1):
        for column in range(1, columns + 1):
            base_cell = baseline_sheet.cell(row, column)
            new_cell = template_sheet.cell(row, column)
            coordinate = f"{get_column_letter(column)}{row}"
            if coordinate in allowed_cells:
                continue
            if _cell_text(base_cell.value) != _cell_text(new_cell.value):
                out.append(Difference(
                    name, "cell_value", coordinate,
                    _cell_text(base_cell.value)[:80], _cell_text(new_cell.value)[:80]))
            if base_cell.number_format != new_cell.number_format:
                out.append(Difference(
                    name, "number_format", coordinate,
                    base_cell.number_format, new_cell.number_format))
            base_locked = bool(base_cell.protection and base_cell.protection.locked)
            new_locked = bool(new_cell.protection and new_cell.protection.locked)
            if base_locked != new_locked:
                out.append(Difference(name, "protection", coordinate,
                                      str(base_locked), str(new_locked)))
    return out


def _validation_signature(sheet) -> list[tuple]:
    out = []
    for validation in sheet.data_validations.dataValidation:
        out.append((str(validation.sqref), validation.type,
                    str(validation.operator), str(validation.formula1),
                    str(validation.formula2)))
    return sorted(out)


def _table_signature(sheet) -> list[tuple]:
    return sorted((name, sheet.tables[name].ref, len(sheet.tables[name].tableColumns))
                  for name in sheet.tables)


def run_checks(baseline_path: Path = DEFAULT_BASELINE,
               template_path: Path = DEFAULT_TEMPLATE) -> Report:
    baseline = openpyxl.load_workbook(baseline_path, data_only=False)
    template = openpyxl.load_workbook(template_path, data_only=False)
    report = Report(baseline=str(baseline_path), template=str(template_path))

    if list(baseline.sheetnames) != list(template.sheetnames):
        report.differences.append(Difference(
            "(workbook)", "sheet_order_or_set",
            "sheetnames", ",".join(baseline.sheetnames),
            ",".join(template.sheetnames)))
        return report

    for name in baseline.sheetnames:
        if name == AUTHORIZED_SHEET:
            continue
        allowed = ALLOWED_CONFIG_CELLS if name == CONFIG_SHEET else frozenset()
        report.checked_sheets.append(name)
        base_sheet = baseline[name]
        new_sheet = template[name]
        report.differences.extend(_compare_sheet(base_sheet, new_sheet, name, allowed))
        if _validation_signature(base_sheet) != _validation_signature(new_sheet):
            report.differences.append(Difference(
                name, "data_validation", "(sheet)", "见基线", "见模板"))
        if _table_signature(base_sheet) != _table_signature(new_sheet):
            report.differences.append(Difference(
                name, "tables", "(sheet)",
                str(_table_signature(base_sheet)), str(_table_signature(new_sheet))))
        if bool(base_sheet.protection.sheet) != bool(new_sheet.protection.sheet):
            report.differences.append(Difference(
                name, "sheet_protection", "(sheet)",
                str(bool(base_sheet.protection.sheet)),
                str(bool(new_sheet.protection.sheet))))
        if sorted(str(r) for r in base_sheet.merged_cells.ranges) != \
                sorted(str(r) for r in new_sheet.merged_cells.ranges):
            report.differences.append(Difference(
                name, "merged_cells", "(sheet)", "见基线", "见模板"))

    # defined names：允许 EV_PUMP_CATEGORY 变化，其余必须完全一致。
    base_names = {k: str(v.value) for k, v in baseline.defined_names.items()}
    new_names = {k: str(v.value) for k, v in template.defined_names.items()}
    for key in sorted(set(base_names) | set(new_names)):
        if key == "EV_PUMP_CATEGORY":
            continue
        if base_names.get(key) != new_names.get(key):
            report.differences.append(Difference(
                "(defined_names)", "defined_name", key,
                str(base_names.get(key)), str(new_names.get(key))))

    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--baseline", type=Path, default=DEFAULT_BASELINE)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--json", type=Path, default=None)
    parser.add_argument("--max-print", type=int, default=20)
    args = parser.parse_args(argv)

    report = run_checks(args.baseline, args.template)
    print(f"baseline : {report.baseline}")
    print(f"template : {report.template}")
    print(f"checked  : {len(report.checked_sheets)} 个非授权 Sheet")
    print(f"diffs    : {len(report.differences)}")
    for diff in report.differences[: args.max_print]:
        print(f"  [{diff.sheet}] {diff.kind} {diff.where}: "
              f"{diff.baseline!r} -> {diff.template!r}")
    print("GATE    : " + ("PASS" if report.passed else "FAIL"))

    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        args.json.write_text(json.dumps(report.as_dict(), ensure_ascii=False, indent=2),
                            encoding="utf-8")
    return 0 if report.passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
