# -*- coding: utf-8 -*-
"""tools/extract_fan.py - 风机标准提取（V1.11风机标准sheet，4张表）
表1离心(0.95≤ψ<1.55) 表2离心(0.25≤ψ<0.95) 表3轴流(轮毂比γ) 表4外转子前向多翼
"""
import json
import re
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")

SRC = r"G:\标准  规范\02_能耗限额_终端产品\用能设备\用能设备能效分析表V1.11.xlsx"
STD_DIR = Path(__file__).resolve().parent.parent / "standards"

# 各表定义：起始行、类型、有无ns列
FAN_TABLES = [
    (1, "离心", True),    # 表1 R1起
    (9, "离心", True),    # 表2 R9起
    (29, "轴流", False),  # 表3 R29起
    (38, "外转子", True), # 表4 R38起
]


def to_float(v):
    if v in (None, "", "—", "-"):
        return None
    try:
        return float(str(v).strip())
    except (TypeError, ValueError):
        return None


def extract_fan(wb) -> dict:
    ws = wb["风机标准"]
    tables = []
    for start_row, ftype, has_ns in FAN_TABLES:
        title = str(ws.cell(start_row, 1).value).strip()
        # 表头行（start+1）：机号区间（C列起，合并单元格仅首列有值）
        hdr_row = start_row + 1
        # 等级行（start+2）
        lv_row = start_row + 2
        # 收集机号区间列起始（表头行起，含"机号"或"No"文本的列）
        no_starts = []
        for c in range(2, 40):
            v = ws.cell(hdr_row, c).value
            if v is not None:
                s = str(v).strip()
                if "机号" in s or s.startswith("No"):
                    no_starts.append(c)
        if not no_starts:
            continue
        # 每区间3列（3级/2级/1级）
        no_ranges = []
        for cs in no_starts:
            lvs = []
            for j in range(3):
                lvs.append(str(ws.cell(lv_row, cs + j).value).strip().replace("级", ""))
            no_ranges.append({"no": str(ws.cell(hdr_row, cs).value).strip(), "levels": lvs})
        # 数据行（ψ/γ列可能是合并单元格，用eff区判定行有效性）
        rows = []
        r = start_row + 3
        max_r = ws.max_row
        n_cols = len(no_starts) * 3
        first_col = no_starts[0]
        prev_psi = ""
        while r <= max_r:
            s_a = str(ws.cell(r, 1).value or "").strip()
            if s_a.startswith("表"):
                break
            eff = [to_float(ws.cell(r, c).value) for c in range(first_col, first_col + n_cols)]
            if all(v is None for v in eff):
                r += 1
                continue
            if s_a:
                prev_psi = s_a
            ns = ""
            if has_ns:
                ns_v = ws.cell(r, 2).value
                ns = str(ns_v).strip() if ns_v is not None else ""
            rows.append({"psi": prev_psi, "ns": ns, "eff": eff})
            r += 1
        tables.append({
            "title": title,
            "type": ftype,
            "has_ns": has_ns,
            "no_ranges": no_ranges,
            "rows": rows,
        })
    return {
        "standard_code": "GB 19761-2020",
        "standard_name": "通风机能效限定值及能效等级",
        "effective_date": "2021-06-01",
        "unit_note": "效率单位%；eff按no_ranges顺序×等级[3,2,1]展开；psi=压力系数区间，ns=比转速区间（轴流表为轮毂比γ）",
        "tables": tables,
    }


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    data = extract_fan(wb)
    wb.close()
    out = STD_DIR / "fan.json"
    out.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    for t in data["tables"]:
        print(f"  {t['title'][:36]:<38} 机号区间{len(t['no_ranges'])}个 数据行{len(t['rows'])}")
    print(f"✅ fan.json -> {out}")


if __name__ == "__main__":
    main()
