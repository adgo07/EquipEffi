"""为V4模板加入工作簿级淘汰目录口径下拉和当前引擎说明。"""
from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "src" / "equipeffi" / "resources" / "templates" / "设备能效分析空白模板_重构版V4_20260825.xlsx"
SCOPES = (
    "仅产业结构调整指导目录",
    "高耗能落后机电设备淘汰目录第一至第四批",
    "产业结构调整指导目录+高耗能落后机电设备淘汰目录第一至第四批",
)


def patch(path: Path = TEMPLATE) -> Path:
    workbook = load_workbook(path)
    notes = workbook["注意事项"]
    config = workbook["配置"]

    # 第10行原为空白，放置工作簿级设置，不改变标准目录的行号。
    notes["A10"] = "淘汰判定口径"
    notes["B10"] = SCOPES[1]
    notes.merge_cells("B10:E10")
    notes["A10"]._style = copy(notes["A9"]._style)
    notes["B10"]._style = copy(notes["B9"]._style)
    notes["A10"].alignment = copy(notes["A9"].alignment)
    notes["B10"].alignment = copy(notes["B9"].alignment)
    notes["B10"].comment = None

    # 将旧的“暂不实现判定”说明改成与当前软件状态一致的表述。
    notes["B9"] = "本空白模板预留标准查询和能效判定结果列；软件可对输入参数进行判定并写回结果。范围异常写入自动备注，不设置“是否在标准适用范围”列。"
    notes["B9"].alignment = copy(notes["B8"].alignment)

    # 配置页新增命名枚举，保留隐藏和保护属性。
    column = config.max_column + 1
    column_letter = get_column_letter(column)
    config.cell(1, column).value = "EV_ELIMINATION_SCOPE"
    for row, value in enumerate(SCOPES, start=2):
        config.cell(row, column).value = value
    defined_name = DefinedName("EV_ELIMINATION_SCOPE", attr_text=f"'配置'!${column_letter}$2:${column_letter}${len(SCOPES) + 1}")
    if "EV_ELIMINATION_SCOPE" in workbook.defined_names:
        del workbook.defined_names["EV_ELIMINATION_SCOPE"]
    workbook.defined_names.add(defined_name)

    validation = DataValidation(type="list", formula1="=EV_ELIMINATION_SCOPE", allow_blank=False)
    validation.errorStyle = "stop"
    validation.errorTitle = "淘汰判定口径无效"
    validation.error = "请从下拉列表选择目录口径。"
    validation.promptTitle = "淘汰判定口径"
    validation.prompt = "选择本工作簿使用的淘汰目录范围。"
    notes.add_data_validation(validation)
    validation.add(notes["B10"])
    unlocked = copy(notes["B10"].protection)
    unlocked.locked = False
    notes["B10"].protection = unlocked

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(path)
    return path


if __name__ == "__main__":
    print(patch())
