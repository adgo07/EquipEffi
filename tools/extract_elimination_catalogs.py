from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Iterable

from docx import Document


CATALOG_NAMES = {
    1: "高耗能机电设备淘汰目录（第一批）",
    2: "高耗能落后机电设备（产品）淘汰目录（第二批）",
    3: "高耗能落后机电设备（产品）淘汰目录（第三批）",
    4: "高耗能落后机电设备（产品）淘汰目录（第四批）",
}

FILE_HINTS = {
    1: "第一批",
    2: "第二批",
    3: "第三批",
    4: "第四批",
}

ITEM_RE = re.compile(r"^\s*(\d+)\s*[-－—–一]\s*(\d+)\s*$")
HEADER_WORDS = ("序号", "产品名称", "淘汰产品名称", "淘汰理由")
MODEL_UNIT_RE = re.compile(
    r"(?:KW|MW|W|KV|V|A|HZ|R/MIN|M3/H|M³/H|MPA|KPA|PA|%|T/H)$",
    re.IGNORECASE,
)


@dataclass
class RawEntry:
    batch: int
    item_no: str
    source_file: str
    source_tables: list[int] = field(default_factory=list)
    product_chunks: list[str] = field(default_factory=list)
    model_chunks: list[str] = field(default_factory=list)
    specification_chunks: list[str] = field(default_factory=list)
    reason_chunks: list[str] = field(default_factory=list)
    scope_chunks: list[str] = field(default_factory=list)
    note_chunks: list[str] = field(default_factory=list)
    raw_rows: list[list[str]] = field(default_factory=list)

    def append_unique(self, target: list[str], value: str) -> None:
        value = clean_text(value)
        if value and value not in target:
            target.append(value)


def clean_text(value: str) -> str:
    value = str(value or "").replace("\u3000", " ").replace("\xa0", " ")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n\s*\n+", "\n", value)
    return value.strip()


def normalized_item_no(value: str) -> str | None:
    match = ITEM_RE.match(clean_text(value))
    return f"{int(match.group(1))}-{int(match.group(2))}" if match else None


def dedupe_adjacent(values: Iterable[str]) -> list[str]:
    result: list[str] = []
    for value in values:
        value = clean_text(value)
        if not result or value != result[-1]:
            result.append(value)
    return result


def paragraph_text(cell) -> str:
    parts = [clean_text(p.text) for p in cell.paragraphs if clean_text(p.text)]
    return "\n".join(parts) if parts else clean_text(cell.text)


def classify_device(product: str, item_no: str, batch: int) -> str:
    text = product.replace(" ", "")
    if "变压器" in text:
        return "transformer"
    if "电动机" in text or "电机" in text:
        return "motor_hv" if "高压" in text else "motor_lv"
    if "空气压缩" in text or "空压机" in text:
        return "compressor"
    if "鼓风机" in text:
        return "blower"
    if any(word in text for word in ("通风机", "引风机", "风机")):
        return "fan"
    if "锅炉" in text:
        return "boiler"
    if "潜水" in text and "泵" in text:
        return "submersible"
    if "泵" in text:
        if any(word in text for word in ("耐腐蚀", "化工", "石油")):
            return "pump_chemical"
        return "pump_water"
    if batch == 1 and item_no.startswith("9-"):
        return "heat_treatment"
    if any(word in text for word in ("热处理", "电阻炉", "盐浴炉", "感应加热")):
        return "heat_treatment"
    if "冷水机" in text or "热泵" in text:
        return "heat_pump_chiller"
    if "空调" in text or "制冷" in text:
        return "unitary_ac"
    return ""


def likely_header(cells: list[str]) -> bool:
    combined = "".join(cells)
    return normalized_item_no(cells[0] if cells else "") is None and sum(w in combined for w in HEADER_WORDS) >= 2


def locate_source_files(source_dir: Path) -> dict[int, Path]:
    files = list(source_dir.glob("*.docx"))
    result: dict[int, Path] = {}
    for batch, hint in FILE_HINTS.items():
        matches = [p for p in files if hint in p.name and not p.name.startswith("~$")]
        if len(matches) != 1:
            raise RuntimeError(f"{hint}应有且仅有一个DOCX，实际找到{len(matches)}个: {[p.name for p in matches]}")
        result[batch] = matches[0]
    return result


def append_row(entry: RawEntry, cells: list[str], table_no: int) -> None:
    if table_no not in entry.source_tables:
        entry.source_tables.append(table_no)
    entry.raw_rows.append(cells)
    payload = cells[1:]
    if entry.batch in (1, 2):
        if payload:
            entry.append_unique(entry.product_chunks, payload[0])
        if len(payload) > 1:
            entry.append_unique(entry.reason_chunks, payload[1])
        if len(payload) > 2:
            entry.append_unique(entry.note_chunks, payload[2])
    elif entry.batch == 3:
        if payload:
            entry.append_unique(entry.product_chunks, payload[0])
        if len(payload) > 1:
            entry.append_unique(entry.model_chunks, payload[1])
        if len(payload) > 2:
            entry.append_unique(entry.specification_chunks, payload[2])
        if len(payload) > 3:
            entry.append_unique(entry.reason_chunks, payload[3])
        if len(payload) > 4:
            entry.append_unique(entry.scope_chunks, payload[4])
    else:
        if payload:
            entry.append_unique(entry.product_chunks, payload[0])
        if len(payload) > 1:
            entry.append_unique(entry.model_chunks, payload[1])
        if len(payload) > 2:
            entry.append_unique(entry.specification_chunks, payload[2])
        if len(payload) > 3:
            entry.append_unique(entry.reason_chunks, payload[3])
        if len(payload) > 4:
            entry.append_unique(entry.scope_chunks, payload[4])


def extract_raw_entries(source_file: Path, batch: int) -> list[RawEntry]:
    document = Document(source_file)
    entries: OrderedDict[str, RawEntry] = OrderedDict()
    current: RawEntry | None = None
    for table_no, table in enumerate(document.tables, 1):
        for row in table.rows:
            cells = dedupe_adjacent(paragraph_text(cell) for cell in row.cells)
            if not cells or likely_header(cells):
                continue
            item_no = normalized_item_no(cells[0])
            if item_no:
                current = entries.setdefault(
                    item_no,
                    RawEntry(batch=batch, item_no=item_no, source_file=source_file.name),
                )
                append_row(current, cells, table_no)
            elif current and any(cells):
                # Only append genuine continuation rows. Decorative titles and footnotes are
                # retained by the DOCX itself but must not be attached to a catalog item.
                if not any(word in "".join(cells) for word in ("附件", "目录", "工业和信息化部")):
                    append_row(current, [current.item_no, *cells], table_no)
    return list(entries.values())


def product_name(entry: RawEntry) -> str:
    if not entry.product_chunks:
        return ""
    first = entry.product_chunks[0]
    parts = [clean_text(p) for p in first.splitlines() if clean_text(p)]
    if entry.batch in (1, 2):
        chinese = [p for p in parts if re.search(r"[\u4e00-\u9fff]", p) and not re.search(r"[:：]", p)]
        return chinese[0] if chinese else (parts[0] if parts else first)
    return " / ".join(parts) if parts else first


def normalize_model_token(value: str) -> str:
    translation = str.maketrans({"－": "-", "—": "-", "–": "-", "‐": "-", "‑": "-", "／": "/", "，": ","})
    value = clean_text(value).translate(translation).upper()
    value = re.sub(r"^（?接上页）?\s*", "", value)
    value = re.sub(r"[,，]\s*(?:3|6|10)\s*KV$", "", value, flags=re.IGNORECASE)
    value = re.sub(r"型$", "", value)
    return value.strip(" ,，、;；。")


def candidate_fragments(text: str) -> Iterable[str]:
    text = text.replace("\t", "\n")
    for line in re.split(r"[\n、；;]+", text):
        line = clean_text(line)
        if not line:
            continue
        # A comma before a voltage suffix belongs to the model; other commas separate models.
        pieces = re.split(r",(?!\s*(?:3|6|10)\s*kV\b)|，(?!\s*(?:3|6|10)\s*kV\b)", line, flags=re.IGNORECASE)
        for piece in pieces:
            piece = clean_text(piece)
            if not piece:
                continue
            if " " in piece and not re.search(r"[（(].*[）)]", piece):
                for token in piece.split():
                    yield token
            else:
                yield piece


def is_model_candidate(value: str) -> bool:
    if not value or len(value) < 3 or len(value) > 60:
        return False
    if re.search(r"[\u4e00-\u9fff]", value):
        return False
    if any(marker in value for marker in ("功率", "效率", "流量", "压力", "温度", "标准", "淘汰", "额定", "系列")):
        return False
    if ":" in value or "：" in value or "%" in value or "~" in value or "～" in value:
        return False
    if MODEL_UNIT_RE.search(value):
        return False
    if re.fullmatch(r"[\d.\-/]+", value):
        return False
    # Most catalog models are Latin-letter/digit combinations. Letter-only codes
    # are accepted only when they contain a separator (for example DJF-C).
    has_latin = bool(re.search(r"[A-Z]", value))
    has_digit = bool(re.search(r"\d", value))
    return has_latin and (has_digit or "-" in value)


def extract_models(entry: RawEntry) -> list[str]:
    if entry.batch == 4:
        sources = entry.specification_chunks
    elif entry.batch == 3:
        sources = entry.model_chunks
    else:
        sources = entry.product_chunks
    result: list[str] = []
    for source in sources:
        parts = source.splitlines()
        if entry.batch in (1, 2) and len(parts) > 1:
            parts = parts[1:]
        for part in parts:
            for raw in candidate_fragments(part):
                model = normalize_model_token(raw)
                if is_model_candidate(model) and model not in result:
                    result.append(model)
    return result


def model_series(entry: RawEntry) -> list[str]:
    text = "\n".join(entry.model_chunks)
    series: list[str] = []
    for match in re.finditer(r"([A-Za-z0-9()]+)\s*系列", text):
        prefix = normalize_model_token(match.group(1))
        if prefix and prefix not in series:
            series.append(prefix)
    return series


def condition_from_scope(entry: RawEntry) -> dict:
    scope = "\n".join(entry.scope_chunks)
    if re.search(r"1997\s*年.*?含.*?前生产", scope):
        return {"production_year": {"op": "<=", "value": 1997}}
    if re.search(r"2003\s*年.*?含.*?前生产", scope):
        return {"production_year": {"op": "<=", "value": 2003}}
    if entry.batch == 4 and entry.item_no in {"2-2", "2-4"}:
        return {"rated_voltage_v": {"op": "==", "value": 6000}}
    return {}


def series_patterns(entry: RawEntry, series: list[str]) -> list[str]:
    # Only the three fourth-batch transformer series have a safely expressible
    # bounded rule. Other catalog series rely on the enumerated exact model list.
    if entry.batch != 4 or entry.item_no not in {"1-1", "1-2", "1-3"}:
        return []
    patterns: list[str] = []
    for prefix in series:
        escaped = re.escape(prefix)
        patterns.append(rf"{escaped}-\d+(?:/\d+(?:\.\d+)?)?")
    return patterns


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def serialize(entry: RawEntry) -> dict:
    name = product_name(entry)
    classification_text = "\n".join(entry.product_chunks)
    device_type = classify_device(classification_text, entry.item_no, entry.batch)
    models = extract_models(entry) if device_type else []
    series = model_series(entry) if device_type else []
    patterns = series_patterns(entry, series)
    conditions = condition_from_scope(entry)
    auto_match = bool(device_type and (models or patterns))
    return {
        "entry_id": f"MOTOR-{entry.batch}-{entry.item_no}",
        "catalog": CATALOG_NAMES[entry.batch],
        "batch": f"第{'一二三四'[entry.batch - 1]}批",
        "item_no": entry.item_no,
        "device_type": device_type,
        "product_name": name,
        "model_exact": models,
        "model_series_original": series,
        "series_patterns": patterns,
        "conditions": conditions,
        "original_condition": "\n".join(entry.scope_chunks),
        "reason": "\n".join(entry.reason_chunks),
        "notes": "\n".join(entry.note_chunks),
        "specifications": "\n".join(entry.specification_chunks),
        "source_file": entry.source_file,
        "source_tables": entry.source_tables,
        "source_status": "DOCX结构化抽取，待人工校对",
        "matching_status": "enabled_unverified" if auto_match else "review_only",
        "raw_rows": entry.raw_rows,
    }


def build_catalog(source_dir: Path) -> dict:
    sources = locate_source_files(source_dir)
    records: list[dict] = []
    source_meta: list[dict] = []
    for batch in range(1, 5):
        path = sources[batch]
        raw_entries = extract_raw_entries(path, batch)
        records.extend(serialize(entry) for entry in raw_entries)
        source_meta.append(
            {
                "batch": batch,
                "file": path.name,
                "sha256": sha256(path),
                "entry_count": len(raw_entries),
            }
        )
    return {
        "schema_version": "1.0",
        "data_version": f"motor_elimination_batches_1_4_docx_{date.today():%Y%m%d}",
        "status": "normalized_pending_human_review",
        "source_priority": "用户提供的第一至第四批DOCX结构化文本；原文行完整保留",
        "matching_policy": "仅精确型号或明确受控系列匹配；附加条件必须全部满足；未校对条目保守启用并保留追溯",
        "sources": source_meta,
        "entries": records,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="从第一至第四批DOCX抽取淘汰目录")
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("src/equipeffi/resources/elimination_catalog_batches_1_4.json"),
    )
    args = parser.parse_args()
    data = build_catalog(args.source_dir)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    enabled = sum(item["matching_status"].startswith("enabled") for item in data["entries"])
    print(
        json.dumps(
            {
                "output": str(args.output.resolve()),
                "entries": len(data["entries"]),
                "match_enabled": enabled,
                "sources": data["sources"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
