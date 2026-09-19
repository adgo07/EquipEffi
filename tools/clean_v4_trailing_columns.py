"""清理V4列插入过程中遗留的空白尾列，保持表格后缀可见且紧凑。"""
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "src" / "equipeffi" / "resources" / "templates" / "设备能效分析空白模板_重构版V4_20260825.xlsx"


def patch(path: Path = TEMPLATE) -> Path:
    workbook = load_workbook(path)
    for sheet in workbook.worksheets[:15]:
        last = max((cell.column for cell in sheet[3] if cell.value not in (None, "")), default=sheet.max_column)
        if sheet.max_column > last:
            sheet.delete_cols(last + 1, sheet.max_column - last)
    workbook.save(path)
    return path


if __name__ == "__main__":
    print(patch())
