"""Extend V4 input validations to the full worksheet data area.

Excel tables grow when rows are pasted below the initial data row, but a
validation range such as ``D4:D103`` does not necessarily grow with the
table.  The V4 contract has no business row limit, so validations are applied
from row 4 through the last Excel row.  This script only changes data
validation ranges; formulas, protection, styles and table definitions are
left untouched.

Usage::

    python tools/expand_v4_validations.py input.xlsx [output.xlsx]

When output is omitted the input file is replaced.  The repository template
is updated explicitly by the build step rather than implicitly on import.
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from equipeffi.infrastructure.excel.template_resource import REQUIRED_V4_SHEETS


MAX_DATA_ROW = 1_048_576


def _extend_sqref(sqref: str) -> str:
    ranges: list[str] = []
    for token in str(sqref).split():
        try:
            min_col, min_row, max_col, max_row = range_boundaries(token)
        except ValueError:
            ranges.append(token)
            continue
        # Only extend ordinary vertical cell ranges beginning at the V4 data
        # row.  Leave any named or non-cell references unchanged.
        if min_row == 4 and max_row >= min_row and min_col == max_col:
            from openpyxl.utils import get_column_letter

            col = get_column_letter(min_col)
            ranges.append(f"{col}4:{col}{MAX_DATA_ROW}")
        else:
            ranges.append(token)
    return " ".join(ranges)


def extend_validations(source: Path, destination: Path | None = None) -> Path:
    source = Path(source)
    destination = Path(destination) if destination else source
    workbook = load_workbook(source, read_only=False, data_only=False)
    try:
        for sheet_name in REQUIRED_V4_SHEETS[3:]:
            sheet = workbook[sheet_name]
            for validation in sheet.data_validations.dataValidation:
                validation.sqref = _extend_sqref(str(validation.sqref))
        workbook.save(destination)
    finally:
        workbook.close()
    return destination


def main(argv: list[str] | None = None) -> int:
    args = list(argv if argv is not None else sys.argv[1:])
    if not args or len(args) > 2:
        print(__doc__.strip())
        return 2
    output = Path(args[1]) if len(args) == 2 else None
    extend_validations(Path(args[0]), output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
