"""整理增加结果列后第1行的标题/数量汇总合并区域，避免重叠合并。"""
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "src" / "equipeffi" / "resources" / "templates" / "设备能效分析空白模板_重构版V4_20260825.xlsx"
BACKUP = ROOT / "tmp" / "设备能效分析空白模板_重构版V4_before_common_results.xlsx"


def patch(path: Path = TEMPLATE, backup: Path = BACKUP) -> Path:
    workbook = load_workbook(path)
    original = load_workbook(backup, data_only=False)
    for name in workbook.sheetnames[:15]:
        sheet = workbook[name]
        old = original[name]
        label = next((cell.column for cell in old[1] if cell.value == "数量合计"), None)
        formula = next((cell.value for cell in old[1] if isinstance(cell.value, str) and cell.value.startswith("=")), None)
        if label is None or formula is None:
            continue
        for merged in list(sheet.merged_cells.ranges):
            if merged.min_row <= 1 <= merged.max_row:
                sheet.unmerge_cells(str(merged))
        # 标题和数量汇总采用不重叠的简单布局；结果字段的分组仍在第2行。
        if label > 1:
            sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=label - 1)
        for cell in sheet[1]:
            if cell.column >= label:
                cell.value = None
        sheet.cell(1, label).value = "数量合计"
        sheet.cell(1, label + 1).value = formula
    workbook.save(path)
    return path


if __name__ == "__main__":
    print(patch())
