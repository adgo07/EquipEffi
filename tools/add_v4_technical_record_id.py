"""为V4设备表增加隐藏的技术记录ID列。

技术记录ID不是设备位号，也不是用户填写字段。它放在固定后缀“铭牌照片、
填表备注、自动备注”之前，锁定并隐藏；读取器会优先使用已写入的ID，
旧工作簿或尚未计算公式的行仍回退为 ``sheet-row``。
"""
from __future__ import annotations

from copy import copy
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.worksheet.table import TableColumn, TableFormula
from openpyxl.utils import get_column_letter, range_boundaries


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "src" / "equipeffi" / "resources" / "templates" / "设备能效分析空白模板_重构版V4_20260825.xlsx"
DEVICE_SHEETS = (
    "变压器", "电动机", "空压机", "离心泵", "离心通风机", "轴流通风机", "鼓风机",
    "潜水电泵", "工业锅炉", "热处理设备", "热泵和冷水机组", "热泵热水机",
    "风管送风式空调", "单元式空调", "多联式空调",
)


def _header_key(value: object) -> str:
    return re.sub(r"[\s\r\n（）()\[\]【】]", "", str(value or "")).strip()


def _shift_merged_ranges(sheet, start_col: int, amount: int) -> None:
    original = [str(item) for item in sheet.merged_cells.ranges]
    for item in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(item))
    for text in original:
        min_col, min_row, max_col, max_row = range_boundaries(text)
        if min_col >= start_col:
            min_col += amount
            max_col += amount
        elif max_col >= start_col:
            max_col += amount
        sheet.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)


def _sync_table_columns(sheet) -> None:
    """同步表对象列名与第3行，避免旧补丁造成TableColumn错位。"""

    for table in sheet.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        names = [str(sheet.cell(min_row, col).value or "") for col in range(min_col, max_col + 1)]
        columns = list(table.tableColumns)
        while len(columns) < len(names):
            columns.append(TableColumn(id=0, name=""))
        columns = columns[: len(names)]
        for index, (column, name) in enumerate(zip(columns, names), start=1):
            column.id = index
            column.name = name
        table.tableColumns = columns


def _set_technical_id_formula(sheet, column: int) -> None:
    """给首个空白行和表对象设置非易失性技术ID计算列。"""

    input_refs = []
    for index in range(2, sheet.max_column + 1):
        if index == column:
            continue
        # 只收集模板中解锁的输入列，避开结果/自动备注公式，且兼容
        # 未来在照片和填表备注之间新增锁定结果列。
        if not sheet.cell(4, index).protection.locked:
            input_refs.append(f"{get_column_letter(index)}4")
    if not input_refs:
        return
    formula = f'IF(COUNTA({",".join(input_refs)})>0,"{sheet.title}-"&ROW(),"")'
    sheet.cell(4, column).value = "=" + formula
    for table in sheet.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        if min_col <= column <= max_col:
            table.tableColumns[column - min_col].calculatedColumnFormula = TableFormula(attr_text=formula)


def _add_to_sheet(sheet) -> bool:
    headers = [sheet.cell(3, col).value for col in range(1, sheet.max_column + 1)]
    if any(_header_key(value) == "技术记录ID" for value in headers):
        _sync_table_columns(sheet)
        existing_col = next(index + 1 for index, value in enumerate(headers) if _header_key(value) == "技术记录ID")
        _set_technical_id_formula(sheet, existing_col)
        return False
    photo_col = next((index + 1 for index, value in enumerate(headers) if "铭牌照片" in _header_key(value)), None)
    if photo_col is None:
        raise ValueError(f"{sheet.title}缺少铭牌照片列")

    # 先移动合并区，再插列；openpyxl不会自动调整已有合并范围。
    _shift_merged_ranges(sheet, photo_col, 1)
    sheet.insert_cols(photo_col, 1)
    inserted_letter = get_column_letter(photo_col)
    source_letter = get_column_letter(photo_col + 1)  # 插入后原“铭牌照片”列
    for row in range(1, sheet.max_row + 1):
        source = sheet[f"{source_letter}{row}"]
        target = sheet[f"{inserted_letter}{row}"]
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
    sheet.cell(3, photo_col).value = "技术记录ID"
    # 技术列永远由系统生成；空模板首行留空，新增行由读取器回退生成。
    for row in range(1, sheet.max_row + 1):
        protection = copy(sheet.cell(row, photo_col).protection)
        protection.locked = True
        sheet.cell(row, photo_col).protection = protection
    sheet.column_dimensions[inserted_letter].hidden = True
    sheet.column_dimensions[inserted_letter].width = 2

    for table in sheet.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        if max_col >= photo_col:
            max_col += 1
        table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
        _sync_table_columns(sheet)
    _set_technical_id_formula(sheet, photo_col)
    return True


def _add_config_rows(workbook) -> int:
    config = workbook["配置"]
    existing = {
        (str(config.cell(row, 1).value or ""), str(config.cell(row, 2).value or ""))
        for row in range(2, config.max_row + 1)
    }
    added = 0
    for sheet_name in DEVICE_SHEETS:
        key = (sheet_name, "technical_record_id")
        if key in existing:
            continue
        row = config.max_row + 1
        values = [
            sheet_name, "technical_record_id", "技术记录ID", "系统字段", "文本", "-", "否",
            "系统生成", "隐藏锁定；空值时由读取器按sheet-行号回退", None, None, None,
            "不清洗、不改写", "用于判定结果和轨迹关联，不作为设备位号", "", "系统字段",
        ]
        for col, value in enumerate(values, start=1):
            config.cell(row, col).value = value
            config.cell(row, col).protection = copy(config.cell(2, col).protection)
        added += 1
    return added


def patch(path: Path = TEMPLATE) -> Path:
    workbook = load_workbook(path)
    changed = sum(_add_to_sheet(workbook[name]) for name in DEVICE_SHEETS)
    config_added = _add_config_rows(workbook)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(path)
    print(f"技术记录ID列：新增{changed}个sheet；配置记录新增{config_added}行；文件：{path}")
    return path


if __name__ == "__main__":
    patch()
