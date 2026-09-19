"""恢复列插入操作中丢失的各sheet数量合计公式。"""
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "src" / "equipeffi" / "resources" / "templates" / "设备能效分析空白模板_重构版V4_20260825.xlsx"
BACKUP = ROOT / "tmp" / "设备能效分析空白模板_重构版V4_before_common_results.xlsx"


def patch(path: Path = TEMPLATE, backup: Path = BACKUP) -> Path:
    current = load_workbook(path)
    original = load_workbook(backup, data_only=False)
    for sheet_name in current.sheetnames[:15]:
        sheet = current[sheet_name]
        old = original[sheet_name]
        old_headers = [old.cell(3, col).value for col in range(1, old.max_column + 1)]
        insertion_col = next((i + 1 for i, value in enumerate(old_headers) if "铭牌照片" in str(value or "").replace("\n", "")), None)
        if insertion_col is None:
            continue
        for cell in old[1]:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                target_col = cell.column + 5 if cell.column >= insertion_col else cell.column
                target = sheet.cell(1, target_col)
                target.value = cell.value
    current.save(path)
    return path


if __name__ == "__main__":
    print(patch())
