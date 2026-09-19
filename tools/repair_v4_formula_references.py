"""修复在V4列插入后，表格计算列公式中未自动平移的单元格引用。"""
from __future__ import annotations

from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "src" / "equipeffi" / "resources" / "templates" / "设备能效分析空白模板_重构版V4_20260825.xlsx"
DEVICE_SHEETS = (
    "变压器", "电动机", "空压机", "离心泵", "离心通风机", "轴流通风机", "鼓风机",
    "潜水电泵", "工业锅炉", "热处理设备", "热泵和冷水机组", "热泵热水机",
    "风管送风式空调", "单元式空调", "多联式空调",
)


CELL_REF = re.compile(r"(?<![A-Z0-9_])(?P<abs>\$?)(?P<col>[A-Z]{1,3})(?P<row>\$?\d+)")


def _shift_formula(formula: str, start_col: int, amount: int) -> str:
    def repl(match: re.Match[str]) -> str:
        col = column_index_from_string(match.group("col"))
        if col < start_col:
            return match.group(0)
        return f"{match.group('abs')}{get_column_letter(col + amount)}{match.group('row')}"

    return CELL_REF.sub(repl, formula)


def patch(path: Path = TEMPLATE) -> Path:
    workbook = load_workbook(path)
    for sheet_name in DEVICE_SHEETS:
        sheet = workbook[sheet_name]
        header = [sheet.cell(3, col).value for col in range(1, sheet.max_column + 1)]
        result_col = next((i + 1 for i, value in enumerate(header) if str(value or "").replace("\n", "") == "采用标准"), None)
        if result_col is None:
            continue
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.value = _shift_formula(cell.value, result_col, 5)
    workbook.save(path)
    return path


if __name__ == "__main__":
    print(patch())
