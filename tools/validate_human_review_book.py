"""验收人工校对册是否与当前标准manifest/机器数据同源。

该工具只读校对册和机器标准包，不会把人工建议修订值写回标准JSON。
"""
from __future__ import annotations

import argparse
from contextlib import closing
from decimal import Decimal, InvalidOperation
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

try:
    from .build_human_review_book import FLAT_HEADERS, MANIFEST, _standard_rows
except ImportError:  # 直接执行 python tools/validate_human_review_book.py
    from build_human_review_book import FLAT_HEADERS, MANIFEST, _standard_rows


REQUIRED_SHEETS = {
    "校对说明", "标准目录", "标准数据平铺", "公式与判定规则", "分类与枚举",
    "单位及换算", "淘汰目录_产业", "淘汰目录_机电四批", "待校对问题", "版本变更记录",
}
EDITABLE_FLAT_HEADERS = {"校对状态", "建议修订值", "校对意见"}


def _canonical_locked_value(value: Any) -> tuple[str, str]:
    """Return a comparison form tolerant of Excel blank and numeric coercion.

    openpyxl may read a machine value written as ``30.0`` back as ``30`` and
    may represent an empty string as ``None``.  These are equivalent for the
    locked standard columns.  Other text remains text so a real standard-value
    edit is still detected.
    """

    if value is None:
        return "blank", ""
    if isinstance(value, bool):
        return "bool", str(value)
    text = str(value).strip()
    if not text:
        return "blank", ""
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError):
        number = None
    if number is not None and number.is_finite():
        # normalize() removes insignificant trailing zeroes while ``f`` keeps
        # the comparison stable for values such as 1E-7.
        return "number", format(number.normalize(), "f")
    return "text", text


def validate(path: Path) -> dict[str, Any]:
    path = Path(path).resolve()
    if not path.is_file():
        return {"is_valid": False, "errors": [f"文件不存在: {path}"]}

    errors: list[str] = []
    warnings: list[str] = []
    expected = _standard_rows()
    expected_flat_rows = len(expected)
    actual_flat_rows = 0
    review_status_counts: dict[str, int] = {status: 0 for status in ("未校对", "正确", "需修改", "存疑")}
    # 机器标准包登记的无数据单元格属于特殊的“正确”记录：其空值是
    # 标准原文明确没有数据，而不是人工漏填。校对册必须保留这一语义，
    # 否则后续导出/激活时可能把空值误修成猜测值。
    expected_no_data_rows = {
        str(row[0]): str(row[18] or "")
        for row in expected
        if row[16] == "正确" and any(token in str(row[18] or "") for token in ("无数据", "无资料", "没有数据"))
    }
    actual_no_data_rows: dict[str, dict[str, Any]] = {}
    activation_blockers: list[str] = []
    actual_rows: list[list[Any]] = []
    locked_value_mismatch_count = 0
    locked_value_mismatch_samples: list[dict[str, Any]] = []
    with closing(load_workbook(path, read_only=False, data_only=False)) as workbook:
        missing = sorted(REQUIRED_SHEETS - set(workbook.sheetnames))
        if missing:
            errors.append(f"缺少sheet: {', '.join(missing)}")
        if "标准数据平铺" not in workbook.sheetnames:
            return {"is_valid": False, "errors": errors, "warnings": warnings}
        # “标准目录”必须与当前manifest的包状态/版本一致；否则即使平铺
        # 数据ID相同，校对册仍可能引用过期的机器数据版本。
        if "标准目录" in workbook.sheetnames:
            try:
                manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
                manifest_by_pack = {
                    str(entry.get("pack_id", "")): entry
                    for entry in manifest.get("packs", [])
                    if entry.get("pack_id")
                }
                directory = workbook["标准目录"]
                for row_number in range(4, directory.max_row + 1):
                    pack_id = directory.cell(row=row_number, column=3).value
                    expected_entry = manifest_by_pack.get(str(pack_id))
                    if expected_entry is None:
                        # 允许校对册中登记的待实施标准（如GB28381-2026）。
                        continue
                    actual_status = directory.cell(row=row_number, column=4).value
                    actual_version = directory.cell(row=row_number, column=5).value
                    if actual_status != expected_entry.get("status"):
                        errors.append(f"标准目录第{row_number}行状态与manifest不一致: {pack_id}")
                    if actual_version != expected_entry.get("data_version", ""):
                        errors.append(f"标准目录第{row_number}行数据版本与manifest不一致: {pack_id}")
            except (OSError, json.JSONDecodeError) as exc:
                errors.append(f"读取标准manifest失败: {exc}")
        sheet = workbook["标准数据平铺"]
        header_row = [cell.value for cell in sheet[3]]
        if tuple(header_row[: len(FLAT_HEADERS)]) != FLAT_HEADERS:
            errors.append("标准数据平铺第3行字段与FLAT_HEADERS不一致")
        header_index = {str(value): index for index, value in enumerate(header_row) if value not in (None, "")}
        for name in FLAT_HEADERS:
            if name not in header_index:
                errors.append(f"标准数据平铺缺少字段: {name}")
        if not errors:
            # 校对册超过5万行；一次顺序遍历同时收集数据和检查保护，避免
            # 对每个单元格反复调用Worksheet.cell造成高CPU和不必要的XML解析。
            for row_number, cells in enumerate(
                sheet.iter_rows(min_row=4, max_col=len(FLAT_HEADERS)), start=4
            ):
                if cells[0].value in (None, ""):
                    continue
                values = [cell.value for cell in cells]
                actual_rows.append(values)
                status = values[16]
                if status in review_status_counts:
                    review_status_counts[status] += 1
                if status not in {"未校对", "正确", "需修改", "存疑"}:
                    errors.append(f"第{row_number}行校对状态无效: {status!r}")
                for column, cell in enumerate(cells[:16], start=1):
                    if not cell.protection.locked:
                        errors.append(f"第{row_number}行标准字段列{column}未锁定")
                for column, cell in enumerate(cells[16:19], start=17):
                    if cell.protection.locked:
                        errors.append(f"第{row_number}行校对字段列{column}未解锁")
                if len(errors) > 100:
                    errors.append("错误超过100项，后续错误省略")
                    break
            expected_ids = [str(row[0]) for row in expected]
            actual_ids = [str(row[0]) for row in actual_rows]
            if len(actual_rows) != len(expected):
                errors.append(f"标准数据平铺行数不一致: 校对册{len(actual_rows)}，机器数据{len(expected)}")
            if len(set(actual_ids)) != len(actual_ids):
                errors.append("标准数据平铺存在重复数据ID")
            if actual_ids != expected_ids:
                expected_set, actual_set = set(expected_ids), set(actual_ids)
                missing_ids = sorted(expected_set - actual_set)
                extra_ids = sorted(actual_set - expected_set)
                errors.append(
                    f"数据ID集合不一致: 缺少{len(missing_ids)}条、额外{len(extra_ids)}条"
                )
            # 比较机器平铺结果的16个锁定字段，确认校对册没有脱离当前
            # manifest/标准JSON而只保留了相同ID。Excel可能把空字符串读成
            # None、把30.0读成30，因此使用上面的规范化比较而不放宽文本值。
            expected_by_id = {
                str(row[0]): row
                for row in expected
                if row and row[0] not in (None, "")
            }
            for actual in actual_rows:
                if not actual:
                    continue
                data_id = str(actual[0])
                expected_row = expected_by_id.get(data_id)
                if expected_row is None:
                    continue
                for index, (expected_value, actual_value) in enumerate(
                    zip(expected_row[:16], actual[:16])
                ):
                    if _canonical_locked_value(expected_value) == _canonical_locked_value(actual_value):
                        continue
                    locked_value_mismatch_count += 1
                    if len(locked_value_mismatch_samples) < 20:
                        locked_value_mismatch_samples.append({
                            "data_id": data_id,
                            "field": FLAT_HEADERS[index],
                            "expected": expected_value,
                            "actual": actual_value,
                        })
            if locked_value_mismatch_count:
                errors.append(
                    "标准数据平铺锁定字段与机器数据不一致: "
                    f"{locked_value_mismatch_count}处"
                )
            actual_by_id = {
                str(row[0]): row
                for row in actual_rows
                if row and row[0] not in (None, "")
            }
            for data_id, expected_opinion in expected_no_data_rows.items():
                actual = actual_by_id.get(data_id)
                if actual is None:
                    # 数据ID集合错误已在上面报告；这里不重复生成一条相同错误。
                    continue
                status, replacement, opinion = actual[16], actual[17], actual[18]
                actual_no_data_rows[data_id] = {
                    "status": status,
                    "replacement": replacement,
                    "opinion": opinion,
                }
                if status != "正确":
                    errors.append(f"无数据记录未标记为正确: {data_id}")
                if replacement not in (None, ""):
                    errors.append(f"无数据记录不得填写建议修订值: {data_id}")
                if not any(token in str(opinion or "") for token in ("无数据", "无资料", "没有数据")):
                    errors.append(f"无数据记录的校对意见未明确说明无数据: {data_id}")
            actual_flat_rows = len(actual_rows)
        for name in REQUIRED_SHEETS:
            if name in workbook.sheetnames and not workbook[name].protection.sheet:
                errors.append(f"sheet未保护: {name}")
        if "淘汰目录_机电四批" in workbook.sheetnames:
            catalog = workbook["淘汰目录_机电四批"]
            if catalog.max_row < 4:
                warnings.append("淘汰目录_机电四批没有目录数据行")
    # 激活门禁是独立于结构验收的可读摘要：只有所有平铺记录已标记“正确”、
    # 且manifest中的标准包均为active时，才允许把校对册建议用于机器数据。
    try:
        manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
        for entry in manifest.get("packs", []):
            if entry.get("status") != "active":
                activation_blockers.append(
                    f"{entry.get('pack_id', entry.get('device_type', ''))}:状态为{entry.get('status', '')}"
                )
    except (OSError, json.JSONDecodeError):
        activation_blockers.append("无法读取manifest，不能确认标准包状态")
    if review_status_counts["未校对"] or review_status_counts["需修改"] or review_status_counts["存疑"]:
        activation_blockers.append(
            f"校对状态未全部为正确（未校对{review_status_counts['未校对']}、"
            f"需修改{review_status_counts['需修改']}、存疑{review_status_counts['存疑']}）"
        )
    return {
        "path": str(path),
        "is_valid": not errors,
        "errors": errors,
        "warnings": warnings,
        "expected_flat_rows": expected_flat_rows,
        "actual_flat_rows": actual_flat_rows,
        "locked_value_mismatch_count": locked_value_mismatch_count,
        "locked_value_mismatch_samples": locked_value_mismatch_samples,
        "review_status_counts": review_status_counts,
        "expected_no_data_rows": len(expected_no_data_rows),
        "actual_no_data_rows": len(actual_no_data_rows),
        "activation_ready": not activation_blockers,
        "activation_blockers": activation_blockers,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("book", type=Path)
    parser.add_argument(
        "--require-activation",
        action="store_true",
        help="除结构校验外，要求所有记录已标记正确且manifest标准包均为active",
    )
    args = parser.parse_args(argv)
    result = validate(args.book)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if not result["is_valid"]:
        return 2
    if args.require_activation and not result["activation_ready"]:
        return 3
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
