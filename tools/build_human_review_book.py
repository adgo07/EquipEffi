"""从当前标准manifest和机器标准包生成可读的人工校对册。

该脚本不修改标准JSON；它把同一份机器数据平铺为可人工核对的Excel，并将
标准值列锁定，只解锁校对状态、建议修订值和校对意见。
"""
from __future__ import annotations

from datetime import date
from copy import copy
import argparse
import json
from pathlib import Path
import re
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from .build_lock import build_lock
except ImportError:  # direct execution: ``python tools/build_human_review_book.py``
    from build_lock import build_lock


ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "src" / "equipeffi" / "standard_manifest.json"
RESOURCE_ROOT = MANIFEST.parent
ELIMINATION = RESOURCE_ROOT / "resources" / "elimination_catalog_batches_1_4.json"
INDUSTRY_ELIMINATION = RESOURCE_ROOT / "resources" / "elimination_catalog_industry_2024.json"
DEFAULT_OUTPUT = ROOT / "outputs" / f"final_{date.today().strftime('%Y%m%d')}" / "设备能效标准数据_人工校对版.xlsx"

FLAT_HEADERS = (
    "数据ID", "标准编号", "表号", "产品类别", "匹配条件", "原文区间", "结构化区间",
    "原文值", "原文单位", "规范化值", "规范化单位", "等级", "比较方向", "插值规则",
    "条款/页码", "来源文件", "校对状态", "建议修订值", "校对意见",
)
REVIEW_STATUSES = ("未校对", "正确", "需修改", "存疑")
RESULT_LEVELS = ("不在范围", "无法判定", "淘汰", "未达标", "1级", "2级", "3级", "4级", "5级", "一等", "二等", "三等", "节能评价值", "能效限定值")
STANDARD_BY_DEVICE = {
    "transformer": "GB 20052-2024", "motor_lv": "GB 18613-2020", "motor_hv": "GB 30254-2024",
    "motor_pmsm": "GB 30253-2024", "compressor": "GB 19153-2019", "pump_water": "GB 19762-2025",
    "pump_chemical": "GB 19762-2025", "fan": "GB 19761-2020", "blower": "GB 28381-2012",
    "submersible": "GB 32030-2022", "boiler": "GB 24500-2020", "heat_treatment": "GB/T 36561-2018",
    "heat_pump_chiller": "GB 19577-2024", "heat_pump_water_heater": "GB 29541-2013",
    "duct_ac": "GB 37479-2019", "unitary_ac": "GB 19576-2019", "multi_split_ac": "GB 21454-2021",
}
INDUSTRY_PDF_PAGE_BY_ID = {
    "IND2024-MOTOR-YB": 134, "IND2024-MOTOR-YBF": 134, "IND2024-MOTOR-YBK": 134,
    "IND2024-PUMP-BA": 134, "IND2024-PUMP-F": 134, "IND2024-PUMP-JD": 134,
    "IND2024-COMPRESSOR-3W": 135,
    "IND2024-PUMP-BOILER-FEED": 136, "IND2024-BOILER-FIXED-GRATE": 136,
    "IND2024-COMPRESSOR-L10": 136, "IND2024-FAN-HIGH-PRESSURE": 136,
    "IND2024-BOILER-COAL-10T": 137, "IND2024-BOILER-BIOMASS-2T": 137,
}

# 模板“注意事项”中的后续标准没有可启用的机器数据包；在人工校对册中
# 以独立目录行保留其发布/实施状态，避免被误认为当前判定依据。
FUTURE_STANDARDS = (
    {
        "device_type": "鼓风机（后续版本）",
        "standard_code": "GB 28381-2026",
        "pack_id": "gb28381_2026_pending",
        "status": "已发布、未实施",
        "data_version": "not-enabled",
        "source_file": "原标准PDF（待实施，未复制为机器标准包）",
        "note": "2027-04-01实施；本版仍采用GB 28381-2012，不启用其字段和指标",
    },
)


def _json_text(value: Any) -> str:
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if value is None:
        return ""
    return str(value)


def _number_like(value: Any) -> bool:
    if isinstance(value, bool) or value is None:
        return False
    if isinstance(value, (int, float)):
        return True
    return bool(re.fullmatch(r"[-+]?\d+(?:\.\d+)?", str(value).strip()))


def _level_from_path(path: str) -> str:
    match = re.search(r"(?:thresholds|eff|offsets)\[(\d+)\]", path)
    if match:
        return f"{int(match.group(1)) + 1}级"
    match = re.search(r"(?:^|[._-])([1-5])(?:级|$)", path)
    return f"{match.group(1)}级" if match else ""


def _table_from_context(path: str, context: dict[str, Any]) -> str:
    if context.get("table"):
        return str(context["table"])
    # 标准表使用 name=“表1”…“表8”；保留真实表号，避免把校对册中的
    # 记录序号误认为标准表号。
    if context.get("name") and re.fullmatch(r"表\d+", str(context["name"]).strip()):
        return str(context["name"]).strip()
    match = re.search(r"(?:^|[._])(?:tables|table8)\[(\d+)\]", path)
    return f"记录{int(match.group(1)) + 1}" if match else ""


def _efficiency_dimension(path: str, context: dict[str, Any]) -> tuple[str, str]:
    """从效率叶节点路径解码等级及表头维度，便于人工复核。"""

    match = re.search(r"\.efficiency\.([1-5])\[(\d+)\]$", path)
    if not match:
        return "", ""
    level = f"{match.group(1)}级"
    index = int(match.group(2))
    dims = context.get("dims")
    if isinstance(dims, list) and index < len(dims):
        return level, str(dims[index])
    return level, str(index)


def _condition_text(context: dict[str, Any], path: str = "") -> str:
    """生成可读的匹配条件，而不是只留下JSON路径。"""

    conditions = context.get("conditions")
    if conditions:
        # 保留标准原有条件，同时追加当前行/叶节点的定位信息。
        result = dict(conditions) if isinstance(conditions, dict) else {"conditions": conditions}
    else:
        result = {}
    keys = ("type", "category", "subtype", "product", "unit_type", "product_standard",
            "voltage_group", "cooling_group", "mode", "range_metric", "min", "max",
            "min_inclusive", "max_inclusive", "power_bins", "source_page")
    for key in keys:
        if key in context and key not in result:
            result[key] = context[key]
    if "power_rule" in context:
        result["额定功率区间"] = context["power_rule"]
    elif "power_min_kw" in context or "power_max_kw" in context:
        result["额定功率区间_kW"] = {
            key: context[key] for key in ("power_min_kw", "power_max_kw", "power_min_inclusive") if key in context
        }
    elif "power_kw" in context:
        result["额定功率_kW"] = context["power_kw"]
    level, dimension = _efficiency_dimension(path, context)
    if level:
        # 单等级表使用表自身的grade；三等级表从叶节点路径解码。
        result["等级"] = level
    if dimension:
        mode = str(context.get("mode", ""))
        result["极数" if mode == "poles" else "转速/维度"] = dimension
    return _json_text(result)


def _raw_interval(context: dict[str, Any]) -> str:
    """提取原文表中的功率/容量区间，保留原始区间文本。"""

    if context.get("power_rule") not in (None, ""):
        return str(context["power_rule"])
    if context.get("power_kw") not in (None, ""):
        return f"{context['power_kw']} kW"
    for key in ("range", "range_text", "original_range", "power_bins"):
        if context.get(key) not in (None, ""):
            return _json_text(context[key])
    return ""


def _product_category(context: dict[str, Any], device_type: str) -> str:
    conditions = context.get("conditions")
    if isinstance(conditions, dict):
        for key in ("category", "type", "unit_type", "product_standard", "subtype"):
            if conditions.get(key) not in (None, ""):
                return str(conditions[key])
    for key in ("category", "type", "product", "unit_type", "subtype"):
        if context.get(key) not in (None, ""):
            return str(context[key])
    return device_type


def _walk(value: Any, path: str = "", context: dict[str, Any] | None = None) -> Iterable[tuple[str, Any, dict[str, Any]]]:
    context = dict(context or {})
    if isinstance(value, dict):
        for key, child in value.items():
            # 这些是机器数据的复核/门禁元数据，不是标准数值记录，
            # 不应在“标准数据平铺”中生成伪数据行。
            if key in {"no_data_cells", "no_data_reason", "no_data_conclusion", "activation_review", "unavailable_reason", "data_version"}:
                continue
            child_path = f"{path}.{key}" if path else str(key)
            if key in {"table", "name", "type", "category", "subtype", "product", "unit_type", "product_standard", "conditions", "range_metric", "min", "max", "min_inclusive", "max_inclusive", "power_bins", "source_page", "source_pages", "pages", "data_id", "power_kw", "power_rule", "power_min_kw", "power_max_kw", "power_min_inclusive", "dims", "mode", "grade", "voltage_group", "cooling_group"}:
                context[key] = child
            yield from _walk(child, child_path, context)
        return
    if isinstance(value, list):
        for index, child in enumerate(value):
            yield from _walk(child, f"{path}[{index}]", context)
        return
    yield path, value, context


def _no_data_paths(payload: dict[str, Any]) -> dict[str, str]:
    """将标准包登记的无数据扁平索引转换为叶节点路径和校对意见。"""
    result: dict[str, str] = {}
    for table_index, table in enumerate(payload.get("tables", []) or []):
        dimensions = len(table.get("dims", []) or [])
        if dimensions <= 0:
            continue
        for row_index, row in enumerate(table.get("rows", []) or []):
            reason = str(row.get("no_data_reason") or "标准原文无数据，保留空值；不参与插值或等级比较")
            efficiency = row.get("efficiency", {})
            grades = [int(key) for key in efficiency if str(key) in {"1", "2", "3"}] if isinstance(efficiency, dict) else []
            for flat_index in row.get("no_data_cells", []) or []:
                index = int(flat_index)
                # 表8～表28每行只保存一个等级，索引从该等级的维度数组
                # 起算；表1～表7/表29才是三个等级数组拼接后的索引。
                if len(grades) == 1:
                    level, dimension = grades[0], index
                else:
                    level, dimension = index // dimensions + 1, index % dimensions
                result[f"tables[{table_index}].rows[{row_index}].efficiency.{level}[{dimension}]"] = reason
    return result


def _standard_rows() -> list[list[Any]]:
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    rows: list[list[Any]] = []
    for entry in manifest.get("packs", []):
        source = (MANIFEST.parent / entry["source"]).resolve()
        if not source.exists():
            continue
        payload = json.loads(source.read_text(encoding="utf-8"))
        device_type = str(entry.get("device_type", ""))
        standard = STANDARD_BY_DEVICE.get(device_type, str(entry.get("standard_code") or payload.get("standard_code") or entry.get("pack_id", "")))
        # 多个HVAC内部类型共用一个JSON文件；只展开当前manifest条目的
        # devices.<device_type>子树，避免把5类设备的数据重复写入校对册。
        scoped_payload = payload.get("devices", {}).get(device_type, payload)
        no_data_paths = _no_data_paths(scoped_payload)
        for path, value, context in _walk(scoped_payload):
            if not path or path.rsplit(".", 1)[-1] in {
                "standard_code", "standard_name", "status", "source_note",
                "suspicious_cells", "no_data_cells", "no_data_reason", "no_data_conclusion", "activation_review", "unavailable_reason", "data_version",
                # 来源页码和页码范围是记录上下文，不是待校对的标准值；
                # 它们已在上面的_walk中写入context，不能再平铺成伪记录。
                "source_page", "source_pages", "pages",
            }:
                continue
            explicit_data_id = context.get("data_id")
            # data_id在不同内部标准包之间也必须全局唯一：pump.json等资源
            # 可能被多个manifest包复用。保留pack_id前缀既避免重复，也能
            # 直接回溯到机器数据包；没有显式ID时继续使用JSON路径。
            base_data_id = (
                f"{entry['pack_id']}:{explicit_data_id}"
                if explicit_data_id
                else f"{entry['pack_id']}:{path}"
            )
            # 一个标准记录会展开为多个叶节点；路径后缀保证校对册中的数据ID逐行唯一。
            data_id = base_data_id if path.endswith(".data_id") else f"{base_data_id}:{path}"
            # 已由机器包登记且经用户确认的无数据叶节点，人工校对册用
            # “—”明确展示其语义；机器JSON仍保留null，避免把占位符当作
            # 可参与计算的数值。其他真正的空值继续显示为空白。
            review_value = "—" if path in no_data_paths and value is None else value
            normalized = review_value if _number_like(review_value) else _json_text(review_value)
            rows.append([
                data_id,
                standard,
                _table_from_context(path, context),
                _product_category(context, device_type),
                _condition_text(context, path),
                _raw_interval(context),
                path,
                _json_text(review_value),
                "",
                normalized,
                "",
                _level_from_path(path),
                ">=" if any(token in path.lower() for token in ("eff", "threshold", "offset")) else "",
                "",
                str(context.get("source_page") or context.get("source_pages") or context.get("pages") or ""),
                str(entry.get("source_file") or source),
                "正确" if path in no_data_paths else "未校对",
                "",
                no_data_paths.get(path, ""),
            ])
    return rows


def _formula_rows() -> list[list[Any]]:
    return [
        ["RULE-FAN-EFF", "GB 19761-2020", "通风机最高效率", "ηr=qvsg1·pF·kp/(1000Pr)×100%", "qvsg1进口滞止容积流量，pF风机压力，kp压缩性修正，Pr叶轮功率", "%", "效率空白时按式(1)计算；不以额定电机功率替代叶轮功率", "第5.2.1条/式(1)"],
        ["RULE-FAN-KP", "GB/T 1236 / GB 19761-2020", "通风机压缩性修正", "kp=ln(1+x)/x×Zp/ln(1+Zp)", "x=(pF/psg1)，Zp=(k-1)/k×Pr/(qvsg1·psg1)", "—", "等熵压缩修正；不缺参数时才计算", "GB/T 1236-2017 14.8.2.2"],
        ["RULE-BLOWER-ETA", "GB 28381-2012", "鼓风机多变效率", "ηp=((k-1)/k×ln(p2/p1))/ln(T2/T1)×100", "p为绝对压力，T为绝对温度，k为绝热指数", "%", "出口压力和温度必须高于进口", ""],
        ["RULE-HT-FUEL", "GB/T 36561-2018", "燃料炉可比单耗", "bk1=QDW×B×α/(29308×Gz)", "QDW燃料热值，B燃料耗量，α为表9燃料系数，Gz折合重量", "kgce/t", "按表9燃料品种取α；无法唯一匹配时不判定", ""],
        ["RULE-BOILER-CAPACITY", "GB 24500-2020", "锅炉容量列选择", "表1：D≤20t/h或Q≤14MW；表3：D≤10t/h或Q≤7MW", "D蒸发量，Q热功率；两者同时填写且落入不同列时无法判定", "t/h或MW", "不把MW数值直接套用t/h阈值", ""],
        ["RULE-BOILER-ELECTRIC", "GB 24500-2020", "电加热锅炉能效限定值", "η≥97%（第5.2条）", "η为额定工况设计热效率；标准未给出电锅炉1级/2级分档，达到限定值映射为3级", "%", "电锅炉/电力只检查固定门槛，不套用燃料锅炉表", ""],
        ["RULE-PUMP-LEVEL", "GB 32030-2022", "潜水电泵三级阈值", "η1=ηDB+偏移1；η2=ηDB+偏移2；η3=ηDB−Δη", "ηDB规定效率，Δη效率容差", "%", "“—”不作要求，不当作0", ""],
        ["RULE-PUMP-CHEM-BOUNDARY", "GB 19762-2025", "石油化工离心泵比转数分档边界", "20≤ns<60；60≤ns<120；120≤ns≤210；210<ns≤300", "ns为按公式(1)计算的比转数", "—", "严格保留原文开闭端点；60和120不沿用前一档", "表2"],
        ["RULE-SUBMERSIBLE-ETA", "GB/T 25409-2010", "小型潜水电泵规定效率", "ηB=ηSP−Δη；ηDB=ηB×ηD/100−1.5", "ηSP按流量/泵型表A1，Δη按比转速表A2，ηD按功率/相数/结构/同步转速表4", "%", "仅标准精确档；不插值、不外推", ""],
        ["RULE-MULTI-EER", "GB 21454-2021", "多联机等级AND条件", "主指标与EERmin/低温COP按同一等级同时满足", "SEER/APF/IPLV(C)/HSPF及固定门槛", "W/W或Wh/Wh", "静压大于0时需先按引用标准修正", ""],
        ["RULE-GB19577-CHANNEL", "GB 19577-2024", "热泵和冷水机组表1～表8选表与比较", "主指标按对应表的1～3级阈值比较；表1/表2/表7/表8的单列指标作为3级固定门槛；表1与表2为二选一指标体系；表5蒸汽型按≤比较", "产品类别、产品标准、机组型式、冷却/热源方式、名义容量和评价指标体系；表2的CSPF/IPLV/ACCOP替代指标", "按表中指标", "表1或表2只选一个指标体系；等于阈值视为达到；仅达到3级主指标时还必须满足同级固定门槛；—不作容量分档或不设置该指标", "PDF第9～12页"],
    ]


def _style_sheet(ws, freeze: str = "A4", filter_ref: str | None = None) -> None:
    ws.freeze_panes = freeze
    if filter_ref:
        ws.auto_filter.ref = filter_ref
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column in range(1, ws.max_column + 1):
        values = [str(ws.cell(row, column).value or "") for row in range(1, min(ws.max_row, 50) + 1)]
        width = min(42, max(12, max((len(value) for value in values), default=10) + 2))
        ws.column_dimensions[get_column_letter(column)].width = width


def _write_book_unlocked(output: Path) -> Path:
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    wb.security.lockStructure = True

    ws = wb.create_sheet("校对说明")
    ws.append(["设备能效标准数据人工校对册"])
    ws.append([])
    ws.append(["用途", "本册与src/equipeffi/standard_manifest.json及resources/standards同源生成；标准值列锁定，校对列供人工填写。"])
    ws.append(["状态流转", "extracted → normalized → verified → active；当前状态以manifest为准。"])
    ws.append(["百分数规则", "效率按1～100本值填写，例如98%填写98；COP/EER/SEER/APF为比值，不套用百分数范围。"])
    ws.append(["生成日期", date.today().isoformat()])
    _style_sheet(ws, freeze="A3")

    ws = wb.create_sheet("标准目录")
    ws.append(["本册包含的标准与机器数据包"])
    ws.append([])
    ws.append(["对应内部类型", "标准编号", "数据包ID", "状态", "数据版本", "来源文件", "说明"])
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    for entry in manifest.get("packs", []):
        device_type = str(entry.get("device_type", ""))
        # manifest以数据包ID为主键，标准编号统一从同源设备映射补齐，
        # 避免人工校对册的“标准目录”出现空白标准编号。
        standard_code = entry.get("standard_code") or STANDARD_BY_DEVICE.get(device_type, "")
        ws.append([device_type, standard_code, entry.get("pack_id", ""), entry.get("status", ""), entry.get("data_version", ""), entry.get("source_file", entry.get("source", "")), entry.get("unavailable_reason", "")])
    for entry in FUTURE_STANDARDS:
        ws.append([entry["device_type"], entry["standard_code"], entry["pack_id"], entry["status"], entry["data_version"], entry["source_file"], entry["note"]])
    _style_sheet(ws, freeze="A4", filter_ref=f"A3:G{ws.max_row}")

    ws = wb.create_sheet("标准数据平铺")
    ws.append(["标准原文/结构化数据平铺（机器数据对应记录）"])
    ws.append([])
    ws.append(list(FLAT_HEADERS))
    for row in _standard_rows():
        ws.append(row)
    _style_sheet(ws, freeze="A4", filter_ref=f"A3:S{ws.max_row}")
    for row in ws.iter_rows(min_row=4, min_col=17, max_col=19):
        for cell in row:
            protection = copy(cell.protection)
            protection.locked = False
            cell.protection = protection
    ws.protection.sheet = True

    ws = wb.create_sheet("公式与判定规则")
    ws.append(["公式、比较方向、变量和示例"])
    ws.append([])
    ws.append(["规则ID", "设备/标准", "中文规则", "数学表达式", "变量说明", "单位", "示例", "备注"])
    for row in _formula_rows():
        ws.append(row)
    _style_sheet(ws, freeze="A4", filter_ref=f"A3:H{ws.max_row}")

    ws = wb.create_sheet("分类与枚举")
    ws.append(["标准分类、设备类别及判定结论枚举"])
    ws.append([])
    ws.append(["枚举ID", "枚举名称", "值", "来源", "备注"])
    for index, value in enumerate(RESULT_LEVELS, start=1):
        ws.append([f"CONCLUSION-{index:02d}", "能效结论", value, "判定引擎", "按设备类型限制可用等级"])
    _style_sheet(ws, freeze="A4", filter_ref=f"A3:E{ws.max_row}")

    ws = wb.create_sheet("单位及换算")
    ws.append(["单位规范及允许换算关系"])
    ws.append([])
    ws.append(["量纲", "原单位", "规范单位", "换算关系", "处理约束", "备注"])
    for row in [
        ["效率", "%", "%", "不换算", "1≤x≤100", "98%填写98"],
        ["功率", "W", "kW", "kW=W/1000", "正数", "输入端按字段单位"],
        ["空调容量", "kW", "W", "W=kW×1000", "用于标准分档", "V4输入kW"],
        ["压力", "kPa", "Pa", "Pa=kPa×1000", "绝对压力需标明", "鼓风机公式使用绝对压力"],
    ]:
        ws.append(row)
    _style_sheet(ws, freeze="A4", filter_ref=f"A3:F{ws.max_row}")

    ws = wb.create_sheet("淘汰目录_产业")
    ws.append(["《产业结构调整指导目录（2024年本）》设备相关淘汰条目（PDF核对版，非全文）"])
    ws.append([])
    ws.append(["目录名称", "条目号", "产品/型号条件", "原文条件", "状态", "来源文件", "PDF页码", "校对意见"])
    if INDUSTRY_ELIMINATION.exists():
        industry_catalog = json.loads(INDUSTRY_ELIMINATION.read_text(encoding="utf-8"))
        for item in industry_catalog.get("entries", []):
            ws.append([item.get("catalog", ""), item.get("item_no", ""), _json_text(item.get("model_exact", []) or item.get("series_patterns", [])), item.get("original_condition", ""), item.get("matching_status", ""), item.get("source_file", ""), INDUSTRY_PDF_PAGE_BY_ID.get(str(item.get("entry_id", "")), ""), item.get("notes", "")])
    _style_sheet(ws, freeze="A4", filter_ref=f"A3:H{ws.max_row}")
    # 产业目录原文和匹配规则锁定；校对意见允许人工补充。
    for row in ws.iter_rows(min_row=4, min_col=8, max_col=8):
        for cell in row:
            protection = copy(cell.protection)
            protection.locked = False
            cell.protection = protection
    ws.protection.sheet = True

    ws = wb.create_sheet("淘汰目录_机电四批")
    ws.append(["高耗能落后机电设备淘汰目录第一至第四批"])
    ws.append([])
    ws.append(["数据ID", "批次", "条目号", "对应设备类型", "产品名称", "精确型号", "原文系列", "受控系列正则", "附加条件", "原文条件", "来源文件", "匹配状态", "校对状态", "校对意见"])
    if ELIMINATION.exists():
        catalog = json.loads(ELIMINATION.read_text(encoding="utf-8"))
        for item in catalog.get("entries", []):
            ws.append([item.get("entry_id", ""), item.get("batch", ""), item.get("item_no", ""), item.get("device_type", ""), item.get("product_name", ""), _json_text(item.get("model_exact", [])), _json_text(item.get("model_series_original", [])), _json_text(item.get("series_patterns", [])), _json_text(item.get("conditions", {})), item.get("original_condition", ""), item.get("source_file", ""), item.get("matching_status", ""), "未校对", ""])
    _style_sheet(ws, freeze="A4", filter_ref=f"A3:N{ws.max_row}")
    for row in ws.iter_rows(min_row=4, min_col=12, max_col=14):
        for cell in row:
            protection = copy(cell.protection)
            protection.locked = False
            cell.protection = protection
    ws.protection.sheet = True

    review_issues: list[list[Any]] = []
    # 将机器包中的明确存疑单元格同步列出，避免用户只能在平铺数据中
    # 通过搜索发现问题。PMSM的空值单元格必须人工对照PDF；确认原文无数据时，
    # 校对意见明确填写“无数据”，建议修订值留空，不得猜改。
    for entry in manifest.get("packs", []):
        source = (MANIFEST.parent / entry["source"]).resolve()
        if not source.exists():
            continue
        payload = json.loads(source.read_text(encoding="utf-8"))
        for table in payload.get("tables", []):
            for row in table.get("rows", []):
                suspicious = row.get("suspicious_cells") or []
                for cell_index in suspicious:
                    pages = table.get("pages") or table.get("source_pages") or []
                    page_text = "、".join(str(page) for page in pages) if pages else "待查"
                    image_hint = (
                        f"GB30253-2024_P{int(pages[0]):02d}.png"
                        if pages and str(pages[0]).isdigit() else ""
                    )
                    review_issues.append([
                        str(entry.get("standard_code") or STANDARD_BY_DEVICE.get(entry.get("device_type", ""), payload.get("standard_code", ""))),
                        f"{table.get('table', table.get('name', ''))}、功率{row.get('power_kw', row.get('power_rule', ''))}的效率数据索引{cell_index}为空，需对照PDF确认；若原文无数据请在校对意见填写‘无数据’",
                        "保留原值并记录待复核；当前标准包按manifest状态运行",
                        "渲染PDF逐格核对；确认原文后才能填写建议修订值并激活标准包",
                        f"{entry.get('pack_id', '')}:{table.get('table', table.get('name', ''))}:效率[{cell_index}]",
                        page_text,
                        image_hint,
                    ])
                for cell_index in row.get("no_data_cells", []) or []:
                    pages = table.get("pages") or table.get("source_pages") or []
                    page_text = "、".join(str(page) for page in pages) if pages else "待查"
                    image_hint = (
                        f"GB30253-2024_P{int(pages[0]):02d}.png"
                        if pages and str(pages[0]).isdigit() else ""
                    )
                    review_issues.append([
                        str(entry.get("standard_code") or STANDARD_BY_DEVICE.get(entry.get("device_type", ""), payload.get("standard_code", ""))),
                        f"{table.get('table', table.get('name', ''))}、功率{row.get('power_kw', row.get('power_rule', ''))}的效率数据索引{cell_index}已按原文无数据处理",
                        "保留空值，不参与插值或等级比较；表1/55 kW/12极三档命中时按用户复核结果判定为不在范围",
                        "如后续取得新的PDF证据，须走人工校对和标准包版本变更流程，不得直接填入标准值",
                        f"{entry.get('pack_id', '')}:{table.get('table', table.get('name', ''))}:效率[{cell_index}]",
                        page_text,
                        image_hint,
                    ])
    for entry in manifest.get("packs", []):
        if entry.get("status") not in {"active", "verified"} and entry.get("unavailable_reason"):
            review_issues.append([
                str(entry.get("standard_code") or STANDARD_BY_DEVICE.get(entry.get("device_type", ""), "")),
                str(entry.get("unavailable_reason")),
                f"当前状态：{entry.get('status', '')}",
                "完成来源复核后更新manifest状态",
                str(entry.get("pack_id", "")),
                "",
                "",
            ])

    for title, headers, rows in [
        ("待校对问题", ["标准编号", "问题", "当前处理", "建议动作", "数据ID/位置", "PDF页码", "渲染图像提示"], review_issues),
        ("版本变更记录", ["日期", "版本/数据包", "变更", "影响", "状态"], [
            [date.today().isoformat(), "gb21454_2021_v1", "EERmin改为逐等级阈值；增加静压修正门禁", "多联式空调判定和V4结果列", "normalized/待人工核对"],
            [date.today().isoformat(), "gb32030_2022_v1", "增加GB/T25409附录A小型潜水电泵ηDB条件式计算；只使用精确档，不插值/外推", "潜水电泵计算指标和判定轨迹", "active/待人工核对引用表4"],
            [date.today().isoformat(), "gb24500_2020_v1", "增加第5.2条电加热锅炉能效限定值≥97%的固定门槛；V4电锅炉类别自动拆分燃烧方式", "工业锅炉电加热产品判定", "active/待真实产品样例验收"],
            [date.today().isoformat(), "gbt36561_2018_v1", "热处理燃料炉按表9燃料品种引入折标系数α；未知燃料不再默认α=1", "热处理设备燃料可比单耗计算", "active/待人工核对表9单位"],
            [date.today().isoformat(), "gb19577_2024_pdf_verified_v1", "按PDF逐表复核并启用表1～表8，共69条记录；增加固定3级门槛和表5反向比较", "热泵和冷水机组8类产品判定", "active/待真实产品样例验收"],
            [date.today().isoformat(), "gb28381_2026_pending", "登记已发布但尚未实施的GB 28381-2026；当前机器判定继续采用GB 28381-2012", "鼓风机后续标准切换准备", "published_not_effective/未启用"]
        ]),
    ]:
        ws = wb.create_sheet(title)
        ws.append([title])
        ws.append([])
        ws.append(headers)
        for row in rows:
            ws.append(row)
        _style_sheet(ws, freeze="A4", filter_ref=f"A3:{get_column_letter(len(headers))}{ws.max_row}")

    # 校对状态和建议列是唯一允许人工改写的内容；其余sheet默认保护。
    for ws in wb.worksheets:
        if not ws.protection.sheet:
            ws.protection.sheet = True
    wb.save(output)
    return output


def _write_book(output: Path) -> Path:
    """在生成大体量校对册时占用项目级互斥锁。"""

    with build_lock(ROOT, operation="人工校对册生成"):
        return _write_book_unlocked(output)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="生成设备能效标准人工校对册")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="输出xlsx路径")
    args = parser.parse_args(argv)
    print(_write_book(Path(args.output).resolve()))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
