# -*- coding: utf-8 -*-
"""验证低压/高压电机及PDF独立重建的永磁同步电机数据。

禁止读取 standards/motor_pmsm.json、standards2/motor_pmsm.json 或任何粗校对
永磁工作簿；GB 30253-2024 只验证生产仓库使用的新29表数据包。
"""
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "standards2"
PMSM_VERIFIED = ROOT / "src/equipeffi/resources/standards/gb30253_2024_pdf_verified_v1.json"


def values(row, level, dims):
    return [row.get("efficiency", {}).get(level, {}).get(str(i)) for i in range(len(dims))]


def main():
    lines = ["# 电机标准数据最终验证", ""]
    total_level = 0
    total_trend = 0
    for key, source in (
        ("motor_lv", OUT / "motor_lv.json"),
        ("motor_hv", OUT / "motor_hv.json"),
        ("motor_pmsm_pdf_verified", PMSM_VERIFIED),
    ):
        data = json.loads(source.read_text(encoding="utf-8"))
        if key == "motor_pmsm_pdf_verified":
            if data.get("pack_id") != "gb30253_2024_pdf_verified_v1" or data.get("verified_table_count") != 29:
                raise RuntimeError("永磁同步电机PDF重建包ID或29表完整性校验失败")
        tables = data.get("tables")
        if tables is None:
            tables = [{"dims": data.get("dims", []), "rows": data.get("rows", [])}]
        level_bad = 0
        trend_bad = 0
        row_count = 0
        for table in tables:
            rows = table.get("rows", [])
            dims = table.get("dims", [])
            row_count += len(rows)
            for row in rows:
                for i in range(len(dims)):
                    level_data = row.get("efficiency", {})
                    def at(level):
                        values = level_data.get(level, {})
                        return values.get(str(i)) if isinstance(values, dict) else values[i] if i < len(values) else None
                    a, b, c = at("1"), at("2"), at("3")
                    if a is not None and b is not None and a < b:
                        level_bad += 1
                    if b is not None and c is not None and b < c:
                        level_bad += 1
            for i in range(len(dims)):
                for level in ("1", "2", "3"):
                    prev = None
                    for row in rows:
                        values = row.get("efficiency", {}).get(level, {})
                        cur = values.get(str(i)) if isinstance(values, dict) else values[i] if i < len(values) else None
                        if cur is not None and prev is not None and cur < prev - 1e-9:
                            trend_bad += 1
                        if cur is not None:
                            prev = cur
        total_level += level_bad
        total_trend += trend_bad
        lines.append(f"- {key}: {row_count}行；等级顺序异常 {level_bad}；按功率下降异常 {trend_bad}")

    wb = openpyxl.load_workbook(OUT / "校对表_全部设备_三来源复核.xlsx", data_only=True)
    for sheet in ("低压电机", "高压电机"):
        ws = wb[sheet]
        nonempty = sum(1 for row in ws.iter_rows(min_row=2) for c in row[1:6] if c.value is not None)
        lines.append(f"- 校对表 {sheet}: {ws.max_row - 1}行，核心数据非空单元格 {nonempty}")
    wb.close()
    pmsm = json.loads(PMSM_VERIFIED.read_text(encoding="utf-8"))
    suspicious = sum(len(row.get("suspicious_cells", [])) for table in pmsm["tables"] for row in table["rows"])
    no_data = sum(len(row.get("no_data_cells", [])) for table in pmsm["tables"] for row in table["rows"])
    table1 = next((table for table in pmsm["tables"] if int(table.get("table_no", -1)) == 1), None)
    confirmed = []
    if table1 is not None:
        row55 = next((row for row in table1.get("rows", []) if float(row.get("power_kw", float("nan"))) == 55.0), None)
        if row55 is not None:
            dims = table1.get("dims", [])
            if 12 in dims:
                dim_index = dims.index(12)
                confirmed = [
                    row55.get("efficiency", {}).get(str(level), [])[dim_index]
                    if isinstance(row55.get("efficiency", {}).get(str(level)), list)
                    and dim_index < len(row55.get("efficiency", {}).get(str(level), []))
                    else None
                    for level in (1, 2, 3)
                ]
    if confirmed != [None, None, None] or not table1 or not any(
        5 in row.get("no_data_cells", []) and 12 in row.get("no_data_cells", []) and 19 in row.get("no_data_cells", [])
        for row in table1.get("rows", [])
        if float(row.get("power_kw", float("nan"))) == 55.0
    ):
        raise RuntimeError("GB 30253-2024表1/55 kW/12极三档无数据登记缺失或被改写")
    lines.append(
        f"- 永磁电机校对源：仅PDF重建包；{len(pmsm['tables'])}张表；存疑单元格 {suspicious}（命中时无法判定）；"
        f"已确认无数据单元格 {no_data}（含表1/55 kW/12极1～3级，命中时不在范围）"
    )
    lines += ["", f"汇总：等级顺序异常 {total_level}，按功率下降异常 {total_trend}。", "", "说明：趋势异常只作为人工复核提示；标准表中功率档位、极数、转速和工况变化可能导致局部非单调，不能仅凭趋势自动改值。"]
    (OUT / "最终验证结果.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("\n".join(lines))


if __name__ == "__main__":
    main()
