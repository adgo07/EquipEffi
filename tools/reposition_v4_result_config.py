"""把新增的统一结果字段移入V4配置字段字典的结束标记之前。"""
from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "src" / "equipeffi" / "resources" / "templates" / "设备能效分析空白模板_重构版V4_20260825.xlsx"
DEVICE_SHEETS = {
    "变压器", "电动机", "空压机", "离心泵", "离心通风机", "轴流通风机", "鼓风机",
    "潜水电泵", "工业锅炉", "热处理设备", "热泵和冷水机组", "热泵热水机",
    "风管送风式空调", "单元式空调", "多联式空调",
}
FIELD_IDS = {"standard_code", "standard_table", "reference_grade", "explanation", "missing_fields"}


def patch(path: Path = TEMPLATE) -> Path:
    workbook = load_workbook(path)
    config = workbook["配置"]
    rows_to_move: list[list[object]] = []
    for row in range(2, config.max_row + 1):
        if row <= 364:
            continue
        sheet = str(config.cell(row, 1).value or "")
        field_id = str(config.cell(row, 2).value or "")
        if sheet in DEVICE_SHEETS and field_id in FIELD_IDS:
            rows_to_move.append([config.cell(row, col).value for col in range(1, 17)])
    if not rows_to_move:
        return path
    first = next((row for row in range(2, config.max_row + 1) if not config.cell(row, 1).value and not config.cell(row, 2).value), None)
    if first is None:
        first = config.max_row + 1
    # 新增行位于辅助映射区末尾，连续删除后再插到字段字典结束标记处。
    last_rows = [row for row in range(2, config.max_row + 1) if str(config.cell(row, 1).value or "") in DEVICE_SHEETS and str(config.cell(row, 2).value or "") in FIELD_IDS and row >= first]
    if last_rows:
        config.delete_rows(min(last_rows), len(last_rows))
    config.insert_rows(first, len(rows_to_move))
    for offset, values in enumerate(rows_to_move):
        row = first + offset
        for col, value in enumerate(values, start=1):
            target = config.cell(row, col)
            target.value = value
            target._style = copy(config.cell(2, col)._style)
            target.alignment = copy(config.cell(2, col).alignment)
            target.protection = copy(config.cell(2, col).protection)
    workbook.save(path)
    return path


if __name__ == "__main__":
    print(patch())
