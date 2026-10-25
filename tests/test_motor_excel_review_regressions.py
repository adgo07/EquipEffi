"""PR 35 regressions for motor percent input formats and sparse V6.2 rows."""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook

from equipeffi.composition import create_batch_evaluation_service, create_motor_analysis_service
from equipeffi.infrastructure.excel.ooxml_reader import OOXMLWorkbook
from equipeffi.infrastructure.excel.unified_workbook_reader import MotorWorkbookReader
from equipeffi.infrastructure.persistence.app_data_paths import AppDataPaths
from equipeffi.infrastructure.persistence.sqlite_batch_record_repository import SqliteBatchRecordRepository


class MotorExcelReviewRegressionTests(unittest.TestCase):
    @staticmethod
    def _label() -> str:
        service = create_motor_analysis_service(with_persistence=False)
        return next(item.label for item in service.input_contract.categories
                    if item.category_id == "three_phase_induction")

    @staticmethod
    def _fill_row(sheet, row: int, *, power=0.14, efficiency=71.43,
                  number_format: str | None = None) -> None:
        for column, value in {
            "D": 1, "F": MotorExcelReviewRegressionTests._label(),
            "G": 0.38, "H": 50, "J": power, "K": 4, "M": efficiency,
        }.items():
            sheet[f"{column}{row}"] = value
        if number_format is not None:
            sheet[f"M{row}"].number_format = number_format

    @staticmethod
    def _batch(root: Path):
        paths = AppDataPaths(root / "data")
        return create_batch_evaluation_service(paths=paths, motor_candidate=True), paths

    def test_percentage_formats_and_literal_suffixes(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            batch, _ = self._batch(root)
            source = root / "formats.xlsx"
            batch.template_resource.download_to(source)
            workbook = load_workbook(source)
            sheet = workbook["电动机"]
            for row, number_format in {
                4: "0%", 5: "0.00%", 6: "0.000%",
                7: '0"%"', 8: r"0\%",
            }.items():
                self._fill_row(sheet, row, efficiency=1, number_format=number_format)
            workbook.save(source)
            workbook.close()

            with OOXMLWorkbook(source) as xlsx:
                metadata = xlsx.cell_metadata("电动机")
            with ZipFile(source) as archive:
                styles = ElementTree.fromstring(archive.read("xl/styles.xml"))
            cell_xfs = styles.find(
                "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}cellXfs")
            self.assertIsNotNone(cell_xfs)
            xfs = list(cell_xfs)
            self.assertEqual(xfs[metadata["M4"].style_index].attrib["numFmtId"], "9")
            self.assertEqual(xfs[metadata["M5"].style_index].attrib["numFmtId"], "10")
            self.assertTrue(metadata["M4"].is_percentage)
            self.assertTrue(metadata["M5"].is_percentage)
            self.assertTrue(metadata["M6"].is_percentage)
            self.assertFalse(metadata["M7"].is_percentage)
            self.assertFalse(metadata["M8"].is_percentage)

            rows = {row.row_number: row for row in MotorWorkbookReader().read(source).rows}
            for number in (4, 5, 6):
                with self.subTest(row=number):
                    self.assertTrue(rows[number].input_errors)
                    self.assertEqual(rows[number].values["measured_efficiency_pct"], "")
                    self.assertTrue(any("普通数值" in message
                                        for message in rows[number].input_errors))
            for number in (7, 8):
                with self.subTest(row=number):
                    self.assertFalse(rows[number].input_errors)
                    self.assertEqual(rows[number].values["measured_efficiency_pct"], 1)

    def test_percent_conflict_and_plain_100_through_batch_and_history(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            batch, paths = self._batch(root)
            source = root / "percent-conflict.xlsx"
            batch.template_resource.download_to(source)
            workbook = load_workbook(source)
            self._fill_row(workbook["电动机"], 4, efficiency=1, number_format="0%")
            self._fill_row(workbook["电动机"], 5, efficiency=100)
            workbook.save(source)
            workbook.close()
            original = source.read_bytes()

            result = batch.evaluate_workbook(source)

            self.assertEqual(source.read_bytes(), original)
            self.assertEqual(result.summary.input_error_quantity, 1)
            self.assertEqual(result.summary.evaluated_quantity, 1)
            percent_outcome = next(
                item["outcome"] for item in result.outcomes
                if item["sheet_name"] == "电动机"
                and item["outcome"].row_number == 4)
            self.assertEqual(percent_outcome.evaluation_status, "INVALID_INPUT")
            self.assertTrue(percent_outcome.is_input_error)
            plain_outcome = next(
                item["outcome"] for item in result.outcomes
                if item["sheet_name"] == "电动机"
                and item["outcome"].row_number == 5)
            self.assertEqual(plain_outcome.evaluation_status, "SUCCESS")
            output = load_workbook(result.result_workbook, data_only=True)
            self.assertIsNone(output["电动机"]["Q4"].value)
            self.assertIn("百分比", output["电动机"]["T4"].value)
            self.assertIn("普通数值", output["电动机"]["T4"].value)
            self.assertIn(output["电动机"]["Q5"].value,
                          {"1级", "2级", "3级", "未达3级", "不在范围"})
            output.close()
            record = SqliteBatchRecordRepository(paths.records_db).load_batch_record(
                result.batch_record_id)
            self.assertEqual(record.summary["input_error_rows"], 1)
            self.assertEqual(record.summary["conclusion_quantities"],
                             result.summary.conclusion_quantities)

    def test_sparse_v62_rows_write_direct_interpolation_and_reimport(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            batch, paths = self._batch(root)
            source = root / "sparse-v62.xlsx"
            batch.template_resource.download_to(source)
            workbook = load_workbook(source)
            sheet = workbook["电动机"]
            self._fill_row(sheet, 104, power=0.12, efficiency=70)
            self._fill_row(sheet, 105, power=0.14, efficiency=71.43)
            result_columns = {"N", "O", "P", "Q", "T"}
            self.assertFalse(any(cell.row in {104, 105}
                                 and cell.column_letter in result_columns
                                 for cell in sheet._cells.values()))
            workbook.save(source)
            workbook.close()
            original = source.read_bytes()
            with OOXMLWorkbook(source) as xlsx:
                retained_parts = {name: xlsx._sheets[name]
                                  for name in ("离心泵", "变压器")}
            with ZipFile(source) as archive:
                retained_bytes = {name: archive.read(part)
                                  for name, part in retained_parts.items()}

            result = batch.evaluate_workbook(source)

            self.assertEqual(source.read_bytes(), original)
            with ZipFile(result.result_workbook) as archive:
                for name, part in retained_parts.items():
                    with self.subTest(retained_sheet=name):
                        self.assertEqual(archive.read(part), retained_bytes[name])
            output = load_workbook(result.result_workbook, data_only=False)
            output_sheet = output["电动机"]
            source_workbook = load_workbook(source, data_only=False)
            source_sheet = source_workbook["电动机"]
            for row in (104, 105):
                for column in ("D", "F", "G", "H", "J", "K", "M"):
                    with self.subTest(input_cell=f"{column}{row}"):
                        self.assertEqual(output_sheet[f"{column}{row}"].value,
                                         source_sheet[f"{column}{row}"].value)
            self.assertIsNotNone(output_sheet["Q104"].value)
            self.assertIsNotNone(output_sheet["Q105"].value)
            self.assertIsNotNone(output_sheet["N104"].value)
            self.assertIsNotNone(output_sheet["P104"].value)
            self.assertEqual(output_sheet["N104"].number_format, "0.###############")
            self.assertEqual(output_sheet["P104"].number_format, "0.###############")
            self.assertEqual(output_sheet["N105"].number_format, "0.00")
            self.assertEqual(output_sheet["P105"].number_format, "0.00")
            self.assertEqual(output_sheet["Q104"]._style, output_sheet["Q103"]._style)
            self.assertEqual(output_sheet["T104"]._style, output_sheet["T103"]._style)
            source_workbook.close()
            output.close()

            record = SqliteBatchRecordRepository(paths.records_db).load_batch_record(
                result.batch_record_id)
            self.assertEqual(record.summary["data_row_count"], 2)
            self.assertEqual(record.summary["conclusion_quantities"],
                             result.summary.conclusion_quantities)
            reimported = batch.evaluate_workbook(Path(result.result_workbook))
            self.assertEqual(reimported.summary.conclusion_quantities,
                             result.summary.conclusion_quantities)
            self.assertEqual(len(SqliteBatchRecordRepository(paths.records_db)
                                 .list_batch_records()), 2)


if __name__ == "__main__":
    unittest.main()
