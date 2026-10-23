"""V6.1 Excel 与单项 Application 入口的变压器电压/阻抗交叉验收。"""
from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from equipeffi.composition import (
    create_batch_evaluation_service,
    create_transformer_analysis_service,
)
from equipeffi.application.services.transformer_analysis_service import (
    TransformerAnalysisRequest,
)
from equipeffi.domain.devices.transformer.input_contract import (
    TRANSFORMER_CATEGORIES,
    condition_options,
)
from equipeffi.infrastructure.excel.ooxml_reader import OOXMLWorkbook
from equipeffi.infrastructure.excel.template_resource import (
    DEFAULT_V61_TEMPLATE,
    TRANSFORMER_SHEET_CONTRACT,
)
from equipeffi.infrastructure.excel.unified_workbook_reader import (
    UnifiedWorkbookReader,
)
from equipeffi.infrastructure.persistence.app_data_paths import AppDataPaths


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_PATH = (ROOT / "src" / "equipeffi" / "resources" / "templates"
                 / DEFAULT_V61_TEMPLATE)
AS_OF = date(2026, 10, 23)


def _category(table_no: int):
    return next(item for item in TRANSFORMER_CATEGORIES
                if item.table_no == table_no)


def _thresholds(no_load_1: str, load_1: str,
                no_load_2: str, load_2: str,
                no_load_3: str, load_3: str) -> dict[str, Decimal]:
    return {
        "no_load_loss_1": Decimal(no_load_1),
        "load_loss_1": Decimal(load_1),
        "no_load_loss_2": Decimal(no_load_2),
        "load_loss_2": Decimal(load_2),
        "no_load_loss_3": Decimal(no_load_3),
        "load_loss_3": Decimal(load_3),
    }


def _column_index(column: str) -> int:
    value = 0
    for char in column:
        value = value * 26 + ord(char.upper()) - ord("A") + 1
    return value - 1


class TransformerVoltageImpedanceFlowTests(unittest.TestCase):
    def test_v61_batch_and_application_agree_on_voltage_and_impedance_cases(self):
        # Expected results below are fixed independently from each Application
        # execution. The loss limits are the selected frozen table values in W.
        cases = [
            {
                "name": "T02 630kVA 4% impedance",
                "table_no": 2, "capacity": "630", "no_load": "800", "load": "5000",
                "cells": {"I": "4"},
                "expected_status": "SUCCESS", "expected_grade": "3",
                "expected_conclusion": "3级",
                "expected_thresholds": _thresholds(
                    "775", "4975", "910", "4975", "1070", "5530"),
                "expected_inputs": {"uk_rated_percent": "4"},
            },
            {
                "name": "T02 630kVA 6% impedance",
                "table_no": 2, "capacity": "630", "no_load": "800", "load": "5000",
                "cells": {"I": "6"},
                "expected_status": "SUCCESS", "expected_grade": "2",
                "expected_conclusion": "2级",
                "expected_thresholds": _thresholds(
                    "750", "5050", "885", "5050", "1040", "5610"),
                "expected_inputs": {"uk_rated_percent": "6"},
            },
            {
                "name": "T15 LV66kV special branch",
                "table_no": 15, "capacity": "120000", "no_load": "44000", "load": "302000",
                "cells": {"W": "66"},
                "expected_status": "SUCCESS", "expected_grade": "1",
                "expected_conclusion": "1级",
                "expected_thresholds": _thresholds(
                    "45000", "303000", "53000", "303000", "65000", "320000"),
                "expected_inputs": {"lv_rated_voltage_kv": "66"},
            },
            {
                "name": "T15 LV69kV special branch",
                "table_no": 15, "capacity": "120000", "no_load": "44000", "load": "302000",
                "cells": {"W": "69"},
                "expected_status": "SUCCESS", "expected_grade": "1",
                "expected_conclusion": "1级",
                "expected_thresholds": _thresholds(
                    "45000", "303000", "53000", "303000", "65000", "320000"),
                "expected_inputs": {"lv_rated_voltage_kv": "69"},
            },
            {
                "name": "T15 LV35kV ordinary branch",
                "table_no": 15, "capacity": "120000", "no_load": "44000", "load": "302000",
                "cells": {"W": "35"},
                "expected_status": "SUCCESS", "expected_grade": "2",
                "expected_conclusion": "2级",
                "expected_thresholds": _thresholds(
                    "43000", "304000", "51000", "304000", "63000", "321000"),
                "expected_inputs": {"lv_rated_voltage_kv": "35"},
            },
            {
                "name": "T27 complete A impedance triplet",
                "table_no": 27, "capacity": "120000", "no_load": "28000", "load": "207000",
                "cells": {"I": "12", "X": "36", "Y": "21"},
                "expected_status": "SUCCESS", "expected_grade": "1",
                "expected_conclusion": "1级",
                "expected_thresholds": _thresholds(
                    "28000", "207000", "33000", "207000", "40000", "219000"),
                "expected_inputs": {
                    "uk_hm_rated_percent": "12",
                    "uk_hl_rated_percent": "36",
                    "uk_ml_rated_percent": "21",
                },
            },
            {
                "name": "T27 incomplete impedance triplet",
                "table_no": 27, "capacity": "120000", "no_load": "28000", "load": "207000",
                "cells": {"I": "12", "X": "36"},
                "expected_status": "INSUFFICIENT_DATA", "expected_grade": None,
                "expected_conclusion": "无法判定", "expected_thresholds": {},
                "expected_inputs": {
                    "uk_hm_rated_percent": "12",
                    "uk_hl_rated_percent": "36",
                    "uk_ml_rated_percent": None,
                },
                "expected_missing": ("uk_ml_rated_percent",),
            },
            {
                "name": "T28 complete C impedance triplet",
                "table_no": 28, "capacity": "120000", "no_load": "33000", "load": "225000",
                "cells": {"I": "15.5", "X": "44", "Y": "29"},
                "expected_status": "SUCCESS", "expected_grade": "2",
                "expected_conclusion": "2级",
                "expected_thresholds": _thresholds(
                    "28000", "225000", "33000", "225000", "40000", "238000"),
                "expected_inputs": {
                    "uk_hm_rated_percent": "15.5",
                    "uk_hl_rated_percent": "44",
                    "uk_ml_rated_percent": "29",
                },
            },
            {
                "name": "T28 B/C multi-hit triplet with equal limits",
                "table_no": 28, "capacity": "120000", "no_load": "35000", "load": "230000",
                "cells": {"I": "13", "X": "44", "Y": "29"},
                "expected_status": "SUCCESS", "expected_grade": "3",
                "expected_conclusion": "3级",
                "expected_thresholds": _thresholds(
                    "28000", "225000", "33000", "225000", "40000", "238000"),
                "expected_inputs": {
                    "uk_hm_rated_percent": "13",
                    "uk_hl_rated_percent": "44",
                    "uk_ml_rated_percent": "29",
                },
                "expected_trace_branches": {
                    "GB20052-T28/uk_group=B", "GB20052-T28/uk_group=C"},
            },
            {
                "name": "T29 closed upper voltage boundary",
                "table_no": 29, "capacity": "500", "no_load": "385", "load": "3900",
                "cells": {"V": "7.2", "W": "1.14"},
                "expected_status": "SUCCESS", "expected_grade": "1",
                "expected_conclusion": "1级",
                "expected_thresholds": _thresholds(
                    "385", "3900", "430", "4330", "480", "5410"),
                "expected_inputs": {
                    "hv_rated_voltage_kv": "7.2",
                    "lv_rated_voltage_kv": "1.14",
                },
            },
            {
                "name": "T29 high voltage above range",
                "table_no": 29, "capacity": "500", "no_load": "385", "load": "3900",
                "cells": {"V": "8", "W": "0.4"},
                "expected_status": "OUT_OF_STANDARD_SCOPE", "expected_grade": None,
                "expected_conclusion": "不在范围", "expected_thresholds": {},
                "expected_inputs": {
                    "hv_rated_voltage_kv": "8",
                    "lv_rated_voltage_kv": "0.4",
                },
            },
            {
                "name": "T29 missing low voltage",
                "table_no": 29, "capacity": "500", "no_load": "385", "load": "3900",
                "cells": {"V": "6.5"},
                "expected_status": "INSUFFICIENT_DATA", "expected_grade": None,
                "expected_conclusion": "无法判定", "expected_thresholds": {},
                "expected_inputs": {
                    "hv_rated_voltage_kv": "6.5",
                    "lv_rated_voltage_kv": None,
                },
                "expected_missing": ("lv_rated_voltage_kv",),
            },
        ]

        with tempfile.TemporaryDirectory(prefix="transformer-voltage-impedance-") as directory:
            root = Path(directory)
            paths = AppDataPaths(root / "app-data")
            source = root / "V6.1-cross-entry-input.xlsx"
            destination = root / "V6.1-cross-entry-output.xlsx"
            source.write_bytes(TEMPLATE_PATH.read_bytes())

            workbook = load_workbook(source)
            sheet = workbook["变压器"]
            for offset, case in enumerate(cases):
                row_number = 4 + offset
                category = _category(case["table_no"])
                cells = {
                    "B": case["name"],
                    "C": f"CROSS-{row_number:02d}",
                    "D": 1,
                    "F": category.label,
                    "G": case["capacity"],
                    "J": case["no_load"],
                    "K": case["load"],
                    **case["cells"],
                }
                if case["table_no"] in {2, 29}:
                    cells["H"] = condition_options(category.category_id)[0].label
                for column, value in cells.items():
                    sheet[f"{column}{row_number}"] = value
            workbook.save(source)
            workbook.close()

            # The direct Application request is made from the exact values read
            # from each Excel row, then checked against explicit per-case truth.
            input_workbook = UnifiedWorkbookReader().read(source)
            rows = {row.row_number: row for row in input_workbook.sheets["变压器"].rows}
            self.assertEqual(len(rows), len(cases))
            application = create_transformer_analysis_service(with_persistence=False)
            expected_by_row = {}
            for offset, case in enumerate(cases):
                row_number = 4 + offset
                with self.subTest(case=case["name"]):
                    source_row = rows[row_number]
                    for field, expected_value in case["expected_inputs"].items():
                        actual_value = source_row.values.get(field)
                        if expected_value is None:
                            self.assertIn(actual_value, (None, ""), field)
                        else:
                            self.assertEqual(str(actual_value), expected_value, field)
                    direct = application.evaluate(TransformerAnalysisRequest(
                        dict(source_row.values), AS_OF))
                    self.assertEqual(direct.evaluation_status, case["expected_status"])
                    self.assertEqual(direct.grade, case["expected_grade"])
                    self.assertEqual(direct.ui_conclusion, case["expected_conclusion"])
                    self.assertEqual(direct.thresholds, case["expected_thresholds"])
                    if "expected_missing" in case:
                        self.assertEqual(direct.missing_fields, case["expected_missing"])
                    if "expected_trace_branches" in case:
                        self.assertEqual(
                            {entry["branch_id"] for entry in direct.trace},
                            case["expected_trace_branches"])
                    expected_by_row[row_number] = (case, direct)

            batch = create_batch_evaluation_service(paths=paths)
            batch_result = batch.evaluate_workbook(
                source, destination=destination, as_of=AS_OF, persist=False)
            self.assertEqual(batch_result.summary.data_row_count, len(cases))
            self.assertEqual(batch_result.summary.conclusion_quantities, {
                "1级": 4, "2级": 3, "3级": 2,
                "不在范围": 1, "无法判定": 2,
            })
            batch_outcomes = {
                item["outcome"].row_number: item["outcome"]
                for item in batch_result.outcomes
                if item["sheet_name"] == "变压器"
            }
            self.assertEqual(set(batch_outcomes), set(expected_by_row))

            result_columns = {
                str(field["threshold_key"]): str(field["column"])
                for field in TRANSFORMER_SHEET_CONTRACT["fields"]
                if field.get("role") == "result" and field.get("threshold_key")
            }
            conclusion_column = next(
                str(field["column"])
                for field in TRANSFORMER_SHEET_CONTRACT["fields"]
                if field.get("result_key") == "conclusion"
            )
            message_column = next(
                str(field["column"])
                for field in TRANSFORMER_SHEET_CONTRACT["fields"]
                if field.get("result_key") == "messages"
            )
            with OOXMLWorkbook(destination) as result_workbook:
                output_rows = result_workbook.rows("变压器")
            for row_number, (case, direct) in expected_by_row.items():
                with self.subTest(excel_output=case["name"]):
                    batch_outcome = batch_outcomes[row_number]
                    self.assertEqual(batch_outcome.evaluation_status,
                                     direct.evaluation_status)
                    self.assertEqual(batch_outcome.grade, direct.grade)
                    self.assertEqual(batch_outcome.conclusion,
                                     direct.ui_conclusion)
                    self.assertEqual(batch_outcome.thresholds,
                                     direct.thresholds)
                    row = output_rows[row_number - 1]
                    self.assertEqual(row[_column_index(conclusion_column)],
                                     case["expected_conclusion"])
                    self.assertTrue(row[_column_index(message_column)])
                    for key, column in result_columns.items():
                        expected_value = direct.thresholds.get(key)
                        expected_text = str(expected_value) if expected_value is not None else ""
                        actual_value = row[_column_index(column)]
                        actual_text = ("" if actual_value in (None, "")
                                       else str(actual_value))
                        self.assertEqual(actual_text, expected_text,
                                         f"{case['name']} {key}")


if __name__ == "__main__":
    unittest.main()
