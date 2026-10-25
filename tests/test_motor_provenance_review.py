"""PR 35 regression checks for trace-free motor references and exact display scaling."""
from __future__ import annotations

from decimal import ROUND_CEILING, localcontext
from datetime import date
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from equipeffi.application.ports.batch_workbook import BatchSourceRow
from equipeffi.application.services.motor_analysis_service import MotorAnalysisRequest
from equipeffi.application.services.motor_basis_formatter import (
    _source_power_kw,
    format_motor_basis_lines,
)
from equipeffi.composition import create_batch_evaluation_service, create_motor_analysis_service
from equipeffi.application.services.pump_batch_evaluation_service import BatchEvaluationSummary
from equipeffi.application.services.workbook_evaluation_service import WorkbookEvaluationService
from equipeffi.infrastructure.persistence.app_data_paths import AppDataPaths
from equipeffi.infrastructure.persistence.sqlite_batch_record_repository import (
    SqliteBatchRecordRepository,
)


class MotorProvenanceReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.paths = AppDataPaths(self.root / "data")
        self.batch = create_batch_evaluation_service(
            paths=self.paths, motor_candidate=True)
        self.source = self.root / "input.xlsx"
        self.batch.template_resource.download_to(self.source)
        self.records = SqliteBatchRecordRepository(self.paths.records_db)
        self.motor = create_motor_analysis_service(with_persistence=False)
        self.category_label = next(
            category.label for category in self.motor.input_contract.categories
            if category.category_id == "three_phase_induction")

    def _write_inputs(self, *, motor: dict[str, object] | None = None,
                      pump: dict[str, object] | None = None) -> None:
        workbook = load_workbook(self.source)
        try:
            for sheet_name, values in (("电动机", motor), ("离心泵", pump)):
                if values:
                    for cell, value in values.items():
                        workbook[sheet_name][cell] = value
            workbook.save(self.source)
        finally:
            workbook.close()

    def _expected_motor_result(self):
        workbook = self.batch.reader.read(self.source)
        row = workbook.sheets["电动机"].rows[0]
        return self.motor.evaluate(MotorAnalysisRequest(row.values))

    def _run(self):
        result = self.batch.evaluate_workbook(
            self.source, destination=self.root / "结果.xlsx")
        self.assertEqual(result.batch_record_error, "")
        record = self.records.load_batch_record(result.batch_record_id)
        return result, record

    @staticmethod
    def _motor_sheet(result):
        return next(sheet for sheet in result.summary.sheets
                    if sheet["sheet_name"] == "电动机")

    @staticmethod
    def _motor_outcome(result, row_number=4):
        return next(
            item["outcome"] for item in result.outcomes
            if item["sheet_name"] == "电动机"
            and item["outcome"].row_number == row_number)

    def _assert_single_motor_provenance(self, result, record, expected):
        motor_sheet = self._motor_sheet(result)
        self.assertEqual(motor_sheet["evaluation_provenance"], [expected.references])
        self.assertEqual(record.summary["evaluation_provenance"], expected.references)
        self.assertEqual(record.canonical_version, expected.references["canonical_version"])
        self.assertEqual(record.numeric_profile_id,
                         expected.references["numeric_profile_id"])

    def test_trace_free_out_of_scope_result_persists_provenance(self):
        self._write_inputs(motor={"D4": 2, "F4": self.category_label,
                                  "H4": "5~120"})
        expected = self._expected_motor_result()
        self.assertEqual(expected.evaluation_status, "OUT_OF_STANDARD_SCOPE")
        self.assertFalse(expected.trace)

        result, record = self._run()
        outcome = self._motor_outcome(result)
        self.assertEqual(outcome.evaluation_status, "OUT_OF_STANDARD_SCOPE")
        self.assertFalse(outcome.is_input_error)
        self.assertFalse(outcome.is_execution_error)
        self.assertEqual(result.summary.execution_error_rows, 0)
        self.assertEqual(result.summary.input_error_rows, 0)
        self.assertEqual(len(self.records.list_batch_records()), 1)
        self._assert_single_motor_provenance(result, record, expected)

    def test_trace_free_uncertain_category_persists_provenance(self):
        self._write_inputs(motor={"D4": 2, "F4": "不确定类别"})
        expected = self._expected_motor_result()
        self.assertEqual(expected.evaluation_status, "INSUFFICIENT_DATA")
        self.assertFalse(expected.trace)

        result, record = self._run()
        outcome = self._motor_outcome(result)
        self.assertEqual(outcome.evaluation_status, "INSUFFICIENT_DATA")
        self.assertEqual(outcome.conclusion, "无法判定")
        self.assertFalse(outcome.is_input_error)
        self.assertFalse(outcome.is_execution_error)
        self.assertEqual(result.summary.execution_error_rows, 0)
        self.assertEqual(result.summary.input_error_rows, 0)
        self._assert_single_motor_provenance(result, record, expected)

    def test_only_name_and_quantity_with_missing_category_remains_legal(self):
        # Exact regression shape requested in review: B has a name, D a quantity,
        # while F and every other motor input stay blank.
        self._write_inputs(motor={"B4": "仅有名称和数量", "D4": 2})
        expected = self._expected_motor_result()
        self.assertEqual(expected.evaluation_status, "INSUFFICIENT_DATA")
        self.assertFalse(expected.trace)

        result, record = self._run()
        outcome = self._motor_outcome(result)
        self.assertEqual(outcome.evaluation_status, "INSUFFICIENT_DATA")
        self.assertEqual(outcome.conclusion, "无法判定")
        self.assertFalse(outcome.is_input_error)
        self.assertFalse(outcome.is_execution_error)
        self.assertEqual(result.summary.execution_error_rows, 0)
        self.assertEqual(result.summary.input_error_rows, 0)
        self.assertEqual(result.summary.evaluated_quantity, 2)
        self._assert_single_motor_provenance(result, record, expected)

        workbook = load_workbook(result.result_workbook, data_only=True)
        try:
            self.assertEqual(workbook["电动机"]["Q4"].value, "无法判定")
        finally:
            workbook.close()

    def test_mixed_batch_keeps_motor_provenance_in_sheet_and_one_batch_record(self):
        self._write_inputs(
            motor={"D4": 2, "F4": "不确定类别"},
            pump={"D4": 3, "F4": "其他类别"},
        )
        expected = self._expected_motor_result()
        self.assertEqual(expected.evaluation_status, "INSUFFICIENT_DATA")
        self.assertFalse(expected.trace)

        result, record = self._run()
        self.assertEqual(record.standard_code, "MULTI")
        self.assertEqual(record.numeric_profile_id, "MULTI")
        self.assertEqual(record.canonical_version, "MULTI")
        self.assertEqual(len(self.records.list_batch_records()), 1)
        recorded_motor = next(
            sheet for sheet in record.summary["sheets"]
            if sheet["sheet_name"] == "电动机")
        self.assertEqual(recorded_motor["evaluation_provenance"],
                         [expected.references])
        outcome = self._motor_outcome(result)
        self.assertEqual(outcome.evaluation_status, "INSUFFICIENT_DATA")
        self.assertFalse(outcome.is_input_error)
        self.assertFalse(outcome.is_execution_error)


class MotorBasisFormatterReviewTests(unittest.TestCase):
    def test_real_300kw_six_pole_trace_is_independent_of_decimal_context(self):
        service = create_motor_analysis_service(with_persistence=False)
        result = service.evaluate(MotorAnalysisRequest({
            "category_id": "three_phase_induction",
            "rated_voltage_kv": "0.38",
            "rated_frequency_hz": "50",
            "rated_output_power_kw": "300",
            "poles": "6",
            "measured_efficiency_pct": "92",
        }))
        self.assertEqual(result.evaluation_status, "SUCCESS")
        self.assertTrue(result.trace)

        normal = format_motor_basis_lines(result.trace)
        with localcontext() as context:
            context.prec = 2
            context.rounding = ROUND_CEILING
            constrained = format_motor_basis_lines(result.trace)
            watt_anchor = _source_power_kw({
                "source_power_text": "315", "source_unit": "W"})

        self.assertEqual(constrained, normal)
        self.assertIn("250～315 kW", constrained[0])
        self.assertEqual(watt_anchor, "0.315")


class MotorRejectedResultProvenanceTests(unittest.TestCase):
    def test_invalid_input_and_execution_failure_do_not_add_provenance(self):
        service = object.__new__(WorkbookEvaluationService)
        row = BatchSourceRow(4, {"quantity": "2"})

        class InvalidMotor:
            def evaluate(self, _request):
                return type("InvalidResult", (), {
                    "evaluation_status": "INVALID_INPUT",
                    "explanation": "效率格式无效",
                    "ui_conclusion": "输入错误",
                    "grade": None,
                    "thresholds": {},
                    "trace": (),
                    "references": {"standard_code": "GB 18613-2020"},
                    "finalizable": False,
                })()

        invalid_summary = BatchEvaluationSummary()
        invalid_references = []
        invalid_outcome = service._device_row(
            row, date.today(), invalid_summary, invalid_references,
            InvalidMotor(), MotorAnalysisRequest, lambda _trace: ())
        self.assertTrue(invalid_outcome.is_input_error)
        self.assertEqual(invalid_summary.input_error_rows, 1)
        self.assertEqual(invalid_references, [])

        class FailingMotor:
            def evaluate(self, _request):
                raise RuntimeError("fixture failure")

        failed_summary = BatchEvaluationSummary()
        failed_references = []
        with patch("equipeffi.application.services.workbook_evaluation_service.logging.getLogger"):
            failed_outcome = service._device_row(
                row, date.today(), failed_summary, failed_references,
                FailingMotor(), MotorAnalysisRequest, lambda _trace: ())
        self.assertTrue(failed_outcome.is_execution_error)
        self.assertEqual(failed_summary.execution_error_rows, 1)
        self.assertEqual(failed_references, [])


if __name__ == "__main__":
    unittest.main()
