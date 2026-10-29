"""Release defaults, real mixed-device flow and immutable older snapshots."""
from copy import deepcopy
from dataclasses import replace
from datetime import date
from hashlib import sha256
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import os

from openpyxl import load_workbook
from equipeffi.composition import (create_batch_evaluation_service,
    create_motor_analysis_service, create_pump_analysis_service,
    create_transformer_analysis_service)
from equipeffi.application.services.centrifugal_pump_analysis_service import PumpAnalysisRequest
from equipeffi.application.services.motor_analysis_service import MotorAnalysisRequest
from equipeffi.application.services.transformer_analysis_service import TransformerAnalysisRequest
from equipeffi.infrastructure.excel.template_resource import V6TemplateResource, V61TransformerTemplateResource
from equipeffi.infrastructure.persistence.app_data_paths import AppDataPaths


class ThreeStandardsReleaseTests(unittest.TestCase):
    def setUp(self):
        environment = patch.dict(os.environ)
        environment.start()
        self.addCleanup(environment.stop)
        os.environ.pop('EQUIPEFFI_ENABLE_MOTOR_CANDIDATE', None)

    def test_default_desktop_assembles_all_three_services_and_v63(self):
        from equipeffi.composition import launch_qt
        with tempfile.TemporaryDirectory() as directory, patch(
                'equipeffi.presentation.qt.app.run', return_value=0) as run:
            self.assertEqual(launch_qt(paths=AppDataPaths(Path(directory))), 0)
            self.assertEqual(run.call_count, 1)
            self.assertIsNotNone(run.call_args.args[2])
            services = run.call_args.kwargs
            self.assertEqual(services['transformer'].standard_overview()['support_status'], 'SUPPORTED')
            self.assertEqual(services['motor'].standard_overview()['support_status'], 'SUPPORTED')
            self.assertEqual(services['batch'].template_identity['template_version'], 'V6.3')
            self.assertIsNotNone(services['batch'].motor)

    def test_formal_singles_mixed_batch_reimport_and_old_snapshots(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            paths = AppDataPaths(root / 'data')
            pump = create_pump_analysis_service(paths=paths)
            transformer = create_transformer_analysis_service(paths=paths)
            motor = create_motor_analysis_service(paths=paths)
            batch = create_batch_evaluation_service(paths=paths)
            self.assertEqual(batch.template_identity['template_version'], 'V6.3')
            self.assertIsNotNone(batch.motor)
            singles = [pump.analyze_and_record(PumpAnalysisRequest(
                '单级单吸清水离心泵', date(2026, 10, 27), QBEP='100', HBEP='50',
                speed='2900', efficiency='90', suction='单吸', stages='1')),
                transformer.analyze_and_record(TransformerAnalysisRequest(dict(
                    category_id='GB20052-T03', capacity_kva='3150', no_load_loss_w='1700', load_loss_w='20700'))),
                motor.analyze_and_record(MotorAnalysisRequest(dict(category_id='three_phase_induction',
                    rated_voltage_v='380', rated_frequency_hz='50', rated_output_power_kw='0.14',
                    poles='4', rated_efficiency_pct='71.43')))]
            for outcome in singles:
                self.assertIsNotNone(outcome.record)
                self.assertEqual(outcome.result.evaluation_status, 'SUCCESS')
                self.assertEqual(outcome.result.support_status, 'SUPPORTED')

            # Explicit pre-release snapshots stay IMPLEMENTED after a new service opens them.
            older = []
            for outcome, service in zip(singles[1:], (transformer, motor)):
                snapshot = deepcopy(outcome.record.result_snapshot)
                snapshot['support_status'] = 'IMPLEMENTED'
                record = replace(outcome.record, record_id='legacy-' + outcome.record.record_id,
                                 result_snapshot=snapshot)
                service._records.append_record(record)
                older.append(record)
            before = deepcopy(pump.list_records())
            database_before = sha256(paths.records_db.read_bytes()).hexdigest()
            reopened = create_pump_analysis_service(paths=paths)
            for record in older:
                self.assertEqual(reopened.open_record(record.record_id), record)
                self.assertEqual(reopened.open_record(record.record_id).result_snapshot['support_status'], 'IMPLEMENTED')
            self.assertEqual(sha256(paths.records_db.read_bytes()).hexdigest(), database_before)

            source, output = root / 'input.xlsx', root / 'output.xlsx'
            batch.export_blank_template(source)
            book = load_workbook(source)
            transformer_label = next(c.label for c in transformer.input_contract.categories if c.category_id == 'GB20052-T03')
            motor_label = next(c.label for c in motor.input_contract.categories if c.category_id == 'three_phase_induction')
            for name, values in (
                ('离心泵', {'D4':2, 'F4':'单级单吸清水离心泵', 'G4':100, 'H4':50, 'I4':2900, 'K4':'单吸', 'L4':1, 'M4':90}),
                ('变压器', {'D4':3, 'F4':transformer_label, 'G4':3150, 'J4':1700, 'K4':20700}),
                ('电动机', {'D4':4, 'F4':motor_label, 'G4':380, 'H4':50, 'J4':0.14, 'K4':4, 'M4':71.43})):
                for cell, value in values.items():
                    book[name][cell] = value
            book.save(source)
            book.close()
            input_hash = sha256(source.read_bytes()).hexdigest()
            result = batch.evaluate_workbook(source, destination=output)
            self.assertTrue(result.batch_record_saved)
            self.assertEqual(len(batch.list_batch_records()), 1)
            self.assertEqual(result.summary.concluded_rows, 3)
            self.assertEqual(result.summary.total_quantity, 9)
            self.assertEqual(result.summary.conclusion_quantities, {'1级':5, '3级':4})
            self.assertEqual(sha256(source.read_bytes()).hexdigest(), input_hash)
            self.assertTrue(output.is_file())
            self.assertEqual(pump.list_records(), before)
            repeated = batch.evaluate_workbook(output)
            self.assertEqual(repeated.summary.conclusion_quantities, result.summary.conclusion_quantities)
            self.assertEqual(len(batch.list_batch_records()), 2)
            self.assertEqual(pump.list_records(), before)

    def test_default_service_evaluates_legacy_v6_and_v61_without_changing_input(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            batch = create_batch_evaluation_service(paths=AppDataPaths(root / 'data'))
            for index, resource in enumerate((V6TemplateResource(), V61TransformerTemplateResource())):
                with self.subTest(version=resource.identity['template_version']):
                    source = root / f'legacy-{index}.xlsx'
                    resource.download_to(source)
                    book = load_workbook(source)
                    book['离心泵']['D4'] = 2
                    book['离心泵']['F4'] = '其他类别'
                    book.save(source)
                    book.close()
                    original = source.read_bytes()
                    result = batch.evaluate_workbook(source)
                    self.assertTrue(result.batch_record_saved)
                    self.assertEqual(result.summary.conclusion_quantities, {'不在范围':2})
                    self.assertEqual(source.read_bytes(), original)
