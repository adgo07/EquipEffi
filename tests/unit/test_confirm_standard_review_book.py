from __future__ import annotations

from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from tools.confirm_standard_review_book import confirm
from tools.build_human_review_book import FLAT_HEADERS


class ConfirmStandardReviewBookTests(unittest.TestCase):
    def _make_book(self, path: Path, statuses: list[str], suggestions: list[object] | None = None) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "标准数据平铺"
        sheet.append([None] * len(FLAT_HEADERS))
        sheet.append([None] * len(FLAT_HEADERS))
        sheet.append(list(FLAT_HEADERS))
        suggestions = suggestions or [None] * len(statuses)
        for index, status in enumerate(statuses, start=1):
            values = [None] * len(FLAT_HEADERS)
            values[0] = f"ID-{index}"
            values[1] = "TEST"
            values[16] = status
            values[17] = suggestions[index - 1]
            sheet.append(values)
        workbook.save(path)
        workbook.close()

    def test_confirm_normalizes_legacy_review_status(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            target = Path(directory) / "target.xlsx"
            self._make_book(source, ["已校对", "未校对", "正确 "])
            result = confirm(source, target, reviewer="测试", review_date="2026-08-30")
            self.assertEqual(result["data_rows"], 3)
            self.assertEqual(result["converted_rows"], 3)
            workbook = load_workbook(target, read_only=True, data_only=False)
            statuses = [workbook["标准数据平铺"].cell(row, 17).value for row in range(4, 7)]
            workbook.close()
            self.assertEqual(statuses, ["正确", "正确", "正确"])

    def test_confirm_rejects_unresolved_or_suggested_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            target = Path(directory) / "target.xlsx"
            self._make_book(source, ["需修改"])
            with self.assertRaisesRegex(ValueError, "需修改"):
                confirm(source, target, review_date="2026-08-30")
            self._make_book(source, ["正确"], suggestions=[99])
            with self.assertRaisesRegex(ValueError, "建议修订值"):
                confirm(source, target, review_date="2026-08-30")


if __name__ == "__main__":
    unittest.main()
