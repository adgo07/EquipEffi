from __future__ import annotations

import hashlib
import json
import os
from dataclasses import replace
from copy import deepcopy
from datetime import date
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from equipeffi.application.lifecycle import (
    WorkspaceSnapshot,
    business_keys_metadata,
)
from equipeffi.application.services.centrifugal_pump_analysis_service import (
    GB19762_STANDARD_CODE,
    PUMP_FINGERPRINT_KEYS,
    PUMP_STANDARD_IDENTITY_KEY,
    PumpAnalysisRequest,
    PumpAnalysisResult,
    CentrifugalPumpAnalysisService,
    _pump_payload,
    user_conclusion_from_snapshot,
    user_conclusion_text,
    user_grade_text,
)
from equipeffi.application.services.pump_calculator_version import (
    PUMP_CALCULATOR_SOURCE_FILES,
    PUMP_CALCULATOR_VERSION,
)
from equipeffi.infrastructure.standards.json_repository import (
    JsonStandardRepository,
    StandardPackError,
)

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


ROOT = Path(__file__).resolve().parents[2]


class _NoPackAccess:
    def get_pack(self, *_args, **_kwargs):
        raise AssertionError("标准身份不匹配时不应读取规则包")


class StandardIdentityTests(unittest.TestCase):
    def test_standard_identity_is_fingerprinted_and_rejected_before_rules(self):
        service = CentrifugalPumpAnalysisService(_NoPackAccess())
        correct = PumpAnalysisRequest(
            "单级单吸清水离心泵", date(2026, 10, 19),
            QBEP="100", HBEP="50", speed="2900", efficiency="80",
            suction="单吸", stages="1")
        wrong = PumpAnalysisRequest(
            correct.product_category, correct.as_of, QBEP=correct.QBEP,
            HBEP=correct.HBEP, speed=correct.speed, efficiency=correct.efficiency,
            suction=correct.suction, stages=correct.stages,
            standard_code="GB 20052-2024")

        self.assertEqual(
            correct.request_fingerprint(include_standard_identity=False),
            wrong.request_fingerprint(include_standard_identity=False))
        self.assertNotEqual(correct.request_fingerprint(), wrong.request_fingerprint())
        result = service.evaluate(wrong)
        self.assertEqual(result.evaluation_status, "INVALID_INPUT")
        self.assertEqual(result.issue_codes, ("STANDARD_IDENTITY_MISMATCH",))
        self.assertEqual(result.standard_code, GB19762_STANDARD_CODE)
        self.assertFalse(result.finalizable)

    def test_new_workspace_persists_standard_identity_in_its_fingerprint(self):
        request = PumpAnalysisRequest(
            "单级单吸清水离心泵", date(2026, 10, 19), standard_code="GB 20052-2024")
        payload = _pump_payload(request)
        self.assertEqual(payload[PUMP_STANDARD_IDENTITY_KEY], "GB 20052-2024")
        self.assertIn(PUMP_STANDARD_IDENTITY_KEY, payload["_business_keys"])
        workspace = WorkspaceSnapshot(
            workspace_id="W-identity",
            standard_code=request.standard_code,
            device_type="centrifugal_pump",
            product_category=request.product_category,
            rule_profile="pump_water",
            as_of=request.as_of.isoformat(),
            payload=payload,
            schema_version=1,
            created_at_utc="2026-10-19T00:00:00+00:00",
            updated_at_utc="2026-10-19T00:00:00+00:00",
        )
        self.assertEqual(workspace.request_fingerprint(), request.request_fingerprint())

    def test_legacy_workspace_fingerprint_remains_reconstructable(self):
        request = PumpAnalysisRequest(
            "单级单吸清水离心泵", date(2026, 10, 19), QBEP="100")
        payload = request.raw_values() | business_keys_metadata(PUMP_FINGERPRINT_KEYS)
        workspace = WorkspaceSnapshot(
            workspace_id="W-legacy",
            standard_code=request.standard_code,
            device_type="centrifugal_pump",
            product_category=request.product_category,
            rule_profile="pump_water",
            as_of=request.as_of.isoformat(),
            payload=payload,
            schema_version=1,
            created_at_utc="2026-10-19T00:00:00+00:00",
            updated_at_utc="2026-10-19T00:00:00+00:00",
        )
        self.assertEqual(
            workspace.request_fingerprint(PUMP_FINGERPRINT_KEYS),
            request.request_fingerprint(include_standard_identity=False),
        )

    def test_finalize_accepts_unchanged_phase4_workspace_shape(self):
        class LegacyWorkspaceRepository:
            def __init__(self, snapshot):
                self.snapshot = snapshot

            def load_workspace(self, _workspace_id):
                return self.snapshot

        class RecordCapture:
            def append_record(self, snapshot):
                self.snapshot = snapshot

        request = PumpAnalysisRequest(
            "单级单吸清水离心泵", date(2026, 10, 19), QBEP="100",
            HBEP="50", speed="2900", efficiency="90", suction="单吸",
            stages="1", workspace_revision=1)
        legacy_workspace = WorkspaceSnapshot(
            workspace_id="W-legacy-finalize",
            standard_code=request.standard_code,
            device_type="centrifugal_pump",
            product_category=request.product_category,
            rule_profile="pump_water",
            as_of=request.as_of.isoformat(),
            payload=request.raw_values() | business_keys_metadata(PUMP_FINGERPRINT_KEYS),
            schema_version=1,
            created_at_utc="2026-10-19T00:00:00+00:00",
            updated_at_utc="2026-10-19T00:00:00+00:00",
            revision=1,
        )
        records = RecordCapture()
        service = CentrifugalPumpAnalysisService(
            JsonStandardRepository(ROOT / "src" / "equipeffi"),
            LegacyWorkspaceRepository(legacy_workspace), records)
        result = service.evaluate(request)
        self.assertTrue(result.finalizable)
        record = service.finalize(
            record_id="R-legacy-workspace",
            workspace_id=legacy_workspace.workspace_id,
            request=request,
            result=result,
        )
        self.assertEqual(record.record_id, "R-legacy-workspace")
        self.assertEqual(records.snapshot.standard_code, request.standard_code)

    def test_standard_pack_failure_propagates_as_system_error(self):
        class FailedPack:
            def get_pack(self, *_args, **_kwargs):
                raise StandardPackError("pack unreadable")

        service = CentrifugalPumpAnalysisService(FailedPack())
        request = PumpAnalysisRequest(
            "单级单吸清水离心泵", date(2026, 10, 19), QBEP="100",
            HBEP="50", speed="2900", efficiency="80", suction="单吸", stages="1")
        with self.assertRaisesRegex(StandardPackError, "pack unreadable"):
            service.evaluate(request)


class StandardPackFailureBoundaryTests(unittest.TestCase):
    def _repository(self, root: Path, *, status: str = "active") -> JsonStandardRepository:
        manifest = root / "manifest.json"
        manifest.write_text(json.dumps({
            "packs": [{
                "device_type": "pump_water",
                "pack_id": "gb19762_test_v1",
                "source": "pump.json",
                "status": status,
            }],
        }), encoding="utf-8")
        return JsonStandardRepository(root, manifest)

    def test_missing_active_pack_is_a_system_pack_error(self):
        with tempfile.TemporaryDirectory() as temp:
            repository = self._repository(Path(temp))
            with self.assertRaisesRegex(StandardPackError, "正式启用.*不存在"):
                repository.get_pack("pump_water")

    def test_malformed_active_pack_is_not_returned_as_business_input_error(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            self._repository(root)
            (root / "pump.json").write_text("{", encoding="utf-8")
            repository = JsonStandardRepository(root, root / "manifest.json")
            with self.assertRaisesRegex(StandardPackError, "无法读取或解析标准包"):
                repository.get_pack("pump_water")

    def test_active_pack_read_error_is_wrapped_at_the_repository_boundary(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            self._repository(root)
            source = root / "pump.json"
            source.write_text("{}", encoding="utf-8")
            repository = JsonStandardRepository(root, root / "manifest.json")
            real_read_text = Path.read_text
            read_failure_injected = False
            resolved_source = source.resolve()

            def fail_pack_read(path, *args, **kwargs):
                nonlocal read_failure_injected
                if path.resolve() == resolved_source:
                    read_failure_injected = True
                    raise PermissionError("test resource access failure")
                return real_read_text(path, *args, **kwargs)

            with patch.object(Path, "read_text", fail_pack_read):
                with self.assertRaisesRegex(StandardPackError, "无法读取或解析标准包"):
                    repository.get_pack("pump_water")
            self.assertTrue(read_failure_injected, "测试必须实际注入标准包读取失败")

    @staticmethod
    def _active_pump_repository(root: Path):
        source = root / "pump.json"
        manifest = root / "manifest.json"
        source.write_text(
            (ROOT / "src" / "equipeffi" / "resources" / "standards" / "pump.json")
            .read_text(encoding="utf-8"),
            encoding="utf-8",
        )
        manifest.write_text(json.dumps({
            "packs": [
                {"device_type": "pump_water", "pack_id": "water-test-v1",
                 "source": "pump.json", "status": "active"},
                {"device_type": "pump_chemical", "pack_id": "chemical-test-v1",
                 "source": "pump.json", "status": "active"},
            ],
        }), encoding="utf-8")
        return JsonStandardRepository(root, manifest), source

    def test_active_water_profile_rejects_an_empty_required_ci_table(self):
        with tempfile.TemporaryDirectory() as temp:
            repository, source = self._active_pump_repository(Path(temp))
            pack = json.loads(source.read_text(encoding="utf-8"))
            pack["water"]["ci"] = []
            source.write_text(json.dumps(pack), encoding="utf-8")

            with self.assertRaisesRegex(StandardPackError, r"pump_water.*water\.ci"):
                repository.get_pack("pump_water")

    def test_active_pump_profiles_accept_the_canonical_required_structures(self):
        with tempfile.TemporaryDirectory() as temp:
            repository, _source = self._active_pump_repository(Path(temp))

            water = repository.get_pack("pump_water")
            chemical = repository.get_pack("pump_chemical")

            self.assertTrue(water["water"]["ci"])
            self.assertTrue(chemical["chemical"]["level_offsets"])

    def test_active_chemical_profile_rejects_a_missing_required_table(self):
        with tempfile.TemporaryDirectory() as temp:
            repository, source = self._active_pump_repository(Path(temp))
            pack = json.loads(source.read_text(encoding="utf-8"))
            pack["chemical"]["level_offsets"] = []
            source.write_text(json.dumps(pack), encoding="utf-8")

            with self.assertRaisesRegex(StandardPackError, r"pump_chemical.*chemical\.level_offsets"):
                repository.get_pack("pump_chemical")

    def test_malformed_chemical_pump_kind_is_wrapped_as_pack_error(self):
        for bad_kind in ([], {}):
            with self.subTest(bad_kind=bad_kind), tempfile.TemporaryDirectory() as temp:
                repository, source = self._active_pump_repository(Path(temp))
                pack = json.loads(source.read_text(encoding="utf-8"))
                pack["chemical"]["level_offsets"][0]["pump"] = bad_kind
                source.write_text(json.dumps(pack), encoding="utf-8")

                with self.assertRaisesRegex(StandardPackError, "pump_chemical.*结构校验失败"):
                    repository.get_pack("pump_chemical")

    def test_shared_pack_source_is_validated_for_each_active_pump_profile(self):
        with tempfile.TemporaryDirectory() as temp:
            repository, source = self._active_pump_repository(Path(temp))
            repository.get_pack("pump_water")
            pack = json.loads(source.read_text(encoding="utf-8"))
            pack["chemical"]["level_offsets"] = []
            source.write_text(json.dumps(pack), encoding="utf-8")

            with self.assertRaisesRegex(StandardPackError, r"pump_chemical.*chemical\.level_offsets"):
                repository.get_pack("pump_chemical")

    def test_missing_inactive_pack_keeps_existing_unavailable_shell(self):
        with tempfile.TemporaryDirectory() as temp:
            repository = self._repository(Path(temp), status="extracted")
            pack = repository.get_pack("pump_water")
            self.assertEqual(pack["status"], "extracted")
            self.assertIn("unavailable_reason", pack)


class UserDisplayProjectionTests(unittest.TestCase):
    def test_live_and_snapshot_conclusions_use_the_same_short_labels(self):
        cases = (
            ({"evaluation_status": "SUCCESS", "grade": "1", "ui_conclusion": "1级"}, "1级"),
            ({"evaluation_status": "SUCCESS", "grade": "BELOW_MINIMUM",
              "ui_conclusion": "未达标"}, "未达3级"),
            ({"evaluation_status": "OUT_OF_STANDARD_SCOPE", "ui_conclusion": "不适用"},
             "不在范围"),
            ({"evaluation_status": "INSUFFICIENT_DATA", "ui_conclusion": "无法判定"},
             "无法判定"),
            ({"evaluation_status": "INVALID_INPUT", "grade": "2",
              "ui_conclusion": "2级"}, "输入错误"),
            ({"evaluation_status": "EXECUTION_ERROR", "grade": "1",
              "ui_conclusion": "1级"}, "执行失败"),
            ({"evaluation_status": None, "requires_category_confirmation": True,
              "ui_conclusion": "无法评价"}, "无法判定"),
        )
        for snapshot, expected in cases:
            with self.subTest(snapshot=snapshot):
                self.assertEqual(user_conclusion_text(snapshot), expected)
                self.assertEqual(user_conclusion_from_snapshot(snapshot), expected)

                live_result = PumpAnalysisResult(
                    rule_profile=None,
                    standard_code=GB19762_STANDARD_CODE,
                    product_category="单级单吸清水离心泵",
                    as_of=date(2026, 10, 19),
                    evaluation_status=snapshot.get("evaluation_status"),
                    category_status=None,
                    support_status=None,
                    ui_conclusion=snapshot.get("ui_conclusion", ""),
                    grade=snapshot.get("grade"),
                    requires_category_confirmation=bool(
                        snapshot.get("requires_category_confirmation")),
                )
                self.assertEqual(user_conclusion_text(live_result), expected)
                self.assertEqual(
                    user_conclusion_from_snapshot(live_result.as_snapshot()), expected)

    def test_internal_grade_codes_never_leak_to_user_text(self):
        self.assertEqual(user_grade_text("1"), "1级")
        self.assertEqual(user_grade_text("BELOW_MINIMUM"), "未达3级")
        self.assertEqual(user_grade_text("PRIVATE_INTERNAL_CODE"), "")


class PumpCalculatorVersionTests(unittest.TestCase):
    def test_source_fingerprint_declares_grading_numeric_and_range_dependencies(self):
        required_sources = {
            "src/equipeffi/application/result_projection.py",
            "src/equipeffi/application/services/pump_release_gate.py",
            "src/equipeffi/domain/evaluation/device_evaluators.py",
            "src/equipeffi/domain/evaluation/evaluators/shared.py",
            "src/equipeffi/domain/evaluation/decimal_math.py",
            "src/equipeffi/domain/evaluation/grading.py",
            "src/equipeffi/domain/common/enums.py",
            "src/equipeffi/domain/common/models.py",
        }
        self.assertTrue(required_sources.issubset(set(PUMP_CALCULATOR_SOURCE_FILES)))

    def test_calculator_version_matches_the_reviewed_source_fingerprint(self):
        digest = hashlib.sha256()
        for relative_path in PUMP_CALCULATOR_SOURCE_FILES:
            digest.update(relative_path.encode("utf-8"))
            digest.update(b"\0")
            digest.update((ROOT / relative_path).read_bytes().replace(b"\r\n", b"\n"))
            digest.update(b"\0")
        self.assertEqual(PUMP_CALCULATOR_VERSION,
                         f"pump-calculator:sha256:{digest.hexdigest()}")


class HomePageHistoricalProjectionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from PySide6.QtWidgets import QApplication

        cls.app = QApplication.instance() or QApplication([])

    def test_recent_history_uses_shared_snapshot_projection_without_mutating_records(self):
        from equipeffi.presentation.qt.pages.home import HomePage
        from PySide6.QtCore import Qt
        from equipeffi.composition import create_pump_analysis_service
        from equipeffi.infrastructure.persistence.app_data_paths import AppDataPaths

        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        actual_service = create_pump_analysis_service(paths=AppDataPaths(Path(temp.name)))
        request = PumpAnalysisRequest("单级单吸清水离心泵", date(2026, 10, 20),
            QBEP="100", HBEP="50", speed="2900", efficiency="80", suction="单吸", stages="1")
        base = actual_service.analyze_and_record(request).record
        self.assertIsNotNone(base)
        # Full real RecordSnapshot contract, including standard and frozen input/
        # result/reference evidence; retain legacy historical conclusion labels.
        records = tuple(replace(base, record_id=rid, grade=grade, evaluation_status=status,
            ui_conclusion=legacy, result_snapshot=dict(base.result_snapshot, grade=grade,
                evaluation_status=status, ui_conclusion=legacy))
            for rid, grade, status, legacy in (
                ("R1", "1", "SUCCESS", "1级"), ("R2", "2", "SUCCESS", "2级"),
                ("R3", "3", "SUCCESS", "3级"), ("R4", "BELOW_MINIMUM", "SUCCESS", "未达标"),
                ("R5", None, "OUT_OF_STANDARD_SCOPE", "不适用")))
        before = deepcopy(records)
        stored_before = deepcopy(actual_service.list_records())

        class HistoryService:
            def standard_overview(self):
                return {"standard_code": "GB 19762-2025", "standard_name": "离心泵能效", "effective_date": "2026-03-01"}

            def list_records(self, *, limit):
                self.limit = limit
                return records

        service = HistoryService()
        page = HomePage(service)
        self.addCleanup(page.deleteLater)
        page.refresh()

        expected = ("1级", "2级", "3级", "未达3级", "不在范围")
        rendered = tuple(page.records.item(index, 2).text() for index in range(page.records.rowCount()))
        self.assertEqual(service.limit, 5)
        self.assertEqual(page.records.rowCount(), 5)
        self.assertEqual(page._record_ids, [record.record_id for record in records])
        for record, label in zip(records, expected):
            with self.subTest(record_id=record.record_id):
                row = records.index(record)
                self.assertEqual(page.records.item(row, 0).data(Qt.ItemDataRole.UserRole), record.record_id)
                self.assertEqual(page.records.item(row, 1).text(), record.standard_code)
                self.assertEqual(page.records.item(row, 2).toolTip(), label)
                self.assertEqual(rendered[row], label)
        self.assertEqual(records, before, "首页展示不得改写历史记录")
        self.assertEqual(actual_service.list_records(), stored_before, "首页展示不得写入原数据库")


class QtBatchDisplayWiringTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from PySide6.QtWidgets import QApplication

        cls.app = QApplication.instance() or QApplication([])

    def _service(self):
        from equipeffi.composition import create_pump_analysis_service
        from equipeffi.infrastructure.persistence.app_data_paths import AppDataPaths

        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        return create_pump_analysis_service(paths=AppDataPaths(Path(temp.name)))

    @staticmethod
    def _write_input(path: Path, *, flow: int, efficiency: int = 80) -> None:
        import openpyxl

        from equipeffi.infrastructure.excel.pump_workbook_reader import PUMP_SHEET
        from equipeffi.infrastructure.excel.template_resource import V6TemplateResource

        V6TemplateResource().download_to(path)
        workbook = openpyxl.load_workbook(path)
        sheet = workbook[PUMP_SHEET]
        row = 4
        for column, value in {
            "B": "测试泵", "C": "C-001", "D": 1, "E": "1号车间",
            "F": "单级单吸清水离心泵", "G": flow, "H": 50, "I": 2900,
            "J": 45, "K": "单吸", "L": 1, "M": efficiency,
        }.items():
            sheet[f"{column}{row}"] = value
        workbook.save(path)
        workbook.close()

    @staticmethod
    def _batch(service):
        from equipeffi.application.services.pump_batch_evaluation_service import (
            PumpBatchEvaluationService,
        )
        from equipeffi.infrastructure.excel.pump_result_writer import (
            PumpResultWorkbookWriter,
        )
        from equipeffi.infrastructure.excel.pump_workbook_reader import (
            V6PumpWorkbookReader,
        )

        return PumpBatchEvaluationService(
            service,
            reader=V6PumpWorkbookReader(),
            writer=PumpResultWorkbookWriter(),
        )

    @staticmethod
    def _page(service):
        from equipeffi.presentation.qt.pages.analysis import AnalysisPage

        page = AnalysisPage(service, as_of=date(2026, 10, 19))
        page.category.setCurrentIndex(
            page.category.findData("单级单吸清水离心泵"))
        for key, value in (
            ("QBEP", "100"), ("HBEP", "50"),
            ("speed", "2900"), ("efficiency", "80"),
        ):
            page.point_inputs[key].setText(value)
        return page

    def test_invalid_input_stays_separate_in_live_qt_and_batch_results(self):
        import openpyxl
        from PySide6.QtTest import QTest

        from equipeffi.infrastructure.excel.pump_workbook_reader import PUMP_SHEET

        service = self._service()
        page = self._page(service)
        self.addCleanup(page.deleteLater)
        request = page._collect_request()
        page.evaluate()
        self.assertIsNotNone(page.last_saved_record_id)
        before = deepcopy(service.list_records())
        edit = page.point_inputs["QBEP"]
        edit.clear()
        QTest.keyClicks(edit, "-1")
        self.assertEqual(edit.text(), "-")
        for text in ("-", "-1"):
            with self.subTest(input=text), patch.object(service, "analyze_and_record", wraps=service.analyze_and_record) as analyze:
                edit.setText(text)
                self.assertIsNone(page.evaluate())
                analyze.assert_not_called()
                self.assertEqual(page.conclusion.text(), "—")
                self.assertIn("流量", edit.numeric_error.text())
                self.assertFalse(edit.numeric_error.isHidden())
                self.assertEqual(page.last_record_status, "NOT_RECORDED")
                self.assertIsNone(page.last_saved_record_id)
                self.assertEqual(service.list_records(), before)
        application = service.analyze_and_record(replace(request, QBEP="-1"))
        self.assertEqual(application.result.evaluation_status, "INVALID_INPUT")
        self.assertEqual(user_conclusion_text(application.result), "输入错误")
        self.assertFalse(application.result.finalizable)
        self.assertEqual(application.record_status, "NOT_RECORDED")
        self.assertIsNone(application.record)
        self.assertEqual(service.list_records(), before)

        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "input.xlsx"
            destination = Path(temp) / "result.xlsx"
            self._write_input(source, flow=-1)
            batch = self._batch(service).evaluate_workbook(
                source, destination=destination,
                as_of=date(2026, 10, 19), persist=False)
            outcome = batch.outcomes[0]
            self.assertTrue(outcome.is_input_error)
            self.assertFalse(outcome.is_execution_error)
            self.assertEqual(outcome.conclusion, user_conclusion_text(application.result))
            self.assertEqual(outcome.conclusion, "输入错误")
            self.assertEqual(batch.summary.input_error_rows, 1)
            self.assertEqual(batch.summary.execution_error_rows, 0)
            self.assertEqual(batch.summary.concluded_rows, 0)
            self.assertEqual(batch.summary.evaluated_quantity, 0)
            self.assertTrue(all(value == 0 for value in batch.summary.conclusion_quantities.values()))
            self.assertEqual(service.list_records(), before)
            workbook = openpyxl.load_workbook(destination, data_only=True)
            self.addCleanup(workbook.close)
            self.assertEqual(workbook[PUMP_SHEET]["X4"].value, "输入错误")

    def test_execution_failure_stays_separate_in_live_qt_and_batch_results(self):
        import openpyxl

        from equipeffi.infrastructure.excel.pump_workbook_reader import PUMP_SHEET
        from equipeffi.presentation.qt.pages.analysis import SYSTEM_FAILURE_TEXT

        service = self._service()
        page = self._page(service)
        self.addCleanup(page.deleteLater)
        page.evaluate()
        self.assertIsNotNone(page.last_saved_record_id)
        before = deepcopy(service.list_records())
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "input.xlsx"
            destination = Path(temp) / "result.xlsx"
            self._write_input(source, flow=100)
            with patch.object(service, "evaluate", side_effect=RuntimeError("simulated failure")) as evaluate:
                self.assertIsNone(page.evaluate())
                evaluate.assert_called_once()
                actual = evaluate.call_args.args[0]
                self.assertEqual((actual.QBEP, actual.HBEP, actual.speed, actual.efficiency,
                                  actual.suction, actual.stages), ("100", "50", "2900", "80", "单吸", "1"))
                self.assertEqual(page.conclusion.text(), "—")
                self.assertEqual(page.summary.text(), SYSTEM_FAILURE_TEXT)
                self.assertIsNone(page._last_result)
                self.assertIsNone(page.last_saved_record_id)
                self.assertEqual(service.list_records(), before)

                evaluate.reset_mock()
                batch = self._batch(service).evaluate_workbook(
                    source, destination=destination,
                    as_of=date(2026, 10, 19), persist=False)
                evaluate.assert_called_once()
            outcome = batch.outcomes[0]
            self.assertTrue(outcome.is_execution_error)
            self.assertFalse(outcome.is_input_error)
            self.assertEqual(outcome.conclusion, "执行失败")
            self.assertNotEqual(outcome.conclusion, "无法判定")
            self.assertEqual(batch.summary.execution_error_rows, 1)
            self.assertEqual(batch.summary.input_error_rows, 0)
            self.assertEqual(batch.summary.concluded_rows, 0)
            self.assertEqual(batch.summary.evaluated_quantity, 0)
            self.assertEqual(service.list_records(), before)
            workbook = openpyxl.load_workbook(destination, data_only=True)
            self.addCleanup(workbook.close)
            self.assertEqual(workbook[PUMP_SHEET]["X4"].value, "执行失败")


if __name__ == "__main__":
    unittest.main()
