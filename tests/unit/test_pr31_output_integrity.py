"""回归 PR31 结果清理与变压器输出区合并检查。"""
from __future__ import annotations

from copy import copy
from contextlib import closing
from decimal import Decimal
from pathlib import Path
import sqlite3
import tempfile
import unittest
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook

from equipeffi.application.ports.batch_workbook import BatchRowOutcome
from equipeffi.composition import create_batch_evaluation_service
from equipeffi.domain.devices.transformer.input_contract import (
    TRANSFORMER_CATEGORIES,
    condition_options,
)
from equipeffi.infrastructure.excel.ooxml_reader import OOXMLReadError, OOXMLWorkbook
from equipeffi.infrastructure.excel.pump_result_writer import _sheet_target
from equipeffi.infrastructure.excel.template_resource import (
    DEFAULT_V6_TEMPLATE,
    DEFAULT_V61_TEMPLATE,
    PUMP_SHEET_CONTRACT,
    TRANSFORMER_SHEET_CONTRACT,
    V61TransformerTemplateResource,
)
from equipeffi.infrastructure.excel.unified_result_writer import (
    TRANSFORMER_RESULT_COLUMNS,
    UnifiedResultWorkbookWriter,
)
from equipeffi.infrastructure.excel.unified_workbook_reader import UnifiedWorkbookReader
from equipeffi.infrastructure.persistence.app_data_paths import AppDataPaths
from equipeffi.infrastructure.excel.worksheet_patch import WorksheetPatch


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE_DIR = ROOT / "src" / "equipeffi" / "resources" / "templates"


def _column_index(column: str) -> int:
    value = 0
    for char in column:
        value = value * 26 + ord(char.upper()) - ord("A") + 1
    return value - 1


def _result_columns(contract: dict) -> tuple[str, ...]:
    return tuple(str(field["column"]) for field in contract["fields"]
                 if field.get("role") == "result")


def _worksheet_part(archive: ZipFile, name: str) -> str:
    part = _sheet_target(archive, name)
    if not part:
        raise AssertionError(f"找不到工作表部件：{name}")
    return part


def _cell_child_names(xml: str, reference: str) -> tuple[str, ...]:
    row_number = int("".join(char for char in reference if char.isdigit()))
    patch = WorksheetPatch(xml)
    row = patch.rows.get(row_number)
    if row is None:
        return ()
    cell = next((item for item in row.cells if item.reference == reference), None)
    if cell is None:
        return ()
    return tuple(child.local for child in cell.element.children
                 if child.ns_uri == "http://schemas.openxmlformats.org/spreadsheetml/2006/main")


class PR31OutputIntegrityTests(unittest.TestCase):
    def setUp(self):
        self._temporary = tempfile.TemporaryDirectory(prefix="equipeffi-pr31-output-")
        self.addCleanup(self._temporary.cleanup)
        self.temp_dir = Path(self._temporary.name)

    def _copy_template(self, name: str, *, v61: bool = True) -> Path:
        source = self.temp_dir / name
        template = DEFAULT_V61_TEMPLATE if v61 else DEFAULT_V6_TEMPLATE
        source.write_bytes((TEMPLATE_DIR / template).read_bytes())
        return source

    def _add_stale_results(self, source: Path, row_number: int = 5) -> None:
        workbook = load_workbook(source)
        for sheet_name, contract in (("离心泵", PUMP_SHEET_CONTRACT),
                                     ("变压器", TRANSFORMER_SHEET_CONTRACT)):
            sheet = workbook[sheet_name]
            for column in _result_columns(contract):
                sheet[f"{column}{row_number}"] = f"旧结果-{column}"
        workbook.save(source)

    def _make_outcome(self, row_number: int) -> BatchRowOutcome:
        return BatchRowOutcome(
            row_number, True, "1级", "SUCCESS", "1级",
            messages=("本次软件结果",),
            thresholds={
                "no_load_loss_1": Decimal("100.25"),
                "no_load_loss_2": Decimal("110.5"),
                "no_load_loss_3": Decimal("120.75"),
                "load_loss_1": Decimal("200.25"),
                "load_loss_2": Decimal("210.5"),
                "load_loss_3": Decimal("220.75"),
            })

    def test_writer_clears_stale_results_for_reader_skipped_rows(self):
        source = self._copy_template("mixed.xlsx")
        self._add_stale_results(source, row_number=5)
        workbook = load_workbook(source)
        workbook["离心泵"]["B4"] = "本次泵输入"
        workbook["变压器"]["S4"] = "铭牌图片引用"
        workbook["变压器"]["T4"] = "用户填写备注"
        workbook.save(source)

        read_result = UnifiedWorkbookReader().read(source)
        self.assertIn(4, [row.row_number for row in read_result.sheets["离心泵"].rows])
        self.assertNotIn(5, [row.row_number for row in read_result.sheets["离心泵"].rows])
        self.assertIn(4, [row.row_number for row in read_result.sheets["变压器"].rows])
        self.assertNotIn(5, [row.row_number for row in read_result.sheets["变压器"].rows])

        source_bytes = source.read_bytes()
        destination = self.temp_dir / "mixed-result.xlsx"
        outcome = self._make_outcome(4)
        UnifiedResultWorkbookWriter().write(
            source,
            {"离心泵": {4: outcome}, "变压器": {4: outcome}},
            destination)

        with OOXMLWorkbook(destination) as output:
            pump_rows = output.rows("离心泵")
            transformer_rows = output.rows("变压器")
        for column in _result_columns(PUMP_SHEET_CONTRACT):
            self.assertEqual(pump_rows[4][_column_index(column)], "", column)
        for column in TRANSFORMER_RESULT_COLUMNS:
            self.assertEqual(transformer_rows[4][_column_index(column)], "", column)
        self.assertEqual(pump_rows[3][_column_index("X")], "1级")
        self.assertEqual(transformer_rows[3][_column_index("R")], "1级")
        self.assertEqual(transformer_rows[3][_column_index("U")], "本次软件结果")
        self.assertEqual(transformer_rows[3][_column_index("S")], "铭牌图片引用")
        self.assertEqual(transformer_rows[3][_column_index("T")], "用户填写备注")
        self.assertEqual(source.read_bytes(), source_bytes, "源工作簿必须保持不变")

        with ZipFile(source) as before, ZipFile(destination) as after:
            target_parts = {
                _worksheet_part(before, "离心泵"),
                _worksheet_part(before, "变压器"),
            }
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if name not in target_parts:
                    self.assertEqual(before.read(name), after.read(name), name)

    def test_writer_clears_results_when_every_connected_sheet_is_empty(self):
        source = self._copy_template("empty.xlsx")
        self._add_stale_results(source, row_number=4)
        destination = self.temp_dir / "empty-result.xlsx"

        UnifiedResultWorkbookWriter().write(
            source, {"离心泵": {}, "变压器": {}}, destination)

        with OOXMLWorkbook(destination) as output:
            pump_rows = output.rows("离心泵")
            transformer_rows = output.rows("变压器")
        for column in _result_columns(PUMP_SHEET_CONTRACT):
            self.assertEqual(pump_rows[3][_column_index(column)], "", column)
        for column in TRANSFORMER_RESULT_COLUMNS:
            self.assertEqual(transformer_rows[3][_column_index(column)], "", column)

    def test_formula_and_style_only_rows_do_not_gain_result_cells(self):
        source = self._copy_template("empty-formatted.xlsx")
        workbook = load_workbook(source)
        pump = workbook["离心泵"]
        transformer = workbook["变压器"]
        pump["N104"] = "=1+1"
        transformer["R104"] = "=1+1"
        transformer["R105"]._style = copy(transformer["R103"]._style)
        workbook.save(source)

        destination = self.temp_dir / "empty-formatted-result.xlsx"
        UnifiedResultWorkbookWriter().write(
            source, {"离心泵": {}, "变压器": {}}, destination)

        with ZipFile(source) as before, ZipFile(destination) as after:
            for sheet_name, formula_ref, style_only_ref in (
                    ("离心泵", "N104", "N105"),
                    ("变压器", "R104", "R105")):
                source_part = _worksheet_part(before, sheet_name)
                result_part = _worksheet_part(after, sheet_name)
                source_xml = before.read(source_part).decode("utf-8")
                result_xml = after.read(result_part).decode("utf-8")
                source_patch = WorksheetPatch(source_xml)
                result_patch = WorksheetPatch(result_xml)
                self.assertEqual(
                    {ref for row in source_patch.rows.values() for ref in
                     (cell.reference for cell in row.cells)},
                    {ref for row in result_patch.rows.values() for ref in
                     (cell.reference for cell in row.cells)})
                self.assertEqual(_cell_child_names(result_xml, formula_ref), ())
                self.assertEqual(_cell_child_names(result_xml, style_only_ref), ())

    def test_legacy_v6_does_not_clear_preserved_transformer_sheet(self):
        source = self._copy_template("legacy-v6.xlsx", v61=False)
        workbook = load_workbook(source)
        workbook["离心泵"]["B4"] = "本次泵输入"
        workbook["变压器"]["R4"] = "保留的旧V6内容"
        workbook.save(source)
        destination = self.temp_dir / "legacy-v6-result.xlsx"

        with ZipFile(source) as before:
            transformer_part = _worksheet_part(before, "变压器")
            original_transformer = before.read(transformer_part)
        UnifiedResultWorkbookWriter().write(
            source, {"离心泵": {4: self._make_outcome(4)}}, destination)
        with ZipFile(destination) as after:
            self.assertEqual(after.read(transformer_part), original_transformer)

    def test_transformer_result_merge_is_rejected_but_title_merges_remain_valid(self):
        source = self._copy_template("merged-transformer.xlsx")
        self.assertTrue(V61TransformerTemplateResource().validate(source).is_valid)

        workbook = load_workbook(source)
        transformer = workbook["变压器"]
        transformer.protection.sheet = False
        transformer.merge_cells("R4:R5")
        workbook.save(source)

        validation = V61TransformerTemplateResource().validate(source)
        self.assertFalse(validation.is_valid, validation.message)
        self.assertIn("R4:R5", validation.message)
        self.assertIn("变压器", validation.message)

    def test_real_transformer_batches_clear_skipped_results_and_reject_result_merge(self):
        paths = AppDataPaths(self.temp_dir / "app-data")
        batch = create_batch_evaluation_service(paths=paths)
        source = self._copy_template("two-transformers.xlsx")
        workbook = load_workbook(source)
        transformer = workbook["变压器"]
        category = TRANSFORMER_CATEGORIES[0]
        condition = condition_options(category.category_id)[0]
        for row_number, name, no_load, load in (
                (4, "一级设备", 60, 400), (5, "二级设备", 68, 480)):
            for column, value in {
                    "B": name,
                    "C": f"T-{row_number}",
                    "D": 1,
                    "F": category.label,
                    "G": 30,
                    "H": condition.label,
                    "J": no_load,
                    "K": load,
            }.items():
                transformer[f"{column}{row_number}"] = value
        workbook.save(source)
        workbook.close()

        first_output = self.temp_dir / "first-evaluation.xlsx"
        first = batch.evaluate_workbook(source, destination=first_output)
        self.assertEqual(first.batch_record_error, "")
        self.assertEqual(first.summary.data_row_count, 2)
        self.assertEqual(first.summary.conclusion_quantities, {"1级": 1, "2级": 1})
        first_outcomes = [item["outcome"] for item in first.outcomes]
        self.assertEqual([(item.row_number, item.conclusion) for item in first_outcomes],
                         [(4, "1级"), (5, "2级")])

        # Reuse the prior result workbook as the next input, then remove every
        # formal input value from row 4 while leaving its old software outputs.
        second_source = self.temp_dir / "second-input.xlsx"
        second_source.write_bytes(first_output.read_bytes())
        workbook = load_workbook(second_source)
        transformer = workbook["变压器"]
        for field in TRANSFORMER_SHEET_CONTRACT["fields"]:
            if field.get("role") == "input":
                transformer[f"{field['column']}4"] = None
        workbook.save(second_source)
        workbook.close()
        with OOXMLWorkbook(second_source) as before_second:
            stale_row = before_second.rows("变压器")[3]
        self.assertEqual(stale_row[_column_index("R")], "1级")

        second_output = self.temp_dir / "second-evaluation.xlsx"
        second = batch.evaluate_workbook(second_source, destination=second_output)
        self.assertEqual(second.batch_record_error, "")
        self.assertEqual(second.summary.data_row_count, 1)
        self.assertEqual(second.summary.conclusion_quantities, {"2级": 1})
        second_outcomes = [item["outcome"] for item in second.outcomes]
        self.assertEqual([(item.row_number, item.conclusion) for item in second_outcomes],
                         [(5, "2级")])
        with OOXMLWorkbook(second_output) as after_second:
            output_rows = after_second.rows("变压器")
        for column in TRANSFORMER_RESULT_COLUMNS:
            self.assertEqual(output_rows[3][_column_index(column)], "", column)
        self.assertEqual(output_rows[4][_column_index("R")], "2级")

        # The same two rows yield different real conclusions. Merging their
        # result cells must fail at the Reader/Service boundary before writing.
        merged_source = self.temp_dir / "merged-input.xlsx"
        merged_source.write_bytes(source.read_bytes())
        workbook = load_workbook(merged_source)
        transformer = workbook["变压器"]
        transformer.protection.sheet = False
        transformer.merge_cells("R4:R5")
        workbook.save(merged_source)
        workbook.close()
        self.assertFalse(V61TransformerTemplateResource().validate(merged_source).is_valid)
        with self.assertRaisesRegex(OOXMLReadError, "R4:R5"):
            UnifiedWorkbookReader().read(merged_source)
        rejected_output = self.temp_dir / "must-not-exist.xlsx"
        with self.assertRaisesRegex(OOXMLReadError, "R4:R5"):
            batch.evaluate_workbook(merged_source, destination=rejected_output)
        self.assertFalse(rejected_output.exists())

        self.assertEqual(len(batch.list_batch_records()), 2)
        self.assertEqual(first.records_created, 0)
        self.assertEqual(second.records_created, 0)
        with closing(sqlite3.connect(paths.records_db)) as connection:
            record_count = connection.execute("SELECT COUNT(*) FROM record").fetchone()[0]
        self.assertEqual(record_count, 0)


if __name__ == "__main__":
    unittest.main()
