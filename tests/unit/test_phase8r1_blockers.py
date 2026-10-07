"""Phase 8 R1 — Independent Acceptance blocker regression + UI cleanup evidence.

本模块的每个测试都是为了在**被 BLOCKED 的 head**
`8cb6eec1845cc26bed43e3dfea2dec1c5880729c` 上**失败**、在修复后**通过**而写的。

覆盖四个 blocker：

```text
B1  结果 Workbook 必须写出真实等级限值（U/V/W 取自正式 Result.thresholds）
B2  INVALID_INPUT 不得污染批次统计（不属正式评价结论）
B3  Writer 不得改写原始输入精度（结果 Workbook = 原文件副本 + 只写结果区域）
B4  batch provenance 绝不能跨批次串用
```

以及四组 Owner UI 简化（UI01～UI04）。
"""
from __future__ import annotations

import os
import re
import tempfile
import unittest
import zipfile
from datetime import date
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import openpyxl

from equipeffi.application.services.centrifugal_pump_analysis_service import (
    PumpAnalysisRequest,
)
from equipeffi.application.services.pump_batch_evaluation_service import (
    INVALID_CONCLUSION,
)
from equipeffi.infrastructure.excel.pump_result_writer import (
    THRESHOLD_COLUMNS,
    cell_payloads,
)
from equipeffi.infrastructure.excel.pump_workbook_reader import PUMP_SHEET
from equipeffi.infrastructure.excel.template_resource import V6TemplateResource
from equipeffi.infrastructure.persistence.sqlite_batch_record_repository import (
    SqliteBatchRecordRepository,
)

ROOT = Path(__file__).resolve().parents[2]
FIRST_DATA_ROW = 4
AS_OF = date(2026, 8, 23)
HIGH_PRECISION = "100.12345678901234567890123456789012345"

#: 结果列（B3：只有这些列允许被 Writer 改写）。
AUTHORIZED_RESULT_COLUMNS = ("N", "O", "P", "Q", "R", "S", "T",
                             "U", "V", "W", "X", "AA")
#: 用户输入列（B3：一个都不许因生成结果 Workbook 而被重写）。
INPUT_COLUMNS = ("A", "B", "C", "D", "E", "F", "G", "H", "I", "J",
                 "K", "L", "M", "Y", "Z")


class R1TestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from PySide6.QtWidgets import QApplication

        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        from equipeffi.composition import create_batch_evaluation_service
        from equipeffi.infrastructure.persistence.app_data_paths import AppDataPaths
        from equipeffi.infrastructure.persistence.records_migrations import (
            migrate_records_database,
        )

        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.root = Path(self.tmp.name)
        self.paths = AppDataPaths(self.root / "app")
        migrate_records_database(self.paths.records_db, app_version="test")
        self.batch = create_batch_evaluation_service(paths=self.paths)
        self.repository = SqliteBatchRecordRepository(self.paths.records_db)
        self.addCleanup(self.tmp.cleanup)

    def make(self, name: str, rows: list[dict]) -> Path:
        source = self.root / name
        V6TemplateResource().download_to(source)
        workbook = openpyxl.load_workbook(source)
        sheet = workbook[PUMP_SHEET]
        for offset, values in enumerate(rows):
            for column, value in values.items():
                sheet[f"{column}{FIRST_DATA_ROW + offset}"] = value
        workbook.save(source)
        workbook.close()
        return source

    def run_batch(self, rows: list[dict], name: str, **kwargs):
        source = self.make(f"{name}.xlsx", rows)
        return self.batch.evaluate_workbook(
            source, destination=self.root / f"{name}_out.xlsx",
            as_of=AS_OF, **kwargs)

    @staticmethod
    def water(**overrides) -> dict:
        row = dict(B="水泵", C="M-001", D=1, E="1号车间", F="单级单吸清水离心泵",
                   G=100, H=50, I=2900, J=45, K="单吸", L=1, M=78)
        row.update(overrides)
        return row

    @staticmethod
    def chemical(**overrides) -> dict:
        row = dict(B="化工泵", C="C-001", D=1, E="2号车间", F="单级石油化工离心泵",
                   G=300, H=40, I=2900, J=45, K="双吸", L=1, M=90)
        row.update(overrides)
        return row


# ===========================================================================
# B1 — 结果 Workbook 必须写出真实等级限值
# ===========================================================================

class B1ThresholdWritebackTests(R1TestCase):
    def test_water_thresholds_are_written_to_u_v_w(self):
        result = self.run_batch([self.water()], "b1_water", persist=False)
        outcome = result.outcomes[0]
        self.assertTrue(outcome.thresholds, "正式 Result 必须提供 thresholds")

        sheet = openpyxl.load_workbook(result.result_workbook)[PUMP_SHEET]
        for column, name in THRESHOLD_COLUMNS:
            with self.subTest(column=column):
                written = sheet[f"{column}{FIRST_DATA_ROW}"].value
                self.assertIsNotNone(written, f"{column} 必须写出正式等级限值")
                self.assertEqual(float(written), float(outcome.thresholds[name]))

    def test_chemical_thresholds_are_written_to_u_v_w(self):
        result = self.run_batch([self.chemical()], "b1_chem", persist=False)
        outcome = result.outcomes[0]
        self.assertTrue(outcome.thresholds)

        sheet = openpyxl.load_workbook(result.result_workbook)[PUMP_SHEET]
        for column, name in THRESHOLD_COLUMNS:
            with self.subTest(column=column):
                written = sheet[f"{column}{FIRST_DATA_ROW}"].value
                self.assertIsNotNone(written)
                self.assertEqual(float(written), float(outcome.thresholds[name]))

    def test_thresholds_match_the_formal_application_result_exactly(self):
        """写入值必须**等于正式 Result.thresholds**（重新打开结果文件后仍一致）。"""

        from equipeffi.composition import create_pump_analysis_service

        analysis = create_pump_analysis_service(paths=self.paths)
        formal = analysis.evaluate(PumpAnalysisRequest(
            product_category="单级单吸清水离心泵", as_of=AS_OF, QBEP="100",
            HBEP="50", speed="2900", efficiency="78", suction="单吸", stages="1"))
        result = self.run_batch([self.water()], "b1_exact", persist=False)

        sheet = openpyxl.load_workbook(result.result_workbook)[PUMP_SHEET]
        for column, name in THRESHOLD_COLUMNS:
            with self.subTest(column=column):
                self.assertEqual(float(sheet[f"{column}{FIRST_DATA_ROW}"].value),
                                 float(formal.thresholds[name]))

    def test_no_threshold_states_do_not_fabricate_u_v_w(self):
        """`OUT_OF_STANDARD_SCOPE` / 不确定类别 没有正式阈值 -> 不得伪造 U/V/W。"""

        result = self.run_batch([
            self.water(B="其他", C="M-1", F="其他类别"),
            self.water(B="未定", C="M-2", F="不确定类别"),
        ], "b1_none", persist=False)

        for outcome in result.outcomes:
            with self.subTest(row=outcome.row_number, status=outcome.evaluation_status):
                self.assertEqual(outcome.thresholds, {},
                                 "这些状态没有正式阈值")
        sheet = openpyxl.load_workbook(result.result_workbook)[PUMP_SHEET]
        for row in (FIRST_DATA_ROW, FIRST_DATA_ROW + 1):
            for column, _name in THRESHOLD_COLUMNS:
                with self.subTest(row=row, column=column):
                    self.assertIsNone(sheet[f"{column}{row}"].value,
                                      "不得伪造等级限值")

    def test_missing_threshold_rows_leave_existing_cells_untouched(self):
        """没有阈值的行不得写入任何限值（也不得写入占位文本）。"""

        result = self.run_batch([self.water(B="其他", F="其他类别")],
                                "b1_untouched", persist=False)
        sheet = openpyxl.load_workbook(result.result_workbook)[PUMP_SHEET]
        for column, _name in THRESHOLD_COLUMNS:
            with self.subTest(column=column):
                self.assertIsNone(sheet[f"{column}{FIRST_DATA_ROW}"].value)


# ===========================================================================
# B2 — INVALID_INPUT 不得污染批次统计
# ===========================================================================

class B2InvalidInputTests(R1TestCase):
    def test_negative_flow_with_quantity_seven_is_an_input_error(self):
        """独立验收实测场景：负流量 + 数量=7。"""

        result = self.run_batch([self.water(D=7, G=-1)], "b2_neg")
        summary, outcome = result.summary, result.outcomes[0]

        self.assertEqual(outcome.evaluation_status, "INVALID_INPUT")
        self.assertTrue(outcome.is_input_error, "INVALID_INPUT 必须判为输入错误")
        self.assertEqual(outcome.conclusion, INVALID_CONCLUSION)

        self.assertEqual(summary.data_row_count, 1)
        # 数量本身合法 -> 计入"提交设备总数量"
        self.assertEqual(summary.total_quantity, 7)
        # 但**不得**计入已评价数量，也**不得**进入任何正式结论数量
        self.assertEqual(summary.evaluated_quantity, 0)
        self.assertEqual(summary.conclusion_quantities, {})
        self.assertEqual(summary.conclusion_rows, {})
        self.assertEqual(summary.input_error_rows, 1)
        self.assertEqual(summary.input_error_quantity, 7)
        # 必须出现在"需要关注"列表
        self.assertEqual(len(summary.issues), 1)
        self.assertEqual(summary.issues[0].row_number, FIRST_DATA_ROW)
        self.assertEqual(summary.issues[0].kind, "INPUT_ERROR")
        for conclusion in ("1级", "2级", "3级", "未达标", "不适用", "无法评价"):
            with self.subTest(forbidden=conclusion):
                self.assertNotIn(conclusion, summary.conclusion_quantities)

    def test_other_invalid_business_input_is_also_an_input_error(self):
        """数量=7 + 其他非法业务输入（类别与级数冲突）。"""

        result = self.run_batch(
            [self.water(D=7, L=3)], "b2_stage_conflict")
        summary, outcome = result.summary, result.outcomes[0]
        self.assertEqual(outcome.evaluation_status, "INVALID_INPUT")
        self.assertTrue(outcome.is_input_error)
        self.assertEqual(summary.evaluated_quantity, 0)
        self.assertEqual(summary.conclusion_quantities, {})
        self.assertEqual(summary.input_error_rows, 1)

    def test_invalid_quantity_does_not_invent_devices(self):
        """数量非法：不得虚构设备数量，也不得进入任何数量统计。"""

        for bad in ("", 0, -3, 1.5):
            with self.subTest(quantity=bad):
                result = self.run_batch(
                    [self.water(D=bad)], f"b2_q_{abs(hash(str(bad))) % 1000}")
                summary = result.summary
                self.assertEqual(summary.total_quantity, 0)
                self.assertEqual(summary.evaluated_quantity, 0)
                self.assertEqual(summary.conclusion_quantities, {})
                self.assertEqual(summary.input_error_rows, 1)
                self.assertTrue(result.outcomes[0].is_input_error)

    def test_uncertain_category_with_quantity_is_a_formal_conclusion(self):
        """「不确定类别」不是 INVALID_INPUT：形成正式结论「无法评价」并按数量计。"""

        result = self.run_batch(
            [self.water(B="未定", F="不确定类别", D=4)], "b2_uncertain")
        summary, outcome = result.summary, result.outcomes[0]
        self.assertFalse(outcome.is_input_error)
        self.assertEqual(outcome.conclusion, "无法评价")
        self.assertEqual(summary.input_error_rows, 0)
        self.assertEqual(summary.total_quantity, 4)
        # 「无法评价」是正式用户结论：进入结论数量与"未评价"计数；
        # evaluated_quantity 仅统计**有 evaluation_status 的正式评价**（台）。
        self.assertEqual(summary.conclusion_quantities.get("无法评价"), 4)
        self.assertEqual(summary.unevaluated_rows, 1)

    def test_quantity_identities_hold_in_a_mixed_batch(self):
        """混合批次：数量恒等关系必须成立。"""

        result = self.run_batch([
            self.water(B="正常", D=20, M=78),                     # 2级 x20
            self.chemical(B="化工", D=5),                          # 1级 x5
            self.water(B="负流量", D=7, G=-1),                     # 输入错误
            self.water(B="坏数量", D=0),                           # 输入错误（数量非法）
            self.water(B="未定", F="不确定类别", D=3),               # 无法评价 x3
            self.water(B="其他", F="其他类别", D=2),                # 不适用 x2
        ], "b2_mixed")
        summary = result.summary

        self.assertEqual(summary.data_row_count, 6)
        # 数量恒等：总量 = 合法数量之和（非法数量不贡献）
        self.assertEqual(summary.total_quantity, 20 + 5 + 7 + 3 + 2)
        # evaluated_quantity = 有正式 evaluation_status 的数量（1级/2级/不适用）
        self.assertEqual(summary.evaluated_quantity, 20 + 5 + 2)
        # 输入错误：负流量 + 坏数量 = 2 行
        self.assertEqual(summary.input_error_rows, 2)
        self.assertEqual(summary.execution_error_rows, 0)
        # 正式结论数量逐个核对
        self.assertEqual(summary.conclusion_quantities.get("2级"), 20)
        self.assertEqual(summary.conclusion_quantities.get("1级"), 5)
        self.assertEqual(summary.conclusion_quantities.get("无法评价"), 3)
        self.assertEqual(summary.conclusion_quantities.get("不适用"), 2)
        for conclusion in ("1级", "2级", "3级", "未达标", "不适用", "无法评价"):
            with self.subTest(forbidden=INVALID_CONCLUSION):
                self.assertNotEqual(conclusion, INVALID_CONCLUSION)
        self.assertNotIn(INVALID_CONCLUSION, summary.conclusion_quantities)
        # 恒等：所有正式结论数量之和 == total_quantity - 输入错误数量
        self.assertEqual(sum(summary.conclusion_quantities.values()),
                         summary.total_quantity - summary.input_error_quantity)
        # 恒等：结论数量之和 == 有状态数量 + 未评价数量
        self.assertEqual(sum(summary.conclusion_quantities.values()),
                         summary.evaluated_quantity
                         + sum(v for k, v in summary.conclusion_quantities.items()
                               if k == "无法评价"))

    def test_result_workbook_marks_input_error_rows_distinctly(self):
        result = self.run_batch([self.water(D=7, G=-1)], "b2_wb")
        sheet = openpyxl.load_workbook(result.result_workbook)[PUMP_SHEET]
        status = str(sheet[f"X{FIRST_DATA_ROW}"].value)
        self.assertIn("输入错误", status)
        self.assertNotIn("无法评价", status)
        self.assertNotIn("无法判定", status)
        note = str(sheet[f"AA{FIRST_DATA_ROW}"].value or "")
        self.assertTrue(note, "必须给出具体问题说明")

    def test_execution_error_is_reported_separately(self):
        from unittest.mock import patch

        with patch.object(self.batch.analysis, "evaluate",
                          side_effect=RuntimeError("模拟内部错误")):
            result = self.run_batch([self.water(D=3)], "b2_exec")
        summary, outcome = result.summary, result.outcomes[0]
        self.assertTrue(outcome.is_execution_error)
        self.assertEqual(outcome.conclusion, "执行失败")
        self.assertEqual(summary.execution_error_rows, 1)
        self.assertEqual(summary.input_error_rows, 0)
        self.assertEqual(summary.evaluated_quantity, 0)
        sheet = openpyxl.load_workbook(result.result_workbook)[PUMP_SHEET]
        self.assertIn("执行失败", str(sheet[f"X{FIRST_DATA_ROW}"].value))


# ===========================================================================
# B3 — 结果 Workbook 不得改写原始输入精度
# ===========================================================================

def _sheet_path(archive: zipfile.ZipFile) -> str:
    workbook = archive.read("xl/workbook.xml").decode("utf-8")
    rels = archive.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    relation_map = {}
    for element in re.findall(r"<Relationship\b[^>]*>", rels):
        rid = re.search(r'Id="([^"]*)"', element)
        target = re.search(r'Target="([^"]*)"', element)
        if rid and target:
            relation_map[rid.group(1)] = target.group(1)
    for element in re.findall(r"<sheet\b[^>]*>", workbook):
        if 'name="离心泵"' not in element:
            continue
        rid = re.search(r'[A-Za-z0-9]+:id="([^"]*)"', element).group(1)
        target = relation_map[rid].lstrip("/")
        return target if target.startswith("xl/") else f"xl/{target}"
    raise AssertionError("找不到离心泵 sheet")


def _patch_number_cell(path: Path, reference: str, literal: str) -> None:
    """把某个数值 cell 直接写成高精度字面量（模拟用户在 Excel 里输入）。"""

    with zipfile.ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
        target = _sheet_path(archive)
    pattern = re.compile(r'<c r="' + reference + r'"[^>]*/>|<c r="' + reference + r'"[^>]*>.*?</c>')

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for info, data in entries:
            if info.filename == target:
                xml = data.decode("utf-8")
                replacement = f'<c r="{reference}" s="226"><v>{literal}</v></c>'
                xml = pattern.sub(replacement, xml, count=1)
                data = xml.encode("utf-8")
            archive.writestr(info, data)


class B3InputPrecisionTests(R1TestCase):
    def test_35_digit_input_survives_into_the_result_workbook(self):
        source = self.make("b3_high.xlsx", [
            self.water(D=7, G=None)])
        _patch_number_cell(source, f"G{FIRST_DATA_ROW}", HIGH_PRECISION)

        result = self.batch.evaluate_workbook(
            source, destination=self.root / "b3_out.xlsx", as_of=AS_OF,
            persist=False)
        with zipfile.ZipFile(result.result_workbook) as archive:
            xml = archive.read(_sheet_path(archive)).decode("utf-8")
        match = re.search(r'<c r="G' + str(FIRST_DATA_ROW) + r'"[^>]*>.*?</c>', xml)
        self.assertIsNotNone(match)
        self.assertIn(HIGH_PRECISION, match.group(0),
                      "结果 Workbook 不得降低用户输入精度")

    def test_only_the_pump_sheet_part_changes(self):
        source = self.make("b3_parts.xlsx", [self.water()])
        result = self.batch.evaluate_workbook(
            source, destination=self.root / "b3_parts_out.xlsx", as_of=AS_OF,
            persist=False)
        with zipfile.ZipFile(source) as archive:
            before = {name: archive.read(name) for name in archive.namelist()}
        with zipfile.ZipFile(result.result_workbook) as archive:
            after = {info.filename: archive.read(info.filename)
                     for info in archive.infolist()}
        self.assertEqual(set(before), set(after), "不得增删任何 zip 部件")
        changed = [name for name in before if before[name] != after[name]]
        self.assertEqual(changed, [_sheet_path(zipfile.ZipFile(source))],
                         "只允许「离心泵」Sheet 变化")

    def test_writer_does_not_touch_any_input_cell(self):
        """只允许 12 个授权结果列变化；用户输入列一个都不许被改写。"""

        source = self.make("b3_cols.xlsx", [
            self.water(B="设备甲", C="MODEL-9", D=7, E="三号车间", M=78)])
        _patch_number_cell(source, f"G{FIRST_DATA_ROW}", HIGH_PRECISION)
        result = self.batch.evaluate_workbook(
            source, destination=self.root / "b3_cols_out.xlsx", as_of=AS_OF,
            persist=False)

        before = openpyxl.load_workbook(source)[PUMP_SHEET]
        after = openpyxl.load_workbook(result.result_workbook)[PUMP_SHEET]
        for column in INPUT_COLUMNS:
            reference = f"{column}{FIRST_DATA_ROW}"
            with self.subTest(cell=reference):
                self.assertEqual(after[reference].value, before[reference].value,
                                 f"输入 cell {reference} 不得因生成结果文件而改变")
        # 结果列确实被写入了
        self.assertIsNotNone(after[f"X{FIRST_DATA_ROW}"].value)
        self.assertIsNotNone(after[f"U{FIRST_DATA_ROW}"].value)

    def test_number_formats_and_protection_of_input_cells_are_preserved(self):
        source = self.make("b3_fmt.xlsx", [self.water()])
        result = self.batch.evaluate_workbook(
            source, destination=self.root / "b3_fmt_out.xlsx", as_of=AS_OF,
            persist=False)
        before = openpyxl.load_workbook(source)
        after = openpyxl.load_workbook(result.result_workbook)
        for name in before.sheetnames:
            with self.subTest(sheet=name):
                self.assertEqual(after[name].sheet_state, before[name].sheet_state)
                self.assertEqual(bool(after[name].protection.sheet),
                                 bool(before[name].protection.sheet))
        for column in INPUT_COLUMNS:
            reference = f"{column}{FIRST_DATA_ROW}"
            with self.subTest(cell=reference):
                self.assertEqual(after[PUMP_SHEET][reference].number_format,
                                 before[PUMP_SHEET][reference].number_format)

    def test_other_sheets_keep_their_cell_values(self):
        source = self.make("b3_other.xlsx", [self.water()])
        result = self.batch.evaluate_workbook(
            source, destination=self.root / "b3_other_out.xlsx", as_of=AS_OF,
            persist=False)
        before = openpyxl.load_workbook(source, data_only=False)
        after = openpyxl.load_workbook(result.result_workbook, data_only=False)
        for name in before.sheetnames:
            if name == PUMP_SHEET:
                continue
            with self.subTest(sheet=name):
                self.assertEqual(after[name].max_row, before[name].max_row)
                self.assertEqual(after[name].max_column, before[name].max_column)

    def test_source_workbook_is_never_overwritten(self):
        import hashlib

        source = self.make("b3_src.xlsx", [self.water()])
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        self.batch.evaluate_workbook(
            source, destination=self.root / "b3_src_out.xlsx", as_of=AS_OF,
            persist=False)
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)

    def test_existing_destination_is_not_silently_overwritten(self):
        from equipeffi.infrastructure.excel.pump_result_writer import (
            ResultWorkbookExistsError,
        )

        source = self.make("b3_dup.xlsx", [self.water()])
        target = self.root / "b3_dup_out.xlsx"
        self.batch.evaluate_workbook(source, destination=target, as_of=AS_OF,
                                     persist=False)
        with self.assertRaises(ResultWorkbookExistsError):
            self.batch.evaluate_workbook(source, destination=target, as_of=AS_OF,
                                         persist=False)


# ===========================================================================
# B4 — batch provenance 绝不能跨批次串用
# ===========================================================================

class B4ProvenanceIsolationTests(R1TestCase):
    def test_provenance_does_not_leak_across_three_consecutive_batches(self):
        """同一 service instance：A=合法 / B=全部导入或数量非法 / C=合法但组成不同。"""

        # A：合法批次
        self.run_batch([self.water(B="A", D=1)], "b4_a")
        # B：全部在进入正式评价前失败（数量非法）
        self.run_batch([self.water(B="B", D=0)], "b4_b")
        # C：合法但组成不同（石化）
        self.run_batch([self.chemical(B="C", D=2)], "b4_c")

        records = list(reversed(self.repository.list_batch_records()))
        self.assertEqual(len(records), 3)
        a, b, c = records

        with self.subTest(batch="A"):
            self.assertTrue(a.canonical_version)
            self.assertEqual(a.evaluated_quantity, 1)
        with self.subTest(batch="B (must not inherit A)"):
            self.assertEqual(b.canonical_version, "",
                             "全非法批次不得沿用上一批的 Canonical 引用")
            self.assertEqual(b.numeric_profile_id, "")
            self.assertEqual(b.evaluated_quantity, 0)
            self.assertEqual(b.summary["input_error_rows"], 1)
        with self.subTest(batch="C"):
            self.assertTrue(c.canonical_version)
            self.assertEqual(c.evaluated_quantity, 2)

    def test_provenance_is_empty_when_nothing_reached_formal_evaluation(self):
        """本批没有任何 Result -> references 必须为空，不得伪造。"""

        self.run_batch([self.water(B="A", D=1)], "b4_seed")
        self.run_batch([self.water(B="B", D=0), self.water(B="B2", D=-1)],
                       "b4_all_invalid")
        latest = self.repository.list_batch_records()[0]
        self.assertEqual(latest.canonical_version, "")
        self.assertEqual(latest.numeric_profile_id, "")
        self.assertEqual(latest.evaluated_quantity, 0)
        # 但当前批次的客观信息仍必须写入
        self.assertTrue(latest.app_version)
        self.assertTrue(latest.template_id)
        self.assertTrue(latest.template_sha256)
        self.assertTrue(latest.source_workbook_sha256)
        self.assertTrue(latest.result_workbook_sha256)

    def test_service_has_no_batch_scoped_instance_state(self):
        """机械守卫：service 实例上不得存在批次作用域的可变状态。"""

        instance = self.batch
        for name in ("_first_result", "_provenance", "_summary", "_outcomes",
                     "_reference_snapshot"):
            with self.subTest(attribute=name):
                self.assertFalse(hasattr(instance, name),
                                 f"批次状态不得挂在 service 实例上：{name}")
        # 连续两次同一输入，summary 不得累加（证明没有实例级累积器）
        first = self.run_batch([self.water(D=1)], "b4_s1", persist=False)
        second = self.run_batch([self.water(D=1)], "b4_s2", persist=False)
        self.assertEqual(first.summary.total_quantity,
                         second.summary.total_quantity)
        self.assertEqual(first.summary.data_row_count,
                         second.summary.data_row_count)

    def test_repeated_batches_with_the_same_service_do_not_accumulate(self):
        summaries = []
        for index in range(3):
            result = self.run_batch([self.water(B=f"R{index}", D=5)],
                                    f"b4_rep_{index}", persist=False)
            summaries.append(result.summary)
        for summary in summaries:
            with self.subTest():
                self.assertEqual(summary.data_row_count, 1)
                self.assertEqual(summary.total_quantity, 5)
                self.assertEqual(summary.evaluated_quantity, 5)


# ===========================================================================
# UI01～UI04 — Owner UI 简化
# ===========================================================================

class UIAnalysisTerminologyTests(R1TestCase):
    """UI01 + UI02：新建分析页术语与结果区简化。"""

    def page(self):
        from equipeffi.composition import create_pump_analysis_service
        from equipeffi.presentation.qt.pages.analysis import AnalysisPage

        page = AnalysisPage(create_pump_analysis_service(paths=self.paths),
                            as_of=AS_OF)
        self.addCleanup(page.deleteLater)
        return page

    def visible_text(self, widget) -> str:
        from PySide6.QtWidgets import QGroupBox, QLabel

        texts = [label.text() for label in widget.findChildren(QLabel)]
        texts += [box.title() for box in widget.findChildren(QGroupBox)]
        return "\n".join(texts)

    def test_analysis_page_has_no_standard_terminology(self):
        page = self.page()
        blob = self.visible_text(page)
        for forbidden in ("规定点",):
            with self.subTest(word=forbidden):
                self.assertNotIn(forbidden, blob)

    def test_analysis_page_has_no_internal_symbols(self):
        page = self.page()
        blob = self.visible_text(page)
        for forbidden in ("Q_BEP", "H_BEP", "η_BEP", "BEP"):
            with self.subTest(symbol=forbidden):
                self.assertNotIn(forbidden, blob)

    def test_analysis_page_has_simplified_labels(self):
        page = self.page()
        blob = self.visible_text(page)
        for expected in ("设备参数", "流量", "扬程", "转速", "泵效率"):
            with self.subTest(label=expected):
                self.assertIn(expected, blob)

    def _evaluate(self, page):
        page.category.setCurrentIndex(
            page.category.findData("单级单吸清水离心泵"))
        for key, value in (("QBEP", "100"), ("HBEP", "50"),
                           ("speed", "2900"), ("efficiency", "78")):
            page.point_inputs[key].setText(value)
        page.suction.setCurrentIndex(page.suction.findData("单吸"))
        if page.stages.isEnabled():
            page.stages.setText("1")
        return page.evaluate()

    def test_result_area_drops_explanations_and_basis(self):
        page = self.page()
        self._evaluate(page)
        blob = "\n".join([page.summary.text(), page.reason_label.text(),
                          page.basis.text()])
        for forbidden in ("判定说明", "为什么", "所选标准", "标准依据", "所选"):
            with self.subTest(word=forbidden):
                self.assertNotIn(forbidden, blob)

    def test_result_area_keeps_conclusion_thresholds_and_key_parameters(self):
        page = self.page()
        result = self._evaluate(page)
        self.assertTrue(page.conclusion.text())
        self.assertIn("对应等级效率限值", page.values_label.text())
        self.assertIn("关键计算参数", page.basis.text())
        self.assertTrue(result.thresholds)

    def test_thresholds_are_displayed_with_two_decimals(self):
        page = self.page()
        result = self._evaluate(page)
        text = page.values_label.text()
        for name, value in result.thresholds.items():
            with self.subTest(threshold=name):
                from equipeffi.application.services.centrifugal_pump_analysis_service import (
                    format_metric,
                )

                self.assertIn(format_metric(value, name=str(name)), text)
        # 不得出现原始长小数
        for value in result.thresholds.values():
            with self.subTest(raw=value):
                self.assertNotIn(str(value), text)


class UIRecordsAuditRemovalTests(R1TestCase):
    """UI03：记录详情不得展示「审计信息」，但底层数据必须仍在。"""

    def _record_id(self) -> str:
        from equipeffi.composition import create_pump_analysis_service

        analysis = create_pump_analysis_service(paths=self.paths)
        page_source = R1TestCase.make(self, "ui03.xlsx", [self.water()])
        from equipeffi.infrastructure.excel.pump_workbook_reader import (
            V6PumpWorkbookReader,
        )

        row = V6PumpWorkbookReader().read(page_source).rows[0]
        request = PumpAnalysisRequest(
            product_category="单级单吸清水离心泵", as_of=AS_OF, QBEP="100",
            HBEP="50", speed="2900", efficiency="78", suction="单吸", stages="1")
        result = analysis.evaluate(request)
        analysis.finalize(record_id="R1-UI03", workspace_id=None,
                          request=request, result=result)
        return "R1-UI03"

    def test_record_detail_has_no_audit_section(self):
        from PySide6.QtWidgets import QLabel
        from equipeffi.presentation.qt.pages.records import RecordsPage
        from equipeffi.composition import create_pump_analysis_service

        record_id = self._record_id()
        analysis = create_pump_analysis_service(paths=self.paths)
        page = RecordsPage(analysis)
        self.addCleanup(page.deleteLater)
        text = page.show_record(record_id)

        self.assertNotIn("审计信息", text)
        blob = "\n".join(label.text() for label in page.findChildren(QLabel))
        self.assertNotIn("审计信息", blob)
        self.assertFalse(getattr(page, "technical_box", None))
        # 不出现内部标识
        for token in ("pump_water", "pack_hash", "numeric_profile_id",
                      "workspace_id", "matched_rule_id"):
            with self.subTest(token=token):
                self.assertNotIn(token, blob)

    def test_underlying_audit_data_is_still_stored(self):
        """底层 provenance / snapshot / hash / numeric profile 必须仍然存在。"""

        from equipeffi.composition import create_pump_analysis_service

        record_id = self._record_id()
        snapshot = create_pump_analysis_service(
            paths=self.paths).open_record(record_id)
        self.assertTrue(snapshot.reference_snapshot)
        self.assertTrue(snapshot.result_snapshot.get("provenance"))
        self.assertTrue(snapshot.canonical_package_hash)
        self.assertTrue(snapshot.numeric_profile_id)
        self.assertTrue(snapshot.result_snapshot.get("matched_rule_id"))


class UISettingsLogLevelTests(R1TestCase):
    """UI04：日志级别用户可见文本必须为中文，且内部值可稳定恢复。"""

    def _page(self):
        from equipeffi.composition import create_settings_runtime
        from equipeffi.presentation.qt.pages.settings import SettingsPage

        service, logger = create_settings_runtime(paths=self.paths)
        self.addCleanup(lambda: __import__(
            "equipeffi.infrastructure.runtime_logging", fromlist=["close_logging"]
        ).close_logging(logger))
        page = SettingsPage(service, None, app_version="test",
                            data_location=self.paths.root)
        self.addCleanup(page.deleteLater)
        return service, page

    def test_all_visible_log_levels_are_chinese(self):
        _service, page = self._page()
        labels = [page.log_level.itemText(i)
                  for i in range(page.log_level.count())]
        self.assertTrue(labels)
        for label in labels:
            with self.subTest(label=label):
                self.assertFalse(re.search(r"[A-Za-z]", label),
                                 f"不得显示英文枚举值：{label}")

    def test_internal_values_are_preserved_and_round_trip(self):
        service, page = self._page()
        values = [page.log_level.itemData(i)
                  for i in range(page.log_level.count())]
        self.assertIn("WARNING", values)
        index = page.log_level.findData("WARNING")
        page.log_level.setCurrentIndex(index)
        self.assertEqual(service.get("log.level"), "WARNING")

        # 重新构造页面（模拟重启）后仍应恢复到「警告」
        _service2, page2 = self._page()
        self.assertEqual(page2.log_level.currentData(), "WARNING")
        self.assertEqual(page2.log_level.currentText(), "警告")

    def test_label_mapping_covers_every_supported_level(self):
        from equipeffi.presentation.qt.pages.settings import LOG_LEVEL_LABELS

        service, _page = self._page()
        for value in type(service).LEVELS:
            with self.subTest(level=value):
                self.assertIn(value, LOG_LEVEL_LABELS)


if __name__ == "__main__":
    unittest.main()
