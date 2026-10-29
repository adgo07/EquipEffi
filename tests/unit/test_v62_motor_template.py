from __future__ import annotations

import hashlib
import posixpath
from pathlib import Path
import runpy
import tempfile
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from equipeffi.application.ports.batch_workbook import BatchRowOutcome
from equipeffi.domain.devices.motor.input_contract import (
    BASIS_OPTIONS,
    MOTOR_CATEGORIES,
    scope_hint,
    efficiency_basis_hint,
    legacy_visible_fields,
)
from equipeffi.domain.devices.motor.calculator import Gb18613MotorCalculator
from equipeffi.infrastructure.excel.ooxml_reader import OOXMLReadError, OOXMLWorkbook
from equipeffi.infrastructure.excel.pump_result_writer import _sheet_target
from equipeffi.infrastructure.excel.template_resource import (
    DEFAULT_V61_TEMPLATE,
    DEFAULT_V62_TEMPLATE,
    MOTOR_SHEET_CONTRACT,
    TEMPLATE_MANIFEST,
    V61_TEMPLATE_ENTRY,
    V62_TEMPLATE_ENTRY,
    V62_TEMPLATE_IDENTITY,
    V61TransformerTemplateResource,
    V62MotorTemplateResource,
    identity_defined_names,
    validate_v6_template,
)
from equipeffi.infrastructure.excel.unified_result_writer import (
    LEGAL_MOTOR_CONCLUSIONS,
    MOTOR_RESULT_COLUMNS,
    ResultWorkbookWriteError,
    UnifiedResultWorkbookWriter,
)
from equipeffi.infrastructure.excel.unified_workbook_reader import UnifiedWorkbookReader
from equipeffi.infrastructure.standards.motor_repository import load_motor_pack


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE_DIR = ROOT / "src" / "equipeffi" / "resources" / "templates"
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _column_index(column: str) -> int:
    result = 0
    for char in column:
        result = result * 26 + ord(char.upper()) - ord("A") + 1
    return result - 1


def _column_name(index: int) -> str:
    result = ""
    number = index + 1
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _sheet_table_target(archive: ZipFile, sheet_name: str) -> str:
    sheet = _sheet_target(archive, sheet_name)
    relation_path = posixpath.join(
        posixpath.dirname(sheet), "_rels", posixpath.basename(sheet) + ".rels")
    root = ET.fromstring(archive.read(relation_path))
    for relation in root.findall(f"{{{REL_NS}}}Relationship"):
        if str(relation.attrib.get("Type", "")).endswith("/table"):
            target = str(relation.attrib["Target"])
            if target.startswith("/"):
                return posixpath.normpath(target.lstrip("/"))
            return posixpath.normpath(posixpath.join(posixpath.dirname(sheet), target))
    raise AssertionError(f"找不到{sheet_name}的Table关系")


def _set_cell(root: ET.Element, reference: str, value) -> None:
    sheet_data = root.find(f"{{{MAIN_NS}}}sheetData")
    assert sheet_data is not None
    row_number = int("".join(char for char in reference if char.isdigit()))
    row = next((item for item in sheet_data.findall(f"{{{MAIN_NS}}}row")
                if item.attrib.get("r") == str(row_number)), None)
    if row is None:
        row = ET.Element(f"{{{MAIN_NS}}}row", {"r": str(row_number)})
        existing = list(sheet_data.findall(f"{{{MAIN_NS}}}row"))
        later = next((item for item in existing
                      if int(item.attrib.get("r", "0")) > row_number), None)
        sheet_data.insert(list(sheet_data).index(later), row) if later is not None else sheet_data.append(row)
    cell = next((item for item in row.findall(f"{{{MAIN_NS}}}c")
                 if item.attrib.get("r") == reference), None)
    if cell is None:
        cell = ET.Element(f"{{{MAIN_NS}}}c", {"r": reference})
        row.append(cell)
    for child in list(cell):
        cell.remove(child)
    if value is None:
        cell.attrib.pop("t", None)
    elif isinstance(value, tuple) and value[0] == "formula":
        cell.attrib.pop("t", None)
        ET.SubElement(cell, f"{{{MAIN_NS}}}f").text = str(value[1])
        ET.SubElement(cell, f"{{{MAIN_NS}}}v")
    elif isinstance(value, tuple) and value[0] == "error":
        cell.set("t", "e")
        ET.SubElement(cell, f"{{{MAIN_NS}}}v").text = str(value[1])
    elif isinstance(value, str):
        cell.set("t", "inlineStr")
        inline = ET.SubElement(cell, f"{{{MAIN_NS}}}is")
        ET.SubElement(inline, f"{{{MAIN_NS}}}t").text = value
    else:
        cell.attrib.pop("t", None)
        ET.SubElement(cell, f"{{{MAIN_NS}}}v").text = str(value)
    cells = list(row.findall(f"{{{MAIN_NS}}}c"))
    for child in cells:
        row.remove(child)
    cells.sort(key=lambda item: _column_index(
        "".join(char for char in item.attrib["r"] if char.isalpha())))
    row.extend(cells)


def _patch_motor_sheet(source: Path, destination: Path, values: dict[str, object],
                       *, capacity: int | None = None) -> Path:
    ET.register_namespace("", MAIN_NS)
    with ZipFile(source) as original:
        motor_part = _sheet_target(original, "电动机")
        entries = [(info, original.read(info.filename)) for info in original.infolist()]
    root = ET.fromstring(dict((info.filename, data) for info, data in entries)[motor_part])
    for reference, value in values.items():
        _set_cell(root, reference, value)
    if capacity is not None:
        last_row = capacity + 3
        dimension = root.find(f"{{{MAIN_NS}}}dimension")
        if dimension is not None:
            dimension.set("ref", f"A1:T{last_row}")
        sheet_data = root.find(f"{{{MAIN_NS}}}sheetData")
        assert sheet_data is not None
        rows_by_number = {
            int(row.attrib["r"]): row
            for row in sheet_data.findall(f"{{{MAIN_NS}}}row")
            if row.attrib.get("r", "").isdigit()
        }
        for index in range(capacity):
            row_number = index + 4
            row = rows_by_number.get(row_number)
            if row is None:
                row = ET.Element(f"{{{MAIN_NS}}}row", {"r": str(row_number)})
                rows_by_number[row_number] = row
                sheet_data.append(row)
            cell = next((item for item in row.findall(f"{{{MAIN_NS}}}c")
                         if item.attrib.get("r") == f"B{row_number}"), None)
            if cell is None:
                cell = ET.SubElement(row, f"{{{MAIN_NS}}}c", {"r": f"B{row_number}"})
            for child in list(cell):
                cell.remove(child)
            cell.set("t", "inlineStr")
            inline = ET.SubElement(cell, f"{{{MAIN_NS}}}is")
            ET.SubElement(inline, f"{{{MAIN_NS}}}t").text = f"motor-{index + 1}"
        for row in list(sheet_data.findall(f"{{{MAIN_NS}}}row")):
            sheet_data.remove(row)
        sheet_data.extend(rows_by_number[number] for number in sorted(rows_by_number))
        for row in rows_by_number.values():
            cells = list(row.findall(f"{{{MAIN_NS}}}c"))
            for cell in cells:
                row.remove(cell)
            cells.sort(key=lambda item: _column_index(
                "".join(char for char in item.attrib["r"] if char.isalpha())))
            row.extend(cells)
    content = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(destination, "w", ZIP_DEFLATED) as patched:
        for info, data in entries:
            patched.writestr(info, content if info.filename == motor_part else data)
    return destination


def _defined_names(path: Path) -> dict[str, tuple[str, ...]]:
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read("xl/workbook.xml"))
    names: dict[str, list[str]] = {}
    for item in root.findall(f"{{{MAIN_NS}}}definedNames/{{{MAIN_NS}}}definedName"):
        names.setdefault(item.attrib.get("name", ""), []).append(item.text or "")
    return {name: tuple(values) for name, values in names.items()}


def _patch_motor_sheet_revision(source: Path, destination: Path, revision: str) -> Path:
    with ZipFile(source) as original:
        entries = [(info, original.read(info.filename)) for info in original.infolist()]
    workbook_xml = "xl/workbook.xml"
    root = ET.fromstring(dict((info.filename, data) for info, data in entries)[workbook_xml])
    marker = TEMPLATE_MANIFEST["identity_markers"]["sheet_contract_prefix"] + "motor"
    node = next(item for item in root.findall(f"{{{MAIN_NS}}}definedNames/{{{MAIN_NS}}}definedName")
                if item.attrib.get("name") == marker)
    node.text = f'"{revision}"'
    content = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with ZipFile(destination, "w", ZIP_DEFLATED) as patched:
        for info, data in entries:
            patched.writestr(info, content if info.filename == workbook_xml else data)
    return destination


def _config_values(path: Path) -> dict[tuple[int, int], object]:
    with OOXMLWorkbook(path) as workbook:
        grid = workbook.rows("配置")
    return {
        (row_number + 1, column_index): value
        for row_number, row in enumerate(grid)
        for column_index, value in enumerate(row)
        if value not in (None, "")
        and not (_column_index("EW") <= column_index <= _column_index("FH"))
    }


class V62MotorTemplateTests(unittest.TestCase):
    def setUp(self):
        self._temporary = tempfile.TemporaryDirectory(prefix="equipeffi-v62-motor-")
        self.addCleanup(self._temporary.cleanup)
        self.temp_dir = Path(self._temporary.name)
        self.template = TEMPLATE_DIR / DEFAULT_V62_TEMPLATE

    def test_manifest_resource_and_domain_contract_are_mechanically_aligned(self):
        self.assertEqual(V62_TEMPLATE_IDENTITY["template_version"], "V6.2")
        self.assertEqual(V62_TEMPLATE_IDENTITY["template_id"],
                         "equipeffi.device-efficiency.V6.2")
        self.assertEqual(V62_TEMPLATE_ENTRY["template_contract_revision"], "unified-v6.v3")
        motor_entry = next(sheet for sheet in V62_TEMPLATE_ENTRY["sheets"]
                           if sheet["sheet_id"] == "motor")
        self.assertEqual(motor_entry["contract_revision"], "motor-input.v2")
        self.assertEqual(V62_TEMPLATE_ENTRY["input_compatibility"], "ACCEPTED")
        self.assertEqual(V62_TEMPLATE_ENTRY["status"], "CANDIDATE")
        digest = hashlib.sha256(self.template.read_bytes()).hexdigest().upper()
        self.assertEqual(digest, V62_TEMPLATE_ENTRY["asset_sha256"])
        self.assertEqual(self.template.stat().st_size, V62_TEMPLATE_ENTRY["asset_size_bytes"])
        validation = V62MotorTemplateResource().validate()
        self.assertTrue(validation.is_valid, validation.message)
        self.assertTrue(validate_v6_template(self.template).is_valid)
        self.assertFalse(V61TransformerTemplateResource().validate(
            self.template).is_valid)
        self.assertEqual(MOTOR_SHEET_CONTRACT["table"],
                         {"name": "TblMotor", "ref": "A3:T103"})

        motor = next(sheet for sheet in V62_TEMPLATE_ENTRY["sheets"]
                     if sheet["sheet_id"] == "motor")
        self.assertEqual(motor["category_ids"],
                         [item.category_id for item in MOTOR_CATEGORIES])
        self.assertEqual(len(motor["categories"]), 8)
        field_map = {field["field_id"]: field for field in motor["fields"]
                     if field.get("field_id")}
        for category in MOTOR_CATEGORIES:
            visible = {field.field_id for field in legacy_visible_fields(category.category_id)}
            self.assertEqual(
                category.category_id in field_map["rated_voltage_kv"]["required_for_categories"],
                "rated_voltage_kv" in visible)
            self.assertEqual(
                category.category_id in field_map["rated_frequency_hz"]["required_for_categories"],
                "rated_frequency_hz" in visible)
            self.assertEqual(
                category.category_id in field_map["poles"]["required_for_categories"],
                "poles" in visible)
            self.assertEqual(
                category.category_id in field_map["max_nameplate_speed_rpm"]["required_for_categories"],
                "max_nameplate_speed_rpm" in visible)
            if category.table_no is not None:
                self.assertTrue(scope_hint(category.category_id))
                self.assertTrue(efficiency_basis_hint(category.category_id))
        self.assertFalse(field_map["poles"].get("defined_name_prefix"))
        self.assertEqual(field_map["poles"]["excel_validation"]["type"], "custom")
        self.assertEqual(field_map["rated_frequency_hz"]["excel_validation"]["type"], "custom")
        self.assertEqual(field_map["efficiency_basis_status"]["defined_name"],
                         "EV_MOTOR_V62_BASIS")
        self.assertEqual([label for _code, label in BASIS_OPTIONS],
                         ["按上述口径填报", "仅有铭牌效率", "已知试验工况不符合",
                          "驱动不可分离但仅有电机本体效率（仅表6）", "效率口径不确定"])

    def test_legacy_v62_builder_replays_registered_asset_after_domain_v3(self):
        builder = runpy.run_path(str(ROOT / "tools" / "build_v62_motor_template.py"))
        with tempfile.TemporaryDirectory(prefix="equipeffi-v62-replay-") as directory:
            digest, size = builder["build"](
                output=Path(directory) / DEFAULT_V62_TEMPLATE,
                update_manifest=False)
        self.assertEqual(digest, V62_TEMPLATE_ENTRY["asset_sha256"])
        self.assertEqual(size, V62_TEMPLATE_ENTRY["asset_size_bytes"])

    def test_frequency_and_poles_use_free_entry_controls(self):
        with ZipFile(self.template) as archive:
            motor_part = _sheet_target(archive, "电动机")
            root = ET.fromstring(archive.read(motor_part))
        validations = {}
        for node in root.findall(f"{{{MAIN_NS}}}dataValidations/{{{MAIN_NS}}}dataValidation"):
            reference = node.attrib.get("sqref", "")
            column = reference.split("4:", 1)[0]
            validations[column] = node

        frequency = validations["H"]
        self.assertEqual(frequency.attrib["type"], "custom")
        frequency_formula = frequency.find(f"{{{MAIN_NS}}}formula1").text
        self.assertIn("ISNUMBER(H4)", frequency_formula)
        self.assertIn("ISTEXT(H4)", frequency_formula)

        poles = validations["K"]
        self.assertEqual(poles.attrib["type"], "custom")
        self.assertEqual(
            poles.find(f"{{{MAIN_NS}}}formula1").text,
            'OR(K4="",AND(ISNUMBER(K4),K4>0,MOD(K4,2)=0))')
        self.assertFalse(any(
            name.startswith("EV_MOTOR_V62_POLES_")
            for name in _defined_names(self.template)))

    def test_previous_motor_input_v1_candidate_is_rejected_before_reading(self):
        old_candidate = self.temp_dir / "motor-input-v1-candidate.xlsx"
        _patch_motor_sheet_revision(self.template, old_candidate, "motor-input.v1")

        validation = V62MotorTemplateResource().validate(old_candidate)
        self.assertFalse(validation.is_valid)
        self.assertIn("重新输出空白模板并复制输入数据", validation.message)
        self.assertIn("旧文件不会被修改", validation.message)
        self.assertIn("motor-input.v1", validation.message)
        self.assertIn("motor-input.v2", validation.message)
        with self.assertRaisesRegex(OOXMLReadError, "motor-input.v1"):
            UnifiedWorkbookReader().read(old_candidate)

    def test_builder_preserves_pump_transformer_and_shared_configuration(self):
        v61_path = TEMPLATE_DIR / DEFAULT_V61_TEMPLATE
        with ZipFile(v61_path) as before, ZipFile(self.template) as after:
            self.assertEqual(before.namelist(), after.namelist())
            pump_part = _sheet_target(before, "离心泵")
            transformer_part = _sheet_target(before, "变压器")
            motor_part = _sheet_target(before, "电动机")
            config_part = _sheet_target(before, "配置")
            instructions_part = _sheet_target(before, "注意事项")
            template_notes_part = _sheet_target(before, "模板说明")
            table_part = _sheet_table_target(before, "电动机")
            self.assertEqual(before.read(pump_part), after.read(pump_part))
            self.assertEqual(before.read(transformer_part), after.read(transformer_part))
            allowed = {motor_part, config_part, instructions_part,
                       template_notes_part, table_part,
                       "xl/styles.xml", "xl/workbook.xml"}
            changed = {name for name in before.namelist()
                       if before.read(name) != after.read(name)}
            self.assertEqual(changed, allowed)

        before_names, after_names = _defined_names(v61_path), _defined_names(self.template)
        old_identity = set(identity_defined_names(V61_TEMPLATE_ENTRY))
        new_identity = set(identity_defined_names(V62_TEMPLATE_ENTRY))
        legacy_motor_names = {
            "EV_MOTOR_CATEGORY", "EV_POLES_LV", "EV_POLES_HV", "EV_POLES_PMSM",
            "EV_MOTOR_VOLTAGE", "EV_COOLING_MOTOR_HV", "EV_COOLING_PMSM",
        }
        keep = lambda name: not (
            name in old_identity | new_identity
            or name in legacy_motor_names
            or name.startswith("EV_MOTOR_") or name.startswith("EV_POLES_"))
        self.assertEqual({name: values for name, values in before_names.items() if keep(name)},
                         {name: values for name, values in after_names.items() if keep(name)})
        for name in ("EV_PUMP_CATEGORY", "EV_TRANSFORMER_CATEGORY"):
            self.assertEqual(before_names[name], after_names[name])
        self.assertEqual(_config_values(v61_path), _config_values(self.template))
        with OOXMLWorkbook(v61_path) as before_book, OOXMLWorkbook(self.template) as after_book:
            before_instructions = before_book.rows("注意事项")
            after_instructions = after_book.rows("注意事项")
            after_notes = after_book.rows("模板说明")
        self.assertEqual(before_instructions[20][1], after_instructions[20][1])
        self.assertIn("GB/T 1032 B法", after_notes[20][1])
        self.assertIn("单速、封闭自扇冷式、N设计、连续工作制",
                      after_notes[20][1])
        self.assertNotIn("三类标准合并到一张表", after_notes[20][1])

        with ZipFile(self.template) as archive:
            motor_part = _sheet_target(archive, "电动机")
            table_part = _sheet_table_target(archive, "电动机")
            root = ET.fromstring(archive.read(motor_part))
            table = ET.fromstring(archive.read(table_part))
        cells = {cell.attrib.get("r"): cell for cell in root.iter(f"{{{MAIN_NS}}}c")}
        self.assertEqual(len([cell for cell in cells if cell.startswith(tuple("ABCDEFGHIJKLMNOPQRST"))]),
                         len(cells))
        self.assertEqual(len([cell for cell in root.iter(f"{{{MAIN_NS}}}c")
                              if cell.find(f"{{{MAIN_NS}}}f") is not None]), 101)
        self.assertFalse(any(
            cell.attrib.get("r", "").startswith("T")
            and cell.find(f"{{{MAIN_NS}}}f") is not None
            for cell in root.iter(f"{{{MAIN_NS}}}c")))
        self.assertEqual(table.attrib.get("ref"), "A3:T103")
        self.assertEqual(MOTOR_RESULT_COLUMNS, ("N", "O", "P", "Q", "T"))
        self.assertEqual(LEGAL_MOTOR_CONCLUSIONS,
                         {"1级", "2级", "3级", "未达3级", "不在范围", "无法判定"})

    def test_reader_preserves_frequency_text_and_arbitrary_poles_and_rejects_formula_input(self):
        category = MOTOR_CATEGORIES[0]
        table_five = MOTOR_CATEGORIES[4]
        source = self.temp_dir / "motor-input.xlsx"
        _patch_motor_sheet(self.template, source, {
            "B4": "风机电机", "C4": "M-750", "D4": 1, "E4": "泵房",
            "F4": category.label, "G4": 0.4, "H4": "5~120",
            "I4": BASIS_OPTIONS[0][1], "J4": 0.75,
            "K4": 10, "M4": 90.25, "R4": "铭牌照片", "S4": "备注",
            "B5": "输入含公式", "F5": category.label,
            "H5": ("formula", '"50"'), "J5": 0.75,
            "B6": "表5留痕", "F6": table_five.label,
            "H6": "50 Hz 铭牌文字", "K6": 12,
        })
        result = UnifiedWorkbookReader().read(source)
        motor = result.sheets["电动机"]
        self.assertEqual(result.template_identity["template_version"], "V6.2")
        self.assertEqual(len(motor.rows), 3)
        first, second, third = motor.rows
        self.assertEqual(first.row_number, 4)
        self.assertEqual(first.values["category_id"], category.category_id)
        self.assertEqual(first.values["rated_frequency_hz"], "5~120")
        self.assertEqual(first.values["poles"], 10)
        self.assertEqual(first.values["efficiency_basis_status"], BASIS_OPTIONS[0][1])
        self.assertEqual(first.values["efficiency_basis_raw"], BASIS_OPTIONS[0][1])
        self.assertEqual(first.values["measured_efficiency_pct"], 90.25)
        self.assertEqual(second.values["rated_frequency_hz"], "")
        self.assertTrue(any("H5" in message and "公式" in message
                            for message in second.input_errors))
        self.assertEqual(third.values["rated_frequency_hz"], "50 Hz 铭牌文字")
        self.assertEqual(third.values["poles"], 12)
        self.assertNotIn("电动机", result.unprocessed_sheets)
        self.assertIn("空压机", result.unprocessed_sheets)

    def test_writer_writes_threshold_states_fixed_summary_and_can_reimport_twice(self):
        category = MOTOR_CATEGORIES[0]
        source = self.temp_dir / "motor-source.xlsx"
        _patch_motor_sheet(self.template, source, {
            "B4": "电机A", "D4": 1, "F4": category.label,
            "G4": 0.4, "H4": 50, "J4": 0.75, "K4": "4极", "M4": 92.25,
            "N4": 999, "O4": 888, "P4": 777, "Q4": "旧等级", "T4": "旧说明",
        })
        outcome = BatchRowOutcome(
            4, True, "1级", "SUCCESS", "1级", messages=("效率达到1级。",),
            thresholds={
                "grade_1": {"state": "AVAILABLE", "display_text": "92.5"},
                "grade_2": {"state": "NO_INDICATOR", "display_text": "—"},
                "grade_3": {"state": "NOT_COMPUTED", "display_text": None},
            })
        summary = {
            "sheets": [{"sheet_name": "电动机", "data_row_count": 1}],
            "unprocessed_sheets": ["空压机"],
            "data_row_count": 1, "total_quantity": 1, "evaluated_quantity": 1,
            "concluded_rows": 1, "unevaluated_rows": 0,
            "input_error_rows": 0, "execution_error_rows": 0,
        }
        first_result = self.temp_dir / "motor-result.xlsx"
        second_result = self.temp_dir / "motor-result-again.xlsx"
        writer = UnifiedResultWorkbookWriter()
        writer.write(source, {"电动机": {4: outcome}}, first_result,
                     batch_summary=summary)
        first_validation = V62MotorTemplateResource().validate(first_result)
        self.assertTrue(first_validation.is_valid, first_validation.message)
        with OOXMLWorkbook(first_result) as result:
            rows = result.rows("电动机")
            notes = result.rows("模板说明")
        self.assertEqual(rows[3][13], 92.5)
        self.assertEqual(rows[3][14], "—")
        self.assertEqual(rows[3][15], "")
        self.assertEqual(rows[3][16], "1级")
        self.assertEqual(rows[3][19], "效率达到1级。")
        self.assertEqual(notes[26][0], "本次批量处理说明")
        self.assertIn("未覆盖", notes[26][1])
        self.assertEqual(notes[27][1], "V6.2")
        self.assertEqual(notes[29][1], "电动机")
        self.assertEqual(notes[30][1], "空压机")
        self.assertIn("执行失败 0 行", notes[31][1])
        self.assertEqual(notes[33][1], "EQUIPEFFI_BATCH_OUTPUT_V1")

        with ZipFile(source) as before, ZipFile(first_result) as after:
            motor_part = _sheet_target(before, "电动机")
            notes_part = _sheet_target(before, "模板说明")
            transformer_part = _sheet_target(before, "变压器")
            pump_part = _sheet_target(before, "离心泵")
            self.assertEqual(before.read(transformer_part), after.read(transformer_part))
            self.assertEqual(before.read(pump_part), after.read(pump_part))
            for name in before.namelist():
                if name not in {motor_part, notes_part}:
                    self.assertEqual(before.read(name), after.read(name), name)

        # The software marker makes this fixed area safe for another run, and
        # the output workbook remains a valid V6.2 input workbook.
        writer.write(first_result, {"电动机": {4: outcome}}, second_result,
                     batch_summary=summary)
        self.assertTrue(V62MotorTemplateResource().validate(second_result).is_valid)
        reimported = UnifiedWorkbookReader().read(second_result)
        self.assertEqual(reimported.sheets["电动机"].rows[0].values["category_id"],
                         category.category_id)

    def test_writer_rejects_user_content_in_fixed_area_and_old_template_motor_write(self):
        user_owned = self.temp_dir / "user-owned-area.xlsx"
        _patch_motor_sheet(self.template, user_owned, {})
        with ZipFile(user_owned) as source_zip:
            notes_part = _sheet_target(source_zip, "模板说明")
            entries = [(info, source_zip.read(info.filename))
                       for info in source_zip.infolist()]
        notes_root = ET.fromstring(dict((info.filename, data) for info, data in entries)[notes_part])
        _set_cell(notes_root, "B28", "用户自填")
        with ZipFile(user_owned, "w", ZIP_DEFLATED) as target:
            for info, data in entries:
                target.writestr(info, ET.tostring(notes_root, encoding="utf-8", xml_declaration=True)
                                if info.filename == notes_part else data)
        destination = self.temp_dir / "should-not-exist.xlsx"
        with self.assertRaisesRegex(ResultWorkbookWriteError, "拒绝覆盖用户内容"):
            UnifiedResultWorkbookWriter().write(
                user_owned, {}, destination, batch_summary={})
        self.assertFalse(destination.exists())

        old_source = self.temp_dir / "old-template.xlsx"
        old_source.write_bytes((TEMPLATE_DIR / DEFAULT_V61_TEMPLATE).read_bytes())
        old_read = UnifiedWorkbookReader().read(old_source)
        self.assertNotIn("电动机", old_read.sheets)
        self.assertIn("电动机", old_read.unprocessed_sheets)
        old_dest = self.temp_dir / "old-template-motor-result.xlsx"
        with self.assertRaisesRegex(ResultWorkbookWriteError, "尚未接通"):
            UnifiedResultWorkbookWriter().write(
                old_source,
                {"电动机": {4: BatchRowOutcome(4, True, "1级", "SUCCESS", "1级")}},
                old_dest)
        self.assertFalse(old_dest.exists())

    def test_real_014_kw_four_pole_interpolation_displays_two_places(self):
        category = MOTOR_CATEGORIES[0]
        pack, _digest = load_motor_pack()
        calculator = Gb18613MotorCalculator()
        interpolated = calculator.evaluate({
            "category_id": category.category_id,
            "rated_output_power_kw": "0.14",
            "rated_voltage_kv": "0.38",
            "rated_frequency_hz": "50",
            "poles": "4",
            "measured_efficiency_pct": "71.43",
        }, pack)
        self.assertEqual(interpolated.thresholds["grade_3"]["display_text"], "66.50")
        self.assertEqual(interpolated.thresholds["grade_3"]["method"],
                         "LINEAR_INTERPOLATION")

        direct = calculator.evaluate({
            "category_id": category.category_id,
            "rated_output_power_kw": "0.12",
            "rated_voltage_kv": "0.38",
            "rated_frequency_hz": "50",
            "poles": "4",
            "measured_efficiency_pct": "70",
        }, pack)
        self.assertTrue(all(value["method"] == "DIRECT"
                            for value in direct.thresholds.values()))

        source = self.temp_dir / "motor-format-source.xlsx"
        _patch_motor_sheet(self.template, source, {
            "B4": "插值工况", "D4": 1, "F4": category.label,
            "G4": 0.38, "H4": 50, "J4": 0.14, "K4": "4极", "M4": 71.43,
            "B5": "直查工况", "D5": 1, "F5": category.label,
            "G5": 0.38, "H5": 50, "J5": 0.12, "K5": "4极", "M5": 70,
        })
        destination = self.temp_dir / "motor-format-result.xlsx"
        outcomes = {
            4: BatchRowOutcome(
                4, True, interpolated.ui_conclusion,
                interpolated.evaluation_status, interpolated.grade,
                messages=(interpolated.explanation,), thresholds=interpolated.thresholds),
            5: BatchRowOutcome(
                5, True, direct.ui_conclusion,
                direct.evaluation_status, direct.grade,
                messages=(direct.explanation,), thresholds=direct.thresholds),
        }
        UnifiedResultWorkbookWriter().write(source, {"电动机": outcomes}, destination)

        from openpyxl import load_workbook
        workbook = load_workbook(destination, data_only=False, read_only=False)
        sheet = workbook["电动机"]
        self.assertEqual(sheet["P4"].value, 66.5)
        self.assertEqual(sheet["P4"].number_format, "0.00")
        self.assertEqual(sheet["N4"].number_format, "0.00")
        self.assertEqual(sheet["N5"].number_format, "0.###############")
        self.assertEqual(sheet["N5"].value,
                         float(direct.thresholds["grade_1"]["display_text"]))
        workbook.close()

        # Reimport and reuse the same row through interpolation -> direct ->
        # interpolation. Each run must choose the format from method, not keep
        # the previous output cell's number format.
        direct_source = self.temp_dir / "motor-format-direct-source.xlsx"
        _patch_motor_sheet(destination, direct_source, {
            "J4": 0.12, "M4": 70, "J5": 0.14, "M5": 71.43})
        direct_result = self.temp_dir / "motor-format-direct-result.xlsx"
        UnifiedResultWorkbookWriter().write(
            direct_source, {"电动机": {4: outcomes[5], 5: outcomes[4]}}, direct_result)
        workbook = load_workbook(direct_result, data_only=False, read_only=False)
        self.assertEqual(workbook["电动机"]["P4"].number_format,
                         "0.###############")
        self.assertEqual(workbook["电动机"]["P5"].value, 66.5)
        self.assertEqual(workbook["电动机"]["P5"].number_format, "0.00")
        workbook.close()
        with ZipFile(direct_source) as source_zip, ZipFile(direct_result) as result_zip:
            self.assertEqual(source_zip.read("xl/styles.xml"),
                             result_zip.read("xl/styles.xml"))

        interpolate_again_source = self.temp_dir / "motor-format-interpolate-source.xlsx"
        _patch_motor_sheet(direct_result, interpolate_again_source,
                           {"J4": 0.14, "M4": 71.43, "J5": 0.12, "M5": 70})
        interpolate_again = self.temp_dir / "motor-format-interpolate-result.xlsx"
        UnifiedResultWorkbookWriter().write(
            interpolate_again_source,
            {"电动机": {4: outcomes[4], 5: outcomes[5]}}, interpolate_again)
        workbook = load_workbook(interpolate_again, data_only=False, read_only=False)
        self.assertEqual(workbook["电动机"]["P4"].value, 66.5)
        self.assertEqual(workbook["电动机"]["P4"].number_format, "0.00")
        self.assertEqual(workbook["电动机"]["P5"].number_format,
                         "0.###############")
        workbook.close()
        with ZipFile(interpolate_again_source) as source_zip, ZipFile(interpolate_again) as result_zip:
            self.assertEqual(source_zip.read("xl/styles.xml"),
                             result_zip.read("xl/styles.xml"))

    def test_reader_writer_have_no_software_row_cap_at_100_1000_10000(self):
        # These synthetic OOXML inputs prove the adapter has no fixed 100-row
        # boundary. Excel/WPS table auto-expansion still requires manual engine
        # evidence and is deliberately reported separately as NOT_RUN.
        writer = UnifiedResultWorkbookWriter()
        for capacity in (100, 1_000, 10_000):
            with self.subTest(capacity=capacity):
                source = self.temp_dir / f"motor-{capacity}.xlsx"
                result = self.temp_dir / f"motor-{capacity}-result.xlsx"
                _patch_motor_sheet(self.template, source, {}, capacity=capacity)
                read = UnifiedWorkbookReader().read(source)
                rows = read.sheets["电动机"].rows
                self.assertEqual(len(rows), capacity)
                last_row = capacity + 3
                outcome = BatchRowOutcome(
                    last_row, True, "1级", "SUCCESS", "1级",
                    thresholds={"grade_1": {"state": "AVAILABLE", "display_text": "88.125"}})
                writer.write(
                    source, {"电动机": {last_row: outcome}}, result,
                    batch_summary={
                        "sheets": [{"sheet_name": "电动机", "data_row_count": capacity}],
                        "unprocessed_sheets": [], "data_row_count": capacity,
                        "total_quantity": capacity, "evaluated_quantity": capacity,
                        "concluded_rows": capacity, "unevaluated_rows": 0,
                        "input_error_rows": 0, "execution_error_rows": 0,
                    })
                with OOXMLWorkbook(result) as workbook:
                    output = workbook.rows("电动机")
                self.assertEqual(output[last_row - 1][13], 88.125)
                self.assertEqual(output[last_row - 1][16], "1级")


if __name__ == "__main__":
    unittest.main()
