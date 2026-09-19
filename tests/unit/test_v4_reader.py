from __future__ import annotations

from pathlib import Path
import shutil
import tempfile
import unittest
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl

from equipeffi.infrastructure.excel.strict_importer import StrictTemplateImporter
from equipeffi.infrastructure.excel.v4_reader import V4WorkbookReaderImpl
from equipeffi.application.services.evaluation_facade import EvaluationFacade
from equipeffi.application.services.evaluation_service import EvaluationService
from equipeffi.infrastructure.standards.json_repository import JsonStandardRepository


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "src" / "equipeffi" / "resources" / "templates" / "设备能效分析空白模板_重构版V4_20260825.xlsx"


class V4ReaderTests(unittest.TestCase):
    def test_reads_full_config_contract_from_bundled_v4(self):
        reader = V4WorkbookReaderImpl()
        contract = reader.read_contract(TEMPLATE)
        self.assertEqual(contract.validate(), [])
        self.assertGreaterEqual(len(contract.fields), 300)
        motor_fields = {field.field_id for field in contract.fields_for_sheet("电动机", editable_only=True)}
        self.assertIn("rated_power", motor_fields)
        self.assertIn("efficiency", motor_fields)
        self.assertIn("EV_MOTOR_CATEGORY", contract.enums)
        self.assertIn("变频调速永磁同步电动机", contract.enums["EV_MOTOR_CATEGORY"])
        cooling = set(contract.enums["EV_COOLING_MOTOR_HV"])
        self.assertTrue({"IC86W", "IC71W(IC3W7)", "IC416", "IC666"}.issubset(cooling))

    def test_reads_workbook_elimination_scope_setting(self):
        settings = V4WorkbookReaderImpl().read_settings(TEMPLATE)
        self.assertEqual(settings["elimination_scope"], "高耗能落后机电设备淘汰目录第一至第四批")

    def test_blank_v4_template_has_no_device_rows(self):
        rows = V4WorkbookReaderImpl().read_rows(TEMPLATE)
        self.assertEqual(rows, ())

    def test_reader_uses_existing_technical_record_id_when_present(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "带技术ID.xlsx"
            shutil.copy2(TEMPLATE, source)
            workbook = openpyxl.load_workbook(source)
            sheet = workbook["电动机"]
            column = next(cell.column for cell in sheet[3] if cell.value == "技术记录ID")
            # 输入最小可识别字段，验证ID读取与普通数据行过滤同时成立。
            for column_index, value in {
                2: "电机1", 3: "M1", 4: 1, 6: "三相异步电动机（一般用途）",
                7: 0.4, 10: 7.5, 11: 4, 12: 1480, 13: 98, column: "TECH-001",
            }.items():
                sheet.cell(4, column_index).value = value
            workbook.save(source)
            rows = V4WorkbookReaderImpl().read_rows(source)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0].record_id, "TECH-001")

    def test_legacy_strict_importer_name_uses_v4_reader(self):
        importer = StrictTemplateImporter()
        self.assertEqual(importer.import_records(str(TEMPLATE)), ())

    def test_copied_v4_row_is_read_and_evaluated(self):
        # 只在临时副本中增加一行，不改动软件内置V4模板。
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "输入.xlsx"
            shutil.copy2(TEMPLATE, source)
            workbook = V4WorkbookReaderImpl()
            from equipeffi.infrastructure.excel.ooxml_reader import OOXMLWorkbook
            with OOXMLWorkbook(source) as package:
                sheet_target = package._sheets["电动机"]
            with ZipFile(source, "r") as archive:
                files = {name: archive.read(name) for name in archive.namelist()}
            cells = [
                '<c r="A4"><v>1</v></c>',
                '<c r="B4" t="inlineStr"><is><t>电机1</t></is></c>',
                '<c r="C4" t="inlineStr"><is><t>M1</t></is></c>',
                '<c r="D4"><v>1</v></c>',
                '<c r="F4" t="inlineStr"><is><t>三相异步电动机（一般用途）</t></is></c>',
                '<c r="G4"><v>0.4</v></c>',
                '<c r="J4"><v>7.5</v></c>',
                '<c r="K4"><v>4</v></c>',
                '<c r="L4"><v>1480</v></c>',
                '<c r="M4"><v>98</v></c>',
            ]
            row_xml = '<row r="4">' + "".join(cells) + "</row>"
            sheet_xml = files[sheet_target].replace(b"</sheetData>", row_xml.encode("utf-8") + b"</sheetData>")
            files[sheet_target] = sheet_xml
            patched = Path(directory) / "patched.xlsx"
            with ZipFile(patched, "w", compression=ZIP_DEFLATED) as archive:
                for name, data in files.items():
                    archive.writestr(name, data)

            rows = workbook.read_rows(patched)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0].sheet_name, "电动机")
            self.assertEqual(rows[0].row_number, 4)
            self.assertEqual(rows[0].values["rated_power"], 7.5)
            result = EvaluationFacade(EvaluationService(JsonStandardRepository(ROOT))).evaluate_v4(
                rows[0].record_id,
                rows[0].sheet_name,
                rows[0].values,
            )
            self.assertEqual(result.conclusion.value, "1级")


if __name__ == "__main__":
    unittest.main()
