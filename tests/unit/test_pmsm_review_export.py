import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from tools.export_pmsm_review_json import _suspicious_cells, export_review
from tools.activate_gb30253_pack import activate


class PmsmReviewExportTests(unittest.TestCase):
    def test_single_grade_table_decodes_row_local_suspicious_index(self):
        pack = {
            "tables": [{
                "table_no": 8,
                "dims": [2, 4],
                "rows": [{
                    "efficiency": {"2": [90.0, 0.0]},
                    "suspicious_cells": [1],
                }],
            }],
        }
        self.assertEqual(_suspicious_cells(pack), [(0, 8, 0, 2, 1)])

    def _fixtures(self, root: Path, status: str = "正确", suggestion=94.9):
        pack = {
            "pack_id": "gb30253_2024_pdf_verified_v1",
            "status": "normalized",
            "source_sha256": "abc",
            "tables": [],
        }
        for table_no in range(1, 30):
            table = {"table_no": table_no, "table": f"表{table_no}", "dims": [2, 4], "pages": [table_no], "rows": []}
            if table_no == 1:
                table["rows"] = [{"power_kw": 55.0, "efficiency": {"1": [96.6, 96.8], "2": [96.2, 0.0], "3": [95.3, 95.7]}, "suspicious_cells": [3]}]
            else:
                table["rows"] = [{"power_kw": float(table_no), "efficiency": {"1": [90.0, 90.1], "2": [89.0, 89.1], "3": [88.0, 88.1]}}]
            pack["tables"].append(table)
        book = Workbook()
        ws = book.active
        ws.title = "标准数据平铺"
        ws.append(["标题"])
        ws.append([])
        ws.append(["数据ID", "标准编号", "表号", "结构化区间", "条款/页码", "校对状态", "建议修订值", "校对意见"])
        ws.append(["id-power", "GB 30253-2024", "表1", "tables[0].rows[0].power_kw", "", status, "", ""])
        ws.append(["id-e1", "GB 30253-2024", "表1", "tables[0].rows[0].efficiency.1[0]", "", status, "", ""])
        ws.append(["id-e2a", "GB 30253-2024", "表1", "tables[0].rows[0].efficiency.1[1]", "", status, "", ""])
        ws.append(["id-e2b", "GB 30253-2024", "表1", "tables[0].rows[0].efficiency.2[0]", "", status, "", ""])
        ws.append(["id-sus", "GB 30253-2024", "表1", "tables[0].rows[0].efficiency.2[1]", "", status, suggestion, "已逐格对照PDF"])
        ws.append(["id-e3a", "GB 30253-2024", "表1", "tables[0].rows[0].efficiency.3[0]", "", status, "", ""])
        ws.append(["id-e3b", "GB 30253-2024", "表1", "tables[0].rows[0].efficiency.3[1]", "", status, "", ""])
        # 其余表各放一个最小完整效率行，覆盖导出器的“29张表/不可删行”门禁。
        for table_index in range(1, 29):
            table_no = table_index + 1
            # GB 30253-2024表8～表28各表只登记一个等级；用表2模拟
            # 这一结构，确保导出器不会假设每张表都有三档数组。
            levels = (1,) if table_no == 2 else (1, 2, 3)
            pack_table = pack["tables"][table_index]
            pack_table["rows"][0]["efficiency"] = {
                str(level): [90.0 + level, 90.1 + level] for level in levels
            }
            for level in levels:
                for dimension in (0, 1):
                    ws.append([
                        f"id-{table_no}-{level}-{dimension}", "GB 30253-2024", f"表{table_no}",
                        f"tables[{table_index}].rows[0].efficiency.{level}[{dimension}]", "", status, "", "",
                    ])
        pack_path = root / "pack.json"
        pack_path.write_text(json.dumps(pack), encoding="utf-8")
        book_path = root / "review.xlsx"
        book.save(book_path)
        return pack_path, book_path

    def test_export_requires_all_pmsm_records_to_be_correct(self):
        with tempfile.TemporaryDirectory() as tmp:
            pack, book = self._fixtures(Path(tmp), status="未校对")
            with self.assertRaisesRegex(ValueError, "未标记为.*正确"):
                export_review(book, Path(tmp) / "review.json", "张三", "2026-08-27", pack)

    def test_export_maps_suspicious_cell_to_activation_record(self):
        with tempfile.TemporaryDirectory() as tmp:
            pack, book = self._fixtures(Path(tmp))
            result = export_review(book, Path(tmp) / "review.json", "张三", "2026-08-27", pack)
            self.assertEqual(len(result["items"]), 1)
            item = result["items"][0]
            self.assertEqual(item["table_no"], 1)
            self.assertEqual(item["power_kw"], 55.0)
            self.assertEqual(item["level"], 2)
            self.assertEqual(item["dimension"], 4)
            self.assertEqual(item["replacement"], 94.9)
            self.assertTrue(item["pdf_checked"])
            self.assertTrue(result["all_cells_reviewed"])
            self.assertEqual(result["reviewed_tables"], list(range(1, 30)))
            activated = activate(pack, Path(tmp) / "review.json", Path(tmp) / "active.json")
            self.assertEqual(activated["status"], "active")

    def test_export_preserves_confirmed_no_data(self):
        with tempfile.TemporaryDirectory() as tmp:
            pack, book = self._fixtures(Path(tmp), suggestion=None)
            from openpyxl import load_workbook
            wb = load_workbook(book)
            ws = wb["标准数据平铺"]
            ws["H8"] = "已逐格对照PDF，原文无数据"
            wb.save(book)
            result = export_review(book, Path(tmp) / "review.json", "张三", "2026-08-27", pack)
            self.assertTrue(result["items"][0]["no_data"])
            self.assertIsNone(result["items"][0]["replacement"])


if __name__ == "__main__":
    unittest.main()
