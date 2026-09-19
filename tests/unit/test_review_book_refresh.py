from __future__ import annotations

from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from tools.refresh_review_book_preserving_edits import refresh


HEADERS = [
    "数据ID", "标准编号", "表号", "产品类别", "匹配条件", "原文区间", "结构化区间",
    "原文值", "原文单位", "规范化值", "规范化单位", "等级", "比较方向", "插值规则",
    "条款/页码", "来源文件", "校对状态", "建议修订值", "校对意见",
]


def _book(path: Path, rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "标准数据平铺"
    sheet.append(["标题"])
    sheet.append([])
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


class ReviewBookRefreshTests(unittest.TestCase):
    def test_refresh_merges_reviewer_columns_by_data_id(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            existing = root / "existing.xlsx"
            output = root / "output.xlsx"
            old_row = ["old-id"] + [""] * 15 + ["正确", 95, "已与PDF核对"]
            _book(existing, [old_row])

            def fake_builder(path: Path) -> Path:
                new_row = ["old-id"] + [""] * 15 + ["未校对", "", ""]
                added_row = ["new-id"] + [""] * 15 + ["未校对", "", ""]
                _book(path, [new_row, added_row])
                return path

            with patch("tools.refresh_review_book_preserving_edits._write_book", side_effect=fake_builder):
                result = refresh(existing, output)
            self.assertEqual(result["previous_review_rows"], 1)
            self.assertEqual(result["rebuilt_rows"], 2)
            self.assertEqual(result["merged_rows"], 1)
            self.assertEqual(result["new_rows_without_previous_review"], 1)

            workbook = load_workbook(output, read_only=True, data_only=False)
            sheet = workbook["标准数据平铺"]
            values = {sheet.cell(row, 1).value: [sheet.cell(row, col).value for col in (17, 18, 19)] for row in range(4, 6)}
            workbook.close()
            self.assertEqual(values["old-id"], ["正确", 95, "已与PDF核对"])
            self.assertEqual(values["new-id"], ["未校对", None, None])


if __name__ == "__main__":
    unittest.main()
