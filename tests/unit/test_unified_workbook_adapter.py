from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
import re
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook

from equipeffi.application.ports.batch_workbook import BatchRowOutcome
from equipeffi.application.services.pump_batch_evaluation_service import BatchEvaluationSummary
from equipeffi.application.services.workbook_evaluation_service import WorkbookEvaluationService
from equipeffi.domain.devices.transformer.input_contract import (
    TRANSFORMER_CATEGORIES,
    condition_options,
)
from equipeffi.infrastructure.excel.ooxml_reader import OOXMLReadError, OOXMLWorkbook
from equipeffi.infrastructure.excel.pump_result_writer import _sheet_target
from equipeffi.infrastructure.excel.pump_workbook_reader import V6PumpWorkbookReader
from equipeffi.infrastructure.excel.template_resource import (
    DEFAULT_V6_TEMPLATE,
    DEFAULT_V61_TEMPLATE,
    TEMPLATE_MANIFEST,
    V6TemplateResource,
    V61TransformerTemplateResource,
)
from equipeffi.infrastructure.excel.unified_result_writer import (
    TRANSFORMER_RESULT_COLUMNS,
    UnifiedResultWorkbookWriter,
)
from equipeffi.infrastructure.excel.unified_workbook_reader import UnifiedWorkbookReader


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE_DIR = ROOT / "src" / "equipeffi" / "resources" / "templates"


class UnifiedWorkbookAdapterTests(unittest.TestCase):
    def setUp(self):
        self._temporary = tempfile.TemporaryDirectory(prefix="equipeffi-unified-")
        self.addCleanup(self._temporary.cleanup)
        self.temp_dir = Path(self._temporary.name)

    def _filled_v61(self, name: str = "source.xlsx") -> Path:
        source = self.temp_dir / name
        source.write_bytes((TEMPLATE_DIR / DEFAULT_V61_TEMPLATE).read_bytes())
        workbook = load_workbook(source)
        sheet = workbook["变压器"]
        first = TRANSFORMER_CATEGORIES[0]
        first_condition = condition_options(first.category_id)[0]
        sheet["B4"] = "配电变压器"
        sheet["C4"] = "S13-630"
        sheet["D4"] = 1
        sheet["F4"] = first.label
        sheet["G4"] = 630
        sheet["H4"] = first_condition.label
        sheet["I4"] = 4.25
        sheet["J4"] = 123.45
        sheet["K4"] = 345.67
        sheet["S4"] = "铭牌"
        sheet["T4"] = "样例"

        table_27 = TRANSFORMER_CATEGORIES[26]
        sheet["B5"] = "自耦变压器"
        sheet["D5"] = 2
        sheet["F5"] = table_27.label
        sheet["G5"] = 120000
        sheet["I5"] = 10.1
        sheet["J5"] = 1000
        sheet["K5"] = 2000
        sheet["V5"] = 500
        sheet["W5"] = 10.5
        sheet["X5"] = 20.2
        sheet["Y5"] = 30.3

        table_23 = TRANSFORMER_CATEGORIES[22]
        sheet["B6"] = "串联绕组末端调压自耦变压器"
        sheet["D6"] = 1
        sheet["F6"] = table_23.label
        sheet["G6"] = 100000
        sheet["I6"] = 11.1
        sheet["J6"] = 900
        sheet["K6"] = 1900

        table_24 = TRANSFORMER_CATEGORIES[23]
        sheet["B7"] = "另一类串联绕组末端调压自耦变压器"
        sheet["D7"] = 1
        sheet["F7"] = table_24.label
        sheet["G7"] = 100000
        sheet["I7"] = 12.1
        sheet["J7"] = 910
        sheet["K7"] = 1910

        table_28 = TRANSFORMER_CATEGORIES[27]
        sheet["B8"] = "三绕组变压器"
        sheet["D8"] = 1
        sheet["F8"] = table_28.label
        sheet["G8"] = 100000
        sheet["I8"] = 13.1
        sheet["J8"] = 920
        sheet["K8"] = 1920
        sheet["X8"] = 21.2
        sheet["Y8"] = 31.3
        workbook.save(source)
        return source

    def _patch_transformer_cells(self, source: Path, replacements) -> Path:
        patched = source.with_name(f"{source.stem}-patched.xlsx")
        main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        ET.register_namespace("", main_ns)
        with ZipFile(source) as original:
            transformer_part = _sheet_target(original, "变压器")
            with ZipFile(patched, "w", ZIP_DEFLATED) as target:
                for info in original.infolist():
                    data = original.read(info.filename)
                    if info.filename == transformer_part:
                        root = ET.fromstring(data)
                        cells = {cell.attrib.get("r"): cell for cell in root.iter(
                            f"{{{main_ns}}}c")}
                        for reference, (cell_type, formula, cached_value) in replacements.items():
                            cell = cells.get(reference)
                            if cell is None:
                                row_number = int(re.search(r"[0-9]+$", reference).group())
                                sheet_data = root.find(
                                    f"{{{main_ns}}}sheetData")
                                row_node = next((row for row in sheet_data.findall(
                                    f"{{{main_ns}}}row")
                                    if row.attrib.get("r") == str(row_number)), None)
                                if row_node is None:
                                    row_node = ET.SubElement(sheet_data,
                                        f"{{{main_ns}}}row", {"r": str(row_number)})
                                cell = ET.SubElement(row_node, f"{{{main_ns}}}c",
                                                     {"r": reference})
                                cells[reference] = cell
                            if cell_type:
                                cell.attrib["t"] = cell_type
                            else:
                                cell.attrib.pop("t", None)
                            formula_node = cell.find(f"{{{main_ns}}}f")
                            if formula is None:
                                if formula_node is not None:
                                    cell.remove(formula_node)
                            elif formula_node is None:
                                ET.SubElement(cell, f"{{{main_ns}}}f").text = formula
                            else:
                                formula_node.text = formula
                            value_node = cell.find(f"{{{main_ns}}}v")
                            if cached_value is None:
                                if value_node is not None:
                                    cell.remove(value_node)
                            elif value_node is None:
                                ET.SubElement(cell, f"{{{main_ns}}}v").text = cached_value
                            else:
                                value_node.text = cached_value
                        data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                    target.writestr(info, data)
        return patched

    def test_v61_resource_has_candidate_identity_and_25_column_contract(self):
        validation = V61TransformerTemplateResource().validate()
        self.assertTrue(validation.is_valid, validation.message)
        self.assertEqual(validation.identity["template_version"], "V6.1")
        entry = next(item for item in TEMPLATE_MANIFEST["templates"]
                     if item["filename"] == DEFAULT_V61_TEMPLATE)
        self.assertEqual(entry["status"], "CANDIDATE")
        contract = next(sheet for sheet in entry["sheets"]
                        if sheet["sheet_id"] == "transformer")
        self.assertEqual(contract["contract_revision"], "transformer-input.v1")
        self.assertEqual(len(contract["fields"]), 25)
        self.assertEqual(contract["table"]["ref"], "A3:Y103")
        self.assertEqual(V6TemplateResource().identity["filename"], DEFAULT_V6_TEMPLATE)
        self.assertEqual(V6TemplateResource().identity["template_version"], "V6-20261005")

    def test_unified_reader_reads_transformer_and_keeps_v6_pump_reader_compatible(self):
        source = self._filled_v61()
        reader = UnifiedWorkbookReader()
        workbook = reader.read(source)
        self.assertEqual(set(workbook.sheets), {"离心泵", "变压器"})
        self.assertEqual(len(workbook.sheets["离心泵"].rows), 0)
        self.assertIn("电动机", workbook.unprocessed_sheets)
        self.assertNotIn("变压器", workbook.unprocessed_sheets)
        self.assertNotIn("配置", workbook.unprocessed_sheets)
        self.assertNotIn("注意事项", workbook.unprocessed_sheets)
        self.assertNotIn("模板说明", workbook.unprocessed_sheets)

        transformer_rows = workbook.sheets["变压器"].rows
        self.assertEqual([row.row_number for row in transformer_rows], [4, 5, 6, 7, 8])
        first, three_winding, two_winding_auto, table_24, table_28 = (
            row.values for row in transformer_rows)
        self.assertEqual(first["category_id"], "GB20052-T01")
        self.assertEqual(first["standard_conditions"],
                         condition_options("GB20052-T01")[0].condition_id)
        self.assertEqual(first["capacity_kva"], 630)
        self.assertEqual(first["no_load_loss_w"], Decimal("123.45"))
        self.assertEqual(first["load_loss_w"], Decimal("345.67"))
        self.assertEqual(first["uk_rated_percent"], Decimal("4.25"))
        self.assertEqual(three_winding["category_id"], "GB20052-T27")
        self.assertNotIn("uk_rated_percent", three_winding)
        self.assertEqual(three_winding["uk_hm_rated_percent"], Decimal("10.1"))
        self.assertEqual(three_winding["uk_hl_rated_percent"], Decimal("20.2"))
        self.assertEqual(three_winding["uk_ml_rated_percent"], Decimal("30.3"))
        self.assertEqual(two_winding_auto["category_id"], "GB20052-T23")
        self.assertEqual(two_winding_auto["uk_hm_rated_percent"], Decimal("11.1"))
        self.assertNotIn("uk_rated_percent", two_winding_auto)
        self.assertEqual(table_24["category_id"], "GB20052-T24")
        self.assertEqual(table_24["uk_hm_rated_percent"], Decimal("12.1"))
        self.assertNotIn("uk_rated_percent", table_24)
        self.assertEqual(table_28["category_id"], "GB20052-T28")
        self.assertEqual(table_28["uk_hm_rated_percent"], Decimal("13.1"))
        self.assertNotIn("uk_rated_percent", table_28)

        # The established pump reader continues to accept the pump sheet in V6.1.
        pump_source = V6PumpWorkbookReader().read(source)
        self.assertEqual(pump_source.rows, ())

        # V6 stays readable as a pump-only source and does not acquire transformer support.
        legacy = self.temp_dir / "legacy-v6.xlsx"
        legacy.write_bytes((TEMPLATE_DIR / DEFAULT_V6_TEMPLATE).read_bytes())
        legacy_result = reader.read(legacy)
        self.assertEqual(tuple(legacy_result.sheets), ("离心泵",))
        self.assertIn("变压器", legacy_result.unprocessed_sheets)
        self.assertEqual(legacy_result.template_identity["template_version"],
                         V6TemplateResource().identity["template_version"])

    def test_reader_rejects_unknown_template_contract_revision(self):
        source = self._filled_v61()
        corrupted = self.temp_dir / "unknown-revision.xlsx"
        with ZipFile(source) as original, ZipFile(corrupted, "w", ZIP_DEFLATED) as target:
            for info in original.infolist():
                data = original.read(info.filename)
                if info.filename == "xl/workbook.xml":
                    xml = data.decode("utf-8")
                    marker = r'(<definedName name="EQUIPEFFI_TEMPLATE_CONTRACT_REVISION">).*?(</definedName>)'
                    xml, replacements = re.subn(
                        marker, r'\1"unknown-revision"\2', xml, count=1)
                    self.assertEqual(replacements, 1)
                    data = xml.encode("utf-8")
                target.writestr(info, data)
        with self.assertRaises(OOXMLReadError):
            UnifiedWorkbookReader().read(corrupted)

    def test_transformer_reader_rejects_formula_cache_date_and_excel_error_inputs(self):
        source = self._filled_v61()
        condition_label = condition_options("GB20052-T01")[0].label
        formula_source = self._patch_transformer_cells(source, {
            "F4": ("str", "\"配电变压器\"", "配电变压器"),
            "H4": ("str", f'"{condition_label}"', condition_label),
            "J4": (None, "1+2", "999.25"),
        })
        with OOXMLWorkbook(formula_source) as workbook:
            cached = workbook.rows("变压器")[3]
        self.assertEqual(cached[5], "配电变压器")
        self.assertEqual(cached[7], condition_label)
        self.assertEqual(cached[9], Decimal("999.25"))

        formula_row = UnifiedWorkbookReader().read(formula_source).sheets[
            "变压器"].rows[0]
        for reference in ("F4", "H4", "J4"):
            self.assertTrue(any(reference in error and "公式" in error
                                for error in formula_row.input_errors), reference)
        self.assertEqual(formula_row.values["category_id"], "")
        self.assertEqual(formula_row.values["standard_conditions"], "")
        self.assertEqual(formula_row.values["no_load_loss_w"], "")
        application = object.__new__(WorkbookEvaluationService)
        application.transformer = Mock()
        summary = BatchEvaluationSummary(data_row_count=1)
        outcome = application._transformer_row(formula_row, date.today(), summary, [])
        self.assertEqual(outcome.evaluation_status, "INVALID_INPUT")
        self.assertTrue(outcome.is_input_error)
        application.transformer.evaluate.assert_not_called()
        self.assertEqual(summary.total_quantity, 1)
        self.assertEqual(summary.input_error_rows, 1)

        date_source = self.temp_dir / "date-input.xlsx"
        date_source.write_bytes((TEMPLATE_DIR / DEFAULT_V61_TEMPLATE).read_bytes())
        date_workbook = load_workbook(date_source)
        date_sheet = date_workbook["变压器"]
        date_sheet["B4"] = "日期误填损耗"
        date_sheet["D4"] = 1
        date_sheet["F4"] = TRANSFORMER_CATEGORIES[0].label
        date_sheet["J4"] = datetime(2026, 10, 23)
        date_sheet["J4"].number_format = "yyyy-mm-dd"
        date_workbook.save(date_source)
        date_row = UnifiedWorkbookReader().read(date_source).sheets[
            "变压器"].rows[0]
        self.assertTrue(any("日期或时间格式" in error for error in date_row.input_errors))
        self.assertEqual(date_row.values["no_load_loss_w"], "")

        blank_date_source = self.temp_dir / "blank-date-style.xlsx"
        blank_date_source.write_bytes((TEMPLATE_DIR / DEFAULT_V61_TEMPLATE).read_bytes())
        blank_date_workbook = load_workbook(blank_date_source)
        blank_date_workbook["变压器"]["J4"].number_format = "yyyy-mm-dd"
        blank_date_workbook.save(blank_date_source)
        blank_date_result = UnifiedWorkbookReader().read(blank_date_source).sheets[
            "变压器"]
        self.assertFalse(any(row.row_number == 4 for row in blank_date_result.rows))

        dated_note_source = self.temp_dir / "dated-note.xlsx"
        dated_note_source.write_bytes((TEMPLATE_DIR / DEFAULT_V61_TEMPLATE).read_bytes())
        dated_note_workbook = load_workbook(dated_note_source)
        dated_note_sheet = dated_note_workbook["变压器"]
        dated_note_sheet["B4"] = "报告日期记录"
        dated_note_sheet["T4"] = datetime(2026, 10, 23)
        dated_note_sheet["T4"].number_format = "yyyy-mm-dd"
        dated_note_workbook.save(dated_note_source)
        dated_note_row = UnifiedWorkbookReader().read(dated_note_source).sheets[
            "变压器"].rows[0]
        self.assertEqual(dated_note_row.input_errors, ())

        error_source = self._patch_transformer_cells(source, {
            "J4": ("e", None, "#DIV/0!"),
        })
        error_row = UnifiedWorkbookReader().read(error_source).sheets[
            "变压器"].rows[0]
        self.assertTrue(any("Excel错误值（#DIV/0!）" in error
                            for error in error_row.input_errors))
        self.assertEqual(error_row.values["no_load_loss_w"], "")

        empty_formula_base = self.temp_dir / "empty-formula-base.xlsx"
        empty_formula_base.write_bytes((TEMPLATE_DIR / DEFAULT_V61_TEMPLATE).read_bytes())
        empty_formula_source = self._patch_transformer_cells(
            empty_formula_base, {"F4": ("str", '""', None)})
        empty_formula_row = UnifiedWorkbookReader().read(empty_formula_source).sheets[
            "变压器"].rows[0]
        self.assertTrue(any("F4" in error and "公式" in error
                            for error in empty_formula_row.input_errors))

    def test_writer_patches_both_sheets_and_preserves_non_target_parts(self):
        source = self._filled_v61()
        workbook = load_workbook(source)
        transformer = workbook["变压器"]
        transformer["L4"] = 999  # stale result must be cleared when unavailable.
        transformer["M4"] = 888
        transformer["N4"] = 777
        transformer["R4"] = "旧结论"
        transformer["U4"] = "旧备注"
        workbook.save(source)

        pump_outcome = BatchRowOutcome(
            4, True, "泵结论", "SUCCESS", "2级", messages=("泵备注",))
        transformer_outcome = BatchRowOutcome(
            4, True, "1级", "SUCCESS", "1级",
            messages=("变压器说明", "第二条"),
            thresholds={
                "no_load_loss_1": Decimal("100.125"),
                "no_load_loss_2": Decimal("110.25"),
                "no_load_loss_3": None,
                "load_loss_1": Decimal("200.125"),
                "load_loss_2": Decimal("210.25"),
                "load_loss_3": Decimal("220.5"),
            })
        input_error = BatchRowOutcome(
            5, False, "输入错误", "INVALID_INPUT", None,
            messages=("请补充必要字段",),
            thresholds={"no_load_loss_1": Decimal("999")},
            is_input_error=True)
        execution_error = BatchRowOutcome(
            6, False, "执行失败", None, None,
            messages=("执行失败：ValueError",),
            thresholds={"no_load_loss_1": Decimal("999")},
            is_execution_error=True)
        destination = self.temp_dir / "result.xlsx"
        UnifiedResultWorkbookWriter().write(
            source,
            {"离心泵": {4: pump_outcome},
             "变压器": {4: transformer_outcome, 5: input_error,
                        6: execution_error}},
            destination)

        with ZipFile(source) as before, ZipFile(destination) as after:
            self.assertEqual(before.namelist(), after.namelist())
            pump_part = _sheet_target(before, "离心泵")
            transformer_part = _sheet_target(before, "变压器")
            self.assertIsNotNone(pump_part)
            self.assertIsNotNone(transformer_part)
            self.assertNotEqual(before.read(pump_part), after.read(pump_part))
            self.assertNotEqual(before.read(transformer_part), after.read(transformer_part))
            for name in before.namelist():
                if name not in {pump_part, transformer_part}:
                    self.assertEqual(before.read(name), after.read(name), name)

        with OOXMLWorkbook(destination) as output:
            pump_rows = output.rows("离心泵")
            transformer_rows = output.rows("变压器")
        self.assertEqual(pump_rows[3][23], "泵结论")  # X4
        self.assertEqual(pump_rows[3][26], "泵备注")  # AA4
        self.assertEqual(transformer_rows[3][11], Decimal("100.125"))
        self.assertEqual(transformer_rows[3][12], Decimal("110.25"))
        self.assertEqual(transformer_rows[3][13], "")
        self.assertEqual(transformer_rows[3][16], Decimal("220.5"))
        self.assertEqual(transformer_rows[3][17], "1级")
        self.assertEqual(transformer_rows[3][20], "变压器说明；第二条")
        self.assertEqual(TRANSFORMER_RESULT_COLUMNS,
                         ("L", "M", "N", "O", "P", "Q", "R", "U"))
        for row_index in (4, 5):
            for column_index in range(11, 17):
                self.assertEqual(transformer_rows[row_index][column_index], "")
            self.assertEqual(transformer_rows[row_index][17], "")
        self.assertIn("输入错误", transformer_rows[4][20])
        self.assertIn("请补充必要字段", transformer_rows[4][20])
        self.assertEqual(transformer_rows[5][20], "执行失败：ValueError")

    def test_writer_keeps_legacy_v6_pump_only_operation(self):
        source = self.temp_dir / "legacy-v6.xlsx"
        source.write_bytes((TEMPLATE_DIR / DEFAULT_V6_TEMPLATE).read_bytes())
        destination = self.temp_dir / "legacy-result.xlsx"
        outcome = BatchRowOutcome(4, True, "泵结论", "SUCCESS", "2级")
        UnifiedResultWorkbookWriter().write(
            source, {"离心泵": {4: outcome}}, destination)
        with ZipFile(source) as before, ZipFile(destination) as after:
            pump = _sheet_target(before, "离心泵")
            transformer = _sheet_target(before, "变压器")
            self.assertIsNotNone(pump)
            self.assertIsNotNone(transformer)
            self.assertEqual(before.read(transformer), after.read(transformer))

    def test_reader_and_writer_support_100_1000_and_10000_rows(self):
        source = self.temp_dir / "ten-thousand-rows.xlsx"
        source.write_bytes((TEMPLATE_DIR / DEFAULT_V61_TEMPLATE).read_bytes())
        workbook = load_workbook(source)
        sheet = workbook["变压器"]
        category = TRANSFORMER_CATEGORIES[0]
        condition = condition_options(category.category_id)[0]
        for index in range(10_000):
            row = index + 4
            sheet.cell(row, 2, f"设备-{index + 1}")
            sheet.cell(row, 4, 1)
            sheet.cell(row, 6, category.label)
            sheet.cell(row, 7, 630)
            sheet.cell(row, 8, condition.label)
            sheet.cell(row, 9, 4.25)
            sheet.cell(row, 10, 123.45)
            sheet.cell(row, 11, 345.67)
        workbook.save(source)

        read_result = UnifiedWorkbookReader().read(source)
        rows = read_result.sheets["变压器"].rows
        self.assertEqual(len(rows), 10_000)
        self.assertEqual(
            [rows[index - 1].row_number for index in (100, 1_000, 10_000)],
            [103, 1_003, 10_003])

        outcomes = {
            row_number: BatchRowOutcome(
                row_number, True, "1级", "SUCCESS", "1级",
                thresholds={"no_load_loss_1": Decimal(row_number)})
            for row_number in (103, 1_003, 10_003)
        }
        destination = self.temp_dir / "ten-thousand-rows-result.xlsx"
        UnifiedResultWorkbookWriter().write(
            source, {"变压器": outcomes}, destination)
        with OOXMLWorkbook(destination) as output:
            result_rows = output.rows("变压器")
        for row_number in (103, 1_003, 10_003):
            self.assertEqual(result_rows[row_number - 1][11], Decimal(row_number))
            self.assertEqual(result_rows[row_number - 1][17], "1级")


if __name__ == "__main__":
    unittest.main()
