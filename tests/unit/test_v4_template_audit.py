from __future__ import annotations

from pathlib import Path
import shutil
import tempfile
import unittest
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
from openpyxl.utils.cell import range_boundaries

from equipeffi.infrastructure.excel.template_resource import REQUIRED_V4_SHEETS
from equipeffi.infrastructure.excel.v4_template_audit import audit_v4_template
from equipeffi.infrastructure.excel.v4_writer import V4WorkbookWriterImpl
from equipeffi.application.ports.v4_workbook import V4WorkbookEvaluationRow
from equipeffi.application.services.evaluation_facade import EvaluationFacade
from equipeffi.application.services.evaluation_service import EvaluationService
from equipeffi.infrastructure.excel.v4_reader import V4WorkbookReaderImpl
from equipeffi.infrastructure.excel.ooxml_reader import OOXMLWorkbook
from equipeffi.infrastructure.excel.v4_writer import _set_cell
from equipeffi.infrastructure.standards.json_repository import JsonStandardRepository
from tools.audit_v4_validations import audit as audit_v4_validations


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "src" / "equipeffi" / "resources" / "templates" / "设备能效分析空白模板_重构版V4_20260825.xlsx"


class V4TemplateAuditTests(unittest.TestCase):
    def test_bundled_template_protection_formula_and_fields(self):
        result = audit_v4_template(TEMPLATE)
        self.assertTrue(result.is_valid)
        self.assertEqual(result.missing_sheets, ())
        self.assertEqual(result.unprotected_sheets, ())
        self.assertEqual(result.forbidden_fields, ())
        self.assertEqual(result.formula_errors, ())
        self.assertEqual(result.volatile_formulas, ())
        self.assertTrue(result.workbook_protected)
        self.assertTrue(result.calc_on_load)
        self.assertGreater(result.formula_count, 0)
        self.assertEqual(result.standard_catalog_count, 16)
        self.assertEqual(result.future_standard_status, "已发布、未实施")
        self.assertEqual(result.standard_catalog_issues, ())
        self.assertEqual(set(result.sheet_names), set(REQUIRED_V4_SHEETS))

    def test_template_explanation_title_is_v4(self):
        workbook = openpyxl.load_workbook(TEMPLATE, read_only=True, data_only=False)
        try:
            title = str(workbook["模板说明"]["A1"].value or "")
            self.assertIn("V4", title)
            self.assertNotIn("V3", title)
        finally:
            workbook.close()

    def test_writer_preserves_existing_media_part(self):
        """写回器保留未知媒体部件，后续可接入图片锚点而不丢附件。"""
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "输入.xlsx"
            shutil.copy2(TEMPLATE, source)
            with ZipFile(source, "r") as archive:
                files = {name: archive.read(name) for name in archive.namelist()}
            media_name = "xl/media/保留附件.bin"
            files[media_name] = b"attachment-preserved"
            with OOXMLWorkbook(source) as workbook:
                sheet_target = workbook._sheets["电动机"]
            sheet_xml = files[sheet_target].decode("utf-8")
            for column, value in {
                0: 1, 1: "电机1", 2: "M1", 3: 1,
                4: "三相异步电动机（一般用途）", 6: 0.4, 9: 7.5,
                10: 4, 11: 1480, 12: 98,
            }.items():
                sheet_xml = _set_cell(sheet_xml, 4, column, value)
            files[sheet_target] = sheet_xml.encode("utf-8")
            patched = Path(directory) / "patched.xlsx"
            with ZipFile(patched, "w", compression=ZIP_DEFLATED) as archive:
                for name, data in files.items():
                    archive.writestr(name, data)

            row = V4WorkbookReaderImpl().read_rows(patched)[0]
            self.assertEqual(row.values["category"], "三相异步电动机（一般用途）")
            self.assertEqual(row.values["rated_voltage"], 0.4)
            self.assertEqual(row.values["rated_power"], 7.5)
            result = EvaluationFacade(EvaluationService(JsonStandardRepository(ROOT))).evaluate_v4(
                row.record_id, row.sheet_name, row.values
            )
            destination = Path(directory) / "结果.xlsx"
            V4WorkbookWriterImpl().write_results(
                patched, destination, [V4WorkbookEvaluationRow(row, result)]
            )
            with ZipFile(destination, "r") as archive:
                self.assertEqual(archive.read(media_name), b"attachment-preserved")

    def test_template_has_workbook_elimination_scope_dropdown(self):
        workbook = openpyxl.load_workbook(TEMPLATE, read_only=False, data_only=False)
        notes = workbook["注意事项"]
        validation = next((item for item in notes.data_validations.dataValidation if str(item.sqref) == "B10"), None)
        self.assertIsNotNone(validation)
        self.assertEqual(validation.formula1, "=EV_ELIMINATION_SCOPE")
        self.assertEqual(notes["B10"].value, "高耗能落后机电设备淘汰目录第一至第四批")
        self.assertFalse(notes["B10"].protection.locked)
        self.assertIn("EV_ELIMINATION_SCOPE", workbook.defined_names)
        workbook.close()

    def test_all_device_sheets_have_common_locked_result_fields(self):
        workbook = openpyxl.load_workbook(TEMPLATE, read_only=False, data_only=False)
        expected = {"采用标准", "匹配表/条款", "参考能效等级", "判定说明", "缺失信息"}
        config_rows = {
            (str(workbook["配置"].cell(row, 1).value or ""), str(workbook["配置"].cell(row, 2).value or ""))
            for row in range(2, workbook["配置"].max_row + 1)
        }
        for sheet in REQUIRED_V4_SHEETS[3:]:
            headers = {str(cell.value or "").replace("\n", "") for cell in workbook[sheet][3]}
            self.assertTrue(expected.issubset(headers), sheet)
            for display in expected:
                column = next(cell.column for cell in workbook[sheet][3] if str(cell.value or "").replace("\n", "") == display)
                self.assertTrue(workbook[sheet].cell(4, column).protection.locked, f"{sheet}:{display}")
            self.assertTrue(all((sheet, field_id) in config_rows for field_id in {"standard_code", "standard_table", "reference_grade", "explanation", "missing_fields"}))
        workbook.close()

    def test_result_columns_keep_fixed_suffix_and_summary_formula(self):
        workbook = openpyxl.load_workbook(TEMPLATE, read_only=False, data_only=False)
        for sheet_name in REQUIRED_V4_SHEETS[3:]:
            sheet = workbook[sheet_name]
            headers = [str(cell.value or "").replace("\n", "") for cell in sheet[3] if cell.value not in (None, "")]
            self.assertEqual(headers[-3:], ["铭牌照片", "填表备注", "自动备注"], sheet_name)
            self.assertTrue(any(isinstance(cell.value, str) and cell.value.startswith("=") for cell in sheet[1]), sheet_name)
            self.assertTrue(isinstance(sheet["A4"].value, str) and sheet["A4"].value.startswith("="), sheet_name)
        workbook.close()

    def test_each_device_sheet_has_hidden_locked_technical_record_id(self):
        workbook = openpyxl.load_workbook(TEMPLATE, read_only=False, data_only=False)
        try:
            for sheet_name in REQUIRED_V4_SHEETS[3:]:
                sheet = workbook[sheet_name]
                column = next(cell.column for cell in sheet[3] if str(cell.value or "").replace("\n", "") == "技术记录ID")
                letter = openpyxl.utils.get_column_letter(column)
                self.assertTrue(sheet.column_dimensions[letter].hidden, sheet_name)
                self.assertTrue(sheet.cell(4, column).protection.locked, sheet_name)
                self.assertTrue(str(sheet.cell(4, column).value or "").startswith("=IF(COUNTA("), sheet_name)
                self.assertLess(column, next(cell.column for cell in sheet[3] if "铭牌照片" in str(cell.value or "")), sheet_name)
                table = next(iter(sheet.tables.values()))
                id_column = next(item for item in table.tableColumns if item.name == "技术记录ID")
                self.assertIsNotNone(id_column.calculatedColumnFormula)
        finally:
            workbook.close()

    def test_common_prefix_order_matches_v4_contract(self):
        workbook = openpyxl.load_workbook(TEMPLATE, read_only=False, data_only=False)
        try:
            expected = ["序号", "设备名称", "型号", "数量", "设备类别", "安装位置"]
            for sheet_name in REQUIRED_V4_SHEETS[3:]:
                headers = [str(cell.value or "").replace("\n", "") for cell in workbook[sheet_name][3]]
                self.assertEqual(headers[:6], expected, sheet_name)
            config = workbook["配置"]
            rows = [
                (row, str(config.cell(row, 1).value or ""), str(config.cell(row, 2).value or ""))
                for row in range(2, config.max_row + 1)
            ]
            for sheet_name in REQUIRED_V4_SHEETS[3:]:
                field_rows = {field_id: row for row, sheet, field_id in rows if sheet == sheet_name}
                self.assertLess(field_rows["category"], field_rows["location"], sheet_name)
                self.assertEqual(config.cell(field_rows["category"], 4).value, "基础信息", sheet_name)
        finally:
            workbook.close()

    def test_input_validations_cover_unbounded_table_rows(self):
        """验证规则覆盖新增/批量粘贴行，而不是只覆盖初始空白行。"""
        workbook = openpyxl.load_workbook(TEMPLATE, read_only=False, data_only=False)
        try:
            for sheet_name in REQUIRED_V4_SHEETS[3:]:
                sheet = workbook[sheet_name]
                self.assertGreater(len(sheet.data_validations.dataValidation), 0, sheet_name)
                for validation in sheet.data_validations.dataValidation:
                    for token in str(validation.sqref).split():
                        try:
                            min_col, min_row, max_col, max_row = range_boundaries(token)
                        except ValueError:
                            continue
                        if min_row == 4 and min_col == max_col:
                            self.assertEqual(max_row, 1_048_576, f"{sheet_name}:{token}")
        finally:
            workbook.close()

    def test_validation_audit_tool_matches_bundled_template(self):
        result = audit_v4_validations(TEMPLATE)
        self.assertTrue(result["is_valid"], result["errors"])
        self.assertEqual(result["checks"]["sheets"], 15)
        self.assertEqual(result["checks"]["efficiency_fields"], 7)

    def test_boiler_efficiency_validation_uses_category_column(self):
        workbook = openpyxl.load_workbook(TEMPLATE, read_only=False, data_only=False)
        try:
            sheet = workbook["工业锅炉"]
            formulas = [
                str(value or "")
                for validation in sheet.data_validations.dataValidation
                if "L4:" in str(validation.sqref)
                for value in (validation.formula1, validation.formula2)
            ]
            self.assertTrue(any('E4="室燃燃烧锅炉（燃气冷凝）"' in formula for formula in formulas))
            self.assertFalse(any('F4="室燃燃烧锅炉（燃气冷凝）"' in formula for formula in formulas))
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
