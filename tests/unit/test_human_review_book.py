from __future__ import annotations

import tempfile
import unittest
import json
from pathlib import Path
from unittest.mock import patch

import openpyxl
from openpyxl.styles import Protection

from tools import build_human_review_book
from tools.validate_human_review_book import validate


ROOT = Path(__file__).resolve().parents[2]


class HumanReviewBookTests(unittest.TestCase):
    def test_no_data_paths_decode_single_grade_tables_by_row_grade(self):
        payload = {
            "tables": [{
                "table_no": 8,
                "dims": ["300", "500"],
                "rows": [{
                    "efficiency": {"2": [90, None]},
                    "no_data_cells": [1],
                }],
            }],
        }
        paths = build_human_review_book._no_data_paths(payload)
        self.assertIn("tables[0].rows[0].efficiency.2[1]", paths)
        self.assertNotIn("tables[0].rows[0].efficiency.1[1]", paths)

    def test_standard_rows_show_registered_no_data_as_dash(self):
        """人可读平铺表以“—”区分标准无数据与普通空值。"""
        payload = {
            "standard_code": "GB 30253-2024",
            "tables": [{
                "table": "表1",
                "product": "异步起动",
                "voltage_group": "≤1140V",
                "cooling_group": "通用",
                "mode": "poles",
                "dims": [12],
                "rows": [{
                    "power_kw": 55,
                    "efficiency": {"1": [None]},
                    "no_data_cells": [0],
                    "no_data_reason": "PDF原文无数据",
                }],
            }],
        }
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pack_path = root / "pack.json"
            manifest_path = root / "manifest.json"
            pack_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            manifest_path.write_text(json.dumps({"packs": [{
                "device_type": "motor_pmsm",
                "pack_id": "test-pmsm",
                "source": "pack.json",
                "standard_code": "GB 30253-2024",
            }]}, ensure_ascii=False), encoding="utf-8")
            with patch.object(build_human_review_book, "MANIFEST", manifest_path):
                rows = build_human_review_book._standard_rows()
        row = next(item for item in rows if item[2] == "表1" and item[6].endswith("efficiency.1[0]"))
        self.assertEqual(row[7], "—")
        self.assertEqual(row[9], "—")
        self.assertEqual(row[16], "正确")
        self.assertIn('"额定功率_kW":55', row[4])
        self.assertIn('"等级":"1级"', row[4])
        self.assertIn('"极数":"12"', row[4])
        self.assertEqual(row[5], "55 kW")

    def test_review_book_keeps_future_blower_standard_and_pmsm_issue(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = build_human_review_book._write_book(Path(tmp) / "review.xlsx")
            book = openpyxl.load_workbook(output, read_only=True, data_only=True)
            standard_rows = [
                [book["标准目录"].cell(row, col).value for col in range(1, 8)]
                for row in range(4, book["标准目录"].max_row + 1)
            ]
            pending = next(row for row in standard_rows if row[1] == "GB 28381-2026")
            self.assertEqual(pending[3], "已发布、未实施")
            self.assertIn("2027-04-01", pending[6])
            changes = [
                book["版本变更记录"].cell(row, 2).value
                for row in range(4, book["版本变更记录"].max_row + 1)
            ]
            self.assertIn("gb28381_2026_pending", changes)
            formula_ids = [
                book["公式与判定规则"].cell(row, 1).value
                for row in range(4, book["公式与判定规则"].max_row + 1)
            ]
            self.assertIn("RULE-PUMP-CHEM-BOUNDARY", formula_ids)
            issue_rows = [
                [book["待校对问题"].cell(row, col).value for col in range(1, 6)]
                for row in range(4, book["待校对问题"].max_row + 1)
            ]
            issue_headers = [book["待校对问题"].cell(3, col).value for col in range(1, 8)]
            self.assertEqual(issue_headers[-2:], ["PDF页码", "渲染图像提示"])
            self.assertTrue(any(str(row[0]) == "GB 30253-2024" for row in issue_rows))
            flat = book["标准数据平铺"]
            # 该校对册超过5万行，只进行一次顺序扫描，避免验收测试重复解压
            # 工作表而长时间占满CPU。
            pmsm_correct = 0
            pmsm_no_data = []
            for values in flat.iter_rows(min_row=4, values_only=True):
                if len(values) < 19 or values[1] != "GB 30253-2024":
                    continue
                if values[16] == "正确":
                    pmsm_correct += 1
                    if "无数据" in str(values[18] or ""):
                        pmsm_no_data.append([values[6], values[7], values[9]])
            self.assertEqual(pmsm_correct, 3)
            self.assertEqual(len(pmsm_no_data), 3)
            self.assertTrue(all(item[1] == "—" and item[2] == "—" for item in pmsm_no_data))
            industry = book["淘汰目录_产业"]
            self.assertEqual(industry.cell(3, 7).value, "PDF页码")
            self.assertFalse(industry.cell(4, 8).protection.locked)
            book.close()

    def test_review_book_validator_checks_directory_version_against_manifest(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = build_human_review_book._write_book(Path(tmp) / "review.xlsx")
            result = validate(output)
            self.assertTrue(result["is_valid"], result["errors"])
            self.assertEqual(result["expected_flat_rows"], result["actual_flat_rows"])
            self.assertEqual(result["expected_no_data_rows"], 3)
            self.assertEqual(result["actual_no_data_rows"], 3)
            self.assertFalse(result["activation_ready"])
            self.assertGreater(result["review_status_counts"]["未校对"], 0)
            self.assertTrue(any("未校对" in item for item in result["activation_blockers"]))

    def test_review_book_validator_rejects_stale_directory_version(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = build_human_review_book._write_book(Path(tmp) / "review.xlsx")
            book = openpyxl.load_workbook(output)
            sheet = book["标准目录"]
            row = next(row for row in range(4, sheet.max_row + 1) if sheet.cell(row, 2).value == "GB 32030-2022")
            sheet.cell(row, 5).value = "2026.08.26-gbt25409-calc-v2"
            book.save(output)
            book.close()
            result = validate(output)
            self.assertFalse(result["is_valid"])
            self.assertTrue(any("数据版本与manifest不一致" in error for error in result["errors"]))

    def test_review_book_validator_rejects_locked_value_drift(self):
        """相同数据ID但标准字段被改写时，校对册必须拒绝通过。"""
        expected_row = [
            "id-locked-value", "TEST", "表1", "测试设备", "{}", "1", "rows[0].value",
            "30.0", "kW", 30.0, "kW", "1级", ">=", "", "1", "pack.json",
            "未校对", "", "",
        ]
        from tools import validate_human_review_book as validator

        with tempfile.TemporaryDirectory() as tmp, patch.object(
            validator, "_standard_rows", return_value=[expected_row]
        ):
            book_path = Path(tmp) / "tampered.xlsx"
            book = openpyxl.Workbook()
            book.active.title = "校对说明"
            for name in sorted(validator.REQUIRED_SHEETS - {"校对说明"}):
                book.create_sheet(name)
            sheet = book["标准数据平铺"]
            sheet.append(["标题"])
            sheet.append([])
            sheet.append(list(validator.FLAT_HEADERS))
            tampered = list(expected_row)
            tampered[7] = "31.0"  # 原文值列属于锁定字段
            sheet.append(tampered)
            for current in book.worksheets:
                current.protection.sheet = True
            for column in (17, 18, 19):
                sheet.cell(row=4, column=column).protection = Protection(locked=False)
            book.save(book_path)
            book.close()

            result = validator.validate(book_path)

        self.assertFalse(result["is_valid"])
        self.assertEqual(result["locked_value_mismatch_count"], 1)
        self.assertTrue(any("锁定字段与机器数据不一致" in error for error in result["errors"]))
        self.assertEqual(result["locked_value_mismatch_samples"][0]["field"], "原文值")

    def test_validator_require_activation_is_a_distinct_exit_gate(self):
        from tools import validate_human_review_book as validator
        with patch.object(validator, "validate", return_value={
            "is_valid": True,
            "activation_ready": False,
            "errors": [],
            "warnings": [],
        }), patch("builtins.print"):
            self.assertEqual(validator.main(["review.xlsx", "--require-activation"]), 3)

    def test_validator_rejects_replacement_for_machine_confirmed_no_data(self):
        """无数据单元格被人工填入数值时必须触发校对册错误。"""
        from tools import validate_human_review_book as validator

        path = "tables[0].rows[0].efficiency.1[5]"
        expected_row = [
            "id-no-data", "GB 30253-2024", "表1", "", "", "", path,
            "", "", "", "", "1级", ">=", "", "", "", "正确", "", "标准原文无数据",
        ]
        with tempfile.TemporaryDirectory() as tmp, patch.object(
            validator, "_standard_rows", return_value=[expected_row]
        ):
            book_path = Path(tmp) / "tampered.xlsx"
            book = openpyxl.Workbook()
            book.active.title = "校对说明"
            for name in sorted(validator.REQUIRED_SHEETS - {"校对说明"}):
                book.create_sheet(name)
            sheet = book["标准数据平铺"]
            sheet.append(["标题"])
            sheet.append([])
            sheet.append(list(validator.FLAT_HEADERS))
            tampered = list(expected_row)
            tampered[17] = 95  # 无数据记录不允许出现建议修订值
            sheet.append(tampered)
            for current in book.worksheets:
                current.protection.sheet = True
            for column in (17, 18, 19):
                sheet.cell(row=4, column=column).protection = Protection(locked=False)
            book.save(book_path)
            book.close()
            result = validator.validate(book_path)
            self.assertFalse(result["is_valid"])
            self.assertTrue(any("不得填写建议修订值" in error for error in result["errors"]))

    def test_builder_accepts_explicit_output_path(self):
        from unittest.mock import patch

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "review.xlsx"
            with patch.object(build_human_review_book, "_write_book", return_value=output) as writer:
                self.assertEqual(build_human_review_book.main(["--output", str(output)]), 0)
            writer.assert_called_once_with(output.resolve())


if __name__ == "__main__":
    unittest.main()
