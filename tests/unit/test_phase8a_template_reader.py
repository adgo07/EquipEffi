"""Phase 8A — V6 模板正式化 / Excel↔Application 一致性 / Reader（专项证据）。

覆盖：
- 模板资产身份（SHA-256 / Sheet 集合与顺序）
- 「离心泵」Sheet 结构契约（27 列 / TblPump / 保护 / editable-locked）
- 类别枚举与 `PUMP_CATEGORIES` 逐项一致（`其他类别` / `不确定类别`）
- 石化泵的「单级/多级」**不**隐含吸入方式
- 空白模板内不存在第二套 GB19762 业务算法
- 其他 17 个设备 Sheet 语义不变
- Reader 十进制保真（`QA-EXCEL-001`）
- Reader 行启用语义（只填安装位置的行不得被静默跳过）
- Reader 不读取旧计算结果作为业务输入
- 容量：少量 / 100 / 1,000 / 10,000 行无截断
"""
from __future__ import annotations

import os
import tempfile
import tracemalloc
import unittest
from decimal import Decimal
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import openpyxl

from equipeffi.application.services.centrifugal_pump_analysis_service import (
    CATEGORY_FIELD_CONSTRAINTS,
    PUMP_CATEGORIES,
    category_field_constraints,
)
from equipeffi.infrastructure.excel.ooxml_reader import OOXMLWorkbook, _parse_number
from equipeffi.infrastructure.excel.pump_workbook_reader import (
    INPUT_COLUMNS,
    PUMP_SHEET,
    RESULT_COLUMNS,
    V6PumpWorkbookReader,
)
from equipeffi.infrastructure.excel.template_resource import (
    REQUIRED_V6_SHEETS,
    V6_TEMPLATE_IDENTITY,
    V6TemplateResource,
)

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = (ROOT / "src" / "equipeffi" / "resources" / "templates"
            / V6_TEMPLATE_IDENTITY["filename"])
OWNER_BASELINE = ROOT / "outputs" / V6_TEMPLATE_IDENTITY["source_filename"]

FORMAL_CATEGORIES = tuple(c.visible_name for c in PUMP_CATEGORIES if c.special is None)
EXPECTED_CATEGORIES = tuple(c.visible_name for c in PUMP_CATEGORIES)

FIRST_DATA_ROW = 4
LAST_DATA_ROW = 103


def _blank_workbook():
    return openpyxl.load_workbook(TEMPLATE, data_only=False)


class TemplateAssetTests(unittest.TestCase):
    """P8-G01：模板资产化与身份。"""

    def test_template_asset_exists_and_hash_matches_identity(self):
        self.assertTrue(TEMPLATE.is_file())
        resource = V6TemplateResource()
        self.assertEqual(resource.validate().sha256.upper(),
                         V6_TEMPLATE_IDENTITY["asset_sha256"])

    def test_owner_baseline_is_untouched(self):
        """构建输入（Owner 指定的 V6 基线）不得被写回或改写。"""

        import hashlib

        if not OWNER_BASELINE.is_file():
            self.skipTest("本地未保留 Owner 指定基线（构建输入为离线资产）")
        actual = hashlib.sha256(OWNER_BASELINE.read_bytes()).hexdigest().upper()
        self.assertEqual(actual, V6_TEMPLATE_IDENTITY["source_sha256"])

    def test_sheet_set_and_order_match_owner_baseline(self):
        self.assertEqual(V6TemplateResource().validate().sheet_names,
                         REQUIRED_V6_SHEETS)

    def test_product_reads_the_repository_asset_not_a_local_dev_path(self):
        """正式产品路径必须使用入仓资产，不得运行时依赖本地开发路径。"""

        resource = V6TemplateResource()
        self.assertTrue(resource.template_path.is_file())
        self.assertEqual(resource.template_path.parent.name, "templates")
        self.assertNotIn("outputs", str(resource.template_path).replace("\\", "/"))


class PumpSheetContractTests(unittest.TestCase):
    """P8-G02：离心泵 Sheet 一致性矩阵。"""

    def setUp(self):
        self.workbook = _blank_workbook()
        self.sheet = self.workbook[PUMP_SHEET]
        self.config = self.workbook["配置"]

    def tearDown(self):
        self.workbook.close()

    def test_pump_table_keeps_27_columns(self):
        table = self.sheet.tables["TblPump"]
        self.assertEqual(len(table.tableColumns), 27)
        self.assertEqual(table.ref, f"A3:AA{LAST_DATA_ROW}")

    def test_excel_category_enum_matches_application_exactly(self):
        excel_values = tuple(
            str(self.config[f"AE{row}"].value).strip()
            for row in range(2, 2 + len(EXPECTED_CATEGORIES)))
        self.assertEqual(excel_values, EXPECTED_CATEGORIES)
        self.assertEqual(len(excel_values), 10)
        self.assertIn("其他类别", excel_values)
        self.assertIn("不确定类别", excel_values)
        self.assertNotIn("其他（请备注说明）", excel_values)

    def test_category_named_range_covers_exactly_the_enum(self):
        defined = self.workbook.defined_names["EV_PUMP_CATEGORY"]
        last = 2 + len(EXPECTED_CATEGORIES) - 1
        self.assertEqual(str(defined.value), f"'配置'!$AE$2:$AE${last}")

    def test_formal_and_special_categories_are_all_present(self):
        excel_values = tuple(
            str(self.config[f"AE{row}"].value).strip()
            for row in range(2, 2 + len(EXPECTED_CATEGORIES)))
        for name in FORMAL_CATEGORIES:
            with self.subTest(category=name):
                self.assertIn(name, excel_values)
        self.assertEqual(set(FORMAL_CATEGORIES) | {"其他类别", "不确定类别"},
                         set(excel_values))

    def test_chemical_suction_is_not_implied_by_stage_count(self):
        """石化类的「单级/多级」不代表吸入方式自动固定。"""

        for name in EXPECTED_CATEGORIES:
            if "石油化工" not in name:
                continue
            with self.subTest(category=name):
                self.assertNotIn("suction", category_field_constraints(name))
        self.assertNotIn("suction",
                         CATEGORY_FIELD_CONSTRAINTS["单级石油化工离心泵"])

    def test_water_suction_is_implied_only_when_category_names_it(self):
        self.assertEqual(category_field_constraints("单级单吸清水离心泵")["suction"], "单吸")
        self.assertEqual(category_field_constraints("单级双吸清水离心泵")["suction"], "双吸")
        self.assertNotIn("suction", category_field_constraints("管道清水离心泵"))

    def test_suction_validation_is_an_enum_not_a_category_offset(self):
        formulas = [str(dv.formula1) for dv in
                    self.sheet.data_validations.dataValidation
                    if str(dv.sqref).startswith("K4:")]
        self.assertTrue(formulas)
        for formula in formulas:
            with self.subTest(formula=formula):
                self.assertIn("EV_SUCTION", formula)
                self.assertNotIn("OFFSET", formula)

    def test_stage_validation_documents_single_stage_categories(self):
        formulas = [str(dv.formula1) for dv in
                    self.sheet.data_validations.dataValidation
                    if str(dv.sqref).startswith("L4:")]
        self.assertTrue(formulas)
        joined = " ".join(formulas)
        for name, constraint in CATEGORY_FIELD_CONSTRAINTS.items():
            if constraint.get("stages") == "1":
                with self.subTest(category=name):
                    self.assertIn(name, joined)
        self.assertIn("其他类别", joined)
        self.assertIn("不确定类别", joined)

    def test_quantity_validation_is_positive_integer(self):
        quantity = [dv for dv in self.sheet.data_validations.dataValidation
                    if str(dv.sqref).startswith("D4:")]
        self.assertTrue(quantity)
        self.assertTrue(any(dv.type == "whole" and str(dv.operator) == "greaterThan"
                            and str(dv.formula1) == "0" for dv in quantity))

    def test_result_columns_are_kept_for_the_writer(self):
        for letter, _field in RESULT_COLUMNS:
            with self.subTest(column=letter):
                self.assertIsNotNone(self.sheet[f"{letter}3"].value)

    def test_result_columns_are_locked_and_input_columns_editable(self):
        for letter, _field in RESULT_COLUMNS:
            with self.subTest(column=letter):
                self.assertTrue(self.sheet[f"{letter}{FIRST_DATA_ROW}"].protection.locked)
        for letter, _field in INPUT_COLUMNS:
            with self.subTest(column=letter):
                cell = self.sheet[f"{letter}{FIRST_DATA_ROW}"]
                self.assertFalse(cell.protection and cell.protection.locked)


class NoSecondBusinessAlgorithmTests(unittest.TestCase):
    """P8-G03：Excel 内不得保留第二套 GB19762 业务算法。"""

    def setUp(self):
        self.workbook = _blank_workbook()
        self.sheet = self.workbook[PUMP_SHEET]

    def tearDown(self):
        self.workbook.close()

    def test_business_formula_columns_are_empty(self):
        offenders = []
        for letter, _field in RESULT_COLUMNS:
            for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
                value = self.sheet[f"{letter}{row}"].value
                if value not in (None, ""):
                    offenders.append(f"{letter}{row}")
        self.assertEqual(offenders, [])

    def test_no_business_algorithm_formula_anywhere_on_pump_sheet(self):
        markers = ("163.33", "168.33", "9.81*", "TEXTJOIN", "比转速", "ns")
        offenders = []
        for row in self.sheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    if any(marker in value for marker in markers):
                        offenders.append(cell.coordinate)
        self.assertEqual(offenders, [])

    def test_no_cross_sheet_reference_to_pump_result_columns(self):
        offenders = []
        for name in self.workbook.sheetnames:
            if name == PUMP_SHEET:
                continue
            for row in self.workbook[name].iter_rows():
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and value.startswith("=") and PUMP_SHEET in value:
                        offenders.append(f"{name}!{cell.coordinate}")
        self.assertEqual(offenders, [])

    def test_other_device_sheets_keep_their_own_structure(self):
        """本阶段未接通其他设备：其 Sheet 与表格必须原样保留。"""

        for name in REQUIRED_V6_SHEETS:
            if name == PUMP_SHEET:
                continue
            with self.subTest(sheet=name):
                self.assertIn(name, self.workbook.sheetnames)
        # 其他设备自身的算法不属于本阶段授权范围，不得在 Phase 8 被删改。
        self.assertTrue(self.workbook["变压器"].tables)


class ReaderNumericFidelityTests(unittest.TestCase):
    """P8-G04：`QA-EXCEL-001` — Reader 必须保持十进制语义。"""

    def test_numbers_never_pass_through_float(self):
        self.assertIsInstance(_parse_number("0.1"), Decimal)
        self.assertNotIsInstance(_parse_number("0.1"), float)

    def test_long_decimal_is_preserved_exactly(self):
        text = "0.12345678901234567890123456789012345"
        parsed = _parse_number(text)
        self.assertEqual(parsed, Decimal(text))
        self.assertNotEqual(str(parsed), str(float(Decimal(text))))

    def test_integers_stay_integers(self):
        self.assertIsInstance(_parse_number("42"), int)
        self.assertEqual(_parse_number("42"), 42)
        self.assertEqual(_parse_number("42.0"), 42)

    def test_scientific_notation_keeps_decimal_semantics(self):
        self.assertEqual(_parse_number("1e-5"), Decimal("0.00001"))
        self.assertEqual(_parse_number("1.5E3"), Decimal("1500"))

    def test_large_integer_is_not_rewritten(self):
        text = "12345678901234567890"
        parsed = _parse_number(text)
        self.assertEqual(parsed, int(text))
        self.assertNotEqual(str(parsed), str(float(Decimal(text))))

    def test_text_and_empty_stay_text(self):
        self.assertEqual(_parse_number(""), "")
        self.assertEqual(_parse_number("abc"), "abc")

    def test_workbook_read_returns_decimal_for_fractional_cells(self):
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "numeric.xlsx"
            workbook = _blank_workbook()
            sheet = workbook[PUMP_SHEET]
            sheet[f"G{FIRST_DATA_ROW}"] = Decimal("0.123456789012345678901234567890")
            sheet[f"B{FIRST_DATA_ROW}"] = "精度测试"
            workbook.save(target)
            workbook.close()
            with OOXMLWorkbook(target) as book:
                grid = book.rows(PUMP_SHEET)
            value = grid[FIRST_DATA_ROW + 2][6]
            self.assertNotIsInstance(value, float)


class ReaderRowEnablementTests(unittest.TestCase):
    """P8-G04：行启用是语义式的，不依赖固定行范围。"""

    def _read(self, mutate) -> tuple:
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "rows.xlsx"
            workbook = _blank_workbook()
            mutate(workbook[PUMP_SHEET])
            workbook.save(target)
            workbook.close()
            return V6PumpWorkbookReader().read(target).rows

    def test_fully_blank_rows_are_skipped(self):
        rows = self._read(lambda sheet: None)
        self.assertEqual(rows, ())

    def test_row_with_only_installation_location_is_read_not_skipped(self):
        """只填「安装位置」的行不得被静默跳过，必须交给软件验证。"""

        def mutate(sheet):
            sheet[f"E{FIRST_DATA_ROW}"] = "1号车间"

        rows = self._read(mutate)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].values["location"], "1号车间")
        self.assertEqual(rows[0].values["category"], "")

    def test_row_beyond_the_table_range_is_still_read(self):
        """超过 `TblPump` 100 行的数据必须被读取（不得有业务行上限）。"""

        def mutate(sheet):
            sheet[f"B{FIRST_DATA_ROW + 500}"] = "表外设备"
            sheet[f"D{FIRST_DATA_ROW + 500}"] = 2

        rows = self._read(mutate)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].row_number, FIRST_DATA_ROW + 500)
        self.assertEqual(rows[0].values["device_name"], "表外设备")

    def test_row_number_is_one_based_excel_coordinate(self):
        def mutate(sheet):
            sheet[f"B{FIRST_DATA_ROW}"] = "首行"
            sheet[f"B{FIRST_DATA_ROW + 9}"] = "第十行"

        rows = self._read(mutate)
        self.assertEqual([row.row_number for row in rows],
                         [FIRST_DATA_ROW, FIRST_DATA_ROW + 9])

    def test_reader_ignores_legacy_result_columns(self):
        """旧计算结果不得成为业务输入（Reader 只读正式输入列）。"""

        def mutate(sheet):
            sheet[f"B{FIRST_DATA_ROW}"] = "含旧结果"
            # 伪造旧结果列（旧比转速 / 旧 C1 / 旧等级 / 旧自动备注）
            sheet[f"N{FIRST_DATA_ROW}"] = 999.99
            sheet[f"O{FIRST_DATA_ROW}"] = 111.11
            sheet[f"X{FIRST_DATA_ROW}"] = "1级"
            sheet[f"AA{FIRST_DATA_ROW}"] = "旧的自动备注"

        rows = self._read(mutate)
        self.assertEqual(len(rows), 1)
        read_fields = set(rows[0].values)
        for letter, field_id in RESULT_COLUMNS:
            with self.subTest(field=field_id):
                self.assertNotIn(field_id, read_fields)
        self.assertNotIn("999.99", str(rows[0].values))
        self.assertNotIn("1级", str(rows[0].values))

    def test_header_contract_is_defensively_checked(self):
        """表头被误改时必须报错，而不是静默读错列。"""

        from equipeffi.infrastructure.excel.ooxml_reader import OOXMLReadError

        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "broken.xlsx"
            workbook = _blank_workbook()
            workbook[PUMP_SHEET]["F3"] = "被改坏的列"
            workbook.save(target)
            workbook.close()
            with self.assertRaises(OOXMLReadError):
                V6PumpWorkbookReader(require_template_structure=False).read(target)


class ReaderCapacityTests(unittest.TestCase):
    """P8-G03：统一容量策略 —— 少量 / 100 / 1,000 / 10,000 行。"""

    CATEGORY_CYCLE = EXPECTED_CATEGORIES

    def _build(self, target: Path, count: int) -> None:
        workbook = _blank_workbook()
        sheet = workbook[PUMP_SHEET]
        for index in range(count):
            row = FIRST_DATA_ROW + index
            sheet[f"B{row}"] = f"设备{index + 1}"
            sheet[f"C{row}"] = f"M-{index + 1:05d}"
            sheet[f"D{row}"] = (index % 5) + 1
            sheet[f"E{row}"] = "1号车间"
            sheet[f"F{row}"] = self.CATEGORY_CYCLE[index % len(self.CATEGORY_CYCLE)]
            sheet[f"G{row}"] = 100 + index
            sheet[f"H{row}"] = 30 + (index % 20)
            sheet[f"I{row}"] = 2900
            sheet[f"J{row}"] = 45 + (index % 30)
            sheet[f"K{row}"] = "单吸"
            sheet[f"L{row}"] = 1
            sheet[f"M{row}"] = 75 + (index % 15)
        workbook.save(target)
        workbook.close()

    def test_capacity_scenarios_read_without_truncation(self):
        for count in (3, 100, 1000, 10000):
            with self.subTest(rows=count):
                with tempfile.TemporaryDirectory() as tmp:
                    target = Path(tmp) / f"pump_{count}.xlsx"
                    self._build(target, count)
                    tracemalloc.start()
                    workbook = V6PumpWorkbookReader().read(target)
                    _, peak = tracemalloc.get_traced_memory()
                    tracemalloc.stop()
                    self.assertEqual(len(workbook.rows), count)
                    self.assertEqual(workbook.rows[0].row_number, FIRST_DATA_ROW)
                    self.assertEqual(workbook.rows[-1].row_number,
                                     FIRST_DATA_ROW + count - 1)
                    # 记录真实代价（不预设阈值；不得 hang / 数据丢失）
                    self.assertLess(peak, 2 * 1024 ** 3)


if __name__ == "__main__":
    unittest.main()
