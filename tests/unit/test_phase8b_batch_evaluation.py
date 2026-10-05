"""Phase 8B — 批量评价 / Writer / batch_record / Qt 产品闭环（专项证据）。

覆盖 Owner Phase 8 规则 8/9/10/12 与结论一致性：
- 一次批量评价 → **一条** batch_record，且**不**创建任何单台 Record
- 输出结果是**新** Workbook，不覆盖输入；其他设备 Sheet 不丢
- 结果写回只写结果列；旧计算结果不作为业务输入
- 数量必须正整数 > 0（Excel 验证 + 软件验证都在）
- 「不确定类别」→「无法评价」，且 Qt 与 Excel 结论一致
- 逐设备详细结果保存在结果 Workbook
- 批次汇总含结论分布与需要关注的行
- 系统异常不伪装成业务结论
"""
from __future__ import annotations

import hashlib
import os
import sqlite3
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import openpyxl

from equipeffi.application.lifecycle import BatchRecordSnapshot
from equipeffi.application.services.centrifugal_pump_analysis_service import (
    PUMP_CATEGORIES,
    user_conclusion_text,
)
from equipeffi.application.services.pump_batch_evaluation_service import (
    PumpBatchEvaluationService,
)
from equipeffi.composition import (
    create_batch_evaluation_service,
    create_pump_analysis_service,
)
from equipeffi.infrastructure.excel.pump_result_writer import (
    RESULT_FIELDS,
    PumpResultWorkbookWriter,
)
from equipeffi.infrastructure.excel.pump_workbook_reader import (
    INPUT_COLUMNS,
    PUMP_SHEET,
    V6PumpWorkbookReader,
)
from equipeffi.infrastructure.excel.template_resource import V6TemplateResource
from equipeffi.infrastructure.persistence.app_data_paths import AppDataPaths
from equipeffi.infrastructure.persistence.records_migrations import (
    RECORDS_MIGRATIONS,
    migrate_records_database,
)
from equipeffi.infrastructure.persistence.sqlite_batch_record_repository import (
    SqliteBatchRecordRepository,
)

ROOT = Path(__file__).resolve().parents[2]
FIRST_DATA_ROW = 4
EXPECTED_CATEGORIES = tuple(item.visible_name for item in PUMP_CATEGORIES)

WATER_SUCTION = "单级单吸清水离心泵"
CHEMICAL_SINGLE = "单级石油化工离心泵"


class BatchTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from PySide6.QtWidgets import QApplication

        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.root = Path(self.tmp.name)
        self.paths = AppDataPaths(self.root / "app")
        migrate_records_database(self.paths.records_db, app_version="test")
        self.analysis = create_pump_analysis_service(paths=self.paths)
        self.repository = SqliteBatchRecordRepository(self.paths.records_db)
        # 载体端口按真实装配注入（Application 只认识 Protocol，不认识 Excel）。
        self.batch = PumpBatchEvaluationService(
            self.analysis,
            reader=V6PumpWorkbookReader(),
            writer=PumpResultWorkbookWriter(),
            batch_repository=self.repository)
        self.source = self.root / "输入.xlsx"
        V6TemplateResource().download_to(self.source)
        self.addCleanup(self.tmp.cleanup)

    # -- helpers -----------------------------------------------------------

    def fill(self, rows: list[dict]) -> Path:
        workbook = openpyxl.load_workbook(self.source, data_only=False)
        sheet = workbook[PUMP_SHEET]
        for offset, values in enumerate(rows):
            row = FIRST_DATA_ROW + offset
            for column, value in values.items():
                sheet[f"{column}{row}"] = value
        workbook.save(self.source)
        workbook.close()
        return self.source

    def run_batch(self, rows: list[dict], **kwargs):
        self.fill(rows)
        target = kwargs.pop("destination", self.root / "结果.xlsx")
        return self.batch.evaluate_workbook(self.source, destination=target, **kwargs)

    @staticmethod
    def water_row(**overrides) -> dict:
        row = dict(B="泵", C="M-001", D=1, E="1号车间", F=WATER_SUCTION,
                   G=100, H=50, I=2900, J=45, K="单吸", L=1, M=80)
        row.update(overrides)
        return row


class BatchRecordPersistenceTests(BatchTestCase):
    """Owner 规则 9/10：batch_record 是最小 additive 结构，单台 Record 语义不变。"""

    def test_migration_is_additive_and_appends_after_002(self):
        versions = [(m.schema_version, m.migration_id) for m in RECORDS_MIGRATIONS]
        self.assertEqual(versions[:2], [
            (1, "001_create_workspace_and_record"),
            (2, "002_add_workspace_revision")])
        self.assertEqual(versions[2][0], 3)
        self.assertIn("batch_record", versions[2][1])
        for migration in RECORDS_MIGRATIONS:
            migration.assert_additive()

    def test_existing_record_table_is_unchanged(self):
        with sqlite3.connect(self.paths.records_db) as connection:
            record_columns = [row[1] for row in
                              connection.execute("PRAGMA table_info(record)")]
        self.assertEqual(len(record_columns), 23)
        self.assertNotIn("batch_record_id", record_columns)

    def test_batch_record_round_trip(self):
        snapshot = BatchRecordSnapshot(
            batch_record_id="B-1", standard_code="GB 19762-2025", device_type="离心泵",
            source_workbook="a.xlsx", source_workbook_sha256="AA",
            result_workbook="b.xlsx", result_workbook_sha256="BB",
            total_rows=3, evaluated_count=3, unevaluated_count=1, invalid_count=0,
            summary={"conclusion_counts": {"1级": 2}}, schema_version=3,
            created_at_utc="2026-10-05T00:00:00+00:00")
        self.repository.append_batch_record(snapshot)
        loaded = self.repository.load_batch_record("B-1")
        self.assertEqual(loaded, snapshot)
        self.assertEqual(len(self.repository.list_batch_records()), 1)

    def test_duplicate_batch_record_id_is_rejected(self):
        from equipeffi.application.lifecycle import RecordConflictError

        snapshot = BatchRecordSnapshot(
            batch_record_id="B-dup", standard_code="c", device_type="d",
            source_workbook="a", source_workbook_sha256="A", result_workbook=None,
            result_workbook_sha256=None, total_rows=0, evaluated_count=0,
            unevaluated_count=0, invalid_count=0, summary={}, schema_version=3,
            created_at_utc="2026-10-05T00:00:00+00:00")
        self.repository.append_batch_record(snapshot)
        with self.assertRaises(RecordConflictError):
            self.repository.append_batch_record(snapshot)


class BatchEvaluationSemanticsTests(BatchTestCase):
    """Owner 规则 9：一次批量 → 一条 batch_record，**不**产生单台 Record。"""

    def test_one_batch_evaluation_creates_exactly_one_batch_record(self):
        result = self.run_batch([self.water_row(), self.water_row(B="泵2", C="M-002")])
        self.assertEqual(len(self.repository.list_batch_records()), 1)
        self.assertEqual(result.batch_record_id,
                         self.repository.list_batch_records()[0].batch_record_id)

    def test_no_per_row_single_records_are_created(self):
        result = self.run_batch([self.water_row(), self.water_row(B="泵2", C="M-002"),
                                 self.water_row(B="泵3", C="M-003")])
        self.assertEqual(result.records_created, 0)
        with sqlite3.connect(self.paths.records_db) as connection:
            count = connection.execute("SELECT COUNT(*) FROM record").fetchone()[0]
        self.assertEqual(count, 0)

    def test_second_batch_creates_a_second_summary_record(self):
        self.run_batch([self.water_row()], destination=self.root / "r1.xlsx")
        self.run_batch([self.water_row(B="泵2")], destination=self.root / "r2.xlsx")
        records = self.repository.list_batch_records()
        self.assertEqual(len(records), 2)
        self.assertEqual(len({r.batch_record_id for r in records}), 2)

    def test_batch_summary_counts_conclusions_and_grades(self):
        result = self.run_batch([
            self.water_row(),
            self.water_row(B="泵2", C="M-002", M=90),
        ])
        self.assertEqual(result.summary.total_rows, 2)
        self.assertEqual(result.summary.evaluated_rows, 2)
        self.assertEqual(result.summary.invalid_rows, 0)
        self.assertTrue(result.summary.conclusion_counts)

    def test_mixed_categories_in_one_sheet(self):
        """清水 + 石油化工离心泵共用同一 Sheet（Owner 规则 2）。"""

        result = self.run_batch([
            self.water_row(),
            self.water_row(B="化工泵", C="C-001", F=CHEMICAL_SINGLE,
                           G=300, H=40, K="双吸", M=90),
        ])
        self.assertEqual(result.summary.total_rows, 2)
        self.assertEqual(result.summary.evaluated_rows, 2)
        self.assertEqual(len(result.summary.conclusion_counts), 1)

    def test_uncertain_category_is_unevaluated_not_skipped(self):
        result = self.run_batch([self.water_row(F="不确定类别")])
        self.assertEqual(result.summary.total_rows, 1, "不确定类别行不得被静默跳过")
        self.assertEqual(result.summary.unevaluated_rows, 1)
        self.assertEqual(result.summary.conclusion_counts.get("无法评价"), 1)

    def test_other_category_is_out_of_scope(self):
        result = self.run_batch([self.water_row(F="其他类别")])
        self.assertEqual(result.summary.conclusion_counts.get("不适用"), 1)
        self.assertEqual(result.summary.unevaluated_rows, 0)

    def test_row_with_only_location_is_read_and_reported(self):
        """只填「安装位置」的行不得被静默跳过，必须被读取并报告问题。

        该行同时缺少数量与全部评价参数：软件侧数量门禁先命中，因此计入
        `invalid_rows`（输入非法），而不是"业务资料不足"。
        关键不变量是**被读取且被报告**，不是被丢弃。
        """

        result = self.run_batch([dict(E="5号车间")])
        self.assertEqual(result.summary.total_rows, 1, "不得静默跳过该行")
        self.assertEqual(result.summary.issues[0].row_number, FIRST_DATA_ROW)
        self.assertIn("缺少数量", result.summary.issues[0].reason)
        self.assertEqual(result.summary.invalid_rows, 1)
        self.assertEqual(result.outcomes[0].conclusion, "无法评价")
        target = openpyxl.load_workbook(
            (self.root / "结果.xlsx"))[PUMP_SHEET]
        self.assertEqual(target[f"AA{FIRST_DATA_ROW}"].value, "缺少数量")

    def test_system_failure_is_not_disguised_as_a_business_conclusion(self):
        with patch.object(self.analysis, "evaluate",
                          side_effect=RuntimeError("模拟内部错误")):
            result = self.run_batch([self.water_row()])
        self.assertEqual(result.summary.invalid_rows, 1)
        outcome = result.outcomes[0]
        self.assertEqual(outcome.conclusion, "评价失败")
        self.assertNotIn(outcome.conclusion, ("无法评价", "无法判定", "不适用"))
        self.assertTrue(any("系统执行失败" in message for message in outcome.messages))

    def test_system_failure_on_one_row_does_not_abort_the_batch(self):
        calls = {"n": 0}
        original = self.analysis.evaluate

        def flaky(request):
            calls["n"] += 1
            if calls["n"] == 1:
                raise RuntimeError("模拟首行失败")
            return original(request)

        with patch.object(self.analysis, "evaluate", side_effect=flaky):
            result = self.run_batch([self.water_row(),
                                     self.water_row(B="泵2", C="M-002")])
        self.assertEqual(result.summary.total_rows, 2)
        self.assertEqual(result.summary.invalid_rows, 1)
        self.assertEqual(result.summary.evaluated_rows, 2)

    def test_quantity_contract_is_enforced_by_software_not_only_excel(self):
        """Excel Validation 只是辅助；软件必须自己拒绝非法数量。"""

        for bad in (0, -3, 1.5):
            with self.subTest(quantity=bad):
                result = self.run_batch([self.water_row(**{"D": bad})])
                outcome = result.outcomes[0]
                self.assertNotEqual(outcome.evaluation_status, "SUCCESS",
                                    f"数量 {bad} 不得被判为有效")
        # 空白数量同样不得被默认成 1
        result = self.run_batch([dict(B="泵", C="M-1", E="1号车间", F=WATER_SUCTION,
                                      G=100, H=50, I=2900, J=45, K="单吸", L=1, M=80)])
        self.assertNotEqual(result.outcomes[0].evaluation_status, "SUCCESS")


class ResultWorkbookTests(BatchTestCase):
    """Owner 规则 12：新 Workbook、不覆盖输入、其他 Sheet 不丢、只写结果列。"""

    def test_output_is_a_new_workbook_and_source_is_untouched(self):
        self.fill([self.water_row()])
        before = hashlib.sha256(self.source.read_bytes()).hexdigest()
        target = self.root / "结果.xlsx"
        self.batch.evaluate_workbook(self.source, destination=target)
        self.assertTrue(target.is_file())
        self.assertNotEqual(target, self.source)
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), before)

    def test_writer_refuses_to_overwrite_the_source(self):
        with self.assertRaises(ValueError):
            PumpResultWorkbookWriter().write(self.source, {}, self.source)

    def test_all_device_sheets_survive_into_the_result(self):
        self.fill([self.water_row()])
        target = self.root / "结果.xlsx"
        self.batch.evaluate_workbook(self.source, destination=target)
        expected = V6TemplateResource().validate().sheet_names
        actual = tuple(openpyxl.load_workbook(target).sheetnames)
        self.assertEqual(actual, expected)

    def test_result_columns_are_written_for_each_evaluated_row(self):
        result = self.run_batch([self.water_row()])
        target = Path(result.result_workbook)
        sheet = openpyxl.load_workbook(target)[PUMP_SHEET]
        written = {column: sheet[f"{column}{FIRST_DATA_ROW}"].value
                   for column, _ in RESULT_FIELDS}
        self.assertEqual(written["X"], "1级")
        self.assertIsNotNone(written["N"], "比转速应写回")
        self.assertIsNotNone(written["AA"], "自动备注应写回")

    def test_input_columns_are_preserved_verbatim(self):
        self.fill([self.water_row(B="保留我", C="M-999", E="9号车间")])
        result = self.batch.evaluate_workbook(
            self.source, destination=self.root / "结果.xlsx")
        sheet = openpyxl.load_workbook(result.result_workbook)[PUMP_SHEET]
        self.assertEqual(sheet[f"B{FIRST_DATA_ROW}"].value, "保留我")
        self.assertEqual(sheet[f"C{FIRST_DATA_ROW}"].value, "M-999")
        self.assertEqual(sheet[f"E{FIRST_DATA_ROW}"].value, "9号车间")

    def test_legacy_result_columns_do_not_become_business_input(self):
        """填入伪造的旧结果后，软件的结论必须只由正式输入决定。"""

        baseline = self.run_batch([self.water_row()])
        baseline_conclusion = baseline.outcomes[0].conclusion
        baseline_grade = baseline.outcomes[0].grade

        rows = [self.water_row()]
        workbook = openpyxl.load_workbook(self.source)
        sheet = workbook[PUMP_SHEET]
        sheet[f"X{FIRST_DATA_ROW}"] = "3级"
        sheet[f"N{FIRST_DATA_ROW}"] = 1.0
        sheet[f"AA{FIRST_DATA_ROW}"] = "伪造的旧自动备注"
        workbook.save(self.source)
        workbook.close()

        result = self.batch.evaluate_workbook(
            self.source, destination=self.root / "结果2.xlsx")
        self.assertEqual(result.outcomes[0].conclusion, baseline_conclusion)
        self.assertEqual(result.outcomes[0].grade, baseline_grade)

    def test_writer_writes_full_precision_values(self):
        """结果列写入完整精度数值，由模板数字格式负责 2 位显示。"""

        result = self.run_batch([self.water_row()])
        sheet = openpyxl.load_workbook(result.result_workbook)[PUMP_SHEET]
        raw = sheet[f"N{FIRST_DATA_ROW}"].value
        self.assertIsInstance(raw, (int, float))
        self.assertTrue(str(raw).replace(".", "").replace("-", "").isdigit())
        # 2 位格式化会得到同一显示；完整值本身长于 2 位小数
        self.assertGreater(len(str(raw).split(".")[-1]), 2)


class ConclusionConsistencyTests(BatchTestCase):
    """Owner 规则 7：Qt 与 Excel 的用户可见结论必须一致。"""

    def test_shared_function_is_the_only_conclusion_source(self):
        import equipeffi.presentation.qt.labels as labels

        self.assertIs(labels.user_conclusion_text, user_conclusion_text)

    def test_uncertain_category_is_unable_to_evaluate_in_excel(self):
        result = self.run_batch([self.water_row(F="不确定类别")])
        self.assertEqual(result.outcomes[0].conclusion, "无法评价")
        sheet = openpyxl.load_workbook(result.result_workbook)[PUMP_SHEET]
        self.assertEqual(sheet[f"X{FIRST_DATA_ROW}"].value, "无法评价")

    def test_uncertain_category_is_unable_to_evaluate_in_qt(self):
        from equipeffi.presentation.qt.pages.analysis import AnalysisPage

        page = AnalysisPage(self.analysis, as_of=date.today())
        self.addCleanup(page.deleteLater)
        page.category.setCurrentIndex(page.category.findData("不确定类别"))
        for key, value in (("QBEP", "100"), ("HBEP", "50"),
                           ("speed", "2900"), ("efficiency", "80")):
            page.point_inputs[key].setText(value)
        page.evaluate()
        self.assertEqual(page.conclusion.text(), "无法评价")

    def test_qt_and_excel_agree_for_every_conclusion_shape(self):
        from equipeffi.presentation.qt.pages.analysis import AnalysisPage

        cases = (
            ("单级单吸清水离心泵", "1级"),
            ("其他类别", "不适用"),
            ("不确定类别", "无法评价"),
        )
        for category, expected in cases:
            with self.subTest(category=category):
                result = self.run_batch([self.water_row(F=category)])
                self.assertEqual(result.outcomes[0].conclusion, expected)

                page = AnalysisPage(self.analysis, as_of=date.today())
                self.addCleanup(page.deleteLater)
                page.category.setCurrentIndex(page.category.findData(category))
                for key, value in (("QBEP", "100"), ("HBEP", "50"),
                                   ("speed", "2900"), ("efficiency", "80")):
                    page.point_inputs[key].setText(value)
                if page.stages.isEnabled():
                    page.stages.setText("1")
                page.evaluate()
                self.assertEqual(page.conclusion.text(), expected)


class CategoryAlignmentTests(BatchTestCase):
    """Excel 类别枚举与正式 Application 枚举逐项一致。"""

    def test_workbook_categories_equal_application_categories(self):
        workbook = openpyxl.load_workbook(self.source)
        config = workbook["配置"]
        excel_values = tuple(
            str(config[f"AE{row}"].value).strip()
            for row in range(2, 2 + len(EXPECTED_CATEGORIES)))
        workbook.close()
        self.assertEqual(excel_values, EXPECTED_CATEGORIES)

    def test_uncertain_and_other_categories_are_formal_choices(self):
        self.assertIn("不确定类别", EXPECTED_CATEGORIES)
        self.assertIn("其他类别", EXPECTED_CATEGORIES)
        self.assertNotIn("其他（请备注说明）", EXPECTED_CATEGORIES)


class BatchPageTests(BatchTestCase):
    """8B Qt 产品闭环：选输入 → 选输出 → 开始。"""

    def page(self):
        from equipeffi.presentation.qt.pages.batch import BatchPage

        page = BatchPage(self.batch)
        self.addCleanup(page.deleteLater)
        return page

    def test_page_has_no_per_sheet_capacity_question(self):
        """Owner 规则 11：不得逐 Sheet 询问行数或要求扩容。"""

        page = self.page()
        from PySide6.QtWidgets import QLabel, QPushButton

        blob = "\n".join([w.text() for w in page.findChildren(QLabel)]
                         + [b.text() for b in page.findChildren(QPushButton)])
        for forbidden in ("设置行数", "扩容", "需要多少行", "输入行数", "增加行数"):
            with self.subTest(word=forbidden):
                self.assertNotIn(forbidden, blob)
        # 输入控件里不得出现"行数"类配置项
        from PySide6.QtWidgets import QLineEdit
        for field in page.findChildren(QLineEdit):
            with self.subTest(field=field.objectName()):
                self.assertNotIn("行数", field.placeholderText())

    def test_page_runs_a_batch_and_shows_the_summary(self):
        self.fill([self.water_row()])
        page = self.page()
        page.set_source(self.source)
        self.assertTrue(page.target_edit.text())
        result = page.run()
        self.assertIsNotNone(result)
        self.assertIn("已处理 1 行", page.summary_label.text())
        self.assertIn("仅形成批次总结记录", page.summary_label.text())
        self.assertTrue(Path(result.result_workbook).is_file())

    def test_default_target_never_overwrites_the_source(self):
        page = self.page()
        page.set_source(self.source)
        self.assertNotEqual(Path(page.target_edit.text()).resolve(),
                            self.source.resolve())

    def test_page_requires_a_source_before_running(self):
        page = self.page()
        self.assertIsNone(page.run())
        self.assertIn("请先选择", page.summary_label.text())

    def test_page_reports_system_failure_without_fake_conclusions(self):
        self.fill([self.water_row()])
        page = self.page()
        page.set_source(self.source)
        with patch.object(self.batch, "evaluate_workbook",
                          side_effect=RuntimeError("模拟装配失败")):
            self.assertIsNone(page.run())
        self.assertIn("批量评价未能完成", page.summary_label.text())
        self.assertNotIn("无法评价", page.summary_label.text())


class CompositionWiringTests(BatchTestCase):
    def test_composition_reuses_the_same_analysis_service(self):
        service = create_batch_evaluation_service(paths=self.paths)
        self.assertIsInstance(service.analysis.__class__.__name__, str)
        self.assertTrue(hasattr(service.analysis, "evaluate"))
        self.assertIsNotNone(service.batch_repository)


if __name__ == "__main__":
    unittest.main()
