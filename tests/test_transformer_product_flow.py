"""真实装配的混合工作簿：单台/批量同核、记录与旧模板兼容。"""
import tempfile
import unittest
from pathlib import Path
from hashlib import sha256

from openpyxl import load_workbook

from equipeffi.composition import create_batch_evaluation_service, create_transformer_analysis_service
from equipeffi.application.services.transformer_analysis_service import TransformerAnalysisRequest
from equipeffi.infrastructure.excel.template_resource import V6TemplateResource
from equipeffi.infrastructure.persistence.app_data_paths import AppDataPaths
from equipeffi.infrastructure.persistence.sqlite_batch_record_repository import SqliteBatchRecordRepository


class TransformerProductFlowTests(unittest.TestCase):
    def test_real_mixed_workbook_matches_single_analysis_and_saves_one_batch(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = AppDataPaths(Path(directory) / "data")
            batch = create_batch_evaluation_service(paths=paths)
            source = Path(directory) / "input.xlsx"
            batch.template_resource.download_to(source)
            # 仅改临时用户输入副本，不修改内置模板。
            workbook = load_workbook(source)
            transformer = create_transformer_analysis_service(with_persistence=False)
            category = transformer.input_contract.categories[2]
            sheet = workbook["变压器"]
            for cell, value in {"D4": 3, "F4": category.label, "G4": 3150,
                                "J4": 1700, "K4": 20700}.items():
                sheet[cell] = value
            workbook["离心泵"]["D4"] = 2
            workbook["离心泵"]["F4"] = "其他类别"
            workbook.save(source)
            workbook.close()
            before = sha256(source.read_bytes()).hexdigest()
            result = batch.evaluate_workbook(source)
            self.assertEqual(result.batch_record_error, "")
            self.assertEqual(result.summary.data_row_count, 2)
            self.assertEqual(result.summary.total_quantity, 5)
            self.assertEqual(result.summary.conclusion_quantities, {"不在范围": 2, "1级": 3})
            self.assertEqual(sha256(source.read_bytes()).hexdigest(), before)
            expected = transformer.evaluate(TransformerAnalysisRequest({
                "category_id": category.category_id, "capacity_kva": "3150",
                "no_load_loss_w": "1700", "load_loss_w": "20700"}))
            output = load_workbook(result.result_workbook, data_only=True)
            try:
                self.assertEqual(output["变压器"]["R4"].value, expected.ui_conclusion)
                self.assertEqual(output["变压器"]["L4"].value, 1700)
            finally:
                output.close()
            record = SqliteBatchRecordRepository(paths.records_db).load_batch_record(result.batch_record_id)
            self.assertEqual(record.standard_code, "MULTI")
            transformer_summary = next(item for item in record.summary["sheets"] if item["sheet_name"] == "变压器")
            self.assertEqual(transformer_summary["evaluation_provenance"][0]["calculator_version"],
                             expected.references["calculator_version"])

    def test_old_v6_pump_only_stays_readable_and_transformer_is_untouched(self):
        with tempfile.TemporaryDirectory() as directory:
            batch = create_batch_evaluation_service(paths=AppDataPaths(Path(directory) / "data"))
            source = Path(directory) / "old.xlsx"
            V6TemplateResource().download_to(source)
            workbook = load_workbook(source)
            workbook["离心泵"]["D4"] = 1
            workbook["离心泵"]["F4"] = "其他类别"
            workbook["变压器"]["B4"] = "旧表用户数据保留"
            workbook.save(source)
            workbook.close()
            result = batch.evaluate_workbook(source)
            self.assertEqual(result.batch_record_error, "")
            self.assertEqual(result.summary.total_quantity, 1)
            self.assertIn("变压器", result.summary.unprocessed_sheets)
            record = SqliteBatchRecordRepository(batch.batch_repository.database).load_batch_record(result.batch_record_id)
            self.assertEqual(record.device_type, "离心泵")
            self.assertEqual(record.standard_code, "GB 19762-2025")
            output = load_workbook(result.result_workbook)
            try:
                self.assertEqual(output["变压器"]["B4"].value, "旧表用户数据保留")
            finally:
                output.close()
