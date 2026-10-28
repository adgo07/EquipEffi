"""真实三设备编排、精确阈值与候选边界；不替代WPS交互验收。"""
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from equipeffi.composition import (create_batch_evaluation_service,
    create_motor_analysis_service, motor_candidate_enabled)
from equipeffi.application.services.motor_analysis_service import MotorAnalysisRequest
from equipeffi.infrastructure.persistence.app_data_paths import AppDataPaths
from equipeffi.infrastructure.persistence.sqlite_batch_record_repository import SqliteBatchRecordRepository


class MotorProductFlowTests(unittest.TestCase):
    def setUp(self):
        environment = patch.dict(os.environ)
        environment.start()
        self.addCleanup(environment.stop)
        os.environ.pop("EQUIPEFFI_ENABLE_MOTOR_CANDIDATE", None)

    def test_environment_disabled_batch_keeps_pump_transformer_and_rejects_motor(self):
        from equipeffi.infrastructure.excel.template_resource import V62MotorTemplateResource

        with tempfile.TemporaryDirectory() as directory, patch.dict(
                os.environ, {"EQUIPEFFI_ENABLE_MOTOR_CANDIDATE": "0"}):
            paths = AppDataPaths(Path(directory) / "data")
            batch = create_batch_evaluation_service(paths=paths)
            source = Path(directory) / "legacy.xlsx"
            batch.template_resource.download_to(source)
            self.assertEqual(batch.template_resource.validate().identity["template_version"], "V6.1")
            workbook = load_workbook(source)
            for sheet, quantity in (("离心泵", 2), ("变压器", 3)):
                workbook[sheet]["D4"] = quantity
                workbook[sheet]["F4"] = "其他类别"
            workbook.save(source)
            workbook.close()
            result = batch.evaluate_workbook(source)
            self.assertEqual(result.summary.total_quantity, 5)
            self.assertEqual(result.summary.conclusion_quantities, {"不在范围": 5})
            self.assertEqual(result.batch_record_error, "")
            output = load_workbook(result.result_workbook, data_only=True)
            self.assertEqual(output["离心泵"]["X4"].value, "不在范围")
            self.assertEqual(output["变压器"]["R4"].value, "不在范围")
            output.close()
            saved = batch.batch_repository.list_batch_records()
            self.assertEqual(len(saved), 1)
            record = batch.batch_repository.load_batch_record(result.batch_record_id)
            self.assertEqual(record.summary["total_quantity"], 5)

            motor_source = Path(directory) / "motor.xlsx"
            V62MotorTemplateResource().download_to(motor_source)
            workbook = load_workbook(motor_source)
            workbook["电动机"]["D4"] = 1
            workbook["电动机"]["F4"] = "其他类别"
            workbook.save(motor_source)
            workbook.close()
            before = motor_source.read_bytes()
            destination = Path(directory) / "rejected.xlsx"
            with self.assertRaisesRegex(ValueError, "未启用"):
                batch.evaluate_workbook(motor_source, destination=destination)
            self.assertFalse(destination.exists())
            self.assertEqual(motor_source.read_bytes(), before)
            self.assertEqual(len(batch.batch_repository.list_batch_records()), 1)

    def test_frequency_range_and_pole_scope_survive_excel_and_history(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = AppDataPaths(Path(directory) / "data")
            batch = create_batch_evaluation_service(paths=paths)
            source = Path(directory) / "feedback.xlsx"
            batch.template_resource.download_to(source)
            workbook = load_workbook(source)
            motor = create_motor_analysis_service(with_persistence=False)
            label = next(c.label for c in motor.input_contract.categories
                         if c.category_id == "three_phase_induction")
            for row, quantity, frequency, poles in ((4, 2, 50, 4), (5, 3, "5~120", 4),
                                                     (6, 4, 50, 10), (7, 5, 50, 3)):
                for column, value in {"D": quantity, "F": label, "G": 0.38,
                                      "H": frequency, "J": 0.14, "K": poles, "M": 71.43}.items():
                    workbook["电动机"][f"{column}{row}"] = value
            workbook.save(source)
            workbook.close()
            before = source.read_bytes()
            result = batch.evaluate_workbook(source)
            self.assertEqual(source.read_bytes(), before)
            self.assertEqual(result.summary.total_quantity, 14)
            self.assertEqual(result.summary.evaluated_quantity, 9)
            self.assertEqual(result.summary.input_error_quantity, 5)
            self.assertEqual(result.summary.conclusion_quantities, {"3级": 2, "不在范围": 7})
            output = load_workbook(result.result_workbook, data_only=True)
            self.assertEqual([output["电动机"][f"Q{r}"].value for r in range(4, 8)],
                             ["3级", "不在范围", "不在范围", None])
            self.assertIn("输入错误", output["电动机"]["T7"].value)
            self.assertIn("正偶数", output["电动机"]["T7"].value)
            self.assertEqual(output["电动机"]["H5"].value, "5~120")
            self.assertEqual(output["电动机"]["K6"].value, 10)
            output.close()
            record = SqliteBatchRecordRepository(paths.records_db).load_batch_record(result.batch_record_id)
            self.assertEqual(record.summary["conclusion_quantities"], result.summary.conclusion_quantities)

    def test_candidate_switch_alone_launches_desktop(self):
        from equipeffi.entrypoint import main
        with patch.dict(os.environ, {"EQUIPEFFI_ENABLE_MOTOR_CANDIDATE": "0"}), patch(
                "equipeffi.composition.launch_qt", return_value=0) as launch:
            self.assertEqual(main(["--motor-candidate"]), 0)
            launch.assert_called_once_with()
            self.assertTrue(motor_candidate_enabled())

    def test_motor_defaults_on_and_legacy_environment_remains_compatible(self):
        self.assertTrue(motor_candidate_enabled())
        with patch.dict(os.environ, {"EQUIPEFFI_ENABLE_MOTOR_CANDIDATE": "0"}):
            self.assertFalse(motor_candidate_enabled())
        with patch.dict(os.environ, {"EQUIPEFFI_ENABLE_MOTOR_CANDIDATE": "1"}):
            self.assertTrue(motor_candidate_enabled())

    def test_only_unconnected_sheet_does_not_create_successful_batch(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = AppDataPaths(Path(directory) / "data")
            batch = create_batch_evaluation_service(paths=paths)
            source = Path(directory) / "unsupported.xlsx"
            destination = Path(directory) / "result.xlsx"
            batch.template_resource.download_to(source)
            workbook = load_workbook(source)
            workbook["空压机"]["B4"] = "保留未接通设备"
            workbook["空压机"]["D4"] = 1
            workbook.save(source)
            workbook.close()
            with self.assertRaisesRegex(ValueError, "没有可评价数据"):
                batch.evaluate_workbook(source, destination=destination)
            self.assertFalse(destination.exists())
            self.assertEqual(batch.batch_repository.list_batch_records(), [])

    def test_three_device_batch_and_reimport_preserve_exact_result(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = AppDataPaths(Path(directory) / "data")
            batch = create_batch_evaluation_service(paths=paths)
            source = Path(directory) / "input.xlsx"
            batch.template_resource.download_to(source)
            workbook = load_workbook(source)
            motor = create_motor_analysis_service(with_persistence=False)
            category = next(c for c in motor.input_contract.categories
                            if c.category_id == "three_phase_induction")
            for cell, value in {"D4": 2, "F4": category.label, "G4": 0.38,
                                "H4": 50, "J4": 0.14, "K4": 4, "M4": 71.43}.items():
                workbook["电动机"][cell] = value
            workbook["离心泵"]["D4"] = 3
            workbook["离心泵"]["F4"] = "其他类别"
            workbook["变压器"]["D4"] = 4
            workbook["变压器"]["F4"] = "其他类别"
            workbook.save(source)
            workbook.close()
            before = source.read_bytes()
            result = batch.evaluate_workbook(source)
            self.assertEqual(source.read_bytes(), before)
            self.assertEqual(result.batch_record_error, "")
            self.assertEqual(result.summary.total_quantity, 9)
            self.assertEqual(result.summary.conclusion_quantities, {"3级": 2, "不在范围": 7})
            output = load_workbook(result.result_workbook, data_only=True)
            self.assertEqual(output["电动机"]["Q4"].value, "3级")
            self.assertEqual(float(output["电动机"]["O4"].value), 71.43)
            self.assertIn("电动机", output["模板说明"]["B30"].value)
            output.close()
            record = SqliteBatchRecordRepository(paths.records_db).load_batch_record(result.batch_record_id)
            self.assertEqual(record.standard_code, "MULTI")
            self.assertEqual(len(record.summary["sheets"]), 3)
            from equipeffi.application.services.batch_summary_formatter import format_device_summary_lines
            device_lines = format_device_summary_lines(record.summary["sheets"])
            self.assertEqual(device_lines, format_device_summary_lines(result.summary.sheets))
            for device, quantity, conclusion in (("离心泵", 3, "不在范围"),
                                                  ("变压器", 4, "不在范围"),
                                                  ("电动机", 2, "3级")):
                index = device_lines.index(f"{device}：共 {quantity} 台（数据 1 行）")
                self.assertIn(f"{conclusion} {quantity} 台", device_lines[index + 1])
                self.assertIn("2级 0 台", device_lines[index + 1])
            repeated = batch.evaluate_workbook(Path(result.result_workbook))
            self.assertEqual(repeated.summary.conclusion_quantities, result.summary.conclusion_quantities)
            with self.assertRaisesRegex(ValueError, "未启用"):
                create_batch_evaluation_service(paths=paths, motor_candidate=False).evaluate_workbook(source)

    def test_efficiency_fraction_is_input_error_not_record(self):
        with tempfile.TemporaryDirectory() as directory:
            motor = create_motor_analysis_service(paths=AppDataPaths(Path(directory)))
            values = dict(category_id="three_phase_induction", rated_voltage_kv="0.38",
                          rated_frequency_hz="50", rated_output_power_kw="0.14", poles="4",
                          measured_efficiency_pct="0.85")
            outcome = motor.analyze_and_record(MotorAnalysisRequest(values))
            self.assertEqual(outcome.result.evaluation_status, "INVALID_INPUT")
            self.assertNotEqual(outcome.record_status, "SAVED")
            values["measured_efficiency_pct"] = "71.43"
            accepted = motor.analyze_and_record(MotorAnalysisRequest(values))
            self.assertEqual(accepted.result.ui_conclusion, "3级")
            self.assertEqual(accepted.record_status, "SAVED")
            threshold = accepted.result.thresholds["grade_2"]
            self.assertEqual(threshold["exact"]["numerator"], "2143")
            self.assertEqual(threshold["exact"]["denominator"], "30")
