"""变压器查表依据面向用户的投影回归。"""
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication, QComboBox, QLineEdit

from equipeffi.application.services.transformer_analysis_service import (
    TransformerAnalysisRequest,
)
from equipeffi.application.services.transformer_basis_formatter import (
    format_transformer_basis_lines,
)
from equipeffi.composition import (
    create_batch_evaluation_service,
    create_pump_analysis_service,
    create_transformer_analysis_service,
)
from equipeffi.domain.devices.transformer.input_contract import (
    TRANSFORMER_INPUT_CONTRACT,
    TRANSFORMER_SCOPE_HINT,
)
from equipeffi.infrastructure.persistence.app_data_paths import AppDataPaths
from equipeffi.presentation.qt.pages.records import RecordsPage
from equipeffi.presentation.qt.pages.transformer_analysis import (
    TransformerAnalysisPage,
)


class _HistoryService:
    def __init__(self, record):
        self.record = record

    def record_categories(self):
        return ()

    def search_records(self, query):
        return SimpleNamespace(records=(), total=0, offset=query.offset)

    def list_records(self, limit=5):
        return ()

    def open_record(self, _record_id):
        return self.record


class TransformerBasisProjectionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.widgets = []

    def tearDown(self):
        for widget in self.widgets:
            widget.close()
            widget.deleteLater()
        self.app.processEvents()

    def _track(self, widget):
        self.widgets.append(widget)
        return widget

    @staticmethod
    def _t32_values(*, losses=True):
        values = {
            "category_id": "GB20052-T32",
            "standard_conditions": {
                "renewable_type": "photovoltaic", "thermal_class": "F"},
            "capacity_kva": "3000",
            "hv_rated_voltage_kv": "10.5",
            "lv_rated_voltage_kv": "0.69",
        }
        if losses:
            values.update(no_load_loss_w="2400", load_loss_w="18000")
        return values

    def test_real_application_interpolation_is_readable_and_hides_trace_ids(self):
        with tempfile.TemporaryDirectory() as directory:
            service = create_transformer_analysis_service(
                paths=AppDataPaths(Path(directory) / "data"), with_persistence=False)
            result = service.evaluate(TransformerAnalysisRequest(self._t32_values()))

        self.assertEqual(result.ui_conclusion, "1级")
        lines = format_transformer_basis_lines(result.trace)
        self.assertEqual(len(lines), 1)
        text = lines[0]
        for expected in ("表32", "分支为F级（120℃）", "3000 kVA",
                         "2500～3150 kVA", "表32第26、29行"):
            self.assertIn(expected, text)
        for internal in ("GB20052-T32", "GB20052-R000479", "source_row_id",
                         "cell_coordinates", "thermal=F", "{", "}"):
            self.assertNotIn(internal, text)

    def test_t27_and_t28_same_value_multimatches_keep_every_source(self):
        with tempfile.TemporaryDirectory() as directory:
            service = create_transformer_analysis_service(
                paths=AppDataPaths(Path(directory) / "data"), with_persistence=False)
            for table_no in (27, 28):
                with self.subTest(table_no=table_no):
                    result = service.evaluate(TransformerAnalysisRequest({
                        "category_id": f"GB20052-T{table_no:02d}",
                        "capacity_kva": "120000",
                        "uk_hm_rated_percent": "12.9",
                        "uk_hl_rated_percent": "44.2",
                        "uk_ml_rated_percent": "29.1",
                        "no_load_loss_w": "0",
                        "load_loss_w": "0",
                    }))
                    self.assertEqual(result.ui_conclusion, "1级")
                    self.assertEqual(len(result.trace), 2)
                    lines = format_transformer_basis_lines(result.trace)
                    text = "\n".join(lines)
                    for expected in (f"表{table_no}", "额定阻抗B组", "额定阻抗C组",
                                     "表" + str(table_no) + "第5行",
                                     "表" + str(table_no) + "第6行"):
                        self.assertIn(expected, text)
                    self.assertNotIn("GB20052-T", text)
                    self.assertNotIn("GB20052-R", text)

    def test_real_excel_result_appends_basis_and_keeps_missing_data_message(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = AppDataPaths(Path(directory) / "data")
            batch = create_batch_evaluation_service(paths=paths)
            source = Path(directory) / "input.xlsx"
            batch.template_resource.download_to(source)
            category = next(option.label for option in
                            TRANSFORMER_INPUT_CONTRACT.categories
                            if option.category_id == "GB20052-T32")
            workbook = load_workbook(source)
            sheet = workbook["变压器"]
            for row, losses in ((4, (2400, 18000)), (5, (None, None))):
                sheet[f"D{row}"] = 1
                sheet[f"F{row}"] = category
                sheet[f"G{row}"] = 3000
                sheet[f"H{row}"] = "光伏，F级（120℃）"
                sheet[f"J{row}"], sheet[f"K{row}"] = losses
                sheet[f"V{row}"] = 10.5
                sheet[f"W{row}"] = 0.69
            workbook.save(source)
            workbook.close()

            outcome = batch.evaluate_workbook(source, persist=False)
            result = load_workbook(outcome.result_workbook, data_only=True)
            try:
                full_message = result["变压器"]["U4"].value
                missing_message = result["变压器"]["U5"].value
                self.assertEqual(result["变压器"]["R4"].value, "1级")
                self.assertEqual(result["变压器"]["R5"].value, "无法判定")
            finally:
                result.close()

        self.assertIn("空载损耗与负载损耗均不高于1级对应限值。", full_message)
        self.assertIn("查表依据：表32", full_message)
        self.assertIn("2500～3150 kVA", full_message)
        self.assertIn("表32第26、29行", full_message)
        self.assertIn("已取得六项限值，但正式判级还需空载损耗和负载损耗。",
                      missing_message)
        self.assertIn("查表依据：表32", missing_message)
        for internal in ("GB20052-T32", "GB20052-R000479", "source_row_id",
                         "cell_coordinates", "thermal=F"):
            self.assertNotIn(internal, full_message)
            self.assertNotIn(internal, missing_message)

    def test_real_qt_and_saved_history_show_basis_without_recalculation(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = AppDataPaths(Path(directory) / "data")
            transformer = create_transformer_analysis_service(paths=paths)
            page = self._track(TransformerAnalysisPage(transformer))
            self.assertEqual(page.category_help.text(), TRANSFORMER_SCOPE_HINT)

            page.category.setCurrentIndex(
                page.category.findData("GB20052-T32"))
            self.assertEqual(page.category_help.text(), TRANSFORMER_SCOPE_HINT)
            condition = next(option for option in
                             transformer.input_contract.condition_options("GB20052-T32")
                             if option.values.get("thermal_class") == "F")
            for field_id, value in self._t32_values().items():
                if field_id in {"category_id", "standard_conditions"}:
                    continue
                control = page.field_inputs[field_id]
                self.assertIsInstance(control, QLineEdit)
                control.setText(value)
            conditions = page.field_inputs["standard_conditions"]
            self.assertIsInstance(conditions, QComboBox)
            conditions.setCurrentIndex(conditions.findData(condition.condition_id))

            result = page.evaluate()
            self.assertEqual(result.ui_conclusion, "1级")
            current_text = page.explanation.text()
            self.assertIn("分支为F级（120℃）", current_text)
            self.assertIn("2500～3150 kVA", current_text)
            self.assertIn("表32第26、29行", current_text)
            self.assertNotIn("GB20052-T32", current_text)
            record_id = page._last_outcome.record.record_id

            pump = create_pump_analysis_service(paths=paths)
            history = self._track(RecordsPage(pump, transformer=transformer))
            with patch.object(transformer, "evaluate",
                              side_effect=AssertionError("历史详情不得重算")):
                historical_text = history.show_record(record_id)
            self.assertIn("分支为F级（120℃）", historical_text)
            self.assertIn("2500～3150 kVA", historical_text)
            self.assertIn("表32第26、29行", historical_text)
            self.assertNotIn("GB20052-T32", historical_text)
            self.assertNotIn("GB20052-R000479", historical_text)

    def test_legacy_history_without_trace_keeps_business_message_without_fabrication(self):
        category = next(option for option in TRANSFORMER_INPUT_CONTRACT.categories
                        if option.category_id == "GB20052-T32")
        snapshot = SimpleNamespace(
            record_id="old-transformer-record",
            standard_code="GB 20052-2024",
            device_type="transformer",
            product_category=category.category_id,
            as_of="2025-01-01",
            grade=None,
            evaluation_status="INSUFFICIENT_DATA",
            ui_conclusion="无法判定",
            input_snapshot={"category_id": category.category_id,
                            "category_label": category.label},
            result_snapshot={
                "missing_fields": ["load_loss_w"],
                "explanation": "请补充负载损耗。",
            },
            reference_snapshot={},
            finalized_at_utc="2025-01-01T00:00:00Z",
        )
        transformer = create_transformer_analysis_service(with_persistence=False)
        history = self._track(RecordsPage(
            _HistoryService(snapshot), transformer=transformer))

        text = history.show_record(snapshot.record_id)
        self.assertIn("缺失信息：负载损耗", text)
        self.assertIn("判定说明：请补充负载损耗。", text)
        self.assertNotIn("查表依据", text)
        self.assertNotIn("第26行", text)
        self.assertNotIn("第29行", text)

    def test_empty_or_missing_trace_does_not_invent_sources(self):
        self.assertEqual(format_transformer_basis_lines(None), ())
        self.assertEqual(format_transformer_basis_lines([]), ())
        self.assertEqual(format_transformer_basis_lines([{"branch_id": "internal"}]), ())

    def test_formatter_preserves_capacity_digits_and_zero_padded_table_number(self):
        line = format_transformer_basis_lines([{
            "table_id": "GB20052-T03",
            "branch_id": "GB20052-T03/default",
            "lookup_method": "interpolated",
            "capacity_kva": "3000.12345678901234567890123456789",
            "lower_capacity_kva": "2500",
            "upper_capacity_kva": "3150",
            "sources": [{"source_sheet": "表03", "source_row": 26}],
        }])[0]
        self.assertIn("表03", line)
        self.assertIn("3000.12345678901234567890123456789 kVA", line)
        self.assertIn("来源表03第26行", line)


if __name__ == "__main__":
    unittest.main()
